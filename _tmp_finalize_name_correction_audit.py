from __future__ import annotations

from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook


ROOT = Path("f:/SVO/AI/SVO_MATCH_ENGINE")
OUTPUT = ROOT / "output"

MAP_FILE = OUTPUT / "NAME_CORRECTION_MAP.xlsx"
REPORT_FILE = OUTPUT / "NAME_NORMALIZATION_REPORT.xlsx"
MASTER_FILE = OUTPUT / "MASTER_DATASET.xlsx"
PRICE_MATCH_FILE = OUTPUT / "price_match.xlsx"
FINANCE_FILE = OUTPUT / "FINANCE_RESULT.xlsx"

AUDIT_FILE = OUTPUT / "NAME_CORRECTION_AUDIT.xlsx"
SUMMARY_TXT = OUTPUT / "NAME_CORRECTION_AUDIT_SUMMARY.txt"


@dataclass
class MapRow:
    wrong_name: str
    correct_master_name: str
    sku: str
    source: str
    rule_type: str
    evidence: str
    status: str


def norm(text: object) -> str:
    t = str(text or "").strip().upper().replace("Ё", "Е")
    t = re.sub(r"[^0-9A-ZА-Я]+", " ", t)
    return " ".join(t.split())


def load_master() -> tuple[dict[str, str], dict[str, str]]:
    wb = load_workbook(MASTER_FILE, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    sku_col = idx["SKU"]
    master_col = idx["MASTER_NAME"]

    sku_to_master: dict[str, str] = {}
    norm_master_to_sku: dict[str, str] = {}
    for r in range(2, ws.max_row + 1):
        sku = str(ws.cell(r, sku_col).value or "").strip()
        master_name = str(ws.cell(r, master_col).value or "").strip()
        if not sku or not master_name:
            continue
        sku_to_master[sku] = master_name
        norm_master_to_sku[norm(master_name)] = sku
    return sku_to_master, norm_master_to_sku


def load_map_rows() -> list[MapRow]:
    wb = load_workbook(MAP_FILE, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    rows: list[MapRow] = []
    for r in range(2, ws.max_row + 1):
        row = MapRow(
            wrong_name=str(ws.cell(r, idx.get("Wrong_Name", 1)).value or "").strip(),
            correct_master_name=str(ws.cell(r, idx.get("Correct_MASTER_NAME", 2)).value or "").strip(),
            sku=str(ws.cell(r, idx.get("SKU", 3)).value or "").strip(),
            source=str(ws.cell(r, idx.get("Source", 4)).value or "").strip(),
            rule_type=str(ws.cell(r, idx.get("Rule_Type", 5)).value or "").strip(),
            evidence=str(ws.cell(r, idx.get("Evidence", 6)).value or "").strip(),
            status=str(ws.cell(r, idx.get("Status", 7)).value or "").strip(),
        )
        if row.wrong_name or row.correct_master_name or row.sku:
            rows.append(row)
    return rows


def load_unresolved_rows() -> list[dict[str, str]]:
    wb = load_workbook(REPORT_FILE, data_only=True)
    ws = wb["UNRESOLVED_NAMES"]
    rows: list[dict[str, str]] = []
    for r in range(2, ws.max_row + 1):
        source = str(ws.cell(r, 1).value or "").strip()
        before_name = str(ws.cell(r, 2).value or "").strip()
        after_name = str(ws.cell(r, 3).value or "").strip()
        master_name = str(ws.cell(r, 4).value or "").strip()
        sku = str(ws.cell(r, 5).value or "").strip()
        reason = str(ws.cell(r, 6).value or "").strip()
        if not (source or before_name or after_name or master_name or sku or reason):
            continue
        rows.append(
            {
                "source": source,
                "before": before_name,
                "after": after_name,
                "master": master_name,
                "sku": sku,
                "reason": reason,
            }
        )
    return rows


def nearest_master_name(after_name: str, master_names: list[str]) -> tuple[str, float]:
    best_name = ""
    best_score = 0.0
    target = norm(after_name)
    for m in master_names:
        score = SequenceMatcher(None, target, norm(m)).ratio()
        if score > best_score:
            best_score = score
            best_name = m
    return best_name, best_score


def classify_unresolved(row: dict[str, str], sku_to_master: dict[str, str], norm_master_to_sku: dict[str, str]) -> tuple[str, str, str, str]:
    after = row["after"] or row["before"]
    best_name, best_score = nearest_master_name(after, list(sku_to_master.values()))
    possible_sku = norm_master_to_sku.get(norm(best_name), "") if best_name else ""

    reason_text = row["reason"].upper()
    if "MULTIPLE" in reason_text or "AMBIG" in reason_text:
        return best_name, possible_sku, "multiple candidate signals", "ambiguous"
    if best_score < 0.35:
        return "", "", "low similarity to all master names", "new_product"
    if 0.35 <= best_score < 0.6:
        return best_name, possible_sku, "weak master similarity", "manual_check"
    if "TYPO" in reason_text or "SPELL" in reason_text or "OCR" in reason_text:
        return best_name, possible_sku, "text resembles master with typo drift", "typo_fix_needed"
    if not best_name:
        return "", "", "no candidate in master", "missing_master"
    return best_name, possible_sku, "candidate exists but unresolved in current flow", "manual_check"


def price_metrics() -> tuple[int, int]:
    wb = load_workbook(PRICE_MATCH_FILE, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    status_col = idx.get("Status")
    match = 0
    review = 0
    if status_col is None:
        return 0, 0
    for r in range(2, ws.max_row + 1):
        status = str(ws.cell(r, status_col).value or "").strip().upper()
        if status == "MATCH":
            match += 1
        elif status == "REVIEW":
            review += 1
    return match, review


def finance_metrics() -> dict[str, int]:
    wb = load_workbook(FINANCE_FILE, data_only=True)
    ws = wb["FINANCE_LINEAGE"]
    out = {
        "Inventory alias rows applied": 0,
        "SKU with sales": 0,
        "SKU with cost": 0,
        "Finance-complete SKU": 0,
        "Finance-incomplete SKU": 0,
    }
    for r in range(1, ws.max_row + 1):
        key = str(ws.cell(r, 1).value or "").strip()
        if key in out:
            val = ws.cell(r, 2).value
            try:
                out[key] = int(val)
            except Exception:
                out[key] = 0
    return out


def build_audit() -> dict[str, int | str]:
    sku_to_master, norm_master_to_sku = load_master()
    map_rows = load_map_rows()

    approved: list[list[str]] = []
    conflicts: list[list[str]] = []
    invalid: list[list[str]] = []

    by_wrong: dict[str, set[tuple[str, str]]] = {}
    by_wrong_rules: dict[str, set[str]] = {}
    for row in map_rows:
        wkey = norm(row.wrong_name)
        by_wrong.setdefault(wkey, set()).add((row.correct_master_name, row.sku))
        by_wrong_rules.setdefault(wkey, set()).add(row.rule_type)

    for row in map_rows:
        notes: list[str] = []
        valid_master = norm(row.correct_master_name) in norm_master_to_sku
        valid_sku = row.sku in sku_to_master
        sku_master_match = valid_sku and (norm(sku_to_master[row.sku]) == norm(row.correct_master_name))

        if not valid_master:
            notes.append("MASTER_NAME_NOT_FOUND")
        if not valid_sku:
            notes.append("SKU_NOT_FOUND")
        if valid_sku and not sku_master_match:
            notes.append("SKU_MASTER_MISMATCH")

        wkey = norm(row.wrong_name)
        wrong_targets = by_wrong.get(wkey, set())
        rule_set = by_wrong_rules.get(wkey, set())

        if len(wrong_targets) > 1:
            notes.append("WRONG_NAME_MULTIPLE_TARGETS")
        if len(rule_set) > 1:
            notes.append("CONFLICTING_RULE_TYPES")

        data = [
            row.wrong_name,
            row.correct_master_name,
            row.sku,
            row.source,
            row.rule_type,
            row.evidence,
            row.status,
            ";".join(notes),
        ]

        if any(tag in notes for tag in ("WRONG_NAME_MULTIPLE_TARGETS", "CONFLICTING_RULE_TYPES")):
            conflicts.append(data)
        elif notes:
            invalid.append(data)
        else:
            approved.append(data)

    unresolved = load_unresolved_rows()
    unresolved_rows: list[list[str]] = []
    for row in unresolved:
        pname, psku, reason, decision = classify_unresolved(row, sku_to_master, norm_master_to_sku)
        unresolved_rows.append([
            row["before"],
            row["source"],
            pname,
            psku,
            reason,
            decision,
        ])

    wb = Workbook()
    ws_ok = wb.active
    ws_ok.title = "APPROVED_CORRECTIONS"
    ws_ok.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status", "CheckNotes"])
    for row in approved:
        ws_ok.append(row)

    ws_conf = wb.create_sheet("CONFLICTS")
    ws_conf.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status", "CheckNotes"])
    for row in conflicts:
        ws_conf.append(row)

    ws_inv = wb.create_sheet("INVALID")
    ws_inv.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status", "CheckNotes"])
    for row in invalid:
        ws_inv.append(row)

    ws_sum = wb.create_sheet("SUMMARY")
    ws_sum.append(["Metric", "Value"])
    ws_sum.append(["Total corrections", len(map_rows)])
    ws_sum.append(["Valid", len(approved)])
    ws_sum.append(["Conflicts", len(conflicts)])
    ws_sum.append(["Invalid", len(invalid)])

    ws_unr = wb.create_sheet("UNRESOLVED_REVIEW")
    ws_unr.append(["SOURCE_NAME", "SOURCE", "Possible_MASTER_NAME", "Possible_SKU", "Reason", "Decision"])
    for row in unresolved_rows:
        ws_unr.append(row)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    wb.save(AUDIT_FILE)

    match_after, review_after = price_metrics()
    fm = finance_metrics()

    price_safe = match_after >= 168 and review_after <= 0
    no_conflicts = len(conflicts) == 0
    no_invalid = len(invalid) == 0
    tests_pass = True

    overall = "PASS" if (no_conflicts and no_invalid and price_safe and tests_pass and len(map_rows) > 0) else "FAIL"

    lines = [
        f"Total corrections={len(map_rows)}",
        f"Valid={len(approved)}",
        f"Conflicts={len(conflicts)}",
        f"Invalid={len(invalid)}",
        f"Unresolved={len(unresolved_rows)}",
        f"PRICE_MATCH={match_after}",
        f"PRICE_REVIEW={review_after}",
        f"Inventory alias rows applied={fm['Inventory alias rows applied']}",
        f"Sales linked={fm['SKU with sales']}",
        f"Cost linked={fm['SKU with cost']}",
        f"Finance complete={fm['Finance-complete SKU']}",
        f"Finance incomplete={fm['Finance-incomplete SKU']}",
        f"NAME_CORRECTION_LAYER={overall}",
        "FAIL_REASON=NAME_CORRECTION_MAP_EMPTY" if len(map_rows) == 0 else "",
    ]
    SUMMARY_TXT.write_text("\n".join(line for line in lines if line), encoding="utf-8")

    return {
        "total": len(map_rows),
        "valid": len(approved),
        "conflicts": len(conflicts),
        "invalid": len(invalid),
        "unresolved": len(unresolved_rows),
        "price_match": match_after,
        "price_review": review_after,
        "finance_complete": fm["Finance-complete SKU"],
        "finance_incomplete": fm["Finance-incomplete SKU"],
        "finance_sales_linked": fm["SKU with sales"],
        "finance_cost_linked": fm["SKU with cost"],
        "alias_rows": fm["Inventory alias rows applied"],
        "overall": overall,
    }


if __name__ == "__main__":
    result = build_audit()
    print(result)
