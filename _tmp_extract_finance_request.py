import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


TARGET_LABELS = {
    "Total SKU",
    "SKU with stock",
    "Finance-complete SKU",
    "Missing MASTER for PRICE SKU",
    "Inventory alias rows applied",
}


def find_metrics(sheet):
    metrics = {}
    nonempty_rows = []
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = list(row)
        if any(value not in (None, "") for value in values):
            nonempty_rows.append((row_idx, values))
        for cell_idx, value in enumerate(values):
            if isinstance(value, str) and value.strip() in TARGET_LABELS:
                label = value.strip()
                next_value = None
                for candidate in values[cell_idx + 1 :]:
                    if candidate not in (None, ""):
                        next_value = candidate
                        break
                metrics[label] = next_value
    return metrics, nonempty_rows


def main():
    base_dir = Path(__file__).resolve().parent
    finance_path = base_dir / "output/FINANCE_RESULT.xlsx"
    price_path = base_dir / "output/price_match.xlsx"
    output_path = base_dir / "_tmp_finance_request_output.json"

    finance_wb = load_workbook(finance_path, data_only=True, read_only=True)
    finance_ws = finance_wb["FINANCE_LINEAGE"]
    metrics, nonempty_rows = find_metrics(finance_ws)
    finance_wb.close()

    price_wb = load_workbook(price_path, data_only=True, read_only=True)
    price_ws = price_wb.active
    rows = price_ws.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    counts = Counter()
    if "Status" in header:
        status_idx = header.index("Status")
        for row in rows:
            if row is None or all(value is None for value in row):
                continue
            value = row[status_idx] if status_idx < len(row) else None
            counts[("" if value is None else str(value).strip().upper())] += 1
    price_wb.close()

    payload = {
        "finance_metrics": {label: metrics.get(label) for label in sorted(TARGET_LABELS)},
        "missing_metric_rows": [],
        "price_status_counts": dict(sorted(counts.items())),
    }

    if any(payload["finance_metrics"][label] is None for label in TARGET_LABELS):
        payload["missing_metric_rows"] = [
            {"row": row_idx, "values": values}
            for row_idx, values in nonempty_rows[:120]
        ]

    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(output_path)


if __name__ == "__main__":
    main()