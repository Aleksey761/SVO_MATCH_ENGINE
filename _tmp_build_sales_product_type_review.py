from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import Workbook, load_workbook

INPUT_REVIEW_PATH = Path("output/SALES_REVIEW_89.xlsx")
INPUT_MASTER_PATH = Path("output/MASTER_DATASET.xlsx")
OUTPUT_PATH = Path("output/SALES_PRODUCT_TYPE_REVIEW.xlsx")

REVIEW_SHEET_NAME = "REVIEW_89"
OUTPUT_REVIEW_SHEET_NAME = "REVIEW_79"

OUTPUT_COLUMNS = [
    "SOURCE_ROW",
    "НАИМЕНОВАНИЕ",
    "MATCH_REASONS",
    "POSSIBLE_MASTER_NAME",
    "POSSIBLE_SKU",
    "MASTER_PRODUCT_TYPE",
    "SOURCE_PRODUCT_TYPE",
    "DIAGNOSIS",
]

DIAGNOSIS_LABELS = [
    "TYPE_REALLY_DIFFERENT",
    "TYPE_NAME_DIFFERENCE",
    "SAME_PRODUCT_DIFFERENT_WORDING",
    "NO_OBVIOUS_MASTER_MATCH",
    "NEED_MANUAL_CHECK",
]

STOPWORDS = {
    "ДЛЯ",
    "И",
    "ИЛИ",
    "С",
    "БЕЗ",
    "НА",
    "В",
    "ПО",
    "ОТ",
    "ДО",
    "К",
    "У",
    "ПРИ",
    "ПОД",
    "НАД",
    "ПРО",
    "СО",
    "А",
    "THE",
    "AND",
    "FOR",
    "WITH",
}

EXAMPLE_BRANDS = {"SVO", "GILAR", "PEARLIS", "BOSSFIX"}

NON_ALNUM_RE = re.compile(r"[^0-9A-ZА-Я]+", re.IGNORECASE)
SPACE_RE = re.compile(r"\s+")
VOLUME_RE = re.compile(r"(\d+(?:[\.,]\d+)?)\s*(МЛ|Л|Г|КГ)\b", re.IGNORECASE)


HeaderMap = Dict[str, int]


def norm_text(text: object) -> str:
    if text is None:
        return ""
    s = str(text).upper().replace("Ё", "Е").strip()
    s = NON_ALNUM_RE.sub(" ", s)
    s = SPACE_RE.sub(" ", s).strip()
    return s


def norm_header(text: object) -> str:
    return norm_text(text)


def split_tokens(text: object) -> List[str]:
    n = norm_text(text)
    return [t for t in n.split(" ") if t]


def normalize_number(num_str: str) -> str:
    x = num_str.replace(",", ".")
    try:
        val = float(x)
    except ValueError:
        return x
    if val.is_integer():
        return str(int(val))
    return ("%f" % val).rstrip("0").rstrip(".")


def extract_volume_tokens(text: object) -> Set[str]:
    if text is None:
        return set()
    s = str(text).upper().replace("Ё", "Е")
    out: Set[str] = set()
    for m in VOLUME_RE.finditer(s):
        num = normalize_number(m.group(1))
        unit = m.group(2).upper()
        out.add(f"{num}{unit}")
    return out


def infer_source_product_type(name: object) -> str:
    s = " " + norm_text(name) + " "
    if " ПОРОШ" in s:
        return "СТИРАЛЬНЫЙ ПОРОШОК"
    if " КОНДИЦИОНЕР" in s:
        return "КОНДИЦИОНЕР"
    if " МЫЛО" in s:
        return "МЫЛО"
    if " ГЕЛЬ" in s and " ДУШ" in s:
        return "ГЕЛЬ ДЛЯ ДУША"
    return ""


def type_family(type_label: object) -> str:
    s = " " + norm_text(type_label) + " "
    if " ПОРОШ" in s:
        return "POWDER"
    if " КОНДИЦИОНЕР" in s or " ОПОЛАСКИВАТЕЛ" in s:
        return "CONDITIONER"
    if " МЫЛО" in s:
        return "SOAP"
    if " ГЕЛЬ" in s and " ДУШ" in s:
        return "SHOWER_GEL"
    return ""


def find_header(headers: HeaderMap, names: Sequence[str], required: bool = True) -> Optional[int]:
    for n in names:
        idx = headers.get(norm_header(n))
        if idx is not None:
            return idx
    if required:
        raise RuntimeError(f"Required header not found. Tried: {names}")
    return None


def meaningful_tokens(
    text: object,
    brand_tokens: Set[str],
    volume_tokens: Set[str],
    drop_product_type_words: bool = True,
) -> Set[str]:
    tokens = split_tokens(text)
    out: Set[str] = set()
    for t in tokens:
        if t in STOPWORDS:
            continue
        if t in brand_tokens:
            continue
        if len(t) <= 2:
            continue
        if t.isdigit():
            continue
        if drop_product_type_words and t in {"ПОРОШОК", "СТИРАЛЬНЫЙ", "КОНДИЦИОНЕР", "МЫЛО", "ГЕЛЬ", "ДУША", "ДУШ"}:
            continue
        out.add(t)
    for v in volume_tokens:
        v_token = norm_text(v)
        if v_token in out:
            out.discard(v_token)
    return out


def worksheet_headers(ws) -> HeaderMap:
    return {norm_header(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}


def cell(ws, row: int, col: Optional[int]) -> object:
    if not col:
        return None
    return ws.cell(row, col).value


def build() -> None:
    review_wb = load_workbook(INPUT_REVIEW_PATH, data_only=True)
    if REVIEW_SHEET_NAME not in review_wb.sheetnames:
        raise RuntimeError(f"Sheet {REVIEW_SHEET_NAME} not found in {INPUT_REVIEW_PATH}")
    review_ws = review_wb[REVIEW_SHEET_NAME]

    master_ws = load_workbook(INPUT_MASTER_PATH, data_only=True).active

    rh = worksheet_headers(review_ws)
    mh = worksheet_headers(master_ws)

    source_row_col = find_header(rh, ["SOURCE_ROW"])
    src_name_col = find_header(rh, ["НАИМЕНОВАНИЕ", "НАИМЕНОВАНИЕ ", "NAME", "PRODUCT_NAME", ""])
    reasons_col = find_header(rh, ["MATCH_REASONS"])

    sku_col = find_header(mh, ["SKU", "MASTER_SKU", "АРТИКУЛ", "КОД"], required=False)
    master_name_col = find_header(mh, ["MASTER_NAME", "NAME", "НАИМЕНОВАНИЕ", "MASTER_PRODUCT_NAME"])
    master_type_col = find_header(
        mh,
        ["CATEGORY", "PRODUCT_TYPE", "MASTER_PRODUCT_TYPE", "TYPE", "КАТЕГОРИЯ", "ТИП", "ВИД"],
        required=False,
    )
    master_brand_col = find_header(mh, ["BRAND", "БРЕНД", "ТОРГОВАЯ МАРКА"], required=False)
    master_aroma_col = find_header(mh, ["VARIANT", "AROMA", "АРОМАТ", "ОТДУШКА", "ВАРИАНТ"], required=False)
    master_volume_col = find_header(mh, ["VOLUME", "ОБЪЕМ", "ОБЬЕМ", "ВЕС", "WEIGHT"], required=False)

    known_brands = set(EXAMPLE_BRANDS)
    if master_brand_col:
        for r in range(2, master_ws.max_row + 1):
            known_brands.update(split_tokens(cell(master_ws, r, master_brand_col)))
    known_brands = {b for b in known_brands if len(b) >= 3}

    masters = []
    for r in range(2, master_ws.max_row + 1):
        master_name = cell(master_ws, r, master_name_col)
        if master_name is None or str(master_name).strip() == "":
            continue
        brand_tokens = set(split_tokens(cell(master_ws, r, master_brand_col)))
        master_volume_tokens = extract_volume_tokens(cell(master_ws, r, master_volume_col))
        if not master_volume_tokens:
            master_volume_tokens = extract_volume_tokens(master_name)

        aroma_base = cell(master_ws, r, master_aroma_col)
        aroma_tokens = meaningful_tokens(
            aroma_base if aroma_base not in (None, "") else master_name,
            brand_tokens=brand_tokens,
            volume_tokens=master_volume_tokens,
        )

        masters.append(
            {
                "sku": cell(master_ws, r, sku_col),
                "name": master_name,
                "type": cell(master_ws, r, master_type_col),
                "brand_tokens": brand_tokens,
                "volume_tokens": master_volume_tokens,
                "aroma_tokens": aroma_tokens,
            }
        )

    filtered_rows = []
    for r in range(2, review_ws.max_row + 1):
        reason = cell(review_ws, r, reasons_col)
        reason_trim = "" if reason is None else str(reason).strip()
        if reason_trim == "PRODUCT_TYPE_MISMATCH":
            filtered_rows.append(r)

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = OUTPUT_REVIEW_SHEET_NAME
    out_ws.append(OUTPUT_COLUMNS)

    diagnosis_counter: Counter[str] = Counter()
    preview_rows: List[List[object]] = []

    for r in filtered_rows:
        source_row = cell(review_ws, r, source_row_col)
        src_name = cell(review_ws, r, src_name_col)
        src_reason = cell(review_ws, r, reasons_col)

        src_brand_tokens = set(t for t in split_tokens(src_name) if t in known_brands)
        src_volume_tokens = extract_volume_tokens(src_name)
        src_aroma_tokens = meaningful_tokens(
            src_name,
            brand_tokens=src_brand_tokens,
            volume_tokens=src_volume_tokens,
        )
        source_type = infer_source_product_type(src_name)

        obvious_candidates = []
        for m in masters:
            brand_ok = True
            if src_brand_tokens:
                brand_ok = bool(src_brand_tokens & m["brand_tokens"])

            volume_ok = bool(src_volume_tokens and m["volume_tokens"] and (src_volume_tokens & m["volume_tokens"]))
            aroma_ok = bool(src_aroma_tokens and m["aroma_tokens"] and (src_aroma_tokens & m["aroma_tokens"]))

            if brand_ok and (volume_ok or aroma_ok):
                obvious_candidates.append(m)

        possible_master_name = ""
        possible_sku = ""
        master_product_type = ""

        if len(obvious_candidates) == 0:
            diagnosis = "NO_OBVIOUS_MASTER_MATCH"
        elif len(obvious_candidates) > 1:
            diagnosis = "NEED_MANUAL_CHECK"
        else:
            m = obvious_candidates[0]
            possible_master_name = "" if m["name"] is None else str(m["name"])
            possible_sku = "" if m["sku"] is None else str(m["sku"])
            master_product_type = "" if m["type"] is None else str(m["type"])

            if not source_type:
                diagnosis = "NEED_MANUAL_CHECK"
            elif not master_product_type.strip():
                diagnosis = "NEED_MANUAL_CHECK"
            else:
                n_source_type = norm_text(source_type)
                n_master_type = norm_text(master_product_type)
                if n_source_type == n_master_type:
                    diagnosis = "SAME_PRODUCT_DIFFERENT_WORDING"
                else:
                    fam_source = type_family(source_type)
                    fam_master = type_family(master_product_type)
                    if fam_source and fam_source == fam_master:
                        diagnosis = "TYPE_NAME_DIFFERENCE"
                    else:
                        diagnosis = "TYPE_REALLY_DIFFERENT"

        if diagnosis not in DIAGNOSIS_LABELS:
            raise RuntimeError(f"Unexpected diagnosis: {diagnosis}")

        row_out = [
            source_row,
            "" if src_name is None else str(src_name),
            "" if src_reason is None else str(src_reason).strip(),
            possible_master_name,
            possible_sku,
            master_product_type,
            source_type,
            diagnosis,
        ]
        out_ws.append(row_out)
        diagnosis_counter[diagnosis] += 1
        if len(preview_rows) < 20:
            preview_rows.append(row_out)

    summary_ws = out_wb.create_sheet("SUMMARY")
    summary_ws.append(["DIAGNOSIS", "COUNT"])
    for label in DIAGNOSIS_LABELS:
        summary_ws.append([label, diagnosis_counter.get(label, 0)])

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(OUTPUT_PATH)

    print(f"REVIEW_TOTAL|{len(filtered_rows)}")
    for label in DIAGNOSIS_LABELS:
        print(f"SUMMARY|{label}|{diagnosis_counter.get(label, 0)}")
    print("PREVIEW_HEADER|" + "|".join(OUTPUT_COLUMNS))
    for row in preview_rows:
        print("PREVIEW_ROW|" + "|".join("" if v is None else str(v) for v in row))
    print(f"OUTPUT_PATH|{OUTPUT_PATH.as_posix()}")


if __name__ == "__main__":
    build()
