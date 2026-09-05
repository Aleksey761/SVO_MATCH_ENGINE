from openpyxl import load_workbook

wb = load_workbook("output/SALES_FINAL_CONTROL_CHECK.xlsx", data_only=True)
s = wb["SUMMARY"]
n = wb["NEW_ROWS"]
r = wb["REVIEW_3"]

print("SHEETS|" + "|".join(wb.sheetnames))
print("SUMMARY_ROWS|" + str(s.max_row))
for i in range(1, s.max_row + 1):
    if s.cell(i, 1).value is None:
        continue
    print("SUMMARY|" + str(s.cell(i, 1).value) + "|" + str(s.cell(i, 2).value))

print("NEW_ROWS_COUNT|" + str(max(n.max_row - 1, 0)))
for i in range(2, n.max_row + 1):
    vals = ["" if n.cell(i, c).value is None else str(n.cell(i, c).value) for c in range(1, n.max_column + 1)]
    print("NEW_ROW|" + "|".join(vals))

print("REVIEW_COUNT|" + str(max(r.max_row - 1, 0)))
for i in range(2, r.max_row + 1):
    vals = ["" if r.cell(i, c).value is None else str(r.cell(i, c).value) for c in range(1, r.max_column + 1)]
    print("REVIEW_ROW|" + "|".join(vals))
