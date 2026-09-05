from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

from svo.business_rules import BusinessRules
from svo.engine import Engine
from svo.matcher import Matcher
from svo.models import effective_product_type

DATA_DIR = Path("data")
FAILURES_FILE = Path("output/EXACT_MATCH_FAILURES.xlsx")
OUT_FILE = Path("output/EXACT_MATCH_PIPELINE_TRACE.xlsx")
AUTO_RESOLVE_MULTIPLE_CANDIDATES = True
AUTO_MATCH_SCORE_DELTA = 20.0


def safe(value: object) -> str:
    return str(value or "").strip()


def norm(matcher: Matcher, value: object) -> str:
    return safe(matcher._normalize_value(safe(value)))


def key_for(candidate) -> str:
    return f"{safe(getattr(candidate, 'sku', None))}|{safe(getattr(candidate, 'normalized_key', None))}"


def row_map_by_number(items: list[object]) -> dict[int, object]:
    mapping: dict[int, object] = {}
    for item in items:
        row_number = int(getattr(item, "row_number", 0) or 0)
        if row_number > 0:
            mapping[row_number] = item
    return mapping


def read_failures(path: Path) -> list[dict[str, object]]:
    wb = load_workbook(path, data_only=True)
    ws = wb["EXACT_MATCH_FAILURES"] if "EXACT_MATCH_FAILURES" in wb.sheetnames else wb.active

    headers = [safe(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)]
    idx = {header.upper(): pos + 1 for pos, header in enumerate(headers) if header}

    row_col = idx.get("ROW")
    sku_col = idx.get("MASTER SKU".upper()) or idx.get("MASTER_SKU") or idx.get("SKU")
    if row_col is None or sku_col is None:
        raise ValueError("Failed to locate 'Row' and 'MASTER SKU' columns in EXACT_MATCH_FAILURES.xlsx")

    failures: list[dict[str, object]] = []
    for r in range(2, ws.max_row + 1):
        row_value = ws.cell(row=r, column=row_col).value
        sku_value = ws.cell(row=r, column=sku_col).value
        row_number = int(row_value or 0)
        sku = safe(sku_value)
        if row_number <= 0 or not sku:
            continue
        failures.append({"row": row_number, "exact_sku": sku})

    return failures


def _review_reason(engine: Engine, item) -> str:
    if safe(getattr(item, "status", None)).upper() != "REVIEW":
        return ""
    high_level, recommendation = engine._build_price_review_reason(item)
    raw_reasons = [safe(r) for r in (getattr(item, "review_reasons", []) or []) if safe(r)]
    detail = ", ".join(raw_reasons) if raw_reasons else "none"
    return f"{high_level} | raw={detail} | recommendation={recommendation}"


def trace_row(*, matcher: Matcher, engine: Engine, master_by_sku: dict[str, object], item, exact_sku: str) -> dict[str, object]:
    exact_master = master_by_sku.get(exact_sku)

    disappeared_stage = ""
    disappeared_why = ""

    # Stage 1. Arrival normalization
    if exact_master is None:
        stage1_present = "NO"
        disappeared_stage = "Arrival normalization"
        disappeared_why = "Exact MASTER SKU not found in loaded MASTER file"
        normalization = "NO | exact MASTER SKU missing in MASTER"
        stage2_present = "NO"
        stage3_present = "NO"
        rank = ""
        selected = safe(getattr(item, "sku", None))
        final_reason = _review_reason(engine, item)
        return {
            "Normalization": normalization,
            "Candidate generated": stage2_present,
            "Candidate survived filtering": stage3_present,
            "Final candidate rank": rank,
            "Selected candidate": selected,
            "Final review reason": final_reason,
            "Pipeline stage where exact candidate disappeared": f"{disappeared_stage}: {disappeared_why}",
            "stage_loss": disappeared_stage,
        }

    arrival_pt = norm(matcher, effective_product_type(item))
    arrival_brand = norm(matcher, getattr(item, "brand", None))
    arrival_volume = norm(matcher, getattr(item, "volume", None))

    master_pt = norm(matcher, effective_product_type(exact_master))
    master_brand = norm(matcher, getattr(exact_master, "brand", None))
    master_volume = norm(matcher, getattr(exact_master, "volume", None))

    stage1_present = "YES"
    if not (arrival_pt and arrival_brand and arrival_volume):
        stage1_present = "NO"
        disappeared_stage = "Arrival normalization"
        disappeared_why = "Arrival normalized ProductType/Brand/Volume is incomplete"
    elif not (master_pt == arrival_pt and master_brand == arrival_brand and master_volume == arrival_volume):
        stage1_present = "NO"
        disappeared_stage = "Arrival normalization"
        disappeared_why = (
            "Normalized triple mismatch: "
            f"arrival=({arrival_pt}, {arrival_brand}, {arrival_volume}) vs "
            f"master=({master_pt}, {master_brand}, {master_volume})"
        )

    normalization = (
        f"{stage1_present} | arrival=({arrival_pt}, {arrival_brand}, {arrival_volume}) "
        f"master=({master_pt}, {master_brand}, {master_volume})"
    )

    # Stage 2. Candidate generation
    candidates = matcher._collect_candidates(item)
    candidate_keys = {key_for(candidate): candidate for candidate in candidates}
    exact_key = key_for(exact_master)
    generated = exact_key in candidate_keys
    stage2_present = "YES" if generated else "NO"
    if not generated and not disappeared_stage:
        disappeared_stage = "Candidate generation"
        disappeared_why = "Exact MASTER SKU did not enter candidate pool from indexed lookup/intersections"

    # Stage 3. Candidate filtering (hard rejection before final decision)
    scored_candidates = [
        {
            **matcher._score_detail(item, candidate),
            "candidate": candidate,
        }
        for candidate in candidates
    ]
    scored_candidates.sort(key=lambda entry: float(entry.get("score") or 0.0), reverse=True)

    exact_scored = None
    for entry in scored_candidates:
        if safe(entry.get("sku")) == exact_sku:
            exact_scored = entry
            break

    survived_filter = False
    if exact_scored is not None:
        rejection_reason = safe(exact_scored.get("rejection_reason"))
        if rejection_reason != "PRODUCT_TYPE_MISMATCH":
            survived_filter = True
        elif not disappeared_stage:
            disappeared_stage = "Candidate filtering"
            disappeared_why = "Exact candidate hard-rejected by PRODUCT_TYPE_MISMATCH"

    stage3_present = "YES" if survived_filter else "NO"
    if exact_scored is None and not disappeared_stage:
        disappeared_stage = "Candidate generation"
        disappeared_why = "Exact candidate absent from scored list because it was not generated"

    # Stage 4. Candidate ordering
    final_rank = ""
    if exact_scored is not None:
        for idx, entry in enumerate(scored_candidates, start=1):
            if safe(entry.get("sku")) == exact_sku:
                final_rank = idx
                break

    # Stage 5. Final selection (production matcher)
    match_result = matcher.match(item)
    selected_candidate = safe(getattr(match_result, "sku", None))

    if not disappeared_stage:
        if safe(getattr(match_result, "status", None)).upper() == "MATCH":
            if selected_candidate != exact_sku:
                disappeared_stage = "Final candidate selection"
                disappeared_why = (
                    "Another candidate selected by threshold/margin/preference rules "
                    f"(selected={selected_candidate or 'none'})"
                )
        else:
            disappeared_stage = "Final candidate selection"
            explanation = getattr(match_result, "review_explanation", {}) or {}
            rejected = explanation.get("best_candidate_rejected_reason")
            if isinstance(rejected, list) and rejected:
                rejected_text = "; ".join(safe(x) for x in rejected if safe(x))
            else:
                rejected_text = "match gate not satisfied"
            disappeared_why = f"No candidate selected for MATCH ({rejected_text})"

    # Stage 6. Review classification (business rules + review reason)
    final_reason = _review_reason(engine, item)

    return {
        "Normalization": normalization,
        "Candidate generated": stage2_present,
        "Candidate survived filtering": stage3_present,
        "Final candidate rank": final_rank,
        "Selected candidate": selected_candidate,
        "Final review reason": final_reason,
        "Pipeline stage where exact candidate disappeared": f"{disappeared_stage}: {disappeared_why}",
        "stage_loss": disappeared_stage,
    }


def main() -> None:
    if not FAILURES_FILE.exists():
        raise FileNotFoundError(f"Missing input workbook: {FAILURES_FILE}")

    engine = Engine()
    master_items = engine.loader.load_master(DATA_DIR / "MASTER.xlsx")
    master_by_sku = {safe(getattr(master, "sku", None)): master for master in master_items if safe(getattr(master, "sku", None))}

    price_candidates = sorted([p for p in DATA_DIR.glob("PRC*.xlsx") if p.is_file()], key=lambda p: p.stat().st_mtime)
    if not price_candidates:
        raise FileNotFoundError("No PRICE workbook found in data/")
    price_file = price_candidates[-1]

    price_items = engine.price_loader.load(price_file)
    for item in price_items:
        engine.normalizer.normalize(item)

    matcher = Matcher(
        master_items,
        auto_resolve_multiple_candidates=AUTO_RESOLVE_MULTIPLE_CANDIDATES,
        auto_match_score_delta=AUTO_MATCH_SCORE_DELTA,
    )

    unresolved_items, _, _ = engine._apply_learning_map_matches(price_items, master_items)
    matcher.match_all(unresolved_items)
    BusinessRules(master_items).apply(price_items)
    try:
        engine._validate_price_name_sku_conflicts(price_items)
        validation_note = "OK"
    except Exception as exc:
        validation_note = f"VALIDATION_SKIPPED: {exc}"

    by_row = row_map_by_number(price_items)
    failures = read_failures(FAILURES_FILE)
    if len(failures) != 15:
        print(f"WARNING: expected 15 rows in EXACT_MATCH_FAILURES.xlsx, found {len(failures)}")

    rows_out: list[list[object]] = []
    disappear_counts = Counter()

    for entry in failures:
        row_number = int(entry["row"])
        exact_sku = safe(entry["exact_sku"])

        item = by_row.get(row_number)
        if item is None:
            stage_loss = "Arrival normalization"
            disappear_counts[stage_loss] += 1
            rows_out.append(
                [
                    row_number,
                    exact_sku,
                    "NO | row is not present in current PRICE dataset",
                    "NO",
                    "NO",
                    "",
                    "",
                    "ROW_NOT_FOUND",
                    "Arrival normalization: PRICE row not found in current run",
                ]
            )
            print(f"Row {row_number} ({exact_sku})")
            print("  1. Arrival normalization: NO | row not found")
            print("  2. Candidate generation: NO")
            print("  3. Candidate filtering: NO")
            print("  4. Candidate ordering: NO")
            print("  5. Final selection: NO")
            print("  6. Review classification: ROW_NOT_FOUND")
            continue

        traced = trace_row(
            matcher=matcher,
            engine=engine,
            master_by_sku=master_by_sku,
            item=item,
            exact_sku=exact_sku,
        )
        disappear_counts[safe(traced["stage_loss"]) or "UNKNOWN"] += 1

        rows_out.append(
            [
                row_number,
                exact_sku,
                traced["Normalization"],
                traced["Candidate generated"],
                traced["Candidate survived filtering"],
                traced["Final candidate rank"],
                traced["Selected candidate"],
                traced["Final review reason"],
                traced["Pipeline stage where exact candidate disappeared"],
            ]
        )

        print(f"Row {row_number} ({exact_sku})")
        print(f"  1. Arrival normalization: {'YES' if str(traced['Normalization']).startswith('YES') else 'NO'}")
        print(f"  2. Candidate generation: {traced['Candidate generated']}")
        print(f"  3. Candidate filtering: {traced['Candidate survived filtering']}")
        print(f"  4. Candidate ordering: {'YES' if traced['Final candidate rank'] else 'NO'}")
        print(f"  5. Final selection: {'YES' if safe(traced['Selected candidate']) == exact_sku else 'NO'}")
        print(f"  6. Review classification: {safe(traced['Final review reason']) or 'NONE'}")
        print(f"  Removed stage: {traced['Pipeline stage where exact candidate disappeared']}")

    wb = Workbook()
    ws = wb.active
    ws.title = "EXACT_MATCH_PIPELINE_TRACE"
    ws.append(
        [
            "Row",
            "Exact MASTER SKU",
            "Normalization",
            "Candidate generated",
            "Candidate survived filtering",
            "Final candidate rank",
            "Selected candidate",
            "Final review reason",
            "Pipeline stage where exact candidate disappeared",
        ]
    )
    for row in rows_out:
        ws.append(row)

    ws_counts = wb.create_sheet("STAGE_DISAPPEAR_COUNTS")
    ws_counts.append(["Stage", "Rows"])
    ordered_stages = [
        "Arrival normalization",
        "Candidate generation",
        "Candidate filtering",
        "Candidate ordering",
        "Final candidate selection",
        "Review classification",
        "UNKNOWN",
    ]
    for stage in ordered_stages:
        if disappear_counts.get(stage, 0):
            ws_counts.append([stage, int(disappear_counts[stage])])

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)

    print("EXACT_MATCH_PIPELINE_TRACE_PATH", OUT_FILE.as_posix())
    print("VALIDATION", validation_note)
    print("DISAPPEAR_COUNTS")
    for stage in ordered_stages:
        count = int(disappear_counts.get(stage, 0))
        if count:
            print(stage, count)


if __name__ == "__main__":
    main()
