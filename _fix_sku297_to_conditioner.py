from openpyxl import load_workbook

TARGET = Path(r"output\MASTER_DATASET.xlsx")
SKU = "SKU-297"

wb = load_workbook(TARGET)
ws = wb.active

found = False

for row in ws.iter_rows(min_row=2):
    if str(row[0].value or "").strip() == SKU:
        row[1].value = "Кондиционер"
        row[2].value = "SVO"
        row[3].value = "Romantik Rose"
        row[4].value = "2,7 л"
        row[5].value = "Romantik Rose"
        row[9].value = "Кондиционер SVO Romantik Rose 2,7 л"
        found = True
        print(
            f"{SKU} | {row[1].value} | {row[2].value} | "
            f"{row[3].value} | {row[4].value} | {row[5].value} | "
            f"{row[9].value}"
        )
        break

if not found:
    raise RuntimeError(f"{SKU} NOT FOUND IN MASTER_DATASET.xlsx")

wb.save(TARGET)
print("SKU-297 CATEGORY CORRECTED: ЖМС 3в1 -> Кондиционер")
