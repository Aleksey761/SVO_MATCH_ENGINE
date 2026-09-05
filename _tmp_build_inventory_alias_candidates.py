from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


GENERIC_EXTRA_WORDS = {
    "ДЛЯ",
    "СТИРКИ",
    "ПАРФЮМИРОВАННЫЙ",
    "ПАРФЮМИРОВАННЫ",
    "ПАРФЮМ",
    "ПАРФЮМНЫЙ",
    "ПАРФЮМ НЫЙ",
    "АРОМАТ",
    "АРОМАТОВ",
    "СРЕДСТВО",
    "МОЮЩЕЕ",
    "ЖИДКОЕ",
    "ДЕТСКИХ",
    "ВЕЩЕЙ",
}

PACKAGING_TAIL_WORDS = {
    "КРАСНЫЙ",
    "СИНИЙ",
    "ЖЕЛТЫЙ",
    "ЧЕРНЫЙ",
    "БЕЛЫЙ",
    "РОЗОВЫЙ",
    "ПЭТ",
    "БУТ",
    "БУТЫЛКА",
    "КАНИСТРА",
    "ДОЙПАК",
    "ПАКЕТ",
    "ШТ",
    "ШТУК",
    "УП",
    "КОР",
    "АКЦИЯ",
}

ABBR_HINTS = {
    "ПАРФЮМ",
    "ПАРФ",
    "КОНД",
    "КОНДИЦ",
    "ЖМС",
    "ГД",
    "ШАМП",
    "БАЛЬЗ",
    "СР",
}


@dataclass
class CandidateRow:
    inventory_name: str
    nearest_master_name: str
    similarity: float
    possible_sku: str
    product_type: str
    brand: str
    volume: str
    variant: str
    confidence: str
    status: str
    group_reason: str
    pattern_key: str


def tokenize(text: str) -> list[str]:
    clean = re.sub(r"[^0-9A-ZА-ЯЁ]+", " ", text.upper().replace("Ё", "Е"))
    return [t for t in clean.split() if t]


def extract_volume(text: str) -> str:
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*(Л|МЛ|KG|КГ|G|Г)\b", text.upper().replace("Ё", "Е"))
    if not m:
        return ""
    num = m.group(1).replace(",", ".")
    unit = m.group(2)
    return f"{num} {unit}"


def has_cyrillic(text: str) -> bool:
    return bool(re.search(r"[А-ЯЁ]", text.upper()))


def has_latin(text: str) -> bool:
    return bool(re.search(r"[A-Z]", text.upper()))


def confidence_from_signals(similarity: float, volume_mismatch: bool, brand_missing: bool, group_reason: str) -> str:
    if group_reason == "реально отсутствующий товар":
        return "LOW"
    if similarity >= 75 and not volume_mismatch and not brand_missing:
        return "HIGH"
    if similarity >= 60 and not volume_mismatch:
        return "MEDIUM"
    return "LOW"


def classify_group(
    inventory_name: str,
    nearest_master_name: str,
    similarity: float,
    brand: str,
) -> tuple[str, str, bool, bool]:
    inv_tokens = tokenize(inventory_name)
    mst_tokens = tokenize(nearest_master_name)
    inv_set = set(inv_tokens)
    mst_set = set(mst_tokens)

    inv_volume = extract_volume(inventory_name)
    mst_volume = extract_volume(nearest_master_name)
    volume_mismatch = bool(inv_volume and mst_volume and inv_volume != mst_volume)

    brand_token = brand.upper().strip()
    brand_missing = bool(brand_token and brand_token not in inv_set)

    if similarity < 40:
        return "реально отсутствующий товар", "low_similarity_no_equivalent", volume_mismatch, brand_missing

    if volume_mismatch:
        return "другой объем", f"volume:{inv_volume}->{mst_volume}", volume_mismatch, brand_missing

    if brand_missing:
        return "отсутствие бренда", f"missing_brand:{brand_token}", volume_mismatch, brand_missing

    if (has_cyrillic(inventory_name) and has_latin(nearest_master_name)) or (has_latin(inventory_name) and has_cyrillic(nearest_master_name)):
        return "другой язык", "mixed_script_language", volume_mismatch, brand_missing

    if any(tok in PACKAGING_TAIL_WORDS for tok in inv_set - mst_set):
        extras = sorted((inv_set - mst_set) & PACKAGING_TAIL_WORDS)
        tail = extras[0] if extras else "tail"
        return "упаковочный хвост", f"packaging_tail:{tail}", volume_mismatch, brand_missing

    if any(tok in ABBR_HINTS for tok in inv_set):
        hit = sorted([tok for tok in inv_set if tok in ABBR_HINTS])[0]
        return "сокращения", f"abbreviation:{hit}", volume_mismatch, brand_missing

    # Same core tokens but sequence/format differences.
    core_inv = [t for t in inv_tokens if t not in GENERIC_EXTRA_WORDS]
    core_mst = [t for t in mst_tokens if t not in GENERIC_EXTRA_WORDS]
    if core_inv and core_mst and set(core_inv) == set(core_mst) and core_inv != core_mst:
        return "другой порядок слов", "same_tokens_different_order", volume_mismatch, brand_missing

    extra_tokens = [t for t in inv_set - mst_set if t not in GENERIC_EXTRA_WORDS]
    if extra_tokens:
        return "дополнительные слова", f"extra:{extra_tokens[0]}", volume_mismatch, brand_missing

    return "дополнительные слова", "generic_extra_descriptors", volume_mismatch, brand_missing


def main() -> None:
    recon_file = Path("output") / "INVENTORY_NAME_RECONCILIATION.xlsx"
    master_dataset_file = Path("output") / "MASTER_DATASET.xlsx"
    out_file = Path("output") / "INVENTORY_ALIAS_CANDIDATES.xlsx"

    if not recon_file.exists():
        raise FileNotFoundError(f"Missing file: {recon_file}")
    if not master_dataset_file.exists():
        raise FileNotFoundError(f"Missing file: {master_dataset_file}")

    wb_recon = load_workbook(recon_file, data_only=True)
    ws_need = wb_recon["NEED_REVIEW"]

    wb_master = load_workbook(master_dataset_file, data_only=True)
    ws_master = wb_master.active
    master_headers = [str(ws_master.cell(1, c).value or "").strip() for c in range(1, ws_master.max_column + 1)]
    midx = {h: i + 1 for i, h in enumerate(master_headers)}

    sku_meta: dict[str, dict[str, str]] = {}
    for r in range(2, ws_master.max_row + 1):
        sku = str(ws_master.cell(r, midx["SKU"]).value or "").strip()
        if not sku:
            continue
        sku_meta[sku] = {
            "product_type": str(ws_master.cell(r, midx["CATEGORY"]).value or "").strip(),
            "brand": str(ws_master.cell(r, midx["BRAND"]).value or "").strip(),
            "volume": str(ws_master.cell(r, midx["VOLUME"]).value or "").strip(),
            "variant": str(ws_master.cell(r, midx["VARIANT"]).value or "").strip(),
            "master_name": str(ws_master.cell(r, midx["MASTER_NAME"]).value or "").strip(),
        }

    candidates: list[CandidateRow] = []
    group_counter: Counter[str] = Counter()
    pattern_counter: Counter[str] = Counter()
    pattern_examples: dict[str, tuple[str, str]] = {}

    for r in range(2, ws_need.max_row + 1):
        inventory_name = str(ws_need.cell(r, 2).value or "").strip()
        possible_sku = str(ws_need.cell(r, 5).value or "").strip()
        nearest_master_name = str(ws_need.cell(r, 6).value or "").strip()
        similarity_raw = ws_need.cell(r, 7).value
        try:
            similarity = float(similarity_raw or 0.0)
        except (TypeError, ValueError):
            similarity = 0.0

        meta = sku_meta.get(possible_sku, {})
        product_type = meta.get("product_type", "")
        brand = meta.get("brand", "")
        volume = meta.get("volume", "")
        variant = meta.get("variant", "")

        group_reason, pattern_key, volume_mismatch, brand_missing = classify_group(
            inventory_name=inventory_name,
            nearest_master_name=nearest_master_name,
            similarity=similarity,
            brand=brand,
        )
        confidence = confidence_from_signals(similarity, volume_mismatch, brand_missing, group_reason)

        row = CandidateRow(
            inventory_name=inventory_name,
            nearest_master_name=nearest_master_name,
            similarity=round(similarity, 2),
            possible_sku=possible_sku,
            product_type=product_type,
            brand=brand,
            volume=volume,
            variant=variant,
            confidence=confidence,
            status="CANDIDATE_ONLY_MANUAL_APPROVAL",
            group_reason=group_reason,
            pattern_key=pattern_key,
        )
        candidates.append(row)
        group_counter[group_reason] += 1
        pattern_counter[pattern_key] += 1
        if pattern_key not in pattern_examples:
            pattern_examples[pattern_key] = (inventory_name, nearest_master_name)

    total = len(candidates)
    deterministic_recoverable = sum(1 for x in candidates if x.confidence == "HIGH")
    no_master_equivalent = group_counter.get("реально отсутствующий товар", 0)
    need_manual_approval = total - deterministic_recoverable

    grouped_rows = []
    ordered_groups = [
        "другой порядок слов",
        "сокращения",
        "другой язык",
        "отсутствие бренда",
        "другой объем",
        "упаковочный хвост",
        "дополнительные слова",
        "реально отсутствующий товар",
    ]
    for group_name in ordered_groups:
        grouped_rows.append((group_name, group_counter.get(group_name, 0), round(group_counter.get(group_name, 0) / total * 100, 2) if total else 0.0))

    # Priority fixes: pattern candidates impacting many rows/SKUs.
    by_pattern_rows: dict[str, list[CandidateRow]] = defaultdict(list)
    for row in candidates:
        by_pattern_rows[row.pattern_key].append(row)

    priority = []
    for pattern, rows in by_pattern_rows.items():
        impacted_skus = sorted({x.possible_sku for x in rows if x.possible_sku})
        example = rows[0]
        priority.append(
            {
                "pattern": pattern,
                "group": example.group_reason,
                "count_rows": len(rows),
                "count_skus": len(impacted_skus),
                "example_inventory": example.inventory_name,
                "example_master": example.nearest_master_name,
                "priority_score": len(rows) * 100 + len(impacted_skus),
            }
        )
    priority.sort(key=lambda x: (x["priority_score"], x["count_rows"], x["count_skus"]), reverse=True)

    top20_patterns = pattern_counter.most_common(20)

    wb_out = Workbook()

    ws_top = wb_out.active
    ws_top.title = "NEED_REVIEW_TOP"
    ws_top.append([
        "InventoryName",
        "Nearest_MASTER_NAME",
        "Similarity",
        "Possible_SKU",
        "ProductType",
        "Brand",
        "Volume",
        "Variant",
        "Confidence",
        "Status",
    ])
    for row in sorted(candidates, key=lambda x: x.similarity, reverse=True):
        ws_top.append(
            [
                row.inventory_name,
                row.nearest_master_name,
                row.similarity,
                row.possible_sku,
                row.product_type,
                row.brand,
                row.volume,
                row.variant,
                row.confidence,
                row.status,
            ]
        )

    ws_groups = wb_out.create_sheet("GROUPED_PATTERNS")
    ws_groups.append(["PatternGroup", "Count", "SharePct", "Description"])
    desc = {
        "другой порядок слов": "Состав ключевых токенов совпадает, но последовательность отличается",
        "сокращения": "Есть сокращенные или усеченные формы слов",
        "другой язык": "Смешение кириллицы/латиницы или языковых форм",
        "отсутствие бренда": "Бренд из MASTER не найден в InventoryName",
        "другой объем": "Объем в InventoryName и ближайшем MASTER отличается",
        "упаковочный хвост": "Цвет/упаковка/служебные хвосты в конце наименования",
        "дополнительные слова": "Лишние описательные слова относительно MASTER",
        "реально отсутствующий товар": "Низкое сходство, эквивалент в MASTER не подтверждается",
    }
    for group_name, count, pct in grouped_rows:
        ws_groups.append([group_name, count, pct, desc[group_name]])

    ws_priority = wb_out.create_sheet("PRIORITY_FIXES")
    ws_priority.append([
        "Pattern",
        "Group",
        "AffectedRows",
        "AffectedSKU",
        "PriorityScore",
        "ExampleInventoryName",
        "ExampleMASTER_NAME",
        "FixType",
    ])
    for item in priority[:50]:
        ws_priority.append(
            [
                item["pattern"],
                item["group"],
                item["count_rows"],
                item["count_skus"],
                item["priority_score"],
                item["example_inventory"],
                item["example_master"],
                "MANUAL_APPROVAL_REQUIRED",
            ]
        )

    ws_summary = wb_out.create_sheet("SUMMARY")
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Inventory rows", total])
    ws_summary.append(["Potential deterministic recoverable", deterministic_recoverable])
    ws_summary.append(["Need manual approval", need_manual_approval])
    ws_summary.append(["No MASTER equivalent", no_master_equivalent])
    ws_summary.append([])
    ws_summary.append(["Top 20 recurring naming patterns", ""])
    ws_summary.append(["Pattern", "Count", "Example InventoryName", "Example MASTER_NAME"])
    for pattern, count in top20_patterns:
        inv_ex, mst_ex = pattern_examples[pattern]
        ws_summary.append([pattern, count, inv_ex, mst_ex])

    for ws in wb_out.worksheets:
        ws.freeze_panes = "A2"

    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(out_file)

    print(f"Generated: {out_file}")
    print(f"Inventory rows: {total}")
    print(f"Potential deterministic recoverable: {deterministic_recoverable}")
    print(f"Need manual approval: {need_manual_approval}")
    print(f"No MASTER equivalent: {no_master_equivalent}")


if __name__ == "__main__":
    main()
