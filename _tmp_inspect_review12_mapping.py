from openpyxl import load_workbook

ws = load_workbook("output/SALES_REVIEW_12_MANUAL_MAPPING.xlsx", data_only=True).active
print("SHEET|" + ws.title)
print("MAX|" + str(ws.max_row) + "|" + str(ws.max_column))
headers = ["" if ws.cell(1, c).value is None else str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
print("HEADERS|" + "|".join(headers))

src_col = None
for i, h in enumerate(headers, start=1):
    if h.strip().upper() == "SOURCE_ROW":
        src_col = i
        break

if src_col is not None:
    found = False
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, src_col).value
        if v is not None and str(v).strip() == "267":
            row_vals = ["" if ws.cell(r, c).value is None else str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            print("ROW267|" + str(r) + "|" + "|".join(row_vals))
            found = True
            break
    if not found:
        print("ROW267|NOT_FOUND")

for r in range(2, min(ws.max_row, 12) + 1):
    row_vals = ["" if ws.cell(r, c).value is None else str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
    print("ROW|" + str(r) + "|" + "|".join(row_vals))
