from pathlib import Path
from openpyxl import load_workbook

TARGET = Path(r"output\MASTER_DATASET.xlsx")

# Confirmed corrections from the full REVIEW pass.
# Existing MASTER rows are updated by SKU; missing SKU-297 is created if needed.
updates = {
    "SKU-057": ("Кондиционер", "SVO", "Black", "2,7 л", "Black",
                "Кондиционер SVO Black 2,7 л"),
    "SKU-060": ("Кондиционер", "SVO", "Rose", "2,7 л", "Rose",
                "Кондиционер SVO Rose 2,7 л"),
    "SKU-061": ("Кондиционер", "SVO", "Papatya", "2,7 л", "Papatya",
                "Кондиционер SVO Papatya 2,7 л"),
    "SKU-062": ("Кондиционер", "SVO", "Midnight", "2,7 л", "Midnight",
                "Кондиционер SVO Midnight 2,7 л"),
    "SKU-063": ("Кондиционер", "SVO", "Dahlia", "2,7 л", "Dahlia",
                "Кондиционер SVO Dahlia 2,7 л"),
    "SKU-143": ("Отбеливатель", "SVO", "White", "750 г", "White",
                "Отбеливатель SVO White 750 г"),
    "SKU-144": ("Отбеливатель", "SVO", "Color", "750 г", "Color",
                "Отбеливатель SVO Color 750 г"),
    "SKU-256": ("Посуда моющее ср-во", "SVO", "Orange", "750 мл", "Orange",
                "Посуда моющее ср-во SVO Orange 750 мл"),
    "SKU-257": ("Посуда моющее ср-во", "SVO", "Grapefruit", "750 мл", "Grapefruit",
                "Посуда моющее ср-во SVO Grapefruit 750 мл"),
    "SKU-259": ("Посуда моющее ср-во", "SVO", "Apple", "750 мл", "Apple",
                "Посуда моющее ср-во SVO Apple 750 мл"),
    "SKU-218": ("Порошок стиральный", "SVO", "Mountain Breeze", "6 кг", "Mountain Breeze",
                "Порошок стиральный SVO Mountain Breeze 6 кг"),
    "SKU-229": ("Порошок стиральный", "SVO", "Mountain Breeze", "9 кг", "Mountain Breeze",
                "Порошок стиральный SVO Mountain Breeze 9 кг"),
    "SKU-295": ("Порошок стиральный", "SVO", "Magina", "5 кг", "Magina",
                "Порошок стиральный SVO Magina 5 кг"),
    "SKU-012": ("Шампунь sport", "GILAR", "Black", "400 мл", "Black",
                "Шампунь sport GILAR Black 400 мл"),
    "SKU-067": None,  # remove obsolete duplicate "Black Gel"
}

# SKU-297: confirmed as CONDITIONER, not ЖМС 3в1.
sku297 = ("Кондиционер", "SVO", "Romantik Rose", "2,7 л", "Romantik Rose",
          "Кондиционер SVO Romantik Rose 2,7 л")

wb = load_workbook(TARGET)
ws = wb.active

sku_rows = {}
for row in ws.iter_rows(min_row=2):
    sku = str(row[0].value or "").strip()
    if sku:
        sku_rows[sku] = row

changed = 0
removed = 0

for sku, values in updates.items():
    row = sku_rows.get(sku)
    if row is None:
        raise RuntimeError(f"{sku} NOT FOUND IN MASTER")

    if values is None:
        # Remove SKU-067 row entirely: Black Gel is not a separate MASTER item.
        ws.delete_rows(row[0].row, 1)
        removed += 1
        # Rebuild row map after deletion.
        sku_rows = {}
        for r in ws.iter_rows(min_row=2):
            s = str(r[0].value or "").strip()
            if s:
                sku_rows[s] = r
        continue

    category, brand, variant, volume, aroma, master_name = values
    row[1].value = category
    row[2].value = brand
    row[3].value = variant
    row[4].value = volume
    row[5].value = aroma
    row[9].value = master_name
    changed += 1

# Add or correct SKU-297.
row = sku_rows.get("SKU-297")
if row is None:
    row_number = ws.max_row + 1
    ws.cell(row_number, 1).value = "SKU-297"
    ws.cell(row_number, 2).value = sku297[0]
    ws.cell(row_number, 3).value = sku297[1]
    ws.cell(row_number, 4).value = sku297[2]
    ws.cell(row_number, 5).value = sku297[3]
    ws.cell(row_number, 6).value = sku297[4]
    ws.cell(row_number, 10).value = sku297[5]
    print("ADDED SKU-297")
else:
    row[1].value = sku297[0]
    row[2].value = sku297[1]
    row[3].value = sku297[2]
    row[4].value = sku297[3]
    row[5].value = sku297[4]
    row[9].value = sku297[5]
    print("UPDATED SKU-297")

wb.save(TARGET)

print(f"MASTER PATCH APPLIED: updated={changed}, removed={removed}")
print("SKU-297 = Кондиционер SVO Romantik Rose 2,7 л")
print("SKU-067 removed = obsolete Black Gel duplicate")
