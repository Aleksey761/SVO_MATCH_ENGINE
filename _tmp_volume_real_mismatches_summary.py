from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

path = Path('output/VOLUME_REAL_MISMATCHS.xlsx')
wb = load_workbook(path, data_only=True)
ws = wb.active

headers = [str(ws.cell(row=1, column=c).value or '').strip() for c in range(1, ws.max_column + 1)]
idx = {h: i + 1 for i, h in enumerate(headers)}

rows_analyzed = max(0, ws.max_row - 1)
equal_count = 0
real_mismatch_count = 0
reasons = Counter()
equal_examples = []

for r in range(2, ws.max_row + 1):
    row_num = int(ws.cell(row=r, column=idx['Row']).value or 0)
    supplier = str(ws.cell(row=r, column=idx['SupplierName']).value or '').strip()
    arrival = str(ws.cell(row=r, column=idx['ArrivalVolume']).value or '').strip()
    sku = str(ws.cell(row=r, column=idx['CandidateSKU']).value or '').strip()
    candidate_volume = str(ws.cell(row=r, column=idx['CandidateVolume']).value or '').strip()
    equal = str(ws.cell(row=r, column=idx['VolumeEqual']).value or '').strip().upper()
    reason = str(ws.cell(row=r, column=idx['MismatchReason']).value or '').strip()

    if equal == 'YES':
        equal_count += 1
        if len(equal_examples) < 20:
            equal_examples.append((row_num, supplier, arrival, sku, candidate_volume))
    else:
        real_mismatch_count += 1
        reasons[reason] += 1

print('Rows analyzed', rows_analyzed)
print('Equal volumes', equal_count)
print('Real volume mismatches', real_mismatch_count)
print('Top mismatch reasons')
for reason, count in reasons.most_common():
    print(f'{reason}: {count}')
if equal_count > 0:
    print('FIRST_20_EQUAL_EXAMPLES')
    for row_num, supplier, arrival, sku, candidate_volume in equal_examples:
        print(f"{row_num} | {supplier.encode('unicode_escape').decode('ascii')} | {arrival.encode('unicode_escape').decode('ascii')} | {sku} | {candidate_volume.encode('unicode_escape').decode('ascii')}")
