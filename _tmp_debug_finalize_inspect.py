from pathlib import Path

from openpyxl import load_workbook

from svo.loader import Loader
from svo.matcher import Matcher
from svo.normalizer import Normalizer
from svo.reporter import Reporter
from svo.sales_loader import SalesLoader


def norm(value: object) -> str:
    return str(value or "").strip().upper()


def main() -> None:
    base = Path(".")
    loader = Loader()
    master_file, _arrival_file, sales_file, _arrival_date, _sales_date = loader.discover_workbooks(
        base / "data",
        require_sales=True,
    )
    master_items = loader.load_master(master_file)

    items = SalesLoader().load(sales_file)
    normalizer = Normalizer()
    for item in items:
        normalizer.normalize(item)
    Matcher(master_items).match_all(items)

    tmp = base / "output" / "_tmp_before_finalize_debug.xlsx"
    Reporter().write_matched_document(sales_file, items, tmp)

    wb = load_workbook(tmp, data_only=False)
    ws = wb.active

    header_cells = [cell.value for cell in ws[1]]
    header_to_col = {}
    for idx, value in enumerate(header_cells, start=1):
        normalized = norm(value)
        if normalized:
            header_to_col[normalized] = idx

    match_status_col = header_to_col.get("MATCH_STATUS")
    match_sku_col = header_to_col.get("MATCH_SKU")
    match_master_name_col = header_to_col.get("MATCH_MASTER_NAME")
    match_reasons_col = header_to_col.get("MATCH_REASONS")
    match_confidence_col = header_to_col.get("MATCH_CONFIDENCE")

    assert match_status_col is not None
    assert match_sku_col is not None

    master_by_sku = {item.sku: item for item in master_items}
    master_name_by_sku = Reporter()._load_master_name_by_sku(master_file)

    master_value_by_result_header = {
        "SKU": lambda item, _name: item.sku,
        "ТИП ТОВАРА": lambda item, _name: item.category,
        "БРЕНД": lambda item, _name: item.brand,
        "VARIANT": lambda item, _name: item.variant,
        "ОБЪЕМ": lambda item, _name: item.volume,
        "НАИМЕНОВАНИЕ": lambda item, name: name,
    }

    sklad1_col = header_to_col.get("СКЛАД 1")
    if sklad1_col is not None and sklad1_col > 7:
        for col, title in enumerate(("№", "SKU", "Наименование", "Тип товара", "Бренд", "Variant", "Объем"), start=1):
            ws.cell(row=1, column=col, value=title)

        header_cells = [cell.value for cell in ws[1]]
        header_to_col = {}
        for idx, value in enumerate(header_cells, start=1):
            normalized = norm(value)
            if normalized:
                header_to_col[normalized] = idx

    actionable_rows = []
    matched_rows_by_sku = {}
    matched_unknown_rows = []
    review_rows = []

    max_col = ws.max_column
    for row_number in range(2, ws.max_row + 1):
        status_value = norm(ws.cell(row=row_number, column=match_status_col).value)
        if status_value not in {"MATCH", "REVIEW"}:
            fallback_match_sku = str(ws.cell(row=row_number, column=match_sku_col).value or "").strip()
            fallback_master_name = (
                str(ws.cell(row=row_number, column=match_master_name_col).value or "").strip()
                if match_master_name_col is not None
                else ""
            )
            fallback_reasons = (
                str(ws.cell(row=row_number, column=match_reasons_col).value or "").strip()
                if match_reasons_col is not None
                else ""
            )
            fallback_confidence = (
                ws.cell(row=row_number, column=match_confidence_col).value
                if match_confidence_col is not None
                else None
            )
            if fallback_match_sku:
                status_value = "MATCH"
            elif fallback_master_name or fallback_reasons or fallback_confidence is not None:
                status_value = "REVIEW"

        if status_value not in {"MATCH", "REVIEW"}:
            continue

        actionable_rows.append(row_number)
        row_values = [ws.cell(row=row_number, column=col).value for col in range(1, max_col + 1)]
        row_values[match_status_col - 1] = status_value

        if status_value == "MATCH":
            sku = str(ws.cell(row=row_number, column=match_sku_col).value or "").strip()
            master = master_by_sku.get(sku)
            if master is not None:
                master_name = master_name_by_sku.get(master.sku)

                for result_header, getter in master_value_by_result_header.items():
                    target_col = header_to_col.get(result_header)
                    if target_col is None:
                        continue
                    row_values[target_col - 1] = getter(master, master_name)

                row_values[match_sku_col - 1] = master.sku
                matched_rows_by_sku.setdefault(master.sku, []).append(row_values)
            else:
                matched_unknown_rows.append(row_values)
        else:
            review_rows.append(row_values)

    ordered_rows = []
    for master in master_items:
        ordered_rows.extend(matched_rows_by_sku.get(master.sku, []))
    ordered_rows.extend(matched_unknown_rows)
    ordered_rows.extend(review_rows)

    print("ws.max_row", ws.max_row)
    print("actionable", len(actionable_rows))
    print("ordered", len(ordered_rows))
    print("review", len(review_rows))
    print("unknown", len(matched_unknown_rows))
    print("last_actionable", actionable_rows[-10:])

    for idx, values in enumerate(ordered_rows[-10:], start=len(ordered_rows) - 9):
        print(
            "ordered",
            idx,
            "status",
            norm(values[match_status_col - 1]),
            "match_sku",
            values[match_sku_col - 1],
            "col1",
            values[0],
            "col2",
            values[1],
            "col3",
            values[2],
        )

    target_rows = list(actionable_rows)
    if len(target_rows) < len(ordered_rows):
        start_row = target_rows[-1] + 1 if target_rows else 2
        target_rows.extend(range(start_row, start_row + (len(ordered_rows) - len(target_rows))))
    elif len(target_rows) > len(ordered_rows):
        target_rows = target_rows[: len(ordered_rows)]

    print("target_rows", len(target_rows), target_rows[-10:])

    tmp.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
