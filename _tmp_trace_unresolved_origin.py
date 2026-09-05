from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"f:/SVO/AI/SVO_MATCH_ENGINE")

TARGETS = [
    ("PRICE", "Unknown product without master"),
    ("PRICE", "mystery foobar 999"),
    ("PRICE", "SVO shampoo 1 l"),
    ("PRICE", "Unknown product 999"),
    ("SALES", "SVO SHAMPUN 1 л"),
    ("SALES", "SVO SHAMPUN 1 л"),
]

PRICE_INPUTS = [
    ROOT / "data" / "PRC.xlsx",
    ROOT / "data" / "PRC_202412150928-Дистрибьюторская Цена.xlsx",
    ROOT / "data" / "Прайс SVO,GILAR,BOSSFIX 10.2024.xlsx",
]

SALES_INPUTS = [
    ROOT / "data" / "ARRIVAL.xlsx",
    ROOT / "input" / "ARRIVAL.xlsx",
    ROOT / "output" / "SALES_MATCH_06.04.2026.xlsx",
    ROOT / "output" / "SALES_MATCH.xlsx",
]


def iter_workbook_rows(path: Path):
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            yield ws.title, r, values


def cell_text(v):
    return "" if v is None else str(v).strip()


def find_in_workbooks(name: str, paths: list[Path]):
    hits = []
    needle = name.strip().lower()
    for p in paths:
        if not p.exists():
            continue
        try:
            for sheet, row, values in iter_workbook_rows(p):
                texts = [cell_text(v) for v in values]
                for idx, t in enumerate(texts, start=1):
                    if t.lower() == needle:
                        hits.append((p, sheet, row, idx, texts))
        except Exception:
            continue
    return hits


def extract_row_context(texts):
    sku = ""
    others = []
    for idx, v in enumerate(texts, start=1):
        s = v.strip()
        if not s:
            continue
        if ("SKU" in s.upper()) and len(s) < 30:
            pass
        if not sku and s.upper().startswith("SKU"):
            sku = s
        if idx in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17):
            others.append(f"C{idx}={s}")
    return sku, "; ".join(others[:8])


def main():
    out_lines = []
    out_lines.append("SOURCE\tINPUT_FILE\tSHEET\tSOURCE_ROW\tORIGINAL_NAME\tSKU_IF_PRESENT\tOTHER_RELEVANT_FIELDS\tWHY_UNRESOLVED")

    for source, name in TARGETS:
        paths = PRICE_INPUTS if source == "PRICE" else SALES_INPUTS
        hits = find_in_workbooks(name, paths)
        if not hits:
            out_lines.append(
                f"{source}\tNOT_FOUND_IN_PRODUCTION_INPUTS\t\t\t{name}\t\t\tNo exact row in current production input files; unresolved entry likely originates from non-production fixture/test data"
            )
            continue

        for p, sheet, row, col, texts in hits:
            sku, others = extract_row_context(texts)
            out_lines.append(
                f"{source}\t{p.relative_to(ROOT)}\t{sheet}\t{row}\t{name}\t{sku}\t{others}\tExact source-name row exists but has no deterministic master resolution in current rules"
            )

    out_path = ROOT / "output" / "_tmp_unresolved_origin_trace.tsv"
    out_path.write_text("\n".join(out_lines), encoding="utf-8")
    print(out_path)


if __name__ == "__main__":
    main()
