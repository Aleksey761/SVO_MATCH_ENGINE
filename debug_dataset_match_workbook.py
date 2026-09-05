from pathlib import Path
from openpyxl import load_workbook
import tempfile

base = Path(tempfile.gettempdir())
files = sorted(base.glob('pytest-of-*/*/pytest-*/test_missing_master_sku_detect*/output/ARRIVAL_MATCH_2026-07-15.xlsx'), key=lambda p: p.stat().st_mtime, reverse=True)
if not files:
    print('MATCH workbook not found under the pytest temp directory.')
    raise SystemExit(1)
path = files[0]
print(f'FILE={path}')
wb = load_workbook(path, data_only=True)
ws = wb.active
print(f'SHEET={ws.title}')
print(f'MAX_ROW={ws.max_row} MAX_COLUMN={ws.max_column}')
headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
print('HEADERS=', headers)
count = 0
samples = []
for r in range(2, ws.max_row+1):
    row = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
    if not row or not row[0]:
        continue
    status = str(row[11]).strip().upper() if len(row) >= 12 and row[11] is not None else ''
    sku = str(row[12]).strip() if len(row) >= 13 and row[12] is not None else ''
    if status == 'MATCH' and sku:
        count += 1
        if len(samples) < 10:
            samples.append((r, row[0], status, sku))
print(f'MATCH_WITH_SKU_ROWS={count}')
print('SAMPLES=', samples)
