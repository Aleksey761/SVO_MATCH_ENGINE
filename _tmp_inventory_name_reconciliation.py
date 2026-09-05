from __future__ import annotations

import re
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook

from svo.loader import Loader


def normalize_name_key(value: object) -> str:
    text = str(value or "").strip().upper()
    if not text:
        return ""
    text = text.replace("Ё", "Е")
    text = re.sub(r"[^0-9A-ZА-Я]+", " ", text)
    return " ".join(text.split())


def resolve_master_file() -> Path:
    preferred = Path("data") / "MASTER.xlsx"
    if preferred.exists():
        return preferred

    candidates = sorted(Path("data").glob("*MASTER*.xlsx"))
    if not candidates:
        raise FileNotFoundError("MASTER workbook not found in data/")
    return candidates[0]


def resolve_inventory_file() -> Path:
    candidates = sorted(Path("data").glob("*инвентар*.xlsx"))
    if not candidates:
        raise FileNotFoundError("Inventory workbook not found in data/ by *инвентар*.xlsx")
    return candidates[0]


def nearest_master(norm_name: str, master_norm_rows: list[dict[str, str]]) -> dict[str, object]:
    best_score = -1.0
    best = None
    for item in master_norm_rows:
        score = SequenceMatcher(None, norm_name, item["norm"]).ratio() if norm_name and item["norm"] else 0.0
        if score > best_score:
            best_score = score
            best = item

    if best is None:
        return {
            "nearest_sku": "",
            "nearest_master_name": "",
            "similarity": 0.0,
        }

    return {
        "nearest_sku": best["sku"],
        "nearest_master_name": best["master_name"],
        "similarity": round(best_score * 100, 2),
    }


def main() -> None:
    master_file = resolve_master_file()
    inventory_file = resolve_inventory_file()
    output_file = Path("output") / "INVENTORY_NAME_RECONCILIATION.xlsx"

    master_items = Loader().load_master(master_file)
    master_norm_to_rows: dict[str, list[dict[str, str]]] = {}
    master_norm_rows: list[dict[str, str]] = []

    for item in master_items:
        sku = str(getattr(item, "sku", "") or "").strip()
        master_name = str(getattr(item, "master_name", "") or "").strip()
        if not sku or not master_name:
            continue
        norm = normalize_name_key(master_name)
        row = {"sku": sku, "master_name": master_name, "norm": norm}
        master_norm_rows.append(row)
        master_norm_to_rows.setdefault(norm, []).append(row)

    wb_inv = load_workbook(filename=inventory_file, data_only=True)
    ws_inv = wb_inv.active

    wb_out = Workbook()
    ws_matched = wb_out.active
    ws_matched.title = "MATCHED_EXACT"
    ws_matched.append([
        "InventoryRow",
        "InventoryName",
        "NormalizedInventoryName",
        "StockQty",
        "SKU",
        "MASTER_NAME",
        "status",
    ])

    ws_review = wb_out.create_sheet("NEED_REVIEW")
    ws_review.append([
        "InventoryRow",
        "InventoryName",
        "NormalizedInventoryName",
        "StockQty",
        "NearestSKU",
        "NearestMASTER_NAME",
        "SimilarityScore",
        "Reason",
        "status",
    ])

    ws_all = wb_out.create_sheet("INVENTORY_ALL")
    ws_all.append([
        "InventoryRow",
        "InventoryName",
        "NormalizedInventoryName",
        "StockQty",
        "status",
    ])

    total_rows = 0
    matched_rows = 0
    review_rows = 0

    for row_idx in range(3, ws_inv.max_row + 1):
        inv_name = str(ws_inv.cell(row=row_idx, column=3).value or "").strip()
        if not inv_name:
            continue

        stock_qty_raw = ws_inv.cell(row=row_idx, column=15).value
        stock_qty = 0 if stock_qty_raw is None else int(float(stock_qty_raw))
        norm = normalize_name_key(inv_name)
        total_rows += 1

        candidates = master_norm_to_rows.get(norm, [])
        if len(candidates) == 1:
            matched = candidates[0]
            matched_rows += 1
            ws_matched.append([
                row_idx,
                inv_name,
                norm,
                stock_qty,
                matched["sku"],
                matched["master_name"],
                "MATCHED_EXACT_NORMALIZED",
            ])
            ws_all.append([row_idx, inv_name, norm, stock_qty, "MATCHED_EXACT_NORMALIZED"])
            continue

        review_rows += 1
        if len(candidates) > 1:
            nearest = candidates[0]
            similarity = 100.0
            reason = "AMBIGUOUS_NORMALIZED_KEY"
        else:
            nearest_info = nearest_master(norm, master_norm_rows)
            nearest = {
                "sku": nearest_info["nearest_sku"],
                "master_name": nearest_info["nearest_master_name"],
            }
            similarity = nearest_info["similarity"]
            reason = "NO_EXACT_NORMALIZED_MATCH"

        ws_review.append([
            row_idx,
            inv_name,
            norm,
            stock_qty,
            nearest["sku"],
            nearest["master_name"],
            similarity,
            reason,
            "NEED_REVIEW",
        ])
        ws_all.append([row_idx, inv_name, norm, stock_qty, "NEED_REVIEW"])

    ws_summary = wb_out.create_sheet("SUMMARY")
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["master_file", str(master_file)])
    ws_summary.append(["inventory_file", str(inventory_file)])
    ws_summary.append(["inventory_rows", total_rows])
    ws_summary.append(["matched_exact_rows", matched_rows])
    ws_summary.append(["need_review_rows", review_rows])

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_file)

    print(f"Generated: {output_file}")
    print(f"Inventory rows: {total_rows}")
    print(f"Matched exact: {matched_rows}")
    print(f"Need review: {review_rows}")


if __name__ == "__main__":
    main()
