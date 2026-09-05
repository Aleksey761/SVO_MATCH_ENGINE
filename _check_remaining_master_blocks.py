from openpyxl import load_workbook
from pathlib import Path

path = Path(r"output\MASTER_DATASET.xlsx")
if not path.exists():
    raise FileNotFoundError(path)

wb = load_workbook(path, data_only=True)
ws = wb.active

targets = (
    "ЖМС",
    "ЖМС 3В1",
    "Кондиционер",
    "Порошок",
    "Пятновыводитель",
    "Средство",
    "Шампунь",
)

print("MASTER SHEET =", ws.title)
print("ROWS =", ws.max_row)
print()

for row in ws.iter_rows(min_row=2, values_only=True):
    sku, category, brand, variant, volume, aroma, c7, c8, c9, master_name = row[:10]
    cat = str(category or "").strip()
    name = str(master_name or "").strip()

    if any(key.upper() in cat.upper() or key.upper() in name.upper()
           for key in targets):
        print(
            f"{sku} | {category} | {brand} | {variant} | "
            f"{volume} | {aroma} | {master_name}"
        )
