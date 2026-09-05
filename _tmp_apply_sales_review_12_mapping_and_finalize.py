from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook


BASE = Path("output")
SOURCE_PATH = BASE / "SALES_MATCH_MANUAL.xlsx"
MAPPING_PATH = BASE / "SALES_REVIEW_12_MANUAL_MAPPING.xlsx"
FINAL_PATH = BASE / "SALES_MATCH_FINAL.xlsx"

TARGET_SOURCE_ROW = 267
TARGET_SKU = "SKU-283"
TARGET_MASTER_NAME = "Подгузники BOSSFIX Size 1 1 уп"
TARGET_DECISION = "APPROVE"
TARGET_COMMENT = "manual confirmed from MASTER"


REQUIRED_MAP_COLS = [
    "SOURCE_ROW",
    "SOURCE_NAME",
    "SKU",
    "MASTER_NAME",
    "DECISION",
    "COMMENT",
]


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def get_header_map(ws) -> Dict[str, int]:
    header_map: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header = normalize_header(ws.cell(1, col).value)
        if header:
            header_map[header] = col
    return header_map


def ensure_mapping_file(source_ws) -> None:
    if MAPPING_PATH.exists():
        return

    source_headers = get_header_map(source_ws)
    if "MATCH_STATUS" not in source_headers:
        raise RuntimeError("MATCH_STATUS column not found in SALES_MATCH_MANUAL.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "MANUAL_MAPPING"

    for idx, col_name in enumerate(REQUIRED_MAP_COLS, start=1):
        ws.cell(1, idx, col_name)

    status_col = source_headers["MATCH_STATUS"]
    source_row_col = source_headers.get("SOURCE_ROW")
    source_name_col = source_headers.get("SOURCE_NAME")

    out_row = 2
    for r in range(2, source_ws.max_row + 1):
        status = normalize_header(source_ws.cell(r, status_col).value).upper()
        if status != "REVIEW":
            continue

        source_row_value = (
            source_ws.cell(r, source_row_col).value if source_row_col is not None else r
        )
        source_name_value = (
            source_ws.cell(r, source_name_col).value if source_name_col is not None else ""
        )

        ws.cell(out_row, 1, source_row_value)
        ws.cell(out_row, 2, source_name_value)
        ws.cell(out_row, 3, "")
        ws.cell(out_row, 4, "")
        ws.cell(out_row, 5, "")
        ws.cell(out_row, 6, "")
        out_row += 1

    wb.save(MAPPING_PATH)


def as_int(value: object) -> Optional[int]:
    if value is None:
        return None
    try:
        if isinstance(value, str) and not value.strip():
            return None
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def find_source_name_for_row(source_ws, source_headers: Dict[str, int], source_row: int) -> str:
    src_row_col = source_headers.get("SOURCE_ROW")
    src_name_col = source_headers.get("SOURCE_NAME")
    if src_name_col is None:
        src_name_col = source_headers.get("НАИМЕНОВАНИЕ")
    if src_name_col is None:
        return ""

    for r in range(2, source_ws.max_row + 1):
        candidate = as_int(source_ws.cell(r, src_row_col).value) if src_row_col is not None else r
        if candidate == source_row:
            val = source_ws.cell(r, src_name_col).value
            return "" if val is None else str(val)
    return ""


def update_mapping_row(mapping_ws, source_ws, source_headers: Dict[str, int]) -> Dict[str, int]:
    mapping_headers = get_header_map(mapping_ws)
    for col in REQUIRED_MAP_COLS:
        if col not in mapping_headers:
            raise RuntimeError(f"{col} column not found in SALES_REVIEW_12_MANUAL_MAPPING.xlsx")

    source_row_col = mapping_headers["SOURCE_ROW"]
    source_name_col = mapping_headers["SOURCE_NAME"]
    sku_col = mapping_headers["SKU"]
    master_col = mapping_headers["MASTER_NAME"]
    decision_col = mapping_headers["DECISION"]
    comment_col = mapping_headers["COMMENT"]

    target_excel_row: Optional[int] = None
    for r in range(2, mapping_ws.max_row + 1):
        if as_int(mapping_ws.cell(r, source_row_col).value) == TARGET_SOURCE_ROW:
            target_excel_row = r
            break

    if target_excel_row is None:
        target_excel_row = mapping_ws.max_row + 1
        mapping_ws.cell(target_excel_row, source_row_col, TARGET_SOURCE_ROW)

    existing_name = mapping_ws.cell(target_excel_row, source_name_col).value
    if existing_name is None or str(existing_name).strip() == "":
        resolved_name = find_source_name_for_row(source_ws, source_headers, TARGET_SOURCE_ROW)
        mapping_ws.cell(target_excel_row, source_name_col, resolved_name)

    mapping_ws.cell(target_excel_row, sku_col, TARGET_SKU)
    mapping_ws.cell(target_excel_row, master_col, TARGET_MASTER_NAME)
    mapping_ws.cell(target_excel_row, decision_col, TARGET_DECISION)
    mapping_ws.cell(target_excel_row, comment_col, TARGET_COMMENT)

    return mapping_headers


def apply_mapping_and_save_final(source_wb, source_ws, mapping_ws, mapping_headers: Dict[str, int]):
    source_headers = get_header_map(source_ws)
    required_source_cols = ["MATCH_STATUS", "MATCH_SKU", "MATCH_MASTER_NAME"]
    for col in required_source_cols:
        if col not in source_headers:
            raise RuntimeError(f"{col} column not found in SALES_MATCH_MANUAL.xlsx")

    src_source_row_col = source_headers.get("SOURCE_ROW")
    src_status_col = source_headers["MATCH_STATUS"]
    src_sku_col = source_headers["MATCH_SKU"]
    src_master_col = source_headers["MATCH_MASTER_NAME"]

    map_source_row_col = mapping_headers["SOURCE_ROW"]
    map_sku_col = mapping_headers["SKU"]
    map_master_col = mapping_headers["MASTER_NAME"]
    map_decision_col = mapping_headers["DECISION"]

    approve_map: Dict[int, Dict[str, str]] = {}
    for r in range(2, mapping_ws.max_row + 1):
        decision = normalize_header(mapping_ws.cell(r, map_decision_col).value).upper()
        srow = as_int(mapping_ws.cell(r, map_source_row_col).value)
        if srow is None:
            continue
        if decision == "APPROVE":
            approve_map[srow] = {
                "SKU": "" if mapping_ws.cell(r, map_sku_col).value is None else str(mapping_ws.cell(r, map_sku_col).value),
                "MASTER_NAME": "" if mapping_ws.cell(r, map_master_col).value is None else str(mapping_ws.cell(r, map_master_col).value),
            }

    final_by_source_row: Dict[int, Dict[str, str]] = {}
    for r in range(2, source_ws.max_row + 1):
        srow = as_int(source_ws.cell(r, src_source_row_col).value) if src_source_row_col is not None else r
        if srow is None:
            continue

        if srow in approve_map:
            source_ws.cell(r, src_status_col, "MATCH")
            source_ws.cell(r, src_sku_col, approve_map[srow]["SKU"])
            source_ws.cell(r, src_master_col, approve_map[srow]["MASTER_NAME"])

        final_by_source_row[srow] = {
            "MATCH_STATUS": normalize_header(source_ws.cell(r, src_status_col).value),
            "MATCH_SKU": "" if source_ws.cell(r, src_sku_col).value is None else str(source_ws.cell(r, src_sku_col).value),
            "MATCH_MASTER_NAME": "" if source_ws.cell(r, src_master_col).value is None else str(source_ws.cell(r, src_master_col).value),
        }

    source_wb.save(FINAL_PATH)

    metrics = {
        "TOTAL_SOURCE_ROWS": 0,
        "MATCH_TOTAL": 0,
        "REVIEW_TOTAL": 0,
        "EMPTY_SKU_MATCH": 0,
        "MANUAL_MAPPING_MISMATCH": 0,
    }

    for r in range(2, source_ws.max_row + 1):
        metrics["TOTAL_SOURCE_ROWS"] += 1
        status = normalize_header(source_ws.cell(r, src_status_col).value).upper()
        sku = normalize_header(source_ws.cell(r, src_sku_col).value)
        if status == "MATCH":
            metrics["MATCH_TOTAL"] += 1
            if not sku:
                metrics["EMPTY_SKU_MATCH"] += 1
        if status == "REVIEW":
            metrics["REVIEW_TOTAL"] += 1

    mismatch = 0
    for srow, vals in approve_map.items():
        final_vals = final_by_source_row.get(srow)
        if final_vals is None:
            mismatch += 1
            continue
        if (
            final_vals["MATCH_STATUS"].upper() != "MATCH"
            or final_vals["MATCH_SKU"] != vals["SKU"]
            or final_vals["MATCH_MASTER_NAME"] != vals["MASTER_NAME"]
        ):
            mismatch += 1
    metrics["MANUAL_MAPPING_MISMATCH"] = mismatch

    row267 = final_by_source_row.get(TARGET_SOURCE_ROW, {})

    return metrics, row267


def main():
    if not SOURCE_PATH.exists():
        raise FileNotFoundError(f"Missing source workbook: {SOURCE_PATH}")

    source_wb = load_workbook(SOURCE_PATH)
    source_ws = source_wb.active
    source_headers = get_header_map(source_ws)

    ensure_mapping_file(source_ws)

    map_wb = load_workbook(MAPPING_PATH)
    map_ws = map_wb.active
    mapping_headers = update_mapping_row(map_ws, source_ws, source_headers)
    map_wb.save(MAPPING_PATH)

    metrics, row267 = apply_mapping_and_save_final(source_wb, source_ws, map_ws, mapping_headers)

    for k in [
        "TOTAL_SOURCE_ROWS",
        "MATCH_TOTAL",
        "REVIEW_TOTAL",
        "EMPTY_SKU_MATCH",
        "MANUAL_MAPPING_MISMATCH",
    ]:
        print(f"{k}|{metrics[k]}")

    print(
        "ROW267|"
        + row267.get("MATCH_STATUS", "")
        + "|"
        + row267.get("MATCH_SKU", "")
        + "|"
        + row267.get("MATCH_MASTER_NAME", "")
    )


if __name__ == "__main__":
    main()
