from pathlib import Path
from openpyxl import load_workbook

path = Path('output/CANDIDATE_RANKING_ANALYSIS.xlsx')
wb = load_workbook(path, data_only=True)
ws = wb['SUMMARY']

summary = {}
for r in range(2, ws.max_row + 1):
    key = str(ws.cell(row=r, column=1).value or '').strip()
    value = ws.cell(row=r, column=2).value
    summary[key] = value

print('Rows analyzed', summary.get('Correct-volume candidate exists', 0) + summary.get('Not found', 0))
print('Rows with correct-volume candidate', summary.get('Correct-volume candidate exists', 0))
print('Rows without correct-volume candidate', summary.get('Not found', 0))
print('Estimated recoverable matches if ranking is improved', summary.get('Correct-volume candidate exists', 0))
print('Average score gap', summary.get('Average score gap', 0))
print('Rank 2', summary.get('Rank 2', 0))
print('Rank 3', summary.get('Rank 3', 0))
print('Rank >3', summary.get('Rank >3', 0))
print('Top score components responsible for losing')
print('Name', summary.get('Top score components: Name', 0))
print('Brand', summary.get('Top score components: Brand', 0))
print('ProductType', summary.get('Top score components: ProductType', 0))
print('Volume', summary.get('Top score components: Volume', 0))
print('Other', summary.get('Top score components: Other', 0))
