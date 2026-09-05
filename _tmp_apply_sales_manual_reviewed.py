from __future__ import annotations

from pathlib import Path
from typing import Dict, Optional

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "output"
SOURCE_FILE = OUTPUT_DIR / "SALES_MATCH.xlsx"
REVIEWED_MAP_FILE = OUTPUT_DIR / "SALES_MANUAL_MAPPING_REVIEWED.xlsx"
TARGET_FILE = OUTPUT_DIR / "SALES_MATCH_MANUAL.xlsx"


def norm_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().upper()


def build_header_index(sheet) -> Dict[str, int]:
    result: Dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        key = norm_header(sheet.cell(1, col).value)
        if key and key not in result:
            result[key] = col
    return result


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_source_row(raw: object) -> Optional[int]:
    text = cell_text(raw)
    if not text:
        return None
    try:
        # Handles values like 123 and 123.0
        return int(float(text))
    except ValueError:
        return None


def find_source_name_col(headers: Dict[str, int]) -> Optional[int]:
    preferred = [
        "SOURCE_NAME",
        "НАИМЕНОВАНИЕ",
        "NAME",
        "PRODUCT_NAME",
        "ITEM_NAME",
    ]
    for key in preferred:
        if key in headers:
            return headers[key]

    # Fallback: first name-like business column, but not match master column.
    for key, col in headers.items():
        if "NAME" in key and "MASTER" not in key and "MATCH" not in key:
            return col
    return None


def main() -> int:
    if not REVIEWED_MAP_FILE.exists():
        print(f"ERROR|Missing reviewed mapping file: {REVIEWED_MAP_FILE}")
        print("STOPPED|No output file was written")
        return 2

    if not SOURCE_FILE.exists():
        print(f"ERROR|Missing source match file: {SOURCE_FILE}")
        print("STOPPED|No output file was written")
        return 2

    sales_wb = load_workbook(SOURCE_FILE)
    if "SALES_MATCH" in sales_wb.sheetnames:
        sales_ws = sales_wb["SALES_MATCH"]
    else:
        sales_ws = sales_wb.active

    map_wb = load_workbook(REVIEWED_MAP_FILE, data_only=True)
    if "MANUAL_MAPPING" not in map_wb.sheetnames:
        print("ERROR|Missing sheet MANUAL_MAPPING in reviewed mapping workbook")
        print("STOPPED|No output file was written")
        return 2
    map_ws = map_wb["MANUAL_MAPPING"]

    sales_headers = build_header_index(sales_ws)
    map_headers = build_header_index(map_ws)

    required_sales = [
        "MATCH_STATUS",
        "MATCH_SKU",
        "MATCH_MASTER_NAME",
        "MATCH_CONFIDENCE",
        "MATCH_REASONS",
    ]
    required_map = ["SOURCE_ROW", "DECISION", "SKU", "MASTER_NAME"]

    missing_sales = [k for k in required_sales if k not in sales_headers]
    missing_map = [k for k in required_map if k not in map_headers]

    if missing_sales:
        print("ERROR|Missing required columns in SALES_MATCH: " + ", ".join(missing_sales))
        print("STOPPED|No output file was written")
        return 2
    if missing_map:
        print("ERROR|Missing required columns in MANUAL_MAPPING: " + ", ".join(missing_map))
        print("STOPPED|No output file was written")
        return 2

    source_name_col = find_source_name_col(sales_headers)
    if source_name_col is None:
        print("ERROR|Could not identify business source name column in SALES_MATCH")
        print("STOPPED|No output file was written")
        return 2

    sales_source_row_col = sales_headers.get("SOURCE_ROW")

    # Build SOURCE_ROW -> worksheet row lookup from SALES_MATCH.
    source_row_to_excel_row: Dict[int, int] = {}
    if sales_source_row_col is not None:
        for r in range(2, sales_ws.max_row + 1):
            sr = parse_source_row(sales_ws.cell(r, sales_source_row_col).value)
            if sr is not None and sr not in source_row_to_excel_row:
                source_row_to_excel_row[sr] = r

    manual_approve_applied = 0
    manual_no_master = 0
    manual_mapping_mismatch = 0

    map_source_row_col = map_headers["SOURCE_ROW"]
    map_decision_col = map_headers["DECISION"]
    map_sku_col = map_headers["SKU"]
    map_master_col = map_headers["MASTER_NAME"]

    sales_status_col = sales_headers["MATCH_STATUS"]
    sales_sku_col = sales_headers["MATCH_SKU"]
    sales_master_col = sales_headers["MATCH_MASTER_NAME"]

    for r in range(2, map_ws.max_row + 1):
        decision = cell_text(map_ws.cell(r, map_decision_col).value).upper()
        if not decision:
            continue

        src_row = parse_source_row(map_ws.cell(r, map_source_row_col).value)

        if decision == "NO_MASTER":
            manual_no_master += 1

        if decision == "REVIEW":
            continue

        if src_row is None:
            manual_mapping_mismatch += 1
            continue

        if sales_source_row_col is not None:
            target_excel_row = source_row_to_excel_row.get(src_row)
        else:
            # Fallback: SOURCE_ROW refers to worksheet row index.
            target_excel_row = src_row if 2 <= src_row <= sales_ws.max_row else None

        if target_excel_row is None:
            manual_mapping_mismatch += 1
            continue

        if decision == "APPROVE":
            sku = cell_text(map_ws.cell(r, map_sku_col).value)
            master_name = cell_text(map_ws.cell(r, map_master_col).value)
            if not sku or not master_name:
                manual_mapping_mismatch += 1
                continue

            sales_ws.cell(target_excel_row, sales_status_col).value = "MATCH"
            sales_ws.cell(target_excel_row, sales_sku_col).value = sku
            sales_ws.cell(target_excel_row, sales_master_col).value = master_name
            # MATCH_CONFIDENCE and MATCH_REASONS remain unchanged by design.
            manual_approve_applied += 1
            continue

        if decision == "NO_MASTER":
            sales_ws.cell(target_excel_row, sales_status_col).value = "REVIEW"
            sales_ws.cell(target_excel_row, sales_sku_col).value = None
            continue

        # Unknown decision: count as mismatch.
        manual_mapping_mismatch += 1

    # Calculate requested metrics on the resulting workbook.
    total_source_rows = 0
    match_total = 0
    review_total = 0
    empty_sku_match = 0

    review_rows_out = []

    for r in range(2, sales_ws.max_row + 1):
        source_name = cell_text(sales_ws.cell(r, source_name_col).value)
        if source_name:
            total_source_rows += 1

        status = cell_text(sales_ws.cell(r, sales_status_col).value).upper()
        sku = cell_text(sales_ws.cell(r, sales_sku_col).value)

        if status == "MATCH":
            match_total += 1
            if not sku:
                empty_sku_match += 1
        elif status == "REVIEW":
            review_total += 1

            if len(review_rows_out) < 12:
                source_row_value = ""
                if sales_source_row_col is not None:
                    source_row_value = cell_text(sales_ws.cell(r, sales_source_row_col).value)
                if not source_row_value:
                    source_row_value = str(r)

                reasons = cell_text(sales_ws.cell(r, sales_headers["MATCH_REASONS"]).value)
                master_name = cell_text(sales_ws.cell(r, sales_master_col).value)

                review_rows_out.append(
                    [
                        source_row_value,
                        source_name,
                        "REVIEW",
                        reasons,
                        sku,
                        master_name,
                    ]
                )

    sales_wb.save(TARGET_FILE)

    print(f"OUTPUT_FILE|{TARGET_FILE}")
    print(f"TOTAL_SOURCE_ROWS|{total_source_rows}")
    print(f"MANUAL_APPROVE_APPLIED|{manual_approve_applied}")
    print(f"MANUAL_NO_MASTER|{manual_no_master}")
    print(f"MATCH_TOTAL|{match_total}")
    print(f"REVIEW_TOTAL|{review_total}")
    print(f"EMPTY_SKU_MATCH|{empty_sku_match}")
    print(f"MANUAL_MAPPING_MISMATCH|{manual_mapping_mismatch}")

    print("REVIEW_REMAINING_SAMPLE|12")
    print("SOURCE_ROW|SOURCE_NAME|MATCH_STATUS|MATCH_REASONS|MATCH_SKU|MATCH_MASTER_NAME")
    for row in review_rows_out:
        print("|".join(row))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
