from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .models import PriceDataset, PriceRecord


class PriceLoader:
    _ALLOWED_STATUSES = {"ACTIVE": "Active", "INACTIVE": "Inactive", "DRAFT": "Draft", "ARCHIVED": "Archived"}
    _REQUIRED_HEADERS = (
        "SKU",
        "UNITCOST",
        "RETAILPRICE",
        "WHOLESALEPRICE",
        "MARKETPLACEPRICE",
        "EFFECTIVEFROM",
        "EFFECTIVETO",
        "STATUS",
    )

    @staticmethod
    def _normalize_header(value: object) -> str:
        return "".join(str(value or "").strip().upper().split())

    @staticmethod
    def _normalize_sku(value: object) -> str:
        return str(value or "").strip().upper()

    @staticmethod
    def _to_decimal(value: object) -> Decimal | None:
        if value is None:
            return None
        if isinstance(value, Decimal):
            return value
        text = str(value).strip().replace(" ", "")
        if not text:
            return None
        text = text.replace(",", ".")
        try:
            return Decimal(text)
        except Exception:
            return None

    @staticmethod
    def _to_date(value: object) -> date | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        if not text:
            return None
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    def _validate_headers(self, ws) -> dict[str, int]:
        headers = [self._normalize_header(cell.value) for cell in ws[1]]
        idx = {header: i + 1 for i, header in enumerate(headers) if header}
        missing = [header for header in self._REQUIRED_HEADERS if header not in idx]
        if missing:
            raise ValueError(f"PRICE workbook is missing required columns: {', '.join(missing)}")
        return idx

    def load(self, price_file: str | Path, valuation_date: date | None = None) -> PriceDataset:
        path = Path(price_file)
        valuation = valuation_date or date.today()

        wb = load_workbook(filename=path, data_only=True)
        ws = wb.active
        idx = self._validate_headers(ws)

        records: list[PriceRecord] = []
        warnings: list[str] = []
        row_count = 0
        rejected_row_count = 0

        for row_num in range(2, ws.max_row + 1):
            sku = self._normalize_sku(ws.cell(row=row_num, column=idx["SKU"]).value)
            unit_cost = self._to_decimal(ws.cell(row=row_num, column=idx["UNITCOST"]).value)
            retail_price = self._to_decimal(ws.cell(row=row_num, column=idx["RETAILPRICE"]).value)
            wholesale_price = self._to_decimal(ws.cell(row=row_num, column=idx["WHOLESALEPRICE"]).value)
            marketplace_price = self._to_decimal(ws.cell(row=row_num, column=idx["MARKETPLACEPRICE"]).value)
            effective_from = self._to_date(ws.cell(row=row_num, column=idx["EFFECTIVEFROM"]).value)
            effective_to = self._to_date(ws.cell(row=row_num, column=idx["EFFECTIVETO"]).value)
            raw_status = str(ws.cell(row=row_num, column=idx["STATUS"]).value or "").strip().upper()
            status = self._ALLOWED_STATUSES.get(raw_status)

            values = [sku, unit_cost, retail_price, wholesale_price, marketplace_price, effective_from, status]
            if all(value in (None, "") for value in values):
                continue

            row_count += 1

            if not sku:
                rejected_row_count += 1
                continue
            if None in (unit_cost, retail_price, wholesale_price, marketplace_price, effective_from):
                rejected_row_count += 1
                continue
            if status is None:
                rejected_row_count += 1
                continue
            if unit_cost < 0 or retail_price < 0 or wholesale_price < 0 or marketplace_price < 0:
                rejected_row_count += 1
                continue
            if effective_to is not None and effective_to < effective_from:
                rejected_row_count += 1
                continue

            records.append(
                PriceRecord(
                    sku=sku,
                    unit_cost=unit_cost,
                    retail_price=retail_price,
                    wholesale_price=wholesale_price,
                    marketplace_price=marketplace_price,
                    effective_from=effective_from,
                    effective_to=effective_to,
                    status=status,
                    source_row=row_num,
                )
            )

        dataset = PriceDataset(
            records=records,
            source_file_name=path.name,
            source_loaded_at=datetime.now(timezone.utc),
            row_count=row_count,
            valid_row_count=len(records),
            rejected_row_count=rejected_row_count,
            warning_count=0,
            warnings=warnings,
        )

        grouped = dataset.by_sku()
        duplicate_active_skus: list[str] = []
        pricing_status_by_sku: dict[str, str] = {}

        for sku, sku_records in grouped.items():
            status = dataset.resolve_pricing_status(sku, valuation)
            pricing_status_by_sku[sku] = status

            active_effective = [
                record
                for record in sku_records
                if record.is_active() and record.is_effective_on(valuation)
            ]
            if len(active_effective) > 1:
                duplicate_active_skus.append(sku)
                warnings.append(f"Duplicate active price overlap for SKU {sku}")

        dataset.duplicate_active_skus = sorted(duplicate_active_skus)
        dataset.pricing_status_by_sku = pricing_status_by_sku
        dataset.warning_count = len(warnings)

        return dataset
