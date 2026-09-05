from openpyxl import load_workbook

ws = load_workbook("output/SALES_MATCH_MANUAL.xlsx", data_only=True).active
print("SHEET|" + ws.title)
print("MAX|" + str(ws.max_row) + "|" + str(ws.max_column))
print(
    "HEADERS|"
    + "|".join(
        "" if ws.cell(1, c).value is None else str(ws.cell(1, c).value)
        for c in range(1, ws.max_column + 1)
    )
)

for r in range(2, min(ws.max_row, 8) + 1):
    print(
        "ROW|"
        + str(r)
        + "|"
        + "|".join(
            "" if ws.cell(r, c).value is None else str(ws.cell(r, c).value)
            for c in range(1, min(ws.max_column, 12) + 1)
        )
    )
