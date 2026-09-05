from openpyxl import load_workbook
from pathlib import Path

base = Path('output')
out = Path('_tmp_manual_inputs_info.txt')

lines = []
wb = load_workbook(base / 'SALES_PRODUCT_TYPE_REVIEW.xlsx', data_only=True)
ws = wb['REVIEW_79']
lines.append(f"REVIEW_SHEET|{ws.title}")
lines.append(f"REVIEW_MAX|{ws.max_row}|{ws.max_column}")
lines.append("REVIEW_HEADERS|" + "|".join("" if ws.cell(1,c).value is None else str(ws.cell(1,c).value) for c in range(1, ws.max_column+1)))

wbm = load_workbook(base / 'MASTER_DATASET.xlsx', data_only=True)
ms = wbm.active
lines.append(f"MASTER_SHEET|{ms.title}")
lines.append(f"MASTER_MAX|{ms.max_row}|{ms.max_column}")
lines.append("MASTER_HEADERS|" + "|".join("" if ms.cell(1,c).value is None else str(ms.cell(1,c).value) for c in range(1, ms.max_column+1)))

out.write_text("\n".join(lines), encoding='utf-8')
print(str(out))
