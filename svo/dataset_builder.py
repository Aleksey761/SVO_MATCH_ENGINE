from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook

from .loader import Loader
from .models import MasterItem


@dataclass
class DatasetRecord:
    """One unified MASTER dataset row keyed by SKU."""

    sku: str
    master_name: str
    category: str
    brand: str
    aroma: str
    arrival_quantity: int = 0
    sales_quantity: int = 0
    last_arrival_date: str = ""
    last_sales_date: str = ""

    @property
    def total_movement(self) -> int:
        return self.arrival_quantity + self.sales_quantity


class DatasetBuilder:
    """Builds and loads MASTER_DATASET as the central analytics dataset."""

    HEADERS = [
        "SKU",
        "MASTER_NAME",
        "CATEGORY",
        "BRAND",
        "AROMA",
        "ARRIVAL_QUANTITY",
        "SALES_QUANTITY",
        "TOTAL_MOVEMENT",
        "LAST_ARRIVAL_DATE",
        "LAST_SALES_DATE",
    ]

    def __init__(self):
        self.loader = Loader()

    def _count_warehouse_quantities(self, row: list, warehouse_columns: int = 11) -> int:
        """Count numeric quantities from columns 2..11."""
        total = 0
        for col_idx in range(1, min(warehouse_columns, len(row))):
            value = row[col_idx]
            if value is None:
                continue
            try:
                if isinstance(value, (int, float)):
                    total += int(round(value))
                elif isinstance(value, str):
                    text = value.strip()
                    if text:
                        total += int(round(float(text)))
            except (ValueError, TypeError):
                pass
        return total

    def _extract_matched_sku(self, row: list) -> Optional[str]:
        if len(row) < 13:
            return None
        sku = row[12]
        if sku and str(sku).strip():
            return str(sku).strip()
        return None

    def _extract_match_status(self, row: list) -> str:
        if len(row) < 12:
            return "UNKNOWN"
        status = row[11]
        if status:
            return str(status).strip().upper()
        return "UNKNOWN"

    def _date_to_key(self, date_str: str) -> tuple[int, int, int]:
        if not date_str:
            return (0, 0, 0)
        try:
            parsed = datetime.strptime(date_str, "%d.%m.%Y")
            return (parsed.year, parsed.month, parsed.day)
        except ValueError:
            return (0, 0, 0)

    def _max_date(self, left: str, right: str) -> str:
        return right if self._date_to_key(right) > self._date_to_key(left) else left

    def _build_validation(
        self,
        *,
        master_records: int,
        dataset_records: int,
        arrival_linked: int,
        sales_linked: int,
        missing_master_sku: int,
        duplicate_sku: int,
    ) -> dict:
        integrity_ok = (
            master_records == dataset_records
            and missing_master_sku == 0
            and duplicate_sku == 0
        )
        return {
            "master_records": master_records,
            "dataset_records": dataset_records,
            "arrival_linked": arrival_linked,
            "sales_linked": sales_linked,
            "missing_master_sku": missing_master_sku,
            "duplicate_sku": duplicate_sku,
            "integrity": "OK" if integrity_ok else "FAILED",
        }

    def _build_summary(self, records: list[DatasetRecord]) -> dict:
        master_sku = len(records)
        arrival_linked = sum(1 for r in records if r.arrival_quantity > 0)
        sales_linked = sum(1 for r in records if r.sales_quantity > 0)
        arrival_qty = sum(r.arrival_quantity for r in records)
        sales_qty = sum(r.sales_quantity for r in records)
        movement_linked = sum(1 for r in records if r.total_movement > 0)
        coverage = (movement_linked / master_sku * 100.0) if master_sku else 0.0
        return {
            "master_sku": master_sku,
            "arrival_linked": arrival_linked,
            "sales_linked": sales_linked,
            "arrival_qty": arrival_qty,
            "sales_qty": sales_qty,
            "coverage_pct": coverage,
        }

    def _print_validation(self, validation: dict) -> None:
        print("=================================")
        print("MASTER DATASET")
        print("=================================")
        print(f"MASTER records      : {validation['master_records']}")
        print(f"Dataset records     : {validation['dataset_records']}")
        print(f"Arrival linked      : {validation['arrival_linked']}")
        print(f"Sales linked        : {validation['sales_linked']}")
        print(f"Missing MASTER SKU  : {validation['missing_master_sku']}")
        print(f"Duplicate SKU       : {validation['duplicate_sku']}")
        print(f"Integrity           : {validation['integrity']}")

    def _print_summary(self, summary: dict) -> None:
        print("=================================")
        print("DATASET SUMMARY")
        print("=================================")
        print(f"MASTER SKU          : {summary['master_sku']}")
        print(f"Arrival linked      : {summary['arrival_linked']}")
        print(f"Sales linked        : {summary['sales_linked']}")
        print(f"Arrival Qty         : {summary['arrival_qty']}")
        print(f"Sales Qty           : {summary['sales_qty']}")
        print(f"Coverage %          : {summary['coverage_pct']:.2f}")
        print("=================================")

    def _collect_match_files(
        self,
        explicit: Optional[str | Path],
        pattern: str,
        output_dir: Path,
    ) -> list[Path]:
        if explicit is None:
            # Respect the caller's output directory. Do not fall back to the
            # repository-wide output directory: isolated builds/tests must not
            # consume unrelated production match workbooks.
            return sorted(output_dir.glob(pattern))
        return [Path(explicit)]

    def _load_authoritative_master(self, master_file: Path) -> list[MasterItem]:
        """Load the explicitly selected MASTER workbook directly.

        DatasetBuilder owns the MASTER source contract for dataset construction: 
        when build() selects input_dir/MASTER.xlsx, that exact workbook must be
        parsed, even if an older Loader implementation contains legacy fallback
        logic that substitutes MASTER_DATASET.xlsx.
        """
        wb = load_workbook(filename=master_file, data_only=True)
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        header_to_index = {
            str(value).strip().upper(): idx
            for idx, value in enumerate(headers)
            if value is not None and str(value).strip()
        }

        sku_idx = header_to_index.get("SKU", 0)
        category_idx = header_to_index.get("CATEGORY", 1)
        brand_idx = header_to_index.get("BRAND", 2)
        variant_idx = header_to_index.get("VARIANT", 3)
        volume_idx = header_to_index.get("VOLUME", 4)
        aroma_idx = header_to_index.get("AROMA", 5)

        def value_at(row: tuple, index: int | None) -> str:
            if index is None or index >= len(row):
                return ""
            return str(row[index] or "").strip()

        items: list[MasterItem] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            sku = value_at(row, sku_idx)
            if not sku:
                continue
            items.append(
                MasterItem(
                    sku=sku,
                    category=value_at(row, category_idx),
                    brand=value_at(row, brand_idx),
                    variant=value_at(row, variant_idx),
                    volume=value_at(row, volume_idx),
                    aroma=value_at(row, aroma_idx),
                    master_name="",
                )
            )
        return items

    def build(
        self,
        input_dir: str | Path = "data",
        arrival_match_file: Optional[str | Path] = None,
        sales_match_file: Optional[str | Path] = None,
        output_file: str | Path = "output/MASTER_DATASET.xlsx",
        print_reports: bool = True,
    ) -> dict:
        """Build MASTER_DATASET and print validation + summary reports."""
        input_dir = Path(input_dir)
        output_file = Path(output_file)

        # MASTER.xlsx inside the requested input directory is authoritative.
        # Do not let Loader.discover_workbooks() substitute another workbook
        # such as a previously generated MASTER_DATASET.xlsx.
        candidates = sorted(input_dir.glob("MASTER.xlsx"))
        if candidates:
            master_file = candidates[0]
        else:
            try:
                master_file, _, _, _, _ = self.loader.discover_workbooks(
                    input_dir, require_sales=False
                )
            except ValueError as exc:
                raise ValueError(f"MASTER.xlsx not found in {input_dir}") from exc

        master_items = self._load_authoritative_master(master_file)
        dataset: dict[str, DatasetRecord] = {}
        duplicate_sku = 0

        for master in master_items:
            if master.sku in dataset:
                duplicate_sku += 1
                continue
            dataset[master.sku] = DatasetRecord(
                sku=master.sku,
                master_name=f"{master.category} {master.brand} {master.variant} {master.volume}".strip(),
                category=master.category,
                brand=master.brand,
                aroma=(master.aroma or master.variant or "").strip(),
            )

        missing_master_sku = 0
        arrival_files = self._collect_match_files(arrival_match_file, "ARRIVAL_MATCH_*.xlsx", output_file.parent)
        for file_path in arrival_files:
            if not file_path.exists():
                continue
            file_date = self.loader.parse_arrival_date_from_filename(file_path.name) or ""
            missing_master_sku += self._merge_matches(
                file_path=file_path,
                dataset=dataset,
                quantity_attr="arrival_quantity",
                date_attr="last_arrival_date",
                file_date=file_date,
            )

        sales_files = self._collect_match_files(sales_match_file, "SALES_MATCH_*.xlsx", output_file.parent)
        for file_path in sales_files:
            if not file_path.exists():
                continue
            file_date = self.loader.parse_arrival_date_from_filename(file_path.name) or ""
            missing_master_sku += self._merge_matches(
                file_path=file_path,
                dataset=dataset,
                quantity_attr="sales_quantity",
                date_attr="last_sales_date",
                file_date=file_date,
            )

        self._write_dataset(dataset, output_file)

        records = list(dataset.values())
        arrival_linked = sum(1 for r in records if r.arrival_quantity > 0)
        sales_linked = sum(1 for r in records if r.sales_quantity > 0)

        validation = self._build_validation(
            master_records=len(master_items),
            dataset_records=len(records),
            arrival_linked=arrival_linked,
            sales_linked=sales_linked,
            missing_master_sku=missing_master_sku,
            duplicate_sku=duplicate_sku,
        )
        summary = self._build_summary(records)

        if print_reports:
            self._print_validation(validation)
            print()
            self._print_summary(summary)

        return {
            "output_file": str(output_file),
            "validation": validation,
            "summary": summary,
        }

    def _merge_matches(
        self,
        *,
        file_path: Path,
        dataset: dict[str, DatasetRecord],
        quantity_attr: str,
        date_attr: str,
        file_date: str,
    ) -> int:
        wb = load_workbook(file_path, data_only=True)
        ws = wb.active
        missing_master_sku = 0

        for row_num in range(2, ws.max_row + 1):
            row = [cell.value for cell in ws[row_num]]
            if not row or not row[0]:
                continue
            if self._extract_match_status(row) != "MATCH":
                continue

            sku = self._extract_matched_sku(row)
            if not sku:
                continue
            record = dataset.get(sku)
            if record is None:
                missing_master_sku += 1
                continue

            quantity = self._count_warehouse_quantities(row)
            if quantity <= 0:
                continue

            setattr(record, quantity_attr, getattr(record, quantity_attr) + quantity)
            if file_date:
                current = getattr(record, date_attr)
                setattr(record, date_attr, self._max_date(current, file_date))

        return missing_master_sku

    def _write_dataset(self, dataset: dict[str, DatasetRecord], output_file: Path) -> None:
        output_file.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "MASTER_DATASET"
        ws.append(self.HEADERS)

        for sku in sorted(dataset.keys()):
            record = dataset[sku]
            ws.append([
                record.sku,
                record.master_name,
                record.category,
                record.brand,
                record.aroma,
                record.arrival_quantity,
                record.sales_quantity,
                record.total_movement,
                record.last_arrival_date,
                record.last_sales_date,
            ])

        wb.save(output_file)

    def load_dataset(self, dataset_file: str | Path = "output/MASTER_DATASET.xlsx") -> list[DatasetRecord]:
        """Load dataset rows from the canonical MASTER_DATASET workbook."""
        workbook = load_workbook(dataset_file, data_only=True)
        ws = workbook.active

        records: list[DatasetRecord] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            records.append(
                DatasetRecord(
                    sku=str(row[0]).strip(),
                    master_name=str(row[1] or "").strip(),
                    category=str(row[2] or "").strip(),
                    brand=str(row[3] or "").strip(),
                    aroma=str(row[4] or "").strip(),
                    arrival_quantity=int(row[5] or 0),
                    sales_quantity=int(row[6] or 0),
                    last_arrival_date=str(row[8] or "").strip(),
                    last_sales_date=str(row[9] or "").strip(),
                )
            )
        return records
