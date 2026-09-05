import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


class InventoryInputReader:
    """Reads Excel input sources. No business calculations here."""

    @staticmethod
    def _normalize_header(value: object) -> str:
        text = str(value or "").strip().upper()
        return " ".join(text.split())

    @staticmethod
    def _is_date_header(value: object) -> bool:
        return isinstance(value, (datetime, date))

    @staticmethod
    def _is_stock_header(normalized_header: str, raw_header: object) -> bool:
        if "ОСТАТОК НА СКЛАДЕ" in normalized_header:
            return True
        if normalized_header in {"ВСЕГО", "ИТОГО"}:
            return True
        # Compact inventory exports may keep only the date in the stock header cell.
        return InventoryInputReader._is_date_header(raw_header)

    @staticmethod
    def _looks_like_product_name(value: object) -> bool:
        text = str(value or "").strip()
        if not text:
            return False
        if text.isdigit():
            return False
        return bool(re.search(r"[A-Za-zА-Яа-я]", text))

    @staticmethod
    def _to_number(value: object) -> float:
        if value is None:
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)

        text = str(value).strip().replace(" ", "")
        if not text:
            return 0.0

        text = text.replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return 0.0

    @staticmethod
    def _to_int(value: float) -> int:
        return int(round(value))

    @staticmethod
    def _parse_date_triplet(value: object) -> tuple[int, int, int] | None:
        if isinstance(value, datetime):
            return value.day, value.month, value.year
        if isinstance(value, date):
            return value.day, value.month, value.year

        text = str(value or "")
        match = re.search(r"(?<!\d)(\d{1,2})\D+(\d{1,2})\D+(\d{2}|\d{4})(?!\d)", text)
        if not match:
            return None
        day = int(match.group(1))
        month = int(match.group(2))
        year = int(match.group(3))
        if year < 100:
            year += 2000
        return day, month, year

    @staticmethod
    def _parse_date_from_filename_triplet(filename: str | Path) -> tuple[int, int, int] | None:
        return InventoryInputReader._parse_date_triplet(Path(filename).stem)

    def _find_opening_stock_column(self, headers: list[object], normalized_headers: list[str]) -> int | None:
        dated_stock_columns: list[tuple[int, tuple[int, int, int]]] = []
        stock_columns: list[int] = []
        for idx, (raw_header, normalized_header) in enumerate(zip(headers, normalized_headers), start=1):
            if not self._is_stock_header(normalized_header, raw_header):
                continue
            stock_columns.append(idx)
            date_triplet = self._parse_date_triplet(raw_header)
            if date_triplet is None:
                date_triplet = self._parse_date_triplet(normalized_header)
            if date_triplet is None:
                continue
            dated_stock_columns.append((idx, date_triplet))

        if not dated_stock_columns:
            if stock_columns:
                return stock_columns[-1]
            return None

        # Opening stock is the earliest dated stock snapshot in header.
        return min(dated_stock_columns, key=lambda item: (item[1][2], item[1][1], item[1][0]))[0]

    def _find_inventory_header_row(self, ws) -> int:
        max_scan_rows = min(ws.max_row, 40)
        max_scan_cols = min(ws.max_column, 120)

        for row in range(1, max_scan_rows + 1):
            raw_values = [ws.cell(row=row, column=col).value for col in range(1, max_scan_cols + 1)]
            normalized = [self._normalize_header(value) for value in raw_values]
            has_name = any(value == "НАИМЕНОВАНИЕ" for value in normalized)
            has_opening = any(
                self._is_stock_header(normalized_value, raw_value)
                for normalized_value, raw_value in zip(normalized, raw_values)
            )
            if has_name and has_opening:
                return row

        return 1

    def _detect_product_column(self, ws, header_row: int, name_header_col: int) -> int:
        candidates = [name_header_col]
        if name_header_col + 1 <= ws.max_column:
            candidates.append(name_header_col + 1)

        best_col = name_header_col
        best_score = -1
        max_scan_row = min(ws.max_row, header_row + 120)
        for col in candidates:
            score = 0
            for row in range(header_row + 1, max_scan_row + 1):
                if self._looks_like_product_name(ws.cell(row=row, column=col).value):
                    score += 1
            if score > best_score:
                best_score = score
                best_col = col

        return best_col

    def _infer_product_column(self, ws, header_row: int) -> int | None:
        best_col = None
        best_score = -1
        max_scan_row = min(ws.max_row, header_row + 120)
        max_scan_col = min(ws.max_column, 12)
        for col in range(1, max_scan_col + 1):
            score = 0
            for row in range(header_row + 1, max_scan_row + 1):
                if self._looks_like_product_name(ws.cell(row=row, column=col).value):
                    score += 1
            if score > best_score:
                best_score = score
                best_col = col
        if best_col is None or best_score <= 0:
            return None
        return best_col

    def _detect_actual_stock_column(
        self,
        headers: list[object],
        normalized_headers: list[str],
        *,
        opening_col: int | None,
        inventory_file: str | Path,
    ) -> int | None:
        stock_columns = [
            idx
            for idx, (raw_header, normalized_header) in enumerate(zip(headers, normalized_headers), start=1)
            if self._is_stock_header(normalized_header, raw_header)
        ]
        if not stock_columns:
            return opening_col

        target_date = self._parse_date_from_filename_triplet(inventory_file)
        if target_date is None:
            non_opening = [col for col in stock_columns if col != opening_col]
            if non_opening:
                return non_opening[-1]
            return stock_columns[-1]

        target_day, target_month, target_year = target_date
        candidates: list[tuple[int, int]] = []
        for col in stock_columns:
            header_date = self._parse_date_triplet(headers[col - 1])
            if header_date is None:
                header_date = self._parse_date_triplet(normalized_headers[col - 1])
            if header_date is None:
                continue
            day, month, year = header_date
            if month == target_month and year == target_year:
                candidates.append((col, day))

        if not candidates:
            non_opening = [col for col in stock_columns if col != opening_col]
            if non_opening:
                return non_opening[-1]
            return stock_columns[-1]

        not_later = [item for item in candidates if item[1] <= target_day]
        if not_later:
            return max(not_later, key=lambda item: item[1])[0]
        return min(candidates, key=lambda item: item[1])[0]

    def read_inventory_quantities(self, inventory_file: str | Path) -> list[dict[str, object]]:
        wb = load_workbook(filename=inventory_file, data_only=True)
        ws = wb.active

        header_row = self._find_inventory_header_row(ws)
        headers = [ws.cell(row=header_row, column=col).value for col in range(1, ws.max_column + 1)]

        header_to_col: dict[str, int] = {}
        normalized_headers: list[str] = []
        for idx, header in enumerate(headers, start=1):
            normalized = self._normalize_header(header)
            normalized_headers.append(normalized)
            if normalized:
                header_to_col[normalized] = idx

        name_header_col = header_to_col.get("НАИМЕНОВАНИЕ")
        product_col = None
        if name_header_col is not None:
            product_col = self._detect_product_column(ws, header_row, name_header_col)
        else:
            # Fallback for compact inventory format: product names are in column C.
            product_col = 3

        opening_col = self._find_opening_stock_column(headers, normalized_headers)
        if opening_col is None:
            # Fallback for compact inventory format: total stock is in column O.
            opening_col = 15

        opening_header = headers[opening_col - 1] if 0 < opening_col <= len(headers) else None
        opening_header_date = self._parse_date_triplet(opening_header)
        if opening_header_date is None:
            opening_header_date = self._parse_date_triplet(
                normalized_headers[opening_col - 1]
            )

        # Compact inventory snapshot format: the revision date is stored
        # in the filename, while the stock column header is just ВСЕГО.
        if opening_header_date is None:
            opening_header_date = self._parse_date_from_filename_triplet(inventory_file)

        if opening_header_date is None:
            raise ValueError(
                "Required REVISION columns not found: "
                "НАИМЕНОВАНИЕ / Остаток на складе (шт) with date header"
            )

        actual_col = self._detect_actual_stock_column(
            headers,
            normalized_headers,
            opening_col=opening_col,
            inventory_file=inventory_file,
        )
        if actual_col is None:
            actual_col = opening_col

        voronezh_col = header_to_col.get("ВОРОНЕЖ")
        krasnodar1_col = header_to_col.get("КРАСНОДАР 1")
        krasnodar2_col = header_to_col.get("КРАСНОДАР 2")
        retail_col = header_to_col.get("РОЗНИЦА")
        sku_col = None
        for candidate in ("SKU", "АРТИКУЛ", "АРТИКУЛ ПОСТАВЩИКА", "КОД"):
            if candidate in header_to_col:
                sku_col = header_to_col[candidate]
                break

        if product_col is None or opening_col is None:
            raise ValueError(
                "Required REVISION columns not found: НАИМЕНОВАНИЕ / Остаток на складе (шт) with date header"
            )

        revision_date_col = 16  # Column P
        first_valid_revision_date: tuple[int, int, int] | None = None
        for probe_row in range(3, ws.max_row + 1):
            parsed = self._parse_date_triplet(ws.cell(row=probe_row, column=revision_date_col).value)
            if parsed is not None:
                first_valid_revision_date = parsed
                break

        quantities: list[dict[str, object]] = []
        # Start immediately below the detected header row. The previous
        # minimum-row=3 rule silently discarded the first data record when
        # the workbook header was on row 1.
        data_start_row = header_row + 1
        for row_idx in range(data_start_row, ws.max_row + 1):
            source_name = str(ws.cell(row=row_idx, column=product_col).value or "").strip()
            if not self._looks_like_product_name(source_name):
                continue

            opening_qty = self._to_number(ws.cell(row=row_idx, column=opening_col).value)
            actual_qty = self._to_number(ws.cell(row=row_idx, column=actual_col).value) if actual_col is not None else 0.0
            # Confirmed v1 mapping:
            # - regional columns are supplier returns by warehouse location
            # - retail column is end-customer sales
            supplier_return_qty = 0.0
            supplier_return_qty += self._to_number(ws.cell(row=row_idx, column=voronezh_col).value) if voronezh_col else 0.0
            supplier_return_qty += self._to_number(ws.cell(row=row_idx, column=krasnodar1_col).value) if krasnodar1_col else 0.0
            supplier_return_qty += self._to_number(ws.cell(row=row_idx, column=krasnodar2_col).value) if krasnodar2_col else 0.0
            sales_qty = self._to_number(ws.cell(row=row_idx, column=retail_col).value) if retail_col else 0.0

            payload = {
                "source_name": source_name,
                "opening": self._to_int(opening_qty),
                "receipt": 0,
                "supplier_return": self._to_int(supplier_return_qty),
                "sales": self._to_int(sales_qty),
                "actual": self._to_int(actual_qty),
            }
            row_revision_date = self._parse_date_triplet(ws.cell(row=row_idx, column=revision_date_col).value)
            effective_revision_date = row_revision_date or first_valid_revision_date
            if effective_revision_date is not None:
                day, month, year = effective_revision_date
                payload["revision_date"] = f"{day:02d}.{month:02d}.{year:04d}"
            if sku_col is not None:
                sku_value = str(ws.cell(row=row_idx, column=sku_col).value or "").strip()
                if sku_value:
                    payload["sku"] = sku_value

            quantities.append(payload)

        return quantities

    @staticmethod
    def _find_result_header_row(ws) -> int:
        max_scan_rows = min(ws.max_row, 30)
        for row in range(1, max_scan_rows + 1):
            normalized = [
                InventoryInputReader._normalize_header(ws.cell(row=row, column=col).value)
                for col in range(1, ws.max_column + 1)
            ]
            has_source = "SOURCE_NAME" in normalized
            has_sku = "SKU" in normalized
            has_status = "STATUS" in normalized
            if has_source and has_sku and has_status:
                return row
        raise ValueError("RESULT workbook header row not found (SOURCE_NAME/SKU/STATUS)")

    def read_result_rows(self, result_file: str | Path) -> list[dict[str, str]]:
        wb = load_workbook(filename=result_file, data_only=True)
        ws = wb.active

        header_row = self._find_result_header_row(ws)
        headers = [
            self._normalize_header(ws.cell(row=header_row, column=col).value)
            for col in range(1, ws.max_column + 1)
        ]
        idx = {name: col for col, name in enumerate(headers, start=1) if name}

        required = ("SOURCE_NAME", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "SKU", "MASTER_NAME", "STATUS")
        missing = [name for name in required if name not in idx]
        if missing:
            raise ValueError(f"RESULT workbook is missing required columns: {', '.join(missing)}")

        rows: list[dict[str, str]] = []
        for row in range(header_row + 1, ws.max_row + 1):
            source_name = str(ws.cell(row=row, column=idx["SOURCE_NAME"]).value or "").strip()
            status = str(ws.cell(row=row, column=idx["STATUS"]).value or "").strip().upper()
            if not source_name and not status:
                continue

            rows.append(
                {
                    "source_name": source_name,
                    "category": str(ws.cell(row=row, column=idx["CATEGORY"]).value or "").strip(),
                    "brand": str(ws.cell(row=row, column=idx["BRAND"]).value or "").strip(),
                    "variant": str(ws.cell(row=row, column=idx["VARIANT"]).value or "").strip(),
                    "volume": str(ws.cell(row=row, column=idx["VOLUME"]).value or "").strip(),
                    "sku": str(ws.cell(row=row, column=idx["SKU"]).value or "").strip(),
                    "master_name": str(ws.cell(row=row, column=idx["MASTER_NAME"]).value or "").strip(),
                    "status": status,
                }
            )

        return rows
