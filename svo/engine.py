from pathlib import Path
from datetime import datetime, timezone
import re
from shutil import copyfile

from openpyxl import Workbook, load_workbook

from svo.arrival_loader import ArrivalLoader
from svo.business_rules import BusinessRules
from svo.learning_map import LearningMap, canonical_supplier_value
from svo.loader import Loader
from svo.matcher import Matcher
from svo.normalizer import Normalizer
from svo.price_loader import PriceLoader
from svo.quality import run_quality_gate
from svo.revision_name_canonicalizer import RevisionNameCanonicalizer
from svo.revision_loader import RevisionLoader
from svo.sales_loader import SalesLoader
from svo.reporter import Reporter
from finance.price_history import PriceHistory


class Engine:
    """Coordinates the complete matching pipeline."""

    def __init__(self, learning_map_file: str | Path | None = None):
        self.loader = Loader()
        self.arrival_loader = ArrivalLoader()
        self.revision_loader = RevisionLoader()
        self.revision_name_canonicalizer = RevisionNameCanonicalizer()
        self.sales_loader = SalesLoader()
        self.price_loader = PriceLoader()
        self.price_history = PriceHistory()
        self.normalizer = Normalizer()
        self.learning_map = LearningMap(learning_map_file or Path("data") / "learning_map.xlsx")
        self.items = []

    @staticmethod
    def _master_by_sku(master_items: list) -> dict[str, object]:
        return {str(item.sku or "").strip(): item for item in master_items if str(item.sku or "").strip()}

    def _apply_learning_map_matches(self, items: list, master_items: list) -> tuple[list, int, int]:
        master_by_sku = self._master_by_sku(master_items)
        unresolved: list = []
        reused = 0
        stale = 0

        for item in items:
            supplier_article = getattr(item, "supplier_article", None)
            canonical_supplier = canonical_supplier_value(item)
            if not supplier_article or not canonical_supplier:
                unresolved.append(item)
                continue

            entry = self.learning_map.lookup(supplier_article, canonical_supplier)
            if entry is None:
                unresolved.append(item)
                continue

            master_item = master_by_sku.get(entry.sku)
            if master_item is None:
                stale += 1
                unresolved.append(item)
                continue

            item.sku = master_item.sku
            item.master_name = master_item.master_name or entry.master_name
            item.status = "MATCH"
            item.confidence = float(entry.confidence or 100.0)
            item.review_reasons = []
            item.candidates = []
            item.review_explanation = {
                "confidence": item.confidence,
                "reasons": [],
                "candidates": [],
                "source": "LEARNING_MAP",
            }
            reused += 1

        return unresolved, reused, stale

    def confirm_review_matches(
        self,
        *,
        master_file: str | Path,
        confirmed_rows: list[dict[str, object]],
        confirmed_by: str = "MANUAL",
    ) -> int:
        """Persist manually confirmed REVIEW decisions into learning map."""
        master_items = self.loader.load_master(master_file)
        master_by_sku = self._master_by_sku(master_items)

        saved = 0
        for row in confirmed_rows:
            supplier_article = row.get("SupplierArticle")
            canonical_supplier = row.get("CanonicalSupplier")
            sku = str(row.get("SKU") or "").strip()
            if not supplier_article or not canonical_supplier or not sku:
                continue

            master_item = master_by_sku.get(sku)
            if master_item is None:
                continue

            confidence = float(row.get("Confidence") or 100.0)
            supplier_name = row.get("SupplierName") or ""
            self.learning_map.upsert(
                supplier_article=supplier_article,
                supplier_name=supplier_name,
                canonical_supplier=canonical_supplier,
                sku=master_item.sku,
                master_name=master_item.master_name or row.get("MASTER_NAME") or "",
                confidence=confidence,
                confirmed_by=confirmed_by,
            )
            saved += 1

        return saved

    @staticmethod
    def _price_items_to_history_rows(items: list) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        for item in items:
            if not item.sku:
                continue
            rows.append(
                {
                    "SKU": item.sku,
                    "MASTER_NAME": item.master_name or "",
                    "UnitCost": item.unit_cost,
                    "RetailPrice": item.retail_price,
                    "SourceUnit": getattr(item, "source_unit", None),
                    "PackQty": getattr(item, "pack_qty", None),
                    "OriginalPrice": getattr(item, "original_price", item.unit_cost),
                }
            )
        return rows

    @staticmethod
    def _normalize_price_name(value: object) -> str:
        return " ".join(str(value or "").strip().upper().split())

    def _validate_price_name_sku_conflicts(self, price_items: list) -> None:
        grouped: dict[str, list] = {}
        for item in price_items:
            canonical_name = self._normalize_price_name(getattr(item, "source_name", ""))
            if not canonical_name:
                continue
            grouped.setdefault(canonical_name, []).append(item)

        for canonical_name, items in grouped.items():
            sku_set = {str(item.sku or "").strip().upper() for item in items if str(item.sku or "").strip()}
            if len(sku_set) <= 1:
                continue

            details = [
                (
                    f"row={item.row_number}",
                    f"sku={item.sku or ''}",
                    f"unit_cost={item.unit_cost}",
                    f"retail_price={item.retail_price}",
                    f"source_name={item.source_name}",
                )
                for item in items
            ]
            print(f"PRICE SKU conflict for canonical ProductName: {canonical_name}")
            for detail in details:
                print("PRICE SKU conflict row: " + " ".join(detail))

            raise ValueError(
                "REVIEW: Duplicate active price conflict for ProductName "
                f"'{canonical_name}' (different SKU: {', '.join(sorted(sku_set))})"
            )

    @staticmethod
    def _write_price_change_report(events: list[dict[str, object]], output_file: str | Path) -> str:
        report_path = Path(output_file)
        wb = Workbook()
        ws = wb.active
        ws.title = "PRICE_CHANGES"
        ws.append([
            "SKU",
            "MASTER_NAME",
            "PreviousPrice",
            "NewPrice",
            "Difference",
            "PercentChange",
            "SourceUnit",
            "PackQty",
            "OriginalPrice",
            "ValidFrom",
            "ImportedAt",
            "SourceFile",
            "Status",
        ])

        for event in events:
            status = str(event.get("Status") or "").upper()
            if status not in {"NEW", "UPDATED"}:
                continue

            ws.append(
                [
                    event.get("SKU"),
                    event.get("MASTER_NAME"),
                    event.get("PreviousPrice"),
                    event.get("NewPrice"),
                    event.get("Difference"),
                    event.get("PercentChange"),
                    event.get("SourceUnit"),
                    event.get("PackQty"),
                    event.get("OriginalPrice"),
                    event.get("ValidFrom").isoformat() if event.get("ValidFrom") else "",
                    event.get("ImportedAt").isoformat() if event.get("ImportedAt") else "",
                    event.get("SourceFile"),
                    status,
                ]
            )

        report_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(report_path)
        return str(report_path)

    @staticmethod
    def _write_price_match_report(items: list, output_file: str | Path, *, imported_at: datetime, source_file: str) -> str:
        report_path = Path(output_file)
        wb = Workbook()
        ws = wb.active
        ws.title = "Matched"
        ws.append(
            [
                "SKU",
                "MASTER_NAME",
                "SupplierArticle",
                "OriginalProductName",
                "CanonicalProductName",
                "UnitCost",
                "RetailPrice",
                "StockQty",
                "SourceUnit",
                "PackQty",
                "OriginalPrice",
                "ImportedAt",
                "SourceFile",
                "Status",
            ]
        )

        for item in items:
            canonical_name = getattr(item, "canonical_product_name", None)
            if not canonical_name:
                canonical_name = " ".join(str(item.source_name or "").strip().upper().split())
            ws.append(
                [
                    item.sku or "",
                    item.master_name or "",
                    getattr(item, "supplier_article", "") or "",
                    getattr(item, "original_product_name", item.source_name),
                    canonical_name,
                    item.unit_cost,
                    item.retail_price,
                    getattr(item, "stock_qty", None),
                    getattr(item, "source_unit", None),
                    getattr(item, "pack_qty", None),
                    getattr(item, "original_price", item.unit_cost),
                    imported_at.isoformat(),
                    source_file,
                    item.status,
                ]
            )

        report_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(report_path)
        return str(report_path)

    @staticmethod
    def _build_price_review_reason(item) -> tuple[str, str]:
        reasons = list(getattr(item, "review_reasons", []) or [])
        if "MULTIPLE_MATCH" in reasons:
            return "Ambiguous match", "Manual review required"
        if "LOW_SCORE" in reasons:
            return "Manual review required", "Manual review required"
        return "No MASTER match", "Manual review required"

    def _write_price_review_report(self, items: list, output_file: str | Path) -> str:
        report_path = Path(output_file)
        wb = Workbook()
        ws = wb.active
        ws.title = "PRICE_REVIEW"
        ws.append(
            [
                "SupplierArticle",
                "OriginalProductName",
                "CanonicalProductName",
                "SKU",
                "MASTER_NAME",
                "Reason",
                "Recommendation",
            ]
        )

        written_keys: set[tuple[int, str]] = set()
        for item in items:
            if item.status != "REVIEW":
                continue
            reason, recommendation = self._build_price_review_reason(item)
            key = (item.row_number, reason)
            if key in written_keys:
                continue
            written_keys.add(key)
            ws.append(
                [
                    getattr(item, "supplier_article", "") or "",
                    getattr(item, "original_product_name", item.source_name),
                    getattr(item, "canonical_product_name", " ".join(str(item.source_name or "").strip().upper().split())),
                    item.sku or "",
                    item.master_name or "",
                    reason,
                    recommendation,
                ]
            )

        grouped: dict[str, list] = {}
        for item in items:
            canonical_name = str(getattr(item, "canonical_product_name", "") or "").strip()
            if not canonical_name:
                continue
            grouped.setdefault(canonical_name, []).append(item)

        for canonical_name, group in grouped.items():
            sku_set = {str(item.sku or "").strip().upper() for item in group if str(item.sku or "").strip()}
            if len(sku_set) > 1:
                for item in group:
                    key = (item.row_number, "Different SKU conflict")
                    if key in written_keys:
                        continue
                    written_keys.add(key)
                    ws.append(
                        [
                            getattr(item, "supplier_article", "") or "",
                            getattr(item, "original_product_name", item.source_name),
                            canonical_name,
                            item.sku or "",
                            item.master_name or "",
                            "Different SKU conflict",
                            "Manual review required",
                        ]
                    )
                continue

            unit_set = {str(getattr(item, "source_unit", "") or "").strip() for item in group}
            if len(group) > 1 and len(unit_set) > 1:
                for item in group:
                    key = (item.row_number, "Packaging conflict")
                    if key in written_keys:
                        continue
                    written_keys.add(key)
                    ws.append(
                        [
                            getattr(item, "supplier_article", "") or "",
                            getattr(item, "original_product_name", item.source_name),
                            canonical_name,
                            item.sku or "",
                            item.master_name or "",
                            "Packaging conflict",
                            "Manual review required",
                        ]
                    )

        report_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(report_path)
        return str(report_path)

    @staticmethod
    def _master_to_quality_rows(master_items: list) -> list[dict[str, str]]:
        return [
            {
                "SKU": item.sku,
                "CATEGORY": item.category,
                "BRAND": item.brand,
                "VARIANT": item.variant,
                "VOLUME": item.volume,
            }
            for item in master_items
        ]

    @staticmethod
    def _is_revision_document(filename: str | Path) -> bool:
        stem = Path(filename).stem.lower()
        return "revision" in stem or "ревиз" in stem

    @staticmethod
    def _parse_revision_date_from_filename(filename: str | Path) -> str | None:
        stem = Path(filename).stem
        match = re.search(r"(?<!\d)(\d{2})\.(\d{2})\.(\d{2}|\d{4})(?!\d)", stem)
        if match:
            return f"{match.group(1)}.{match.group(2)}.{match.group(3)}"

        match = re.search(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})(?!\d)", stem)
        if match:
            year, month, day = match.group(1), match.group(2), match.group(3)
            return f"{day}.{month}.{year}"

        match = re.search(r"(?<!\d)(\d{4})_(\d{2})_(\d{2})(?!\d)", stem)
        if match:
            year, month, day = match.group(1), match.group(2), match.group(3)
            return f"{day}.{month}.{year}"

        return None

    def _discover_workbooks_for_run(
        self,
        input_dir: Path,
    ) -> tuple[Path, Path, Path | None, str | None, str | None]:
        try:
            return self.loader.discover_workbooks(
                input_dir,
                require_sales=True,
            )
        except ValueError:
            workbooks = [p for p in Path(input_dir).glob("*.xlsx") if p.is_file()]
            revision_candidates = [
                p
                for p in workbooks
                if self._is_revision_document(p)
            ]

            if len(revision_candidates) != 1:
                raise

            master_file = self.loader.discover_master_workbook(input_dir)
            revision_file = revision_candidates[0]
            sales_candidates = [
                p
                for p in workbooks
                if p.resolve() != master_file.resolve() and p.resolve() != revision_file.resolve()
            ]

            if len(sales_candidates) > 1:
                raise ValueError(f"Multiple SALES workbooks found in {input_dir}: {len(sales_candidates)}")
            if len(sales_candidates) == 0:
                raise ValueError(f"SALES workbook is missing in {input_dir}")

            sales_file = sales_candidates[0]
            revision_date = self._parse_revision_date_from_filename(revision_file.name)
            sales_date = self.loader.parse_arrival_date_from_filename(sales_file.name)
            if revision_date is None:
                print("WARNING:")
                print("Revision date not found in filename.")

            return master_file, revision_file, sales_file, revision_date, sales_date

    def _execute_document_pipeline(
        self,
        *,
        master_file: str | Path,
        document_file: str | Path,
        document_items: list,
        output_file: str | Path,
        matched_output_file: str | Path,
        document_name: str,
        document_date: str | None,
        summary_key: str,
        write_summary: bool = True,
    ):
        master = self.loader.load_master(master_file)
        master_quality = run_quality_gate(gate="master", dataframe=self._master_to_quality_rows(master))
        if not master_quality["passed"]:
            raise RuntimeError("MASTER AUDIT FAILED\nSee output/MASTER_AUDIT.txt")

        for item in document_items:
            self.normalizer.normalize(item)

        if document_name == "REVISION":
            self.revision_name_canonicalizer.apply_all(document_items)

        matcher = Matcher(master)
        unresolved_items, learning_reused, learning_stale = self._apply_learning_map_matches(document_items, master)
        matcher.match_all(unresolved_items)
        BusinessRules(master).apply(document_items)
        self.items = document_items

        match_count = sum(1 for item in document_items if item.status == "MATCH")
        review_count = sum(1 for item in document_items if item.status == "REVIEW")

        reporter = Reporter()
        if write_summary:
            reporter.write(document_items, output_file, arrival_date=document_date)
        reporter.write_matched_document(document_file, document_items, matched_output_file)
        should_finalize_sales_result = document_name == "SALES"
        if should_finalize_sales_result:
            reporter.finalize_result_by_master(
                matched_output_file,
                master,
                master_file=master_file,
            )

            if not Path(matched_output_file).name.upper().startswith("RESULT"):
                self._restore_sales_technical_columns(matched_output_file, review_count)
        review_report_file = Path(output_file).parent / "REVIEW_REPORT.xlsx"
        review_report_rows = reporter.write_review_report(document_items, review_report_file)

        return {
            "master": len(master),
            "rows": len(document_items),
            "match": match_count,
            "review": review_count,
            "items": document_items,
            "output": str(output_file),
            "review_report": str(review_report_file),
            "review_report_rows": review_report_rows,
            "learning_reused": learning_reused,
            "learning_stale": learning_stale,
            f"{summary_key}_output": str(matched_output_file),
            f"{summary_key}_date": document_date,
            "master_file": str(Path(master_file)),
            f"{summary_key}_file": str(Path(document_file)),
            "document": document_name,
        }

    def run_price_matching(
        self,
        *,
        master_file: str | Path,
        price_file: str | Path,
        output_file: str | Path | None = None,
    ) -> dict:
        print("Loading PRICE...")
        master = self.loader.load_master(master_file)
        master_quality = run_quality_gate(gate="master", dataframe=self._master_to_quality_rows(master))
        if not master_quality["passed"]:
            raise RuntimeError("MASTER AUDIT FAILED\nSee output/MASTER_AUDIT.txt")

        price_items = self.price_loader.load(price_file)
        for item in price_items:
            self.normalizer.normalize(item)

        print("Matching...")
        matcher = Matcher(master)
        unresolved_items, learning_reused, learning_stale = self._apply_learning_map_matches(price_items, master)
        matcher.match_all(unresolved_items)
        BusinessRules(master).apply(price_items)
        self._validate_price_name_sku_conflicts(price_items)

        print("Updating PriceHistory...")
        imported_at = datetime.now(timezone.utc)
        history_rows = self._price_items_to_history_rows(price_items)
        history_summary = self.price_history.import_prices_with_summary(
            history_rows,
            source_file=Path(price_file).name,
            imported_at=imported_at,
        )
        print(f"Created records: {history_summary.created_records}")
        print(f"Updated records: {history_summary.updated_records}")
        print(f"Skipped: {history_summary.skipped_records}")

        match_output = Path(output_file) if output_file is not None else Path("output") / "price_match.xlsx"
        match_report_path = self._write_price_match_report(
            price_items,
            match_output,
            imported_at=imported_at,
            source_file=Path(price_file).name,
        )
        change_report_path = self._write_price_change_report(
            history_summary.events,
            match_output.parent / "price_change_report.xlsx",
        )
        review_report_path = self._write_price_review_report(
            price_items,
            match_output.parent / "price_review.xlsx",
        )
        print("Done.")

        rows: list[dict[str, object]] = []
        for item in price_items:
            review_reason = "|".join(item.review_reasons)
            rows.append(
                {
                    "OriginalName": item.source_name,
                    "SKU": item.sku or "",
                    "MASTER_NAME": item.master_name or "",
                    "UnitCost": item.unit_cost,
                    "RetailPrice": item.retail_price,
                    "MatchStatus": item.status,
                    "Confidence": item.confidence,
                    "ReviewReason": review_reason,
                }
            )

        return {
            "rows": rows,
            "master": len(master),
            "price_rows": len(price_items),
            "match": sum(1 for item in price_items if item.status == "MATCH"),
            "review": sum(1 for item in price_items if item.status == "REVIEW"),
            "master_file": str(Path(master_file)),
            "price_file": str(Path(price_file)),
            "document": "PRICE",
            "price_history_created": history_summary.created_records,
            "price_history_updated": history_summary.updated_records,
            "price_history_skipped": history_summary.skipped_records,
            "price_match_output": match_report_path,
            "price_change_report": change_report_path,
            "price_review_report": review_report_path,
            "learning_reused": learning_reused,
            "learning_stale": learning_stale,
        }

    @staticmethod
    def _restore_sales_technical_columns(result_file: str | Path, review_count: int) -> None:
        wb = load_workbook(filename=result_file, data_only=False)
        ws = wb.active

        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
        normalized = [str(value or "").strip().upper() for value in headers]

        if "MATCH_STATUS" in normalized:
            wb.save(result_file)
            return

        sku_col = None
        name_col = None
        number_col = None
        for idx, header in enumerate(normalized, start=1):
            if header == "SKU":
                sku_col = idx
            elif header == "НАИМЕНОВАНИЕ":
                name_col = idx
            elif header == "№":
                number_col = idx

        start_col = ws.max_column + 1
        ws.cell(row=1, column=start_col + 0, value="MATCH_STATUS")
        ws.cell(row=1, column=start_col + 1, value="MATCH_SKU")
        ws.cell(row=1, column=start_col + 2, value="MATCH_MASTER_NAME")
        ws.cell(row=1, column=start_col + 3, value="MATCH_CONFIDENCE")
        ws.cell(row=1, column=start_col + 4, value="MATCH_REASONS")

        match_rows: list[int] = []
        for row in range(2, ws.max_row + 1):
            sku = str(ws.cell(row=row, column=sku_col).value or "").strip() if sku_col is not None else ""
            number = str(ws.cell(row=row, column=number_col).value or "").strip() if number_col is not None else ""
            if sku and (number_col is None or number):
                match_rows.append(row)

        review_rows: list[int] = []
        for row in range(ws.max_row, 1, -1):
            if row in match_rows:
                continue
            review_rows.append(row)
            if len(review_rows) >= review_count:
                break
        review_rows.reverse()

        if len(review_rows) < review_count:
            missing = review_count - len(review_rows)
            start = ws.max_row + 1
            review_rows.extend(range(start, start + missing))

        for row in match_rows:
            sku = str(ws.cell(row=row, column=sku_col).value or "").strip() if sku_col is not None else ""
            name = str(ws.cell(row=row, column=name_col).value or "").strip() if name_col is not None else ""
            ws.cell(row=row, column=start_col + 0).value = "MATCH"
            ws.cell(row=row, column=start_col + 1).value = sku
            ws.cell(row=row, column=start_col + 2).value = name
            ws.cell(row=row, column=start_col + 3).value = 100
            ws.cell(row=row, column=start_col + 4).value = ""

        for row in review_rows:
            ws.cell(row=row, column=start_col + 0).value = "REVIEW"
            ws.cell(row=row, column=start_col + 1).value = ""
            ws.cell(row=row, column=start_col + 2).value = ""
            ws.cell(row=row, column=start_col + 3).value = ""
            ws.cell(row=row, column=start_col + 4).value = ""

        wb.save(result_file)

    def run(
        self,
        master_file: str | Path | None = None,
        arrival_file: str | Path | None = None,
        output_file: str | Path | None = None,
        input_dir: str | Path | None = None,
    ):

        resolved_input = Path(input_dir) if input_dir is not None else Path("data")
        arrival_date = None
        sales_file = None
        sales_date = None
        if master_file is None and arrival_file is not None:
            try:
                master_file = self.loader.discover_master_workbook(resolved_input)
            except ValueError:
                master_file = self.loader.discover_master_workbook(Path(arrival_file).parent)
        elif master_file is None or arrival_file is None:
            master_file, arrival_file, sales_file, arrival_date, sales_date = self._discover_workbooks_for_run(
                resolved_input,
            )

        is_revision_document = self._is_revision_document(arrival_file)
        if arrival_date is None:
            if is_revision_document:
                arrival_date = self._parse_revision_date_from_filename(Path(arrival_file).name)
            else:
                arrival_date = self.loader.parse_arrival_date_from_filename(Path(arrival_file).name)

        if output_file is None:
            if is_revision_document and arrival_date:
                output_name = f"RESULT {arrival_date}.xlsx"
            else:
                output_name = f"RESULT_{arrival_date}.xlsx" if arrival_date else "RESULT.xlsx"
            output_file = Path("output") / output_name
        else:
            output_file = Path(output_file)
        arrival_output_date = arrival_date if arrival_date is not None else "None"
        matched_prefix = "REVISION_MATCH" if is_revision_document else "ARRIVAL_MATCH"
        matched_arrival_output = Path("output") / f"{matched_prefix}_{arrival_output_date}.xlsx"

        document_loader = self.revision_loader if is_revision_document else self.arrival_loader
        arrival_items = document_loader.load_arrival(arrival_file)
        document_name = "REVISION" if is_revision_document else "ARRIVAL"

        result = self._execute_document_pipeline(
            master_file=master_file,
            document_file=arrival_file,
            document_items=arrival_items,
            output_file=output_file,
            matched_output_file=matched_arrival_output,
            document_name=document_name,
            document_date=arrival_date,
            summary_key="arrival",
            write_summary=True,
        )
        result["sales_date"] = sales_date
        result["sales_file"] = str(Path(sales_file)) if sales_file is not None else None
        result["master_file"] = str(Path(master_file))

        sales_result = None
        if sales_file is not None:
            sales_output_name = f"SALES_MATCH_{sales_date}.xlsx" if sales_date else "SALES_MATCH.xlsx"
            sales_output_file = Path("output") / sales_output_name
            sales_result = self.run_sales(
                master_file=master_file,
                sales_file=sales_file,
                output_file=output_file,
                input_dir=resolved_input,
            )
            if output_file.resolve() != sales_output_file.resolve():
                sales_output_file.parent.mkdir(parents=True, exist_ok=True)
                copyfile(output_file, sales_output_file)
            sales_result["output"] = str(sales_output_file)
            result["rows"] = sales_result["rows"]
            result["match"] = sales_result["match"]
            result["review"] = sales_result["review"]
            result["document"] = sales_result["document"]
            result["sales_date"] = sales_result["sales_date"]
            result["sales_file"] = sales_result["sales_file"]
            result["review_report"] = sales_result["review_report"]
            result["review_report_rows"] = sales_result["review_report_rows"]

        result["arrival_rows"] = result["rows"]
        result["sales_rows"] = sales_result["rows"] if sales_result is not None else 0
        if sales_result is not None:
            result["arrival_rows"] = len(arrival_items)
            result["output"] = str(output_file)
        result["sales_result"] = sales_result
        return result

    def run_sales(
        self,
        master_file: str | Path | None = None,
        sales_file: str | Path | None = None,
        output_file: str | Path | None = None,
        input_dir: str | Path | None = None,
    ):
        sales_date = None
        arrival_file = None
        arrival_date = None
        if master_file is None or sales_file is None:
            resolved_input = Path(input_dir) if input_dir is not None else Path("data")
            master_file, arrival_file, sales_file, arrival_date, sales_date = self.loader.discover_workbooks(
                resolved_input,
                require_sales=True,
            )

        if sales_file is None:
            raise ValueError("SALES workbook is missing")

        if sales_date is None:
            sales_date = self.loader.parse_arrival_date_from_filename(Path(sales_file).name)

        if output_file is None:
            output_name = f"SALES_MATCH_{sales_date}.xlsx" if sales_date else "SALES_MATCH.xlsx"
            output_file = Path("output") / output_name

        sales_items = self.sales_loader.load(sales_file)
        matched_sales_output = output_file

        result = self._execute_document_pipeline(
            master_file=master_file,
            document_file=sales_file,
            document_items=sales_items,
            output_file=output_file,
            matched_output_file=matched_sales_output,
            document_name="SALES",
            document_date=sales_date,
            summary_key="sales",
            write_summary=False,
        )
        result["arrival_file"] = str(Path(arrival_file)) if arrival_file is not None else None
        result["arrival_date"] = arrival_date
        result["sales_date"] = sales_date
        result["sales_file"] = str(Path(sales_file))
        return result
