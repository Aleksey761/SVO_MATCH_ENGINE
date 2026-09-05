from pathlib import Path
from openpyxl import load_workbook

from .arrival_loader import ArrivalLoader
from .models import ArrivalItem


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().upper().split())


def _looks_like_product_name(value: object) -> bool:
    if not isinstance(value, str):
        return False

    text = value.strip()
    if len(text) < 15:
        return False

    has_letters = any(char.isalpha() for char in text)
    has_separator = any(separator in text for separator in (",", " "))
    return has_letters and has_separator


class SalesLoader(ArrivalLoader):
    """Loads SALES workbook rows into the shared arrival item model."""

    def _find_quantity_column_indices(self, worksheet) -> tuple[int, int, int]:
        required_headers = ("ВОРОНЕЖ", "КРАСНОДАР 1", "КРАСНОДАР 2")
        found: dict[str, int] = {}

        for row in worksheet.iter_rows():
            for column_index, cell in enumerate(row):
                header = _normalize_header(cell.value)
                if header in required_headers:
                    found[header] = column_index
            if all(header in found for header in required_headers):
                return tuple(found[header] for header in required_headers)

        # A SALES workbook may legitimately contain only product names
        # (for example, a review-only source). In that case quantities are
        # treated as zero; the three warehouse columns remain mandatory when
        # the file is not a recognised SALES contract.
        if any(
            _normalize_header(
                worksheet.cell(row=row_idx, column=column_idx).value
            ) in {"НАИМЕНОВАНИЕ", "SOURCE_NAME", "NAME", "PRODUCT", "PRODUCT NAME"}
            for row_idx in range(1, min(worksheet.max_row, 12) + 1)
            for column_idx in range(1, worksheet.max_column + 1)
        ):
            return ()
        raise ValueError("Required SALES quantity columns not found")

    @staticmethod
    def _quantity(value: object) -> int:
        if value is None or value == "":
            return 0
        try:
            return int(float(str(value).strip().replace(" ", "").replace(",", ".")))
        except (TypeError, ValueError):
            return 0

    def _find_name_column_index(self, worksheet) -> int:
        # Prefer the explicit SALES contract header. This supports files with
        # a title/blank row before the actual header and does not depend on
        # product-name heuristics.
        explicit_headers = {
            "НАИМЕНОВАНИЕ",
            "NAME",
            "PRODUCT",
            "PRODUCT NAME",
            "SOURCE_NAME",
        }
        for row_idx in range(1, min(worksheet.max_row, 12) + 1):
            for column_index, cell in enumerate(worksheet[row_idx]):
                if _normalize_header(cell.value) in explicit_headers:
                    return column_index

        # Fallback for legacy/unlabelled SALES exports.
        for column_index in range(worksheet.max_column):
            product_name_count = 0

            for row in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row,
                values_only=True,
            ):
                if column_index >= len(row):
                    continue

                value = row[column_index]
                if _looks_like_product_name(value):
                    product_name_count += 1

            if product_name_count >= 2:
                return column_index

        raise ValueError(
            "Could not determine SALES product-name column from worksheet"
        )

    def _find_explicit_header_row(self, worksheet) -> int | None:
        explicit_headers = {
            "НАИМЕНОВАНИЕ",
            "NAME",
            "PRODUCT",
            "PRODUCT NAME",
            "SOURCE_NAME",
        }
        for row_idx in range(1, min(worksheet.max_row, 12) + 1):
            for cell in worksheet[row_idx]:
                if _normalize_header(cell.value) in explicit_headers:
                    return row_idx
        return None

    def load(self, filename: str | Path) -> list[ArrivalItem]:
        wb = load_workbook(filename=filename, data_only=True)
        worksheet = wb.active
        explicit_header_row = self._find_explicit_header_row(worksheet)

        if explicit_header_row is not None:
            explicit_headers = {
                "НАИМЕНОВАНИЕ",
                "NAME",
                "PRODUCT",
                "PRODUCT NAME",
                "SOURCE_NAME",
            }
            name_column_index = None
            for column_index, cell in enumerate(worksheet[explicit_header_row]):
                if _normalize_header(cell.value) in explicit_headers:
                    candidate = column_index

                    # Some SALES workbooks have "Наименование" above a
                    # numbering column, while actual product names are one
                    # column to the right.
                    if candidate + 1 < worksheet.max_column:
                        next_values = [
                            worksheet.cell(
                                row=r,
                                column=candidate + 2,
                            ).value
                            for r in range(
                                explicit_header_row + 1,
                                min(worksheet.max_row, explicit_header_row + 10) + 1,
                            )
                        ]
                        if sum(
                            1 for value in next_values
                            if _looks_like_product_name(value)
                        ) >= 2:
                            candidate += 1

                    name_column_index = candidate
                    break

            if name_column_index is None:
                raise ValueError(
                    "Could not determine SALES product-name column from explicit header"
                )
        else:
            name_column_index = self._find_name_column_index(worksheet)

        quantity_column_indices = self._find_quantity_column_indices(worksheet)
        data_start_row = explicit_header_row + 1 if explicit_header_row is not None else 2

        items: list[ArrivalItem] = []
        current_context = ""
        current_context_type = ""

        import re

        powder_pattern = re.compile(
            r"^\u041f\u043e\u0440\u043e\u0448\u043e\u043a\s+(\d+(?:[.,]\d+)?)\s+(.+)$",
            re.IGNORECASE,
        )

        magina_pattern = re.compile(
            r"^\u0413\u0435\u043b\u044c\s+\u043a\u043e\u043d\u0434\u0438\u0446\u0438\u043e\u043d\u0435\u0440\s+1400\s+MAGINA$",
            re.IGNORECASE,
        )

        elegant_pattern = re.compile(
            r"^SVO\s+1500\s+ELEGANT\s+WHITE$",
            re.IGNORECASE,
        )

        powder_aromas = {
            "\u041b\u0410\u0412\u0410\u041d\u0414\u0410": "Lavender",
            "\u0427\u0415\u0420\u041d\u042b\u0419": "Black",
            "\u0420\u041e\u0417\u0410": "Rose",
            "\u041f\u041e\u0414\u0421\u041d\u0415\u0416\u041d\u0418\u041a": "Snowdrop",
            "\u041f\u041e\u0414\u0421\u041d\u0415\u0416\u041d\u041c\u041a": "Snowdrop",
            "\u0413\u041e\u0420\u041d\u042b\u0419": "Mountain Breeze",
            "\u0413\u041e\u0420\u041d\u0412\u0419": "Mountain Breeze",
            "\u0413\u041e\u0420\u0413\u042b\u0419": "Mountain Breeze",
            "\u041b\u0415\u041c\u041e\u041d": "Lemon",
            "\u0414\u0415\u0422\u0421\u041a\u0418\u0419": "Baby",
        }

        elegant_aromas = {
            "WHITE": "White",
            "MANGO": "Mango",
            "FLORIAL MIST": "Floral Mist",
            "MIDNIGHT": "Midnight",
            "SPRING": "Spring",
            "SWEET TROPIK": "Sweet Tropic",
            "VELVET": "Velvet",
            "DREAM": "Dream",
            "BLACK": "Black",
        }

        for excel_row, row in enumerate(
            worksheet.iter_rows(min_row=data_start_row, values_only=True),
            start=data_start_row,
        ):
            name = row[name_column_index] if name_column_index < len(row) else None
            if name is None:
                continue

            text = str(name).strip()
            if not text:
                continue

            powder_match = powder_pattern.match(text)

            if powder_match:
                size_raw = powder_match.group(1).replace(",", ".")
                size = str(float(size_raw)).rstrip("0").rstrip(".")
                source_aroma = powder_match.group(2).strip()
                aroma = powder_aromas.get(
                    source_aroma.upper(),
                    source_aroma,
                )

                current_context = (
                    f"\u041f\u043e\u0440\u043e\u0448\u043e\u043a \u0441\u0442\u0438\u0440\u0430\u043b\u044c\u043d\u044b\u0439 SVO {size} \u043a\u0433"
                )
                current_context_type = "powder"
                match_name = f"{current_context} {aroma}"

            elif magina_pattern.match(text):
                current_context = (
                    "\u041a\u043e\u043d\u0434\u0438\u0446\u0438\u043e\u043d\u0435\u0440 SVO 1,44 \u043b"
                )
                current_context_type = "magina"
                match_name = f"{current_context} MAGINA"

            elif elegant_pattern.match(text):
                current_context = (
                    "\u0416\u041c\u0421 SVO Elegant 1,5 \u043b"
                )
                current_context_type = "elegant_15"
                match_name = f"{current_context} White"

            elif current_context_type == "powder":
                if text.upper().startswith("OXGEEN"):
                    current_context = ""
                    current_context_type = ""
                    match_name = text
                elif len(text) < 40:
                    aroma = powder_aromas.get(text.upper(), text)
                    match_name = f"{current_context} {aroma}"
                else:
                    current_context = ""
                    current_context_type = ""
                    match_name = text

            elif current_context_type == "elegant_15" and len(text) < 40:
                aroma = elegant_aromas.get(text.upper(), text)
                match_name = f"{current_context} {aroma}"

            elif current_context_type == "magina" and len(text) < 40:
                match_name = f"{current_context} {text}"

            else:
                match_name = text
                current_context = ""
                current_context_type = ""

            shipped_qty = sum(
                self._quantity(row[column_index] if column_index < len(row) else None)
                for column_index in quantity_column_indices
            )

            item_kwargs = {
                "row_number": excel_row,
                "source_name": text,
                "shipped_qty": shipped_qty,
            }
            # Contextual SALES rows (for example aroma-only continuation
            # lines like "Лаванда" after "Порошок 1.3") need the expanded
            # product name for normalization and matching.
            if match_name != text:
                item_kwargs["source_name"] = match_name
                item_kwargs["match_name"] = match_name

            items.append(ArrivalItem(**item_kwargs))

        return items
