from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook

from .models import ArrivalItem
from .price_cleaner import PriceCleaner


class PriceLoader:
    """Loads PRICE workbook rows into arrival-style items with price payload fields."""

    _PACKAGE_UNIT_RE = re.compile(r"^\s*упак\s*\(\s*(\d+)\s*шт\s*\)\s*$", re.IGNORECASE)

    _HEADER_CANDIDATES = {
        "ProductName": ("PRODUCTNAME", "ПОЛНОЕНАИМЕНОВАНИЕ"),
        "Price": ("PRICE", "DISTRIBUTORPRICE", "ДИСТРИБЬЮТОРСКАЯЦЕНА"),
        "UnitCost": ("UNITCOST",),
        "RetailPrice": ("RETAILPRICE",),
        "StockQty": ("STOCKQTY", "FREESTOCK", "СВОБОДНЫЙОСТАТОК"),
        "SupplierArticle": ("SUPPLIERARTICLE", "ARTICLE", "АРТИКУЛ", "АРТИКУЛПОСТАВЩИКА"),
        "SourceUnit": ("SOURCEUNIT", "UNIT", "ЕДИНИЦА", "ЕДИЗМ", "ЕД.ИЗМ"),
    }

    _OPTIONAL_HEADERS = {
        "WHOLESALEPRICE": "WholesalePrice",
        "MARKETPLACEPRICE": "MarketplacePrice",
    }

    def __init__(self) -> None:
        self.cleaner = PriceCleaner()

    @staticmethod
    def _normalize_header(value: object) -> str:
        return "".join(str(value or "").strip().upper().split())

    @staticmethod
    def _normalize_name(value: object) -> str:
        return " ".join(str(value or "").strip().upper().split())

    @staticmethod
    def _find_column(header_idx: dict[str, int], candidates: tuple[str, ...]) -> int | None:
        for candidate in candidates:
            found = header_idx.get(candidate)
            if found is not None:
                return found
        return None

    @staticmethod
    def _parse_numeric(value: object) -> Decimal | None:
        if value is None:
            return None
        text = str(value).strip().replace(" ", "")
        if not text:
            return None
        text = text.replace(",", ".")
        try:
            return Decimal(text)
        except (InvalidOperation, ValueError):
            return None

    def _detect_header_row(self, ws) -> int:
        max_scan_row = min(ws.max_row, 25)
        best_row = 1
        best_score = -1

        for row in range(1, max_scan_row + 1):
            row_values = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
            non_empty = [value for value in row_values if value is not None and str(value).strip()]
            if not non_empty:
                continue

            normalized = [self._normalize_header(value) for value in row_values]
            header_idx = {header: idx + 1 for idx, header in enumerate(normalized) if header}

            mapped_hits = 0
            for candidates in self._HEADER_CANDIDATES.values():
                if self._find_column(header_idx, candidates) is not None:
                    mapped_hits += 1

            string_like = sum(1 for value in non_empty if isinstance(value, str))
            numeric_like = sum(1 for value in non_empty if self._parse_numeric(value) is not None)

            score = mapped_hits * 100 + string_like * 3 + len(non_empty)
            if numeric_like > string_like:
                score -= 10

            if score > best_score:
                best_score = score
                best_row = row

        return best_row

    def _build_header_index(self, ws, header_row: int) -> tuple[dict[str, int], dict[int, str]]:
        original_header_by_col: dict[int, str] = {}
        normalized_header_idx: dict[str, int] = {}

        for idx, cell in enumerate(ws[header_row], start=1):
            raw_header = str(cell.value or "").strip()
            if not raw_header:
                continue
            original_header_by_col[idx] = raw_header
            normalized = self._normalize_header(raw_header)
            if normalized:
                normalized_header_idx[normalized] = idx

        return normalized_header_idx, original_header_by_col

    def _find_data_start_row(self, ws, header_row: int) -> int:
        for row in range(header_row + 1, ws.max_row + 1):
            values = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
            if any(value is not None and str(value).strip() for value in values):
                return row
        return header_row + 1

    def _infer_columns_from_data(self, ws, header_row: int, mapped_columns: dict[str, int | None]) -> dict[str, int | None]:
        data_start_row = self._find_data_start_row(ws, header_row)
        sample_end_row = min(ws.max_row, data_start_row + 30)

        def sample_values(column: int) -> list[object]:
            values: list[object] = []
            for row in range(data_start_row, sample_end_row + 1):
                value = ws.cell(row=row, column=column).value
                if value is None:
                    continue
                if isinstance(value, str) and not value.strip():
                    continue
                values.append(value)
            return values

        columns = list(range(1, ws.max_column + 1))
        numeric_scores: dict[int, float] = {}
        integer_scores: dict[int, float] = {}
        text_scores: dict[int, float] = {}

        for column in columns:
            values = sample_values(column)
            if not values:
                continue

            numeric_count = sum(1 for value in values if self._parse_numeric(value) is not None)
            integer_count = sum(
                1
                for value in values
                if (parsed := self._parse_numeric(value)) is not None and parsed == parsed.to_integral_value()
            )
            text_count = sum(1 for value in values if isinstance(value, str) and self._parse_numeric(value) is None)
            avg_text_len = 0.0
            text_values = [str(value).strip() for value in values if isinstance(value, str)]
            if text_values:
                avg_text_len = sum(len(text) for text in text_values) / len(text_values)

            total = len(values)
            numeric_scores[column] = numeric_count / total
            integer_scores[column] = integer_count / total
            text_scores[column] = (text_count / total) + (avg_text_len / 100.0)

        inferred = dict(mapped_columns)

        if numeric_scores:
            best_numeric_col = max(numeric_scores, key=numeric_scores.get)
            best_numeric_score = numeric_scores[best_numeric_col]
            current_price_col = inferred.get("Price")
            current_price_score = numeric_scores.get(current_price_col, 0.0) if current_price_col is not None else 0.0
            if best_numeric_score >= 0.6 and best_numeric_score > current_price_score + 0.25:
                inferred["Price"] = best_numeric_col

        if integer_scores:
            current_price_col = inferred.get("Price")
            stock_candidates = {
                column: score
                for column, score in integer_scores.items()
                if column != current_price_col
            }
            best_integer_col = max(stock_candidates, key=stock_candidates.get) if stock_candidates else None
            best_integer_score = stock_candidates.get(best_integer_col, 0.0) if best_integer_col is not None else 0.0
            current_stock_col = inferred.get("StockQty")
            current_stock_score = integer_scores.get(current_stock_col, 0.0) if current_stock_col is not None else 0.0
            if (
                best_integer_col is not None
                and best_integer_score >= 0.6
                and best_integer_score > current_stock_score + 0.25
            ):
                inferred["StockQty"] = best_integer_col

        if inferred.get("Price") is None and numeric_scores:
            inferred["Price"] = max(numeric_scores, key=numeric_scores.get)

        if inferred.get("StockQty") is None and integer_scores:
            candidates = {
                column: score
                for column, score in integer_scores.items()
                if column != inferred.get("Price")
            }
            if candidates:
                inferred["StockQty"] = max(candidates, key=candidates.get)

        if inferred.get("ProductName") is None and text_scores:
            excluded = {value for key, value in inferred.items() if key in {"Price", "StockQty", "SupplierArticle"} and value is not None}
            candidates = {
                column: score
                for column, score in text_scores.items()
                if column not in excluded
            }
            if candidates:
                inferred["ProductName"] = max(candidates, key=candidates.get)

        if inferred.get("SourceUnit") is None and text_scores:
            excluded = {
                value
                for key, value in inferred.items()
                if key in {"Price", "StockQty", "SupplierArticle", "ProductName"} and value is not None
            }
            candidates = {
                column: score
                for column, score in text_scores.items()
                if column not in excluded
            }
            if candidates:
                inferred["SourceUnit"] = min(candidates, key=candidates.get)

        return inferred

    @staticmethod
    def _debug_print_mapping(
        *,
        worksheet_name: str,
        header_row: int,
        original_header_by_col: dict[int, str],
        mapped_columns: dict[str, int | None],
    ) -> None:
        original_columns = [
            {"column": col, "name": name}
            for col, name in sorted(original_header_by_col.items())
        ]

        mapped_column_names: dict[str, str | None] = {}
        for field, column in mapped_columns.items():
            if column is None:
                mapped_column_names[field] = None
                continue
            mapped_name = original_header_by_col.get(column)
            if mapped_name:
                mapped_column_names[field] = mapped_name
            else:
                mapped_column_names[field] = f"<inferred: column {column}>"

        print(f"PRICE loader worksheet: {worksheet_name}")
        print(f"PRICE loader header row: {header_row}")
        print(f"PRICE loader original columns: {original_columns}")
        print(f"PRICE loader mapped columns: {mapped_column_names}")

    @staticmethod
    def _stock_signature(value: object) -> object:
        parsed = PriceLoader._parse_numeric(value)
        if parsed is not None:
            return parsed
        return str(value or "").strip()

    @classmethod
    def _extract_pack_qty(cls, source_unit: object) -> int | None:
        unit_text = str(source_unit or "").strip()
        if not unit_text:
            return None
        match = cls._PACKAGE_UNIT_RE.match(unit_text)
        if match is None:
            return None
        qty = int(match.group(1))
        if qty <= 0:
            return None
        return qty

    @staticmethod
    def _print_duplicate_rows(canonical_name: str, rows: list[dict[str, object]]) -> None:
        print(f"PRICE duplicate investigation for canonical ProductName: {canonical_name}")
        for row in rows:
            print(
                "PRICE duplicate row "
                f"source_row={row['row_number']} "
                f"original_product_name={row['original_name']} "
                f"canonical_product_name={canonical_name} "
                f"unit_cost={row['unit_cost']} "
                f"retail_price={row['retail_price']} "
                f"stock_qty={row['stock_qty']}"
            )

    def _to_canonical_rows(self, ws) -> list[dict[str, object]]:
        header_row = self._detect_header_row(ws)
        header_idx, original_header_by_col = self._build_header_index(ws, header_row)

        mapped_columns: dict[str, int | None] = {
            field: self._find_column(header_idx, candidates)
            for field, candidates in self._HEADER_CANDIDATES.items()
        }
        mapped_columns = self._infer_columns_from_data(ws, header_row, mapped_columns)
        self._debug_print_mapping(
            worksheet_name=ws.title,
            header_row=header_row,
            original_header_by_col=original_header_by_col,
            mapped_columns=mapped_columns,
        )

        name_col = mapped_columns["ProductName"]
        if name_col is None:
            raise ValueError("PRICE workbook is missing required mapped column ProductName")

        source_price_col = mapped_columns["Price"]
        unit_cost_col = mapped_columns["UnitCost"]
        retail_price_col = mapped_columns["RetailPrice"]

        if source_price_col is None and unit_cost_col is None:
            raise ValueError("PRICE workbook is missing required mapped column Price")
        if source_price_col is None and retail_price_col is None:
            raise ValueError("PRICE workbook is missing required mapped column Price")

        stock_qty_col = mapped_columns["StockQty"]
        supplier_article_col = mapped_columns["SupplierArticle"]
        source_unit_col = mapped_columns["SourceUnit"]
        wholesale_price_col = header_idx.get("WHOLESALEPRICE")
        marketplace_price_col = header_idx.get("MARKETPLACEPRICE")

        # Supplier article is intentionally ignored: it is not part of matching keys.
        _ = mapped_columns["SupplierArticle"]

        canonical_rows: list[dict[str, object]] = []
        for excel_row in range(header_row + 1, ws.max_row + 1):
            canonical_rows.append(
                {
                    "RowNumber": excel_row,
                    "ProductName": ws.cell(row=excel_row, column=name_col).value,
                    "UnitCost": ws.cell(
                        row=excel_row,
                        column=unit_cost_col or source_price_col,
                    ).value,
                    "RetailPrice": ws.cell(
                        row=excel_row,
                        column=retail_price_col or source_price_col,
                    ).value,
                    "StockQty": ws.cell(row=excel_row, column=stock_qty_col).value if stock_qty_col is not None else None,
                    "SupplierArticle": ws.cell(row=excel_row, column=supplier_article_col).value if supplier_article_col is not None else None,
                    "SourceUnit": ws.cell(row=excel_row, column=source_unit_col).value if source_unit_col is not None else None,
                    "WholesalePrice": ws.cell(row=excel_row, column=wholesale_price_col).value if wholesale_price_col is not None else None,
                    "MarketplacePrice": ws.cell(row=excel_row, column=marketplace_price_col).value if marketplace_price_col is not None else None,
                }
            )

        return canonical_rows

    @staticmethod
    def _to_decimal(value: object, field_name: str, row_number: int, *, required: bool) -> Decimal | None:
        if value is None:
            if required:
                raise ValueError(f"Missing required {field_name} at row {row_number}")
            return None

        text = str(value).strip().replace(" ", "")
        if not text:
            if required:
                raise ValueError(f"Missing required {field_name} at row {row_number}")
            return None

        text = text.replace(",", ".")
        try:
            parsed = Decimal(text)
        except (InvalidOperation, ValueError):
            raise ValueError(f"Invalid Decimal in {field_name} at row {row_number}")

        if parsed < Decimal("0"):
            raise ValueError(f"Negative {field_name} at row {row_number}")

        return parsed

    def load(self, filename: str | Path) -> list[ArrivalItem]:
        wb = load_workbook(filename=filename, data_only=True)
        ws = wb.active
        canonical_rows = self._to_canonical_rows(ws)

        items: list[ArrivalItem] = []
        seen_names: dict[str, dict[str, object]] = {}

        for row in canonical_rows:
            excel_row = int(row["RowNumber"])
            if all(
                value is None or (isinstance(value, str) and not value.strip())
                for value in (
                    row["ProductName"],
                    row["UnitCost"],
                    row["RetailPrice"],
                    row["StockQty"],
                    row["WholesalePrice"],
                    row["MarketplacePrice"],
                )
            ):
                continue

            raw_name = row["ProductName"]
            cleaned_name = self.cleaner.clean(raw_name)
            normalized_name = self._normalize_name(cleaned_name)
            if not normalized_name:
                continue

            unit_cost = self._to_decimal(
                row["UnitCost"],
                "UnitCost",
                excel_row,
                required=True,
            )
            retail_price = self._to_decimal(
                row["RetailPrice"],
                "RetailPrice",
                excel_row,
                required=True,
            )
            wholesale_price = None
            if row["WholesalePrice"] is not None:
                wholesale_price = self._to_decimal(
                    row["WholesalePrice"],
                    "WholesalePrice",
                    excel_row,
                    required=False,
                )
            marketplace_price = None
            if row["MarketplacePrice"] is not None:
                marketplace_price = self._to_decimal(
                    row["MarketplacePrice"],
                    "MarketplacePrice",
                    excel_row,
                    required=False,
                )

            item = ArrivalItem(
                row_number=excel_row,
                source_name=cleaned_name,
            )
            item.unit_cost = unit_cost
            item.retail_price = retail_price
            item.stock_qty = row["StockQty"]
            item.supplier_article = str(row["SupplierArticle"] or "").strip() or None
            item.original_product_name = str(raw_name or "").strip()
            item.canonical_product_name = normalized_name
            item.source_unit = str(row["SourceUnit"] or "").strip() or None
            item.pack_qty = self._extract_pack_qty(item.source_unit)
            item.original_price = unit_cost
            item.wholesale_price = wholesale_price
            item.marketplace_price = marketplace_price

            duplicate_row = {
                "row_number": excel_row,
                "original_name": str(raw_name or "").strip(),
                "unit_cost": unit_cost,
                "retail_price": retail_price,
                "stock_qty": row["StockQty"],
            }

            existing = seen_names.get(normalized_name)
            if existing is None:
                seen_names[normalized_name] = {
                    "rows": [duplicate_row],
                }
                items.append(item)
                continue

            duplicate_rows = [*existing["rows"], duplicate_row]
            self._print_duplicate_rows(normalized_name, duplicate_rows)
            seen_names[normalized_name]["rows"].append(duplicate_row)
            print(
                "PRICE duplicate resolution deferred: "
                "continuing import for SKU-aware history processing"
            )

            items.append(item)

        return items

    @staticmethod
    def to_canonical_price_dataset(items: list[ArrivalItem]) -> list[dict[str, object]]:
        dataset: list[dict[str, object]] = []
        for item in items:
            reasons = [str(reason).strip() for reason in getattr(item, "review_reasons", []) if str(reason).strip()]
            dataset.append(
                {
                    "SKU": item.sku or "",
                    "MASTER_NAME": item.master_name or "",
                    "UnitCost": getattr(item, "unit_cost", None),
                    "RetailPrice": getattr(item, "retail_price", None),
                    "Status": item.status,
                    "Confidence": item.confidence,
                    "ReviewReason": "|".join(reasons),
                }
            )
        return dataset
