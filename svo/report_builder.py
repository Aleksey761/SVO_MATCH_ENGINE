from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill


class ReportBuilder:
    """Builds and formats Excel workbooks. No business calculations here."""

    def build_inventory_reconciliation_workbook(
        self,
        *,
        output_file: str | Path,
        headers: list[str],
        stock_rows: list[list[object]],
        variance_rows: list[list[object]],
        variance_col: int = 9,
        numeric_start_col: int = 3,
        numeric_end_col: int = 9,
    ) -> None:
        output_path = Path(output_file)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws_stock = wb.active
        ws_stock.title = "STOCK_RESULT"
        ws_stock.append(headers)
        for row in stock_rows:
            ws_stock.append(row)

        ws_variance = wb.create_sheet("VARIANCE_REPORT")
        ws_variance.append(headers)
        for row in variance_rows:
            ws_variance.append(row)

        self._apply_sheet_formatting(ws_stock, numeric_start_col=numeric_start_col, numeric_end_col=numeric_end_col)
        self._apply_sheet_formatting(ws_variance, numeric_start_col=numeric_start_col, numeric_end_col=numeric_end_col)
        self._apply_variance_conditional_formatting(ws_stock, variance_col=variance_col)
        self._apply_variance_conditional_formatting(ws_variance, variance_col=variance_col)

        wb.save(output_path)

    @staticmethod
    def _apply_sheet_formatting(ws, *, numeric_start_col: int, numeric_end_col: int) -> None:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        bold_font = Font(bold=True)
        for col in range(1, ws.max_column + 1):
            ws.cell(row=1, column=col).font = bold_font

        center = Alignment(horizontal="center", vertical="center")
        for row in range(2, ws.max_row + 1):
            for col in range(numeric_start_col, numeric_end_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.alignment = center
                cell.number_format = "#,##0"

        for col in range(1, ws.max_column + 1):
            max_len = 0
            for row in range(1, ws.max_row + 1):
                value = ws.cell(row=row, column=col).value
                value_len = len(str(value)) if value is not None else 0
                if value_len > max_len:
                    max_len = value_len
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(max_len + 2, 60)

    @staticmethod
    def _apply_variance_conditional_formatting(ws, *, variance_col: int) -> None:
        if ws.max_row < 2:
            return

        col_letter = ws.cell(row=1, column=variance_col).column_letter
        range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"

        light_red = PatternFill(fill_type="solid", fgColor="FFC7CE")
        light_yellow = PatternFill(fill_type="solid", fgColor="FFF2CC")

        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="lessThan", formula=["0"], fill=light_red),
        )
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="greaterThan", formula=["0"], fill=light_yellow),
        )
