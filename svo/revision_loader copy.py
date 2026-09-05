from pathlib import Path
import re

from openpyxl import load_workbook

from .models import ArrivalItem


class RevisionLoader:
    """Loads REVISION workbook rows into the shared arrival item model."""

    @staticmethod
    def _normalize_header(value: object) -> str:
        text = str(value or "").strip().upper()
        return " ".join(text.split())

    @staticmethod
    def _parse_date_triplet(text: str) -> tuple[int, int, int] | None:
        match = re.search(r"(?<!\d)(\d{1,2})\D+(\d{1,2})\D+(\d{2}|\d{4})(?!\d)", text)
        if not match:
            return None

        day = int(match.group(1))
        month = int(match.group(2))
        year = int(match.group(3))
        if year < 100:
            year += 2000
        return day, month, year

    @staticmethod
    def _parse_revision_date_from_filename(filename: str | Path) -> tuple[int, int, int] | None:
        stem = Path(filename).stem
        return RevisionLoader._parse_date_triplet(stem)

    @staticmethod
    def _looks_like_product_name(value: object) -> bool:
        text = str(value or "").strip()
        if not text:
            return False
        if text.isdigit():
            return False
        return bool(re.search(r"[A-Za-zА-Яа-я]", text))

    def _find_header_row(self, worksheet) -> int:
        max_scan_rows = min(worksheet.max_row, 30)
        max_scan_cols = min(worksheet.max_column, 80)

        for row in range(1, max_scan_rows + 1):
            normalized = [
                self._normalize_header(worksheet.cell(row=row, column=col).value)
                for col in range(1, max_scan_cols + 1)
            ]
            has_name = any(value == "НАИМЕНОВАНИЕ" for value in normalized)
            has_stock = any("ОСТАТОК НА СКЛАДЕ" in value for value in normalized)
            if has_name and has_stock:
                return row

        return 1

    def _detect_product_column(self, worksheet, header_row: int, name_header_col: int) -> int:
        candidates = [name_header_col]
        if name_header_col + 1 <= worksheet.max_column:
            candidates.append(name_header_col + 1)

        best_col = name_header_col
        best_score = -1
        max_scan_row = min(worksheet.max_row, header_row + 120)
        for col in candidates:
            score = 0
            for row in range(header_row + 1, max_scan_row + 1):
                if self._looks_like_product_name(worksheet.cell(row=row, column=col).value):
                    score += 1
            if score > best_score:
                best_score = score
                best_col = col

        return best_col

    def _detect_closing_column(
        self,
        normalized_headers: list[str],
        *,
        opening_col: int | None,
        filename: str | Path,
    ) -> int | None:
        stock_columns = [
            idx
            for idx, header in enumerate(normalized_headers, start=1)
            if header.startswith("СКЛАД")
        ]
        if stock_columns:
            revision_date = self._parse_revision_date_from_filename(filename)
            if revision_date is not None:
                revision_day, revision_month, revision_year = revision_date
                dated_candidates: list[tuple[int, int]] = []
                for col in stock_columns:
                    header_date = self._parse_date_triplet(normalized_headers[col - 1])
                    if header_date is None:
                        continue
                    day, month, year = header_date
                    if month == revision_month and year == revision_year:
                        dated_candidates.append((col, day))

                if dated_candidates:
                    not_later = [item for item in dated_candidates if item[1] <= revision_day]
                    if not_later:
                        return max(not_later, key=lambda item: item[1])[0]
                    return min(dated_candidates, key=lambda item: item[1])[0]

            return stock_columns[-1]

        stock_balance_columns = [
            idx
            for idx, header in enumerate(normalized_headers, start=1)
            if header.startswith("ОСТАТОК НА СКЛАДЕ")
        ]
        if stock_balance_columns:
            return stock_balance_columns[-1]

        return opening_col

    @staticmethod
    def _to_number(value: object) -> float:
        if value is None:
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)

        text = str(value).strip().replace(" ", "")
        if not text:
            return 0.0

        text = text.replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return 0.0

    @staticmethod
    def _to_int(value: float) -> int:
        return int(round(value))

    def load(self, filename: str | Path) -> list[ArrivalItem]:
        wb = load_workbook(filename=filename, data_only=True)
        ws = wb.active

        header_row = self._find_header_row(ws)
        headers = [ws.cell(row=header_row, column=col).value for col in range(1, ws.max_column + 1)]
        header_to_col: dict[str, int] = {}
        normalized_headers: list[str] = []
        for idx, header in enumerate(headers, start=1):
            normalized = self._normalize_header(header)
            normalized_headers.append(normalized)
            if normalized:
                header_to_col[normalized] = idx

        name_header_col = header_to_col.get("НАИМЕНОВАНИЕ")
        product_col = None
        if name_header_col is not None:
            product_col = self._detect_product_column(ws, header_row, name_header_col)

        opening_col = None
        for idx, normalized in enumerate(normalized_headers, start=1):
            if normalized.startswith("ОСТАТОК НА СКЛАДЕ"):
                if opening_col is None:
                    opening_col = idx
        closing_col = self._detect_closing_column(
            normalized_headers,
            opening_col=opening_col,
            filename=filename,
        )

        voronezh_col = header_to_col.get("ВОРОНЕЖ")
        krasnodar1_col = header_to_col.get("КРАСНОДАР 1")
        krasnodar2_col = header_to_col.get("КРАСНОДАР 2")
        retail_col = header_to_col.get("РОЗНИЦА")

        items: list[ArrivalItem] = []

        for row_idx in range(header_row + 1, ws.max_row + 1):
            product_name = str(ws.cell(row=row_idx, column=product_col).value or "").strip() if product_col is not None else ""
            if not self._looks_like_product_name(product_name):
                continue

            opening_qty = self._to_number(ws.cell(row=row_idx, column=opening_col).value) if opening_col is not None else 0.0
            closing_qty = self._to_number(ws.cell(row=row_idx, column=closing_col).value) if closing_col is not None else 0.0

            shipped_qty = (
                self._to_number(ws.cell(row=row_idx, column=voronezh_col).value) if voronezh_col is not None else 0.0
            )
            shipped_qty += (
                self._to_number(ws.cell(row=row_idx, column=krasnodar1_col).value) if krasnodar1_col is not None else 0.0
            )
            shipped_qty += (
                self._to_number(ws.cell(row=row_idx, column=krasnodar2_col).value) if krasnodar2_col is not None else 0.0
            )
            shipped_qty -= (
                self._to_number(ws.cell(row=row_idx, column=retail_col).value) if retail_col is not None else 0.0
            )

            item = ArrivalItem(
                row_number=row_idx,
                source_name=product_name,
            )
            item.ProductName = product_name
            item.OpeningQty = self._to_int(opening_qty)
            item.ShippedQty = self._to_int(shipped_qty)
            item.ClosingQty = self._to_int(closing_qty)

            item.product_name = item.ProductName
            item.opening_qty = item.OpeningQty
            item.shipped_qty = item.ShippedQty
            item.closing_qty = item.ClosingQty

            items.append(item)

        return items

    def load_arrival(self, filename: str | Path) -> list[ArrivalItem]:
        return self.load(filename)
