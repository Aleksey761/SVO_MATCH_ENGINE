from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from difflib import SequenceMatcher

from openpyxl import Workbook, load_workbook

from .business_rules import BusinessRules
from .engine import Engine
from .matcher import Matcher


AUTO_RESOLVE_MULTIPLE_CANDIDATES = False
AUTO_MATCH_SCORE_DELTA = 20.0


ROOT_CAUSES = [
    "MULTIPLE_CANDIDATES",
    "NO_MATCH",
    "PRODUCT_TYPE_MISMATCH",
    "BRAND_MISMATCH",
    "VOLUME_MISMATCH",
    "AROMA_MISMATCH",
    "MISSING_ALIAS",
    "OCR_OR_TYPOS",
    "NEW_PRODUCT",
    "OTHER",
]

SUGGESTED_FIXES = {
    "MULTIPLE_CANDIDATES": "Manual review",
    "NO_MATCH": "Manual review",
    "PRODUCT_TYPE_MISMATCH": "Manual review",
    "BRAND_MISMATCH": "Add alias",
    "VOLUME_MISMATCH": "Improve volume parser",
    "AROMA_MISMATCH": "Improve aroma parser",
    "MISSING_ALIAS": "Add alias",
    "OCR_OR_TYPOS": "Add alias",
    "NEW_PRODUCT": "Add MASTER item",
    "OTHER": "Manual review",
}


@dataclass(frozen=True)
class ReviewClassificationEntry:
    source_row: int
    supplier_article: str
    original_name: str
    canonical_name: str
    matched_sku: str
    review_reason: str
    root_cause: str
    best_candidate_sku: str
    best_candidate_score: float | None
    suggested_fix: str


@dataclass(frozen=True)
class MultipleCandidatesAnalysisEntry:
    supplier_article: str
    original_name: str
    candidate1_sku: str
    candidate1_score: float | None
    candidate2_sku: str
    candidate2_score: float | None
    candidate3_sku: str
    candidate3_score: float | None
    score_difference_12: float | None
    auto_match_possible: str
    recommendation: str


def _display_review_reason(engine: Engine, item) -> str:
    reason, _ = engine._build_price_review_reason(item)
    return reason


def _safe_text(value: object) -> str:
    return str(value or "").strip()


def _candidate_details(item, matcher: Matcher) -> list[dict[str, object]]:
    explanation = getattr(item, "review_explanation", {}) or {}
    detailed = [
        candidate
        for candidate in explanation.get("candidates", [])
        if isinstance(candidate, dict)
    ]
    if not detailed and getattr(item, "candidates", None):
        detailed = [matcher._score_detail(item, candidate) for candidate in item.candidates]
    detailed.sort(key=lambda entry: float(entry.get("score") or 0.0), reverse=True)
    return detailed


def _best_candidate(item, matcher: Matcher) -> dict[str, object] | None:
    explanation = getattr(item, "review_explanation", {}) or {}
    best = explanation.get("best_candidate")
    if isinstance(best, dict):
        return best
    details = _candidate_details(item, matcher)
    if details:
        return details[0]
    return None


def _best_source_similarity(item, matcher: Matcher) -> tuple[object | None, float]:
    best_master = None
    best_similarity = 0.0
    for master in matcher.master_items:
        similarity = matcher._source_master_similarity(item.source_name, master)
        if similarity > best_similarity:
            best_similarity = similarity
            best_master = master
    return best_master, best_similarity


def _classify_root_cause(item, matcher: Matcher) -> tuple[str, dict[str, object] | None, float]:
    reasons = {str(reason or "").strip().upper() for reason in getattr(item, "review_reasons", [])}
    best = _best_candidate(item, matcher)
    breakdown = dict(best.get("breakdown") or {}) if isinstance(best, dict) else {}
    best_score = float(best.get("score") or 0.0) if isinstance(best, dict) else 0.0

    if "PRODUCT_TYPE_MISMATCH" in reasons or (best and best.get("rejection_reason") == "PRODUCT_TYPE_MISMATCH"):
        return "PRODUCT_TYPE_MISMATCH", best, 0.0

    if "MULTIPLE_MATCH" in reasons:
        return "MULTIPLE_CANDIDATES", best, 0.0

    if breakdown.get("Brand", 0.0) <= 0.0 and _safe_text(getattr(item, "brand", None)):
        return "BRAND_MISMATCH", best, 0.0

    if "UNKNOWN_BRAND" in reasons:
        return "MISSING_ALIAS", best, 0.0

    if "NO_VOLUME" in reasons or (breakdown.get("Volume", 0.0) <= 0.0 and _safe_text(getattr(item, "volume", None))):
        return "VOLUME_MISMATCH", best, 0.0

    arrival_aroma = _safe_text(getattr(item, "aroma", None) or getattr(item, "variant", None))
    if arrival_aroma and breakdown and float(breakdown.get("Aroma", 0.0)) <= 0.0:
        return "AROMA_MISMATCH", best, 0.0

    best_master, best_similarity = _best_source_similarity(item, matcher)
    if best is None:
        if best_similarity >= 0.55:
            payload = None
            if best_master is not None:
                payload = {"sku": best_master.sku, "score": round(best_similarity * 100.0, 2)}
            return "OCR_OR_TYPOS", payload, best_similarity
        if best_similarity >= 0.30:
            payload = None
            if best_master is not None:
                payload = {"sku": best_master.sku, "score": round(best_similarity * 100.0, 2)}
            return "MISSING_ALIAS", payload, best_similarity
        return "NEW_PRODUCT", None, best_similarity

    if best_score < matcher.confidence_threshold:
        if best_similarity >= 0.75:
            return "OCR_OR_TYPOS", best, best_similarity
        if best_similarity >= 0.35:
            return "MISSING_ALIAS", best, best_similarity
        if best_similarity < 0.20:
            return "NEW_PRODUCT", best, best_similarity
        return "NO_MATCH", best, best_similarity

    source_name = _safe_text(getattr(item, "source_name", None)).upper()
    best_name = _safe_text(best.get("master_name") if isinstance(best, dict) else "").upper()
    if source_name and best_name:
        name_similarity = SequenceMatcher(None, source_name, best_name).ratio()
        if name_similarity >= 0.72 and best_score < 100.0:
            return "OCR_OR_TYPOS", best, best_similarity

    return "OTHER", best, best_similarity


def _write_classification_report(output_path: Path, entries: list[ReviewClassificationEntry]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "REVIEW_CLASSIFICATION"
    ws.append([
        "SourceRow",
        "SupplierArticle",
        "OriginalName",
        "CanonicalName",
        "MatchedSKU",
        "ReviewReason",
        "RootCause",
        "BestCandidateSKU",
        "BestCandidateScore",
        "SuggestedFix",
    ])
    for entry in entries:
        ws.append([
            entry.source_row,
            entry.supplier_article,
            entry.original_name,
            entry.canonical_name,
            entry.matched_sku,
            entry.review_reason,
            entry.root_cause,
            entry.best_candidate_sku,
            entry.best_candidate_score,
            entry.suggested_fix,
        ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_statistics_report(output_path: Path, entries: list[ReviewClassificationEntry]) -> list[tuple[str, int, str]]:
    total = len(entries)
    counts = Counter(entry.root_cause for entry in entries)
    rows = sorted(counts.items(), key=lambda item: (-item[1], item[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "REVIEW_STATISTICS"
    ws.append(["RootCause", "Count", "Percent"])
    report_rows: list[tuple[str, int, str]] = []
    for root_cause, count in rows:
        percent = f"{(count / total * 100.0):.2f}%" if total else "0.00%"
        ws.append([root_cause, count, percent])
        report_rows.append((root_cause, count, percent))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return report_rows


def _score_diff_12(details: list[dict[str, object]]) -> float | None:
    if len(details) < 2:
        return None
    first = details[0].get("score")
    second = details[1].get("score")
    try:
        return round(float(first) - float(second), 2)
    except (TypeError, ValueError):
        return None


def _find_master_by_sku(master_items, sku: str):
    normalized_sku = _safe_text(sku)
    if not normalized_sku:
        return None
    for master_item in master_items:
        if _safe_text(getattr(master_item, "sku", "")) == normalized_sku:
            return master_item
    return None


def _try_auto_resolve_multiple_candidate(
    *,
    item,
    matcher: Matcher,
    master_items,
    auto_match_threshold: float,
) -> bool:
    details = _candidate_details(item, matcher)
    score_diff_12 = _score_diff_12(details)
    if score_diff_12 is None or score_diff_12 < float(auto_match_threshold):
        return False

    top_candidate = details[0] if details else {}
    top_sku = _safe_text(top_candidate.get("sku"))
    if not top_sku:
        return False

    master_item = _find_master_by_sku(master_items, top_sku)
    if master_item is None:
        return False

    item.status = "MATCH"
    item.sku = master_item.sku
    item.master_name = master_item.master_name or item.master_name or item.source_name
    item.review_reasons = []
    item.review_explanation = {
        "auto_resolved_multiple_candidates": True,
        "score_difference_12": score_diff_12,
        "threshold": float(auto_match_threshold),
        "candidate_sku": master_item.sku,
    }
    return True


def _write_empty_reports(output_dir: Path) -> tuple[Path, Path]:
    classification_path = output_dir / "REVIEW_CLASSIFICATION.xlsx"
    statistics_path = output_dir / "REVIEW_STATISTICS.xlsx"
    _write_classification_report(classification_path, [])
    _write_statistics_report(statistics_path, [])
    return classification_path, statistics_path


def _write_multiple_candidates_analysis_report(
    output_path: Path,
    rows: list[MultipleCandidatesAnalysisEntry],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "MULTIPLE_CANDIDATES_ANALYSIS"
    ws.append([
        "SupplierArticle",
        "OriginalName",
        "Candidate1SKU",
        "Candidate1Score",
        "Candidate2SKU",
        "Candidate2Score",
        "Candidate3SKU",
        "Candidate3Score",
        "ScoreDifference12",
        "AutoMatchPossible",
        "Recommendation",
    ])

    for row in rows:
        ws.append([
            row.supplier_article,
            row.original_name,
            row.candidate1_sku,
            row.candidate1_score,
            row.candidate2_sku,
            row.candidate2_score,
            row.candidate3_sku,
            row.candidate3_score,
            row.score_difference_12,
            row.auto_match_possible,
            row.recommendation,
        ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_multiple_candidates_summary_report(
    output_path: Path,
    rows: list[MultipleCandidatesAnalysisEntry],
) -> list[tuple[str, int]]:
    ranges = [
        ("0-5", 0.0, 5.0),
        ("5-10", 5.0, 10.0),
        ("10-20", 10.0, 20.0),
        ("20-30", 20.0, 30.0),
        ("30+", 30.0, None),
    ]

    counters: dict[str, int] = {name: 0 for name, _, _ in ranges}
    for row in rows:
        diff = row.score_difference_12
        if diff is None:
            continue
        for name, start, end in ranges:
            if end is None and diff >= start:
                counters[name] += 1
                break
            if end is not None and start <= diff < end:
                counters[name] += 1
                break

    ordered = [(name, counters[name]) for name, _, _ in ranges]

    wb = Workbook()
    ws = wb.active
    ws.title = "MULTIPLE_CANDIDATES_SUMMARY"
    ws.append(["ScoreDifference range", "Count"])
    for name, count in ordered:
        ws.append([name, count])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return ordered


def _load_multiple_candidate_keys(classification_path: Path) -> set[tuple[str, str]]:
    if not classification_path.exists() or not classification_path.is_file():
        return set()

    wb = load_workbook(classification_path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(row=1, column=col).value or "").strip().upper() for col in range(1, ws.max_column + 1)]

    supplier_idx = headers.index("SUPPLIERARTICLE") + 1 if "SUPPLIERARTICLE" in headers else None
    original_idx = headers.index("ORIGINALNAME") + 1 if "ORIGINALNAME" in headers else None
    root_idx = headers.index("ROOTCAUSE") + 1 if "ROOTCAUSE" in headers else None
    if supplier_idx is None or original_idx is None or root_idx is None:
        return set()

    keys: set[tuple[str, str]] = set()
    for row in range(2, ws.max_row + 1):
        root = _safe_text(ws.cell(row=row, column=root_idx).value).upper()
        if root != "MULTIPLE_CANDIDATES":
            continue
        supplier = _safe_text(ws.cell(row=row, column=supplier_idx).value)
        original = _safe_text(ws.cell(row=row, column=original_idx).value)
        keys.add((supplier, original))
    return keys


def _render_multiple_candidates_summary(rows: list[MultipleCandidatesAnalysisEntry]) -> str:
    total = len(rows)
    auto = sum(1 for row in rows if row.auto_match_possible == "YES")
    manual = total - auto
    lines = [
        f"Total MULTIPLE_CANDIDATES: {total}",
        f"Resolvable automatically: {auto}",
        f"Still require REVIEW: {manual}",
    ]
    return "\n".join(lines)


def _render_console_summary(entries: list[ReviewClassificationEntry]) -> str:
    counts = Counter(entry.root_cause for entry in entries)
    lines = [
        "========== REVIEW ANALYSIS ==========",
        "",
        f"Total REVIEW: {len(entries)}",
        "",
    ]
    for root_cause in ROOT_CAUSES:
        lines.append(f"{root_cause}: {counts.get(root_cause, 0)}")
    return "\n".join(lines)


def build_review_classification_reports(
    *,
    master_file: str | Path | None,
    price_file: str | Path | None,
    output_dir: str | Path = "output",
    learning_map_file: str | Path | None = None,
    auto_resolve_multiple_candidates: bool = AUTO_RESOLVE_MULTIPLE_CANDIDATES,
    auto_match_threshold: float = AUTO_MATCH_SCORE_DELTA,
) -> dict[str, object]:
    output_path = Path(output_dir)
    classification_path = output_path / "REVIEW_CLASSIFICATION.xlsx"
    statistics_path = output_path / "REVIEW_STATISTICS.xlsx"
    multiple_candidates_path = output_path / "MULTIPLE_CANDIDATES_ANALYSIS.xlsx"
    multiple_candidates_summary_path = output_path / "MULTIPLE_CANDIDATES_SUMMARY.xlsx"

    if not master_file or not price_file:
        empty_classification, empty_statistics = _write_empty_reports(output_path)
        _write_multiple_candidates_analysis_report(multiple_candidates_path, [])
        multiple_candidates_summary = _write_multiple_candidates_summary_report(multiple_candidates_summary_path, [])
        console_text = _render_console_summary([])
        multiple_candidates_console_text = _render_multiple_candidates_summary([])
        return {
            "classification_path": empty_classification,
            "statistics_path": empty_statistics,
            "multiple_candidates_path": multiple_candidates_path,
            "multiple_candidates_summary_path": multiple_candidates_summary_path,
            "entries": [],
            "statistics": [],
            "multiple_candidates_summary": multiple_candidates_summary,
            "console_text": console_text,
            "multiple_candidates_console_text": multiple_candidates_console_text,
        }

    engine = Engine(learning_map_file=learning_map_file)
    master_items = engine.loader.load_master(master_file)
    price_items = engine.price_loader.load(price_file)
    for item in price_items:
        engine.normalizer.normalize(item)

    matcher = Matcher(
        master_items,
        auto_resolve_multiple_candidates=bool(auto_resolve_multiple_candidates),
        auto_match_score_delta=float(auto_match_threshold),
    )
    unresolved, _, _ = engine._apply_learning_map_matches(price_items, master_items)
    matcher.match_all(unresolved)

    BusinessRules(master_items).apply(price_items)
    engine._validate_price_name_sku_conflicts(price_items)

    entries: list[ReviewClassificationEntry] = []
    for item in price_items:
        if getattr(item, "status", "") != "REVIEW":
            continue

        root_cause, best_candidate, _ = _classify_root_cause(item, matcher)
        best_candidate_sku = ""
        best_candidate_score = None
        if isinstance(best_candidate, dict):
            best_candidate_sku = _safe_text(best_candidate.get("sku"))
            raw_score = best_candidate.get("score")
            if raw_score is not None:
                try:
                    best_candidate_score = round(float(raw_score), 2)
                except (TypeError, ValueError):
                    best_candidate_score = None

        entries.append(
            ReviewClassificationEntry(
                source_row=int(getattr(item, "row_number", 0) or 0),
                supplier_article=_safe_text(getattr(item, "supplier_article", None)),
                original_name=_safe_text(getattr(item, "original_product_name", None) or getattr(item, "source_name", None)),
                canonical_name=_safe_text(getattr(item, "canonical_product_name", None)),
                matched_sku=_safe_text(getattr(item, "sku", None)),
                review_reason=_display_review_reason(engine, item),
                root_cause=root_cause,
                best_candidate_sku=best_candidate_sku,
                best_candidate_score=best_candidate_score,
                suggested_fix=SUGGESTED_FIXES[root_cause],
            )
        )

    _write_classification_report(classification_path, entries)
    statistics = _write_statistics_report(statistics_path, entries)

    multiple_candidate_keys = _load_multiple_candidate_keys(classification_path)

    multiple_candidates_rows: list[MultipleCandidatesAnalysisEntry] = []
    for item in price_items:
        key = (
            _safe_text(getattr(item, "supplier_article", None)),
            _safe_text(getattr(item, "original_product_name", None) or getattr(item, "source_name", None)),
        )
        if key not in multiple_candidate_keys:
            continue

        details = _candidate_details(item, matcher)

        def candidate_at(index: int) -> tuple[str, float | None]:
            if index >= len(details):
                return "", None
            candidate = details[index]
            sku = _safe_text(candidate.get("sku"))
            score = candidate.get("score")
            try:
                return sku, round(float(score), 2)
            except (TypeError, ValueError):
                return sku, None

        c1_sku, c1_score = candidate_at(0)
        c2_sku, c2_score = candidate_at(1)
        c3_sku, c3_score = candidate_at(2)

        score_diff_12 = None
        auto_match_possible = "NO"
        if c1_score is not None and c2_score is not None:
            score_diff_12 = round(c1_score - c2_score, 2)
            if score_diff_12 >= float(auto_match_threshold):
                auto_match_possible = "YES"

        recommendation = (
            "Increase auto-match threshold candidate"
            if auto_match_possible == "YES"
            else "Keep in REVIEW"
        )

        multiple_candidates_rows.append(
            MultipleCandidatesAnalysisEntry(
                supplier_article=key[0],
                original_name=key[1],
                candidate1_sku=c1_sku,
                candidate1_score=c1_score,
                candidate2_sku=c2_sku,
                candidate2_score=c2_score,
                candidate3_sku=c3_sku,
                candidate3_score=c3_score,
                score_difference_12=score_diff_12,
                auto_match_possible=auto_match_possible,
                recommendation=recommendation,
            )
        )

    _write_multiple_candidates_analysis_report(multiple_candidates_path, multiple_candidates_rows)
    multiple_candidates_summary = _write_multiple_candidates_summary_report(
        multiple_candidates_summary_path,
        multiple_candidates_rows,
    )
    console_text = _render_console_summary(entries)
    multiple_candidates_console_text = _render_multiple_candidates_summary(multiple_candidates_rows)

    return {
        "classification_path": classification_path,
        "statistics_path": statistics_path,
        "multiple_candidates_path": multiple_candidates_path,
        "multiple_candidates_summary_path": multiple_candidates_summary_path,
        "entries": entries,
        "statistics": statistics,
        "multiple_candidates_summary": multiple_candidates_summary,
        "console_text": console_text,
        "multiple_candidates_console_text": multiple_candidates_console_text,
    }