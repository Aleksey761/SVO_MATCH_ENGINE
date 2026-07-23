from pathlib import Path

from openpyxl import load_workbook

from .models import ArrivalItem


def load_document_rows(filename: str | Path) -> list[ArrivalItem]:
    """Load first-column document rows into the shared arrival-style model."""
    wb = load_workbook(filename=filename, data_only=True)
    ws = wb.active

    items: list[ArrivalItem] = []

    for excel_row, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        name = row[0]

        if name is None:
            continue

        text = str(name).strip()

        if not text:
            continue

        items.append(
            ArrivalItem(
                row_number=excel_row,
                source_name=text,
            )
        )

    return items