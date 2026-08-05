from pathlib import Path

from openpyxl import load_workbook

from .arrival_loader import ArrivalLoader
from .models import ArrivalItem


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().upper().split())


def _looks_like_product_name(value: object) -> bool:
    if not isinstance(value, str):
        return False

    text = value.strip()
    if len(text) < 15:
        return False

    has_letters = any(char.isalpha() for char in text)
    has_separator = any(separator in text for separator in (",", " "))
    return has_letters and has_separator


class SalesLoader(ArrivalLoader):
    """Loads SALES workbook rows into the shared arrival item model."""

    def _find_name_column_index(self, worksheet) -> int:
        for row in worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True):
            normalized_headers = [_normalize_header(value) for value in row]

            for index, header in enumerate(normalized_headers):
                if header == "НАИМЕНОВАНИЕ":
                    return index

            non_empty_indexes = [index for index, header in enumerate(normalized_headers) if header]
            if len(non_empty_indexes) == 1:
                return non_empty_indexes[0]

        for row in worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True):
            text_candidates = [
                index
                for index, value in enumerate(row)
                if _looks_like_product_name(value)
            ]
            has_non_text_data = any(
                value is not None and not isinstance(value, str)
                for value in row
            )
            if len(text_candidates) == 1 and has_non_text_data:
                return text_candidates[0]

        raise ValueError("Could not determine SALES name column from worksheet header")

    def load(self, filename: str | Path) -> list[ArrivalItem]:
        wb = load_workbook(filename=filename, data_only=True)
        worksheet = wb.active
        name_column_index = self._find_name_column_index(worksheet)
        items: list[ArrivalItem] = []

        for excel_row, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            name = row[name_column_index] if name_column_index < len(row) else None
            if name is None:
                continue

            text = str(name).strip()
            if not text:
                continue

            items.append(
                ArrivalItem(
                    row_number=excel_row,
                    source_name=text,
                )
            )

        return items