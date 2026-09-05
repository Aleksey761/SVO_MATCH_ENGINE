from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook


@dataclass
class NameCorrectionDecision:
    source: str
    before_name: str
    after_name: str
    change_reason: str
    applied: bool


@dataclass
class CorrectionMapRow:
    wrong_name: str
    correct_master_name: str
    sku: str
    source: str
    rule_type: str
    evidence: str
    status: str


class NameCorrectionLayer:
    """Applies deterministic typo normalization before matching and records diagnostics."""

    _RULES = [
        ("COCOLADE", "CHOCOLATE", "SPELLING", "confirmed typo from inventory validation"),
        ("WILDBERES", "WILD BERRIES", "SPELLING", "confirmed typo from inventory validation"),
        ("LAVANDER", "LAVENDER", "SPELLING", "confirmed typo from inventory validation"),
        ("MEDIN", "MEDINA", "TYPO", "confirmed typo from business examples"),
    ]

    def __init__(self, master_items: list) -> None:
        self._master_by_normalized_name: dict[str, list[object]] = {}
        for item in master_items:
            master_name = str(getattr(item, "master_name", "") or "").strip()
            sku = str(getattr(item, "sku", "") or "").strip()
            if not master_name or not sku:
                continue
            self._master_by_normalized_name.setdefault(self._normalize(master_name), []).append(item)

        self._diagnostics: list[dict[str, str]] = []
        self._map_rows: list[CorrectionMapRow] = []

    @staticmethod
    def _normalize(value: object) -> str:
        text = str(value or "").strip().upper().replace("Ё", "Е")
        text = re.sub(r"[^0-9A-ZА-Я]+", " ", text)
        return " ".join(text.split())

    @staticmethod
    def _apply_rule(text: str, wrong: str, correct: str) -> tuple[str, bool]:
        pattern = r"\b" + re.escape(wrong) + r"\b"
        replaced = re.sub(pattern, correct, text)
        return replaced, replaced != text

    def apply_to_name(self, source_name: str, *, source: str) -> NameCorrectionDecision:
        before = str(source_name or "").strip()
        if not before:
            return NameCorrectionDecision(source=source, before_name="", after_name="", change_reason="", applied=False)

        working = self._normalize(before)
        reasons: list[str] = []
        for wrong, correct, rule_type, evidence in self._RULES:
            replaced, changed = self._apply_rule(working, wrong, correct)
            if changed:
                reasons.append(f"{rule_type}:{wrong}->{correct} ({evidence})")
                working = replaced

        after = " ".join(working.split())
        return NameCorrectionDecision(
            source=source,
            before_name=before,
            after_name=after,
            change_reason="; ".join(reasons),
            applied=bool(reasons),
        )

    def resolve_master_exact(self, corrected_name: str) -> tuple[str, str] | None:
        candidates = self._master_by_normalized_name.get(self._normalize(corrected_name), [])
        if len(candidates) != 1:
            return None
        item = candidates[0]
        return str(getattr(item, "sku", "") or "").strip(), str(getattr(item, "master_name", "") or "").strip()

    def apply_to_item(self, item, *, source: str) -> NameCorrectionDecision:
        working_name = str(getattr(item, "match_name", None) or getattr(item, "source_name", "") or "")
        decision = self.apply_to_name(working_name, source=source)
        setattr(item, "source_name_before_correction", decision.before_name)
        setattr(item, "source_name_after_correction", decision.after_name)
        setattr(item, "name_correction_reason", decision.change_reason)
        if decision.applied:
            item.match_name = decision.after_name
        return decision

    def record_diagnostic(
        self,
        *,
        source: str,
        before_name: str,
        after_name: str,
        master_name: str,
        sku: str,
        change_reason: str,
        status: str,
    ) -> None:
        self._diagnostics.append(
            {
                "Source": source,
                "BEFORE_NAME": before_name,
                "AFTER_NAME": after_name,
                "MASTER_NAME": master_name,
                "SKU": sku,
                "CHANGE_REASON": change_reason,
                "Status": status,
            }
        )

    def record_correction_map_row(self, *, source: str, wrong_name: str, corrected_name: str, sku: str, master_name: str, rule_type: str, evidence: str) -> None:
        if not sku or not master_name:
            return
        self._map_rows.append(
            CorrectionMapRow(
                wrong_name=wrong_name,
                correct_master_name=master_name,
                sku=sku,
                source=source,
                rule_type=rule_type,
                evidence=evidence,
                status="APPROVED",
            )
        )

    @staticmethod
    def _rule_type_from_reason(reason: str) -> str:
        text = str(reason or "").strip()
        if not text:
            return "SPACE_ERROR"
        token = text.split(";", 1)[0].strip()
        if ":" in token:
            token = token.split(":", 1)[0].strip()
        token = token.upper()
        allowed = {
            "TYPO",
            "TRANSLITERATION",
            "SPACE_ERROR",
            "SPELLING",
            "ABBREVIATION",
            "WORD_ORDER",
        }
        return token if token in allowed else "TYPO"

    def _map_rows_from_diagnostics(self) -> list[CorrectionMapRow]:
        rows: list[CorrectionMapRow] = []
        for record in self._diagnostics:
            before_name = str(record.get("BEFORE_NAME", "") or "").strip()
            after_name = str(record.get("AFTER_NAME", "") or "").strip()
            master_name = str(record.get("MASTER_NAME", "") or "").strip()
            sku = str(record.get("SKU", "") or "").strip()
            source = str(record.get("Source", "") or "").strip()
            reason = str(record.get("CHANGE_REASON", "") or "").strip()
            status = str(record.get("Status", "") or "").strip().upper()

            if not before_name or not after_name or before_name == after_name:
                continue
            if status != "MATCH":
                continue
            if not sku or not master_name:
                continue

            evidence = reason or "applied name normalization before MATCH"
            rows.append(
                CorrectionMapRow(
                    wrong_name=before_name,
                    correct_master_name=master_name,
                    sku=sku,
                    source=source,
                    rule_type=self._rule_type_from_reason(reason),
                    evidence=evidence,
                    status="APPROVED",
                )
            )

        return rows

    def write_name_correction_map(self, output_file: str | Path = "output/NAME_CORRECTION_MAP.xlsx") -> str:
        path = Path(output_file)
        existing_rows: list[tuple[str, str, str, str, str, str, str]] = []
        if path.exists():
            wb_existing = load_workbook(path, data_only=True)
            ws_existing = wb_existing["NAME_CORRECTION_MAP"] if "NAME_CORRECTION_MAP" in wb_existing.sheetnames else wb_existing.active
            for row in range(2, ws_existing.max_row + 1):
                values = tuple(str(ws_existing.cell(row=row, column=col).value or "").strip() for col in range(1, 8))
                if any(values):
                    existing_rows.append(values)

        wb = Workbook()
        ws = wb.active
        ws.title = "NAME_CORRECTION_MAP"
        ws.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status"])

        seen: set[tuple[str, str, str, str]] = set()
        for values in existing_rows:
            ws.append(list(values))
            seen.add((values[0], values[1], values[2], values[3]))

        all_rows = [*self._map_rows, *self._map_rows_from_diagnostics()]
        for row in all_rows:
            key = (row.wrong_name, row.correct_master_name, row.sku, row.source)
            if key in seen:
                continue
            seen.add(key)
            ws.append([
                row.wrong_name,
                row.correct_master_name,
                row.sku,
                row.source,
                row.rule_type,
                row.evidence,
                row.status,
            ])

        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return str(path)

    def write_name_normalization_report(self, output_file: str | Path = "output/NAME_NORMALIZATION_REPORT.xlsx") -> str:
        path = Path(output_file)
        wb = Workbook()

        ws_corrected = wb.active
        ws_corrected.title = "CORRECTED_NAMES"
        ws_corrected.append(["Source", "BEFORE_NAME", "AFTER_NAME", "MASTER_NAME", "SKU", "CHANGE_REASON"])

        ws_unresolved = wb.create_sheet("UNRESOLVED_NAMES")
        ws_unresolved.append(["Source", "BEFORE_NAME", "AFTER_NAME", "MASTER_NAME", "SKU", "CHANGE_REASON"])

        corrected_rows = 0
        unresolved_rows = 0
        review_removed_by_correction = 0
        for row in self._diagnostics:
            if row["BEFORE_NAME"] == row["AFTER_NAME"]:
                continue

            if row["Status"] == "MATCH":
                corrected_rows += 1
                review_removed_by_correction += 1
                ws_corrected.append([
                    row["Source"],
                    row["BEFORE_NAME"],
                    row["AFTER_NAME"],
                    row["MASTER_NAME"],
                    row["SKU"],
                    row["CHANGE_REASON"],
                ])
            else:
                unresolved_rows += 1
                ws_unresolved.append([
                    row["Source"],
                    row["BEFORE_NAME"],
                    row["AFTER_NAME"],
                    row["MASTER_NAME"],
                    row["SKU"],
                    row["CHANGE_REASON"],
                ])

        ws_summary = wb.create_sheet("SUMMARY")
        ws_summary.append(["Metric", "Value"])
        ws_summary.append(["Corrected names", corrected_rows])
        ws_summary.append(["Unresolved names", unresolved_rows])
        ws_summary.append(["Review removed only by name correction", review_removed_by_correction])

        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return str(path)
