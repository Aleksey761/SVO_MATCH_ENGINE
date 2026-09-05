from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from time import perf_counter

from openpyxl import Workbook, load_workbook

from svo.dashboard_builder import DashboardBuilder, RunContext
from svo.engine import Engine
from svo.loader import Loader
from svo.inventory_input_reader import InventoryInputReader
from svo.name_correction import NameCorrectionLayer
from svo.inventory_reconciliation import build_inventory_reconciliation
from svo.review_classification import build_review_classification_reports
from svo.analytics import generate_match_analytics_report
from svo.validator import ResultValidationReport
from finance.calculator import FinanceCalculator
from finance.price_history import PriceHistoryRecord


AUTO_RESOLVE_MULTIPLE_CANDIDATES = False
AUTO_MATCH_SCORE_DELTA = 20.0


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="SVO Match Engine CLI")
    parser.add_argument(
        "--mode",
        required=True,
        choices=("price", "sales", "inventory", "finance", "full"),
        help="Execution mode",
    )
    parser.add_argument("--master", help="Path to MASTER workbook")
    parser.add_argument("--price", help="Path to PRICE workbook")
    parser.add_argument("--sales", help="Path to SALES workbook or RESULT workbook (inventory mode)")
    parser.add_argument("--inventory", help="Path to inventory revision workbook")
    parser.add_argument("--stock-match", help="Path to STOCK_MATCH workbook (finance mode)")
    parser.add_argument("--output", help="Path to output workbook where supported")
    parser.add_argument(
        "--auto-resolve-multiple-candidates",
        action="store_true",
        default=AUTO_RESOLVE_MULTIPLE_CANDIDATES,
        help="Auto-assign Candidate1 for MULTIPLE_CANDIDATES when score delta condition is met",
    )
    parser.add_argument(
        "--auto-match-score-delta",
        type=float,
        default=AUTO_MATCH_SCORE_DELTA,
        help="Minimum Candidate1-Candidate2 score difference required for auto-resolve",
    )
    return parser


def _ensure_file(path_value: str | None, *, name: str) -> Path:
    if not path_value:
        raise ValueError(f"Missing required argument --{name}")

    path = Path(path_value)
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"File not found for --{name}: {path}")
    return path


def _ensure_runtime_dirs() -> None:
    for folder in ("output", "logs", "reports"):
        Path(folder).mkdir(parents=True, exist_ok=True)


def _run_price(engine: Engine, args: argparse.Namespace) -> dict:
    master_file = _ensure_file(args.master, name="master")
    price_file = _ensure_file(args.price, name="price")
    output_file = Path(args.output) if args.output else None

    print("[price] Running PRICE matching pipeline")
    result = engine.run_price_matching(master_file=master_file, price_file=price_file, output_file=output_file)
    print("=====================================")
    print("PRICE IMPORT SUMMARY")
    print("=====================================")
    print(f"MASTER rows   : {result['master']}")
    print(f"PRICE rows    : {result['price_rows']}")
    print(f"Matched       : {result['match']}")
    print(f"Review        : {result['review']}")
    print(f"New prices    : {result['price_history_created'] - result['price_history_updated']}")
    print(f"Updated prices: {result['price_history_updated']}")
    print(f"Skipped       : {result['price_history_skipped']}")
    print("Output files")
    print(f"price_match.xlsx: {result['price_match_output']}")
    print(f"price_change_report.xlsx: {result['price_change_report']}")
    print(f"price_review.xlsx: {result['price_review_report']}")
    print("Completed successfully")
    print("=====================================")
    return result


def _run_sales(engine: Engine, args: argparse.Namespace) -> dict:
    master_file = _ensure_file(args.master, name="master")
    sales_file = _ensure_file(args.sales, name="sales")
    output_file = Path(args.output) if args.output else None

    print("[sales] Running SALES matching pipeline")
    result = engine.run_sales(master_file=master_file, sales_file=sales_file, output_file=output_file)
    print(f"[sales] Rows   : {result['rows']}")
    print(f"[sales] MATCH  : {result['match']}")
    print(f"[sales] REVIEW : {result['review']}")
    print(f"[sales] OUTPUT : {result['output']}")
    return result


def _run_inventory(args: argparse.Namespace) -> dict:
    master_file = _ensure_file(args.master, name="master")
    result_file = _ensure_file(args.sales, name="sales")
    inventory_file = _ensure_file(args.inventory, name="inventory")
    output_file = Path(args.output) if args.output else None

    print("[inventory] Running inventory reconciliation pipeline")
    result = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )
    print(f"[inventory] Matched rows : {result['matched_rows']}")
    print(f"[inventory] Variance rows: {result['variance_rows']}")
    print(f"[inventory] OUTPUT       : {result['output']}")
    return result


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().upper().split())


def _to_decimal_or_none(value: object) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "")
    if not text:
        return None
    text = text.replace(",", ".")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _to_int_or_zero(value: object) -> int:
    parsed = _to_decimal_or_none(value)
    if parsed is None:
        return 0
    return int(parsed.to_integral_value())


def _load_price_match_rows(price_match_file: Path) -> list[dict[str, object]]:
    wb = load_workbook(filename=price_match_file, data_only=True)
    ws = wb.active

    headers = [str(ws.cell(row=1, column=col).value or "").strip() for col in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
    required = ("SKU", "MASTER_NAME", "UnitCost", "RetailPrice", "StockQty", "Status")
    missing = [name for name in required if name not in idx]
    if missing:
        raise ValueError(f"PRICE match workbook is missing required columns: {', '.join(missing)}")

    rows: list[dict[str, object]] = []
    for row in range(2, ws.max_row + 1):
        status = str(ws.cell(row=row, column=idx["Status"]).value or "").strip().upper()
        if not status:
            continue
        rows.append(
            {
                "SKU": str(ws.cell(row=row, column=idx["SKU"]).value or "").strip(),
                "MASTER_NAME": str(ws.cell(row=row, column=idx["MASTER_NAME"]).value or "").strip(),
                "UnitCost": ws.cell(row=row, column=idx["UnitCost"]).value,
                "RetailPrice": ws.cell(row=row, column=idx["RetailPrice"]).value,
                "StockQty": ws.cell(row=row, column=idx["StockQty"]).value,
                "Status": status,
            }
        )

    return rows


def _find_sales_columns(ws) -> tuple[int | None, int | None]:
    sku_col = None
    qty_col = None

    max_scan_row = min(ws.max_row, 4)
    max_scan_col = min(ws.max_column, 40)
    for row in range(1, max_scan_row + 1):
        normalized = [_normalize_header(ws.cell(row=row, column=col).value) for col in range(1, max_scan_col + 1)]
        if sku_col is None:
            for idx, header in enumerate(normalized, start=1):
                if header == "SKU":
                    sku_col = idx
                    break
        if qty_col is None:
            for idx, header in enumerate(normalized, start=1):
                if header == "ВСЕГО":
                    qty_col = idx
                    break

    if sku_col is not None and qty_col is None and ws.max_column >= 14:
        qty_col = 14

    return sku_col, qty_col


def _load_sales_qty_by_sku(sales_file: Path | None) -> dict[str, int]:
    if sales_file is None or not sales_file.exists():
        return {}

    wb = load_workbook(filename=sales_file, data_only=True)
    ws = wb.active
    sku_col, qty_col = _find_sales_columns(ws)
    if sku_col is None or qty_col is None:
        return {}

    sales_qty_by_sku: dict[str, int] = {}
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row=row, column=sku_col).value or "").strip()
        if not sku:
            continue
        qty = _to_int_or_zero(ws.cell(row=row, column=qty_col).value)
        sales_qty_by_sku[sku] = sales_qty_by_sku.get(sku, 0) + qty
    return sales_qty_by_sku


def _load_source_to_skus_map(bridge_file: Path | None) -> dict[str, set[str]]:
    if bridge_file is None or not bridge_file.exists():
        return {}

    wb = load_workbook(filename=bridge_file, data_only=True)
    ws = wb.active
    headers = [_normalize_header(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
    source_col = idx.get("SOURCE_NAME")
    sku_col = idx.get("SKU")
    status_col = idx.get("STATUS")
    if source_col is None or sku_col is None:
        return {}

    source_to_skus: dict[str, set[str]] = {}
    for row in range(2, ws.max_row + 1):
        if status_col is not None:
            status = str(ws.cell(row=row, column=status_col).value or "").strip().upper()
            if status and status != "MATCH":
                continue

        source = " ".join(str(ws.cell(row=row, column=source_col).value or "").strip().upper().split())
        sku = str(ws.cell(row=row, column=sku_col).value or "").strip()
        if not source or not sku:
            continue
        source_to_skus.setdefault(source, set()).add(sku)

    return source_to_skus


def _load_source_to_skus_map_with_correction(
    bridge_file: Path | None,
    *,
    correction_layer: NameCorrectionLayer,
) -> dict[str, set[str]]:
    if bridge_file is None or not bridge_file.exists():
        return {}

    wb = load_workbook(filename=bridge_file, data_only=True)
    ws = wb.active
    headers = [_normalize_header(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
    source_col = idx.get("SOURCE_NAME")
    sku_col = idx.get("SKU")
    status_col = idx.get("STATUS")
    if source_col is None or sku_col is None:
        return {}

    source_to_skus: dict[str, set[str]] = {}
    for row in range(2, ws.max_row + 1):
        if status_col is not None:
            status = str(ws.cell(row=row, column=status_col).value or "").strip().upper()
            if status and status != "MATCH":
                continue

        source = str(ws.cell(row=row, column=source_col).value or "").strip()
        decision = correction_layer.apply_to_name(source, source="SALES")
        normalized_source = " ".join(decision.after_name.strip().upper().split())

        sku = str(ws.cell(row=row, column=sku_col).value or "").strip()
        if not normalized_source or not sku:
            continue
        source_to_skus.setdefault(normalized_source, set()).add(sku)

    return source_to_skus


def _load_shipped_by_sku(
    inventory_file: Path | None,
    *,
    source_to_skus: dict[str, set[str]],
    correction_layer: NameCorrectionLayer | None = None,
) -> tuple[dict[str, int], int, str | None]:
    if inventory_file is None or not inventory_file.exists():
        return {}, 0, None

    try:
        quantities = InventoryInputReader().read_inventory_quantities(inventory_file)
    except ValueError as exc:
        return {}, 0, str(exc)

    shipped_by_sku: dict[str, int] = {}
    ambiguous_fallback_rows = 0

    for quantity in quantities:
        explicit_sku = str(quantity.get("sku", "")).strip()
        if explicit_sku:
            sku = explicit_sku
        else:
            raw_source_name = str(quantity.get("source_name", "")).strip()
            corrected_source_name = raw_source_name
            reason = ""
            if correction_layer is not None:
                decision = correction_layer.apply_to_name(raw_source_name, source="INVENTORY")
                corrected_source_name = decision.after_name
                reason = decision.change_reason

            source_key = " ".join(corrected_source_name.strip().upper().split())
            candidate_skus = source_to_skus.get(source_key)
            if not source_key or not candidate_skus:
                if correction_layer is not None and raw_source_name:
                    correction_layer.record_diagnostic(
                        source="INVENTORY",
                        before_name=raw_source_name,
                        after_name=corrected_source_name,
                        master_name="",
                        sku="",
                        change_reason=reason,
                        status="REVIEW",
                    )
                continue
            if len(candidate_skus) != 1:
                ambiguous_fallback_rows += 1
                if correction_layer is not None and raw_source_name:
                    correction_layer.record_diagnostic(
                        source="INVENTORY",
                        before_name=raw_source_name,
                        after_name=corrected_source_name,
                        master_name="",
                        sku="",
                        change_reason=reason,
                        status="REVIEW",
                    )
                continue
            sku = next(iter(candidate_skus))

            if correction_layer is not None and raw_source_name:
                correction_layer.record_diagnostic(
                    source="INVENTORY",
                    before_name=raw_source_name,
                    after_name=corrected_source_name,
                    master_name="",
                    sku=sku,
                    change_reason=reason,
                    status="MATCH",
                )

                if raw_source_name != corrected_source_name:
                    master_name = ""
                    resolved = correction_layer.resolve_master_exact(corrected_source_name)
                    if resolved is not None:
                        _, master_name = resolved
                    correction_layer.record_correction_map_row(
                        source="INVENTORY",
                        wrong_name=raw_source_name,
                        corrected_name=corrected_source_name,
                        sku=sku,
                        master_name=master_name,
                        rule_type="TYPO",
                        evidence=reason or "inventory source exact correction",
                    )

        shipped_qty = int(quantity.get("supplier_return", 0)) - int(quantity.get("sales", 0))
        shipped_by_sku[sku] = shipped_by_sku.get(sku, 0) + shipped_qty

    return shipped_by_sku, ambiguous_fallback_rows, None


def _load_inventory_alias_map(alias_map_file: Path | None) -> dict[str, str]:
    if alias_map_file is None or not alias_map_file.exists():
        return {}

    wb = load_workbook(filename=alias_map_file, data_only=True)
    ws = wb["APPROVED_CANDIDATES"] if "APPROVED_CANDIDATES" in wb.sheetnames else wb.active
    headers = [str(ws.cell(row=1, column=col).value or "").strip() for col in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
    inventory_name_col = idx.get("InventoryName")
    sku_col = idx.get("SKU")
    status_col = idx.get("Status")
    if inventory_name_col is None or sku_col is None:
        return {}

    alias_by_inventory_name: dict[str, str] = {}
    allowed_statuses = {"APPROVED_FOR_REVIEW", "APPROVED_READY"}

    for row in range(2, ws.max_row + 1):
        status = str(ws.cell(row=row, column=status_col).value or "").strip().upper() if status_col else ""
        if status_col and status and status not in allowed_statuses:
            continue

        inventory_name = str(ws.cell(row=row, column=inventory_name_col).value or "").strip()
        sku = str(ws.cell(row=row, column=sku_col).value or "").strip()
        if not inventory_name or not sku:
            continue
        alias_by_inventory_name[inventory_name] = sku

    return alias_by_inventory_name


def _load_inventory_stock_by_sku_exact_alias(
    inventory_file: Path | None,
    *,
    alias_by_inventory_name: dict[str, str],
) -> tuple[dict[str, int], int]:
    if inventory_file is None or not inventory_file.exists() or not alias_by_inventory_name:
        return {}, 0

    wb = load_workbook(filename=inventory_file, data_only=True)
    ws = wb.active

    stock_by_sku: dict[str, int] = {}
    applied_alias_rows = 0
    for row_idx in range(3, ws.max_row + 1):
        inventory_name = str(ws.cell(row=row_idx, column=3).value or "").strip()
        if not inventory_name:
            continue

        sku = alias_by_inventory_name.get(inventory_name)
        if not sku:
            continue

        stock_qty = _to_int_or_zero(ws.cell(row=row_idx, column=15).value)
        stock_by_sku[sku] = stock_by_sku.get(sku, 0) + stock_qty
        applied_alias_rows += 1

    return stock_by_sku, applied_alias_rows


def _load_stock_match_by_sku(stock_match_file: Path | None) -> dict[str, int]:
    if stock_match_file is None or not stock_match_file.exists():
        raise FileNotFoundError(
            "STOCK_MATCH workbook not found: "
            f"{stock_match_file}. Build it first via '--mode inventory' or '--mode full'."
        )

    wb = load_workbook(filename=stock_match_file, data_only=True)
    ws = wb["STOCK_RESULT"] if "STOCK_RESULT" in wb.sheetnames else wb.active

    headers = [_normalize_header(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}

    sku_col = idx.get("SKU")
    stock_col = idx.get("ФАКТИЧЕСКИЙ ОСТАТОК")
    if sku_col is None:
        raise ValueError(
            "STOCK_MATCH.xlsx: column 'SKU' not found in STOCK_RESULT"
        )

    if stock_col is None:
        raise ValueError(
            "STOCK_MATCH.xlsx: column 'ФАКТИЧЕСКИЙ ОСТАТОК' not found in STOCK_RESULT"
        )

    stock_by_sku: dict[str, int] = {}
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row=row, column=sku_col).value or "").strip()
        if not sku:
            continue
        stock_qty = _to_int_or_zero(ws.cell(row=row, column=stock_col).value)
        stock_by_sku[sku] = stock_qty

    return stock_by_sku


def _execute_finance_pipeline(args: argparse.Namespace) -> dict:
    master_file = _ensure_file(args.master, name="master")
    output_file = Path(args.output) if args.output else Path("output") / "FINANCE_RESULT.xlsx"
    stock_match_arg = getattr(args, "stock_match", None)
    stock_match_file = Path(stock_match_arg) if stock_match_arg else Path("output") / "STOCK_MATCH.xlsx"

    price_match_file = Path(args.price) if args.price else Path("output") / "price_match.xlsx"
    if not price_match_file.exists():
        raise FileNotFoundError(f"PRICE match workbook not found: {price_match_file}")

    sales_file = Path(args.sales) if args.sales else None
    if sales_file is None:
        for candidate in (
            Path("output") / "SALES_MATCH_06.04.2026.xlsx",
            Path("output") / "SALES_MATCH.xlsx",
        ):
            if candidate.exists():
                sales_file = candidate
                break

    inventory_file = Path(args.inventory) if args.inventory else None
    if inventory_file is None:
        data_candidates = sorted(Path("data").glob("*инвентар*.xlsx"))
        if data_candidates:
            inventory_file = data_candidates[0]

    bridge_file = Path("output") / "_FULL_RESULT_FOR_STOCK.xlsx"
    price_rows = [row for row in _load_price_match_rows(price_match_file) if row["Status"] == "MATCH"]
    master_items = Loader().load_master(master_file)
    master_by_sku = {str(item.sku or "").strip(): item for item in master_items if str(item.sku or "").strip()}
    correction_layer = NameCorrectionLayer(master_items)

    sales_qty_by_sku = _load_sales_qty_by_sku(sales_file)
    if stock_match_arg is None and not stock_match_file.exists():
        stock_by_sku = {}
    else:
        stock_by_sku = _load_stock_match_by_sku(stock_match_file)
    source_to_skus = _load_source_to_skus_map_with_correction(
        bridge_file,
        correction_layer=correction_layer,
    )
    shipped_by_sku, ambiguous_fallback_rows, inventory_parse_warning = _load_shipped_by_sku(
        inventory_file,
        source_to_skus=source_to_skus,
        correction_layer=correction_layer,
    )
    correction_layer.write_name_correction_map()
    correction_layer.write_name_normalization_report()

    calculator = FinanceCalculator()
    finance_rows: list[dict[str, object]] = []
    missing_master = 0
    for row in price_rows:
        sku = str(row["SKU"] or "").strip()
        if not sku:
            continue

        master_item = master_by_sku.get(sku)
        if master_item is None:
            missing_master += 1

        master_name = str(row["MASTER_NAME"] or "").strip() or (master_item.master_name if master_item is not None else "")
        unit_cost = _to_decimal_or_none(row["UnitCost"])
        retail_price = _to_decimal_or_none(row["RetailPrice"])
        stock_qty = int(stock_by_sku.get(sku, 0))
        sales_qty = int(sales_qty_by_sku.get(sku, 0))
        shipped_qty = int(shipped_by_sku.get(sku, 0))

        revenue = Decimal("0")
        cogs = Decimal("0")
        gross_profit = Decimal("0")
        margin_pct = Decimal("0")
        stock_value = Decimal("0")
        if unit_cost is not None and retail_price is not None:
            revenue = calculator.calculate_revenue(sales_qty=sales_qty, retail_price=retail_price, pricing_status="OK")
            cogs = calculator.calculate_cogs(sales_qty=sales_qty, unit_cost=unit_cost, pricing_status="OK")
            gross_profit = calculator.calculate_gross_profit(revenue=revenue, cogs=cogs)
            margin_pct = calculator.calculate_gross_margin_percent(revenue=revenue, gross_profit=gross_profit)
            valuation = calculator.calculate_inventory_valuation(
                calculated_stock_qty=stock_qty,
                actual_stock_qty=stock_qty,
                variance_qty=0,
                supplier_return_qty=0,
                unit_cost=unit_cost,
                pricing_status="OK",
            )
            stock_value = valuation["StockValue"]

        finance_rows.append(
            {
                "SKU": sku,
                "MASTER_NAME": master_name,
                "SalesQty": sales_qty,
                "RetailPrice": retail_price,
                "Revenue": revenue,
                "UnitCost": unit_cost,
                "COGS": cogs,
                "GrossProfit": gross_profit,
                "MarginPct": margin_pct,
                "StockQty": stock_qty,
                "StockValue": stock_value,
                "Отгружено": shipped_qty,
                "HasSales": sales_qty > 0,
                "HasCost": unit_cost is not None,
                "HasStock": stock_qty > 0,
            }
        )

    total_sku = len(finance_rows)
    with_sales = sum(1 for row in finance_rows if row["HasSales"])
    with_cost = sum(1 for row in finance_rows if row["HasCost"])
    with_stock = sum(1 for row in finance_rows if row["HasStock"])
    complete = sum(1 for row in finance_rows if row["HasSales"] and row["HasCost"] and row["HasStock"])

    wb = Workbook()
    ws = wb.active
    ws.title = "FINANCE_RESULT"
    ws.append([
        "SKU",
        "MASTER_NAME",
        "SalesQty",
        "RetailPrice",
        "Revenue",
        "UnitCost",
        "COGS",
        "GrossProfit",
        "MarginPct",
        "StockQty",
        "StockValue",
        "Отгружено",
    ])
    for row in finance_rows:
        ws.append(
            [
                row["SKU"],
                row["MASTER_NAME"],
                row["SalesQty"],
                row["RetailPrice"],
                row["Revenue"],
                row["UnitCost"],
                row["COGS"],
                row["GrossProfit"],
                row["MarginPct"],
                row["StockQty"],
                row["StockValue"],
                row["Отгружено"],
            ]
        )

    ws_lineage = wb.create_sheet("FINANCE_LINEAGE")
    ws_lineage.append(["COLUMN", "SOURCE", "MODULE", "NOTES"])
    ws_lineage.append(["SKU", str(price_match_file), "PRICE_MATCH", "Exact MATCH SKU from output/price_match.xlsx"])
    ws_lineage.append(["MASTER_NAME", str(price_match_file), "PRICE_MATCH/MASTER", "price_match MASTER_NAME; fallback from Loader master_by_sku"])
    ws_lineage.append(["SalesQty", str(sales_file) if sales_file else "N/A", "SALES_MATCH", "Column ВСЕГО by SKU when available, else 0"])
    ws_lineage.append(["RetailPrice", str(price_match_file), "PRICE_MATCH", "RetailPrice column"])
    ws_lineage.append(["Revenue", "computed", "FinanceCalculator", "sales_qty * retail_price"])
    ws_lineage.append(["UnitCost", str(price_match_file), "PRICE_MATCH", "UnitCost column"])
    ws_lineage.append(["COGS", "computed", "FinanceCalculator", "sales_qty * unit_cost"])
    ws_lineage.append(["GrossProfit", "computed", "FinanceCalculator", "revenue - cogs"])
    ws_lineage.append(["MarginPct", "computed", "FinanceCalculator", "(gross_profit / revenue) * 100; revenue=0 => 0"])
    ws_lineage.append([
        "StockQty",
        str(stock_match_file),
        "STOCK_MATCH",
        "Fact stock by SKU from STOCK_MATCH.xlsx (STOCK_RESULT: SKU + Фактический остаток)",
    ])
    ws_lineage.append(["StockValue", "computed", "FinanceCalculator", "stock_qty * unit_cost"])
    ws_lineage.append(["Отгружено", str(inventory_file) if inventory_file else "N/A", "InventoryInputReader", "supplier_return - sales, SKU-first mapping with deterministic fallback"])
    if inventory_parse_warning:
        ws_lineage.append(["Отгружено_WARNING", str(inventory_file), "InventoryInputReader", inventory_parse_warning])

    ws_lineage.append([])
    ws_lineage.append(["SUMMARY", "VALUE"])
    ws_lineage.append(["Total SKU", total_sku])
    ws_lineage.append(["SKU with sales", with_sales])
    ws_lineage.append(["SKU without sales", total_sku - with_sales])
    ws_lineage.append(["SKU with cost", with_cost])
    ws_lineage.append(["SKU without cost", total_sku - with_cost])
    ws_lineage.append(["SKU with stock", with_stock])
    ws_lineage.append(["SKU without stock", total_sku - with_stock])
    ws_lineage.append(["Finance-complete SKU", complete])
    ws_lineage.append(["Finance-incomplete SKU", total_sku - complete])
    ws_lineage.append(["Missing MASTER for PRICE SKU", missing_master])
    ws_lineage.append(["Ambiguous inventory fallback rows", ambiguous_fallback_rows])
    ws_lineage.append(["Stock rows loaded from STOCK_MATCH", len(stock_by_sku)])

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    print("[finance] FINANCE_RESULT generated")
    print(f"[finance] OUTPUT      : {output_file}")
    print(f"[finance] Total SKU   : {total_sku}")
    print(f"[finance] With sales  : {with_sales}")
    print(f"[finance] With cost   : {with_cost}")
    print(f"[finance] With stock  : {with_stock}")
    print(f"[finance] Incomplete  : {total_sku - complete}")
    print(f"[finance] STOCK rows  : {len(stock_by_sku)}")
    if inventory_parse_warning:
        print(f"[finance] Inventory warning: {inventory_parse_warning}")

    return {
        "output": str(output_file),
        "total_sku": total_sku,
        "with_sales": with_sales,
        "without_sales": total_sku - with_sales,
        "with_cost": with_cost,
        "without_cost": total_sku - with_cost,
        "with_stock": with_stock,
        "without_stock": total_sku - with_stock,
        "finance_complete": complete,
        "finance_incomplete": total_sku - complete,
        "missing_master": missing_master,
        "ambiguous_inventory_fallback_rows": ambiguous_fallback_rows,
        "price_match_file": str(price_match_file),
        "sales_file": str(sales_file) if sales_file else None,
        "inventory_file": str(inventory_file) if inventory_file else None,
        "stock_match_file": str(stock_match_file),
        "inventory_parse_warning": inventory_parse_warning,
        "stock_rows_loaded": len(stock_by_sku),
    }


def _run_finance(args: argparse.Namespace) -> dict:
    return _execute_finance_pipeline(args)


def _touch_xlsx(path: Path, headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    if headers:
        ws.append(headers)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _copy_xlsx(src: str | Path, dst: Path) -> None:
    wb = load_workbook(filename=src, data_only=False)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)


def _copy_sheet_values(*, source_file: Path, source_sheet: str | None, target_ws, target_title: str) -> None:
    wb = load_workbook(filename=source_file, data_only=True)
    ws = wb[source_sheet] if source_sheet else wb.active
    target_ws.title = target_title
    for row in ws.iter_rows(values_only=True):
        target_ws.append(list(row))


def _read_text_lines(path: Path) -> list[str]:
    if not path.exists():
        return []
    return path.read_text(encoding="utf-8").splitlines()


def _extract_new_products(lines: list[str]) -> list[str]:
    marker = "NEW PRODUCTS FOR MASTER"
    start = None
    for idx, line in enumerate(lines):
        if line.strip().upper() == marker:
            start = idx + 1
            break
    if start is None:
        return []

    extracted: list[str] = []
    for line in lines[start:]:
        value = line.strip()
        if value.startswith("----------------------------------------") or value.startswith("VALIDATION") or value.startswith("========================================"):
            break
        if value:
            extracted.append(value)
    return extracted


def _build_final_report(
    *,
    output_dir: Path,
    summary_payload: dict[str, object],
    analytics_file: Path,
    audit_file: Path,
    review_file: Path,
    sales_file: Path,
    stock_file: Path,
    price_change_file: Path,
) -> Path:
    final_path = output_dir / "FINAL_REPORT.xlsx"

    analytics_lines = _read_text_lines(analytics_file)
    audit_lines = _read_text_lines(audit_file)
    new_products = _extract_new_products(analytics_lines)

    wb = Workbook()
    ws_dashboard = wb.active
    ws_dashboard.title = "Dashboard"
    ws_dashboard.append(["KPI", "Value"])

    sales_total = int(summary_payload.get("sales_matched", 0)) + int(summary_payload.get("sales_review", 0))
    coverage = (int(summary_payload.get("sales_matched", 0)) / sales_total * 100.0) if sales_total else 0.0

    ws_dashboard.append(["MASTER SKU", int(summary_payload.get("master_records", 0))])
    ws_dashboard.append(["PRICE Match", int(summary_payload.get("price_matched", 0))])
    ws_dashboard.append(["PRICE Review", int(summary_payload.get("price_review", 0))])
    ws_dashboard.append(["SALES Match", int(summary_payload.get("sales_matched", 0))])
    ws_dashboard.append(["STOCK Match", int(summary_payload.get("stock_matched", 0))])
    ws_dashboard.append(["Coverage", round(coverage, 2)])
    ws_dashboard.append(["Price increases", int(summary_payload.get("price_inc", 0))])
    ws_dashboard.append(["Price decreases", int(summary_payload.get("price_dec", 0))])
    ws_dashboard.append(["New products", len(new_products)])
    ws_dashboard.append(["Execution time", f"{int(summary_payload.get('duration_sec', 0))} sec"])

    ws_price = wb.create_sheet("Price Changes")
    ws_price.append(["SKU", "MASTER_NAME", "OldPrice", "NewPrice", "Difference", "ChangedAt"])
    if price_change_file.exists():
        wb_changes = load_workbook(filename=price_change_file, data_only=True)
        ws_changes = wb_changes.active
        headers = [str(ws_changes.cell(row=1, column=col).value or "").strip().upper() for col in range(1, ws_changes.max_column + 1)]
        idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
        for row in range(2, ws_changes.max_row + 1):
            ws_price.append(
                [
                    ws_changes.cell(row=row, column=idx.get("SKU", 1)).value,
                    ws_changes.cell(row=row, column=idx.get("MASTER_NAME", 2)).value,
                    ws_changes.cell(row=row, column=idx.get("PREVIOUSPRICE", 3)).value,
                    ws_changes.cell(row=row, column=idx.get("NEWPRICE", 4)).value,
                    ws_changes.cell(row=row, column=idx.get("DIFFERENCE", 5)).value,
                    ws_changes.cell(row=row, column=idx.get("IMPORTEDAT", idx.get("VALIDFROM", 1))).value,
                ]
            )

    ws_review = wb.create_sheet("Review")
    if review_file.exists():
        _copy_sheet_values(source_file=review_file, source_sheet=None, target_ws=ws_review, target_title="Review")
    else:
        ws_review.append(["No review output found"])

    ws_new = wb.create_sheet("New Products")
    ws_new.append(["Entry"])
    if new_products:
        for value in new_products:
            ws_new.append([value])
    else:
        ws_new.append(["none"])

    ws_sales = wb.create_sheet("Sales")
    if sales_file.exists():
        _copy_sheet_values(source_file=sales_file, source_sheet=None, target_ws=ws_sales, target_title="Sales")
    else:
        ws_sales.append(["No sales output found"])

    ws_stock = wb.create_sheet("Stock")
    if stock_file.exists():
        stock_sheet_name = None
        wb_stock = load_workbook(filename=stock_file, data_only=True)
        if "STOCK_RESULT" in wb_stock.sheetnames:
            stock_sheet_name = "STOCK_RESULT"
        _copy_sheet_values(source_file=stock_file, source_sheet=stock_sheet_name, target_ws=ws_stock, target_title="Stock")
    else:
        ws_stock.append(["No stock output found"])

    ws_analytics = wb.create_sheet("Analytics")
    ws_analytics.append(["Line"])
    if analytics_lines:
        for line in analytics_lines:
            ws_analytics.append([line])
    else:
        ws_analytics.append(["MATCH_ANALYTICS not found"])

    ws_audit = wb.create_sheet("Audit")
    ws_audit.append(["Line"])
    if audit_lines:
        for line in audit_lines:
            ws_audit.append([line])
    else:
        ws_audit.append(["MASTER_AUDIT not found"])

    final_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(final_path)
    return final_path


def _write_master_match(master_file: Path, output_path: Path) -> int:
    wb_src = load_workbook(filename=master_file, data_only=True)
    ws_src = wb_src.active

    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER_MATCH"

    count = 0
    for row in ws_src.iter_rows(values_only=True):
        ws.append(list(row))
        if count > 0:
            sku = str((row[0] if row and len(row) > 0 else "") or "").strip()
            if sku:
                count += 1
            continue
        count += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    # subtract header row if present
    return max(0, count - 1)


def _write_master_dataset_from_master(master_file: Path, output_path: Path) -> int:
    """Create canonical MASTER_DATASET structure from MASTER source workbook."""
    wb_src = load_workbook(filename=master_file, data_only=True)
    ws_src = wb_src.active

    headers = [str(ws_src.cell(row=1, column=col).value or "").strip().upper() for col in range(1, ws_src.max_column + 1)]
    header_idx = {name: pos + 1 for pos, name in enumerate(headers) if name}

    def _col(*candidates: str) -> int | None:
        for candidate in candidates:
            idx = header_idx.get(candidate)
            if idx is not None:
                return idx
        return None

    sku_col = _col("SKU")
    category_col = _col("CATEGORY", "КАТЕГОРИЯ")
    brand_col = _col("BRAND", "БРЕНД")
    variant_col = _col("VARIANT", "АРОМАТ")
    volume_col = _col("VOLUME", "ОБЪЕМ")
    aroma_col = _col("AROMA", "АРОМАТ", "VARIANT")
    master_name_col = _col("MASTER_NAME", "НАИМЕНОВАНИЕ")

    if sku_col is None:
        raise ValueError("MASTER workbook is missing required SKU column")

    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER_DATASET"
    ws.append(["SKU", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "AROMA", "C7", "C8", "C9", "MASTER_NAME"])

    def _cell(row: int, col: int | None) -> str:
        if col is None:
            return ""
        return str(ws_src.cell(row=row, column=col).value or "").strip()

    count = 0
    for row in range(2, ws_src.max_row + 1):
        sku = _cell(row, sku_col)
        if not sku:
            continue

        category = _cell(row, category_col)
        brand = _cell(row, brand_col)
        variant = _cell(row, variant_col)
        volume = _cell(row, volume_col)
        aroma = _cell(row, aroma_col) or variant
        master_name = _cell(row, master_name_col)
        if not master_name:
            master_name = f"{category} {brand} {variant} {volume}".strip()

        ws.append([sku, category, brand, variant, volume, aroma, "", "", "", master_name])
        count += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return count


def _write_price_history(records: list[PriceHistoryRecord], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE_HISTORY"
    ws.append([
        "SKU",
        "MASTER_NAME",
        "UnitCost",
        "RetailPrice",
        "ValidFrom",
        "ValidTo",
        "Status",
        "SourceFile",
        "ImportedAt",
        "SourceUnit",
        "PackQty",
        "OriginalPrice",
    ])

    for record in records:
        ws.append([
            record.SKU,
            record.MASTER_NAME,
            record.UnitCost,
            record.RetailPrice,
            record.ValidFrom.isoformat() if record.ValidFrom else "",
            record.ValidTo.isoformat() if record.ValidTo else "",
            record.Status,
            record.SourceFile,
            record.ImportedAt.isoformat() if record.ImportedAt else "",
            record.SourceUnit,
            record.PackQty,
            record.OriginalPrice,
        ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_errors(errors: list[dict[str, object]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ERRORS"
    ws.append(["Module", "Row", "Object", "Reason", "Details"])
    for error in errors:
        ws.append([
            error.get("Module", ""),
            error.get("Row", ""),
            error.get("Object", ""),
            error.get("Reason", ""),
            error.get("Details", ""),
        ])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_stock_result_source_from_sales(items: list, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESULT"
    ws.append(["SOURCE_NAME", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "SKU", "MASTER_NAME", "STATUS"])
    for item in items:
        ws.append([
            getattr(item, "source_name", ""),
            getattr(item, "category", "") or "",
            getattr(item, "brand", "") or "",
            getattr(item, "variant", "") or "",
            getattr(item, "volume", "") or "",
            getattr(item, "sku", "") or "",
            getattr(item, "master_name", "") or "",
            getattr(item, "status", "") or "",
        ])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_run_summary(output_path: Path, payload: dict[str, object]) -> None:
    text = "\n".join(
        [
            "====================================",
            "",
            "MASTER",
            f"Records: {payload.get('master_records', 0)}",
            "",
            "PRICE",
            f"Matched : {payload.get('price_matched', 0)}",
            f"Review  : {payload.get('price_review', 0)}",
            f"New SKU : {payload.get('price_new_sku', 0)}",
            "",
            "SALES",
            f"Matched : {payload.get('sales_matched', 0)}",
            f"Review  : {payload.get('sales_review', 0)}",
            "",
            "STOCK",
            f"Matched : {payload.get('stock_matched', 0)}",
            f"Review  : {payload.get('stock_review', 0)}",
            "",
            "Price changes",
            f"Increased : {payload.get('price_inc', 0)}",
            f"Decreased : {payload.get('price_dec', 0)}",
            "",
            "Execution time",
            f"{payload.get('duration_sec', 0)} sec",
            "",
            "====================================",
            "",
        ]
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(text, encoding="utf-8")


def _write_master_audit_json(output_path: Path, *, passed: bool, errors_count: int) -> None:
    payload = {
        "gate": "master",
        "status": "PASSED" if passed else "FAILED",
        "errors": int(errors_count),
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    import json
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_business_metrics_xlsx(output_path: Path, *, summary_payload: dict[str, object]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "BUSINESS_METRICS"
    ws.append(["Metric", "Value"])
    ws.append(["MASTER rows", int(summary_payload.get("master_records", 0))])
    ws.append(["PRICE matched", int(summary_payload.get("price_matched", 0))])
    ws.append(["PRICE review", int(summary_payload.get("price_review", 0))])
    ws.append(["SALES matched", int(summary_payload.get("sales_matched", 0))])
    ws.append(["SALES review", int(summary_payload.get("sales_review", 0))])
    ws.append(["STOCK matched", int(summary_payload.get("stock_matched", 0))])
    ws.append(["STOCK review", int(summary_payload.get("stock_review", 0))])
    ws.append(["Duration sec", int(summary_payload.get("duration_sec", 0))])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _run_full(engine: Engine, args: argparse.Namespace) -> dict:
    started = perf_counter()
    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)

    master_match_file = output_dir / "MASTER_MATCH.xlsx"
    master_dataset_file = output_dir / "MASTER_DATASET.xlsx"
    price_match_file = output_dir / "price_match.xlsx"
    price_history_file = output_dir / "PRICE_HISTORY.xlsx"
    price_change_file = output_dir / "price_change_report.xlsx"
    price_review_file = output_dir / "price_review.xlsx"
    sales_match_file = output_dir / "SALES_MATCH.xlsx"
    stock_match_file = output_dir / "STOCK_MATCH.xlsx"
    errors_file = output_dir / "ERRORS.xlsx"
    summary_file = output_dir / "RUN_SUMMARY.txt"
    final_report_file = output_dir / "FINAL_REPORT.xlsx"
    stock_source_result_file = output_dir / "_FULL_RESULT_FOR_STOCK.xlsx"
    analytics_file = output_dir / "MATCH_ANALYTICS.txt"
    audit_file = output_dir / "MASTER_AUDIT.txt"
    audit_json_file = output_dir / "MASTER_AUDIT.json"
    business_metrics_file = output_dir / "BUSINESS_METRICS.xlsx"
    review_file = output_dir / "REVIEW_REPORT.xlsx"
    review_classification_file = output_dir / "REVIEW_CLASSIFICATION.xlsx"
    review_statistics_file = output_dir / "REVIEW_STATISTICS.xlsx"
    multiple_candidates_analysis_file = output_dir / "MULTIPLE_CANDIDATES_ANALYSIS.xlsx"
    multiple_candidates_summary_file = output_dir / "MULTIPLE_CANDIDATES_SUMMARY.xlsx"

    errors: list[dict[str, object]] = []

    master_file = _ensure_file(args.master, name="master")
    master_records = _write_master_match(master_file, master_match_file)
    _write_master_dataset_from_master(master_file, master_dataset_file)

    price_matched = 0
    price_review = 0
    price_new_sku = 0
    price_inc = 0
    price_dec = 0

    sales_matched = 0
    sales_review = 0
    sales_items: list = []

    stock_matched = 0
    stock_review = 0

    price_file = Path(args.price) if args.price else None
    if price_file is not None and price_file.exists():
        try:
            price_result = engine.run_price_matching(
                master_file=master_file,
                price_file=price_file,
                output_file=price_match_file,
            )
            price_matched = int(price_result.get("match", 0))
            price_review = int(price_result.get("review", 0))
            created = int(price_result.get("price_history_created", 0))
            updated = int(price_result.get("price_history_updated", 0))
            price_new_sku = max(0, created - updated)

            _write_price_history(engine.price_history.records, price_history_file)
            _copy_xlsx(price_result.get("price_change_report", price_change_file), price_change_file)
            _copy_xlsx(price_result.get("price_review_report", price_review_file), price_review_file)

            for event in price_result.get("rows", []):
                _ = event
            for change_event in getattr(engine.price_history, "records", []):
                _ = change_event

            # Count increased/decreased from import events if available.
            price_events = price_result.get("price_change_report")
            if price_events:
                pass

            # Use price history events from engine output by reopening the generated report.
            wb = load_workbook(price_change_file, data_only=True)
            ws = wb.active
            headers = [str(ws.cell(row=1, column=col).value or "").strip().upper() for col in range(1, ws.max_column + 1)]
            diff_col = headers.index("DIFFERENCE") + 1 if "DIFFERENCE" in headers else None
            if diff_col is not None:
                for row in range(2, ws.max_row + 1):
                    value = ws.cell(row=row, column=diff_col).value
                    if value is None:
                        continue
                    try:
                        number = float(value)
                    except Exception:
                        continue
                    if number > 0:
                        price_inc += 1
                    elif number < 0:
                        price_dec += 1
        except Exception as exc:
            errors.append(
                {
                    "Module": "PRICE",
                    "Row": "",
                    "Object": str(price_file),
                    "Reason": type(exc).__name__,
                    "Details": str(exc),
                }
            )
            _touch_xlsx(price_match_file, ["SKU", "MASTER_NAME", "MatchStatus", "Confidence", "ReviewReason"])
            _touch_xlsx(price_history_file, ["SKU", "MASTER_NAME", "UnitCost", "RetailPrice", "Status"])
            _touch_xlsx(price_change_file, ["SKU", "MASTER_NAME", "PreviousPrice", "NewPrice", "Difference", "Status"])
            _touch_xlsx(price_review_file, ["SupplierArticle", "OriginalProductName", "CanonicalProductName", "SKU", "MASTER_NAME", "Reason", "Recommendation"])
    else:
        errors.append(
            {
                "Module": "PRICE",
                "Row": "",
                "Object": str(price_file or ""),
                "Reason": "MissingInput",
                "Details": "PRICE file is missing for full pipeline",
            }
        )
        _touch_xlsx(price_match_file, ["SKU", "MASTER_NAME", "MatchStatus", "Confidence", "ReviewReason"])
        _touch_xlsx(price_history_file, ["SKU", "MASTER_NAME", "UnitCost", "RetailPrice", "Status"])
        _touch_xlsx(price_change_file, ["SKU", "MASTER_NAME", "PreviousPrice", "NewPrice", "Difference", "Status"])
        _touch_xlsx(price_review_file, ["SupplierArticle", "OriginalProductName", "CanonicalProductName", "SKU", "MASTER_NAME", "Reason", "Recommendation"])

    sales_file = Path(args.sales) if args.sales else None
    if sales_file is not None and sales_file.exists():
        try:
            sales_result = engine.run_sales(
                master_file=master_file,
                sales_file=sales_file,
                output_file=sales_match_file,
            )
            sales_matched = int(sales_result.get("match", 0))
            sales_review = int(sales_result.get("review", 0))
            sales_items = list(sales_result.get("items", []))
        except Exception as exc:
            errors.append(
                {
                    "Module": "SALES",
                    "Row": "",
                    "Object": str(sales_file),
                    "Reason": type(exc).__name__,
                    "Details": str(exc),
                }
            )
            _touch_xlsx(sales_match_file, ["MATCH_STATUS", "MATCH_SKU", "MATCH_MASTER_NAME", "MATCH_CONFIDENCE", "MATCH_REASONS"])
    else:
        errors.append(
            {
                "Module": "SALES",
                "Row": "",
                "Object": str(sales_file or ""),
                "Reason": "MissingInput",
                "Details": "SALES file is missing for full pipeline",
            }
        )
        _touch_xlsx(sales_match_file, ["MATCH_STATUS", "MATCH_SKU", "MATCH_MASTER_NAME", "MATCH_CONFIDENCE", "MATCH_REASONS"])

    inventory_file = Path(args.inventory) if args.inventory else None
    if inventory_file is not None and inventory_file.exists() and sales_items:
        try:
            _write_stock_result_source_from_sales(sales_items, stock_source_result_file)
            stock_result = build_inventory_reconciliation(
                master_file=master_file,
                result_file=stock_source_result_file,
                inventory_file=inventory_file,
                output_file=stock_match_file,
            )
            stock_matched = int(stock_result.get("matched_rows", 0))
            stock_review = int(stock_result.get("unmatched_rows", 0))
        except Exception as exc:
            errors.append(
                {
                    "Module": "STOCK",
                    "Row": "",
                    "Object": str(inventory_file),
                    "Reason": type(exc).__name__,
                    "Details": str(exc),
                }
            )
            _touch_xlsx(stock_match_file, ["SKU", "Наименование", "Начальный остаток", "Приход", "Возврат поставщику", "Продажи", "Расчетный остаток", "Фактический остаток", "Расхождение"])
    else:
        errors.append(
            {
                "Module": "STOCK",
                "Row": "",
                "Object": str(inventory_file or ""),
                "Reason": "MissingInput",
                "Details": "STOCK inventory file is missing or SALES import was not available",
            }
        )
        _touch_xlsx(stock_match_file, ["SKU", "Наименование", "Начальный остаток", "Приход", "Возврат поставщику", "Продажи", "Расчетный остаток", "Фактический остаток", "Расхождение"])

    duration = int(round(perf_counter() - started))
    summary_payload = {
        "master_records": master_records,
        "price_matched": price_matched,
        "price_review": price_review,
        "price_new_sku": price_new_sku,
        "sales_matched": sales_matched,
        "sales_review": sales_review,
        "stock_matched": stock_matched,
        "stock_review": stock_review,
        "price_inc": price_inc,
        "price_dec": price_dec,
        "duration_sec": duration,
    }
    _write_run_summary(summary_file, summary_payload)
    _write_business_metrics_xlsx(business_metrics_file, summary_payload=summary_payload)

    validation_report = ResultValidationReport(
        passed=True,
        match_count=sales_matched,
        review_count=sales_review,
        empty_sku_count=0,
        empty_name_count=0,
        duplicate_sku_count=0,
        issues=[],
    )
    try:
        generate_match_analytics_report(
            report_date=None,
            source_file=sales_file,
            total_rows=sales_matched + sales_review,
            match_count=sales_matched,
            review_count=sales_review,
            validation_report=validation_report,
            items=sales_items,
            output_file=analytics_file,
        )
    except Exception as exc:
        errors.append(
            {
                "Module": "ANALYTICS",
                "Row": "",
                "Object": str(analytics_file),
                "Reason": type(exc).__name__,
                "Details": str(exc),
            }
        )

    final_report_path = _build_final_report(
        output_dir=output_dir,
        summary_payload=summary_payload,
        analytics_file=analytics_file,
        audit_file=audit_file,
        review_file=review_file,
        sales_file=sales_match_file,
        stock_file=stock_match_file,
        price_change_file=price_change_file,
    )
    _write_errors(errors, errors_file)
    _write_master_audit_json(audit_json_file, passed=(len(errors) == 0), errors_count=len(errors))

    review_analysis = build_review_classification_reports(
        master_file=master_file,
        price_file=price_file,
        output_dir=output_dir,
        auto_resolve_multiple_candidates=bool(args.auto_resolve_multiple_candidates),
        auto_match_threshold=float(args.auto_match_score_delta),
    )
    review_classification_file = Path(review_analysis["classification_path"])
    review_statistics_file = Path(review_analysis["statistics_path"])
    multiple_candidates_analysis_file = Path(review_analysis["multiple_candidates_path"])
    multiple_candidates_summary_file = Path(review_analysis["multiple_candidates_summary_path"])

    print("====================================")
    print("MASTER")
    print(f"Records: {master_records}")
    print("")
    print("PRICE")
    print(f"Matched : {price_matched}")
    print(f"Review  : {price_review}")
    print(f"New SKU : {price_new_sku}")
    print("")
    print("SALES")
    print(f"Matched : {sales_matched}")
    print(f"Review  : {sales_review}")
    print("")
    print("STOCK")
    print(f"Matched : {stock_matched}")
    print(f"Review  : {stock_review}")
    print("")
    print("Price changes")
    print(f"Increased : {price_inc}")
    print(f"Decreased : {price_dec}")
    print("")
    print("Execution time")
    print(f"{duration} sec")
    print("====================================")
    print(f"Summary file: {summary_file}")
    print(f"Errors file : {errors_file}")
    print(f"Final report: {final_report_path}")
    print(review_analysis["console_text"])
    print(review_analysis["multiple_candidates_console_text"])

    return {
        "document": "FULL",
        "summary": str(summary_file),
        "errors": str(errors_file),
        "outputs": {
            "MASTER_DATASET": str(master_dataset_file),
            "MASTER_MATCH": str(master_match_file),
            "PRICE_MATCH": str(price_match_file),
            "PRICE_HISTORY": str(price_history_file),
            "PRICE_CHANGE_REPORT": str(price_change_file),
            "PRICE_REVIEW": str(price_review_file),
            "SALES_MATCH": str(sales_match_file),
            "STOCK_MATCH": str(stock_match_file),
            "BUSINESS_METRICS": str(business_metrics_file),
            "MATCH_ANALYTICS": str(analytics_file),
            "MASTER_AUDIT": str(audit_json_file),
            "REVIEW_CLASSIFICATION": str(review_classification_file),
            "REVIEW_STATISTICS": str(review_statistics_file),
            "MULTIPLE_CANDIDATES_ANALYSIS": str(multiple_candidates_analysis_file),
            "MULTIPLE_CANDIDATES_SUMMARY": str(multiple_candidates_summary_file),
            "FINAL_REPORT": str(final_report_file),
        },
        "errors_count": len(errors),
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }


def main(argv: list[str] | None = None) -> int:
    parser = _build_parser()
    parsed_argv = list(sys.argv[1:] if argv is None else argv)
    if parsed_argv and str(parsed_argv[0]).strip().lower().endswith((".py", ".pyw")):
        parsed_argv = parsed_argv[1:]
    args = parser.parse_args(parsed_argv)
    run_started = datetime.now(timezone.utc)

    try:
        _ensure_runtime_dirs()
        print("[runner] Runtime directories ensured: output/, logs/, reports/")

        engine = Engine()
        print(f"[runner] Mode selected: {args.mode}")

        run_result: dict[str, object] | None = None
        if args.mode == "price":
            run_result = _run_price(engine, args)
        elif args.mode == "sales":
            run_result = _run_sales(engine, args)
        elif args.mode == "inventory":
            run_result = _run_inventory(args)
        elif args.mode == "finance":
            run_result = _run_finance(args)
        elif args.mode == "full":
            run_result = _run_full(engine, args)
        else:
            raise ValueError(f"Unsupported mode: {args.mode}")

        if args.mode == "full":
            elapsed_sec = int(round((datetime.now(timezone.utc) - run_started).total_seconds()))
            outputs = dict(run_result.get("outputs", {})) if run_result else {}
            run_context = RunContext(
                master_dataset_path=Path(outputs["MASTER_DATASET"]),
                price_match_path=Path(outputs["PRICE_MATCH"]),
                price_review_path=Path(outputs["PRICE_REVIEW"]),
                price_change_report_path=Path(outputs["PRICE_CHANGE_REPORT"]),
                business_metrics_path=Path(outputs["BUSINESS_METRICS"]),
                match_analytics_path=Path(outputs["MATCH_ANALYTICS"]),
                master_audit_path=Path(outputs["MASTER_AUDIT"]),
                output_directory=Path("output"),
                run_date=run_started.strftime("%Y-%m-%d"),
                run_time=run_started.strftime("%H:%M:%S"),
                duration=f"{elapsed_sec} sec",
                source_master=str(Path(outputs["MASTER_DATASET"]).resolve()),
                source_price=str(Path(outputs["PRICE_MATCH"]).resolve()),
                version="runner-full",
                status="SUCCESS",
            )
            dashboard_path = DashboardBuilder(run_context).build()
            print(f"[runner] Dashboard generated: {dashboard_path}")

        print("[runner] Completed successfully")
        return 0
    except (ValueError, FileNotFoundError, NotImplementedError) as exc:
        print(f"[runner] ERROR: {exc}")
        return 1 if args.mode == "full" else 2
    except Exception as exc:  # pragma: no cover - defensive top-level guard
        print(f"[runner] UNEXPECTED ERROR: {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
