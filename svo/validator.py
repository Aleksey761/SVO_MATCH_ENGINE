from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ResultValidationReport:
    passed: bool
    match_count: int
    review_count: int
    empty_sku_count: int
    empty_name_count: int
    duplicate_sku_count: int
    issues: list[str]


def _normalize_header(value: object) -> str:
    return str(value or "").strip().upper()


def validate_result(
    result_file: str | Path,
    *,
    pipeline_match_count: int | None = None,
    pipeline_review_count: int | None = None,
) -> ResultValidationReport:
    wb = load_workbook(result_file, data_only=True)
    ws = wb.active

    headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
    normalized_headers = [_normalize_header(h) for h in headers]

    issues: list[str] = []

    required_headers = ("№", "SKU")
    header_to_col = {
        normalized: idx + 1
        for idx, normalized in enumerate(normalized_headers)
        if normalized
    }
    for required in required_headers:
        if required not in header_to_col:
            issues.append(f"Required column is missing: {required}")

    technical_headers = [h for h in normalized_headers if h.startswith("MATCH_")]
    if technical_headers:
        unique_headers = sorted({h for h in technical_headers})
        issues.append(f"Technical columns must be removed: {', '.join(unique_headers)}")

    match_count = pipeline_match_count if pipeline_match_count is not None else 0
    review_count = pipeline_review_count if pipeline_review_count is not None else 0
    empty_sku_count = 0
    empty_name_count = 0
    sku_values: list[str] = []

    number_col = header_to_col.get("№")
    sku_col = header_to_col.get("SKU")
    name_col = header_to_col.get("НАИМЕНОВАНИЕ")

    max_col = ws.max_column
    for row_idx in range(2, ws.max_row + 1):
        row_values = [ws.cell(row=row_idx, column=col).value for col in range(1, max_col + 1)]
        if all(value is None or str(value).strip() == "" for value in row_values):
            continue

        number_value = str(ws.cell(row=row_idx, column=number_col).value or "").strip() if number_col else ""
        sku = str(ws.cell(row=row_idx, column=sku_col).value or "").strip() if sku_col else ""
        name = str(ws.cell(row=row_idx, column=name_col).value or "").strip() if name_col else ""

        # Treat a row as a product row only when at least one required field is present.
        # This skips template/support rows that contain only warehouse-unit metadata.
        if not (number_value or sku or name):
            continue

        if not number_value:
            issues.append(f"Row {row_idx}: required value is empty in column '№'")
        if not sku:
            empty_sku_count += 1
            issues.append(f"Row {row_idx}: required value is empty in column 'SKU'")
        if name_col is not None and not name:
            empty_name_count += 1
            issues.append(f"Row {row_idx}: required value is empty in column 'Наименование'")

        if sku:
            sku_values.append(sku)

    duplicate_sku_count = sum(count - 1 for count in Counter(sku_values).values() if count > 1)

    passed = len(issues) == 0

    return ResultValidationReport(
        passed=passed,
        match_count=match_count,
        review_count=review_count,
        empty_sku_count=empty_sku_count,
        empty_name_count=empty_name_count,
        duplicate_sku_count=duplicate_sku_count,
        issues=issues,
    )
