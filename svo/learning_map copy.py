from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook


_HEADERS = [
    "SupplierArticle",
    "SupplierName",
    "CanonicalSupplier",
    "SKU",
    "MASTER_NAME",
    "Confidence",
    "CreatedAt",
    "UpdatedAt",
    "ConfirmedBy",
]


def _normalize_text(value: object) -> str:
    return " ".join(str(value or "").strip().upper().split())


def canonical_supplier_value(item: object) -> str:
    canonical = getattr(item, "canonical_product_name", None)
    if canonical:
        return _normalize_text(canonical)
    return _normalize_text(getattr(item, "source_name", None))


@dataclass
class LearningMapEntry:
    supplier_article: str
    supplier_name: str
    canonical_supplier: str
    sku: str
    master_name: str
    confidence: float
    created_at: str
    updated_at: str
    confirmed_by: str


class LearningMap:
    """Persistent map of manually confirmed REVIEW decisions."""

    def __init__(self, file_path: str | Path = "data/learning_map.xlsx") -> None:
        self.file_path = Path(file_path)
        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        if not self.file_path.exists():
            self._create_empty_workbook()

    def _create_empty_workbook(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "LEARNING_MAP"
        ws.append(_HEADERS)
        wb.save(self.file_path)

    @staticmethod
    def _key(supplier_article: object, canonical_supplier: object) -> tuple[str, str]:
        return _normalize_text(supplier_article), _normalize_text(canonical_supplier)

    @staticmethod
    def _now_iso() -> str:
        return datetime.now(timezone.utc).isoformat()

    def _load_workbook(self):
        if not self.file_path.exists():
            self._create_empty_workbook()
        wb = load_workbook(filename=self.file_path)
        ws = wb.active
        if ws.max_row < 1:
            ws.append(_HEADERS)
        else:
            header = [str(ws.cell(row=1, column=i).value or "").strip() for i in range(1, len(_HEADERS) + 1)]
            if header != _HEADERS:
                ws.delete_rows(1, ws.max_row)
                ws.append(_HEADERS)
        return wb, ws

    def all_entries(self) -> list[LearningMapEntry]:
        wb, ws = self._load_workbook()
        entries: list[LearningMapEntry] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            supplier_article = _normalize_text(row[0])
            canonical_supplier = _normalize_text(row[2])
            sku = str(row[3] or "").strip()
            if not supplier_article or not canonical_supplier or not sku:
                continue
            entries.append(
                LearningMapEntry(
                    supplier_article=supplier_article,
                    supplier_name=str(row[1] or "").strip(),
                    canonical_supplier=canonical_supplier,
                    sku=sku,
                    master_name=str(row[4] or "").strip(),
                    confidence=float(row[5] or 0.0),
                    created_at=str(row[6] or ""),
                    updated_at=str(row[7] or ""),
                    confirmed_by=str(row[8] or "").strip(),
                )
            )
        wb.close()
        return entries

    def lookup(self, supplier_article: object, canonical_supplier: object) -> LearningMapEntry | None:
        key = self._key(supplier_article, canonical_supplier)
        if not key[0] or not key[1]:
            return None

        wb, ws = self._load_workbook()
        found: LearningMapEntry | None = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_key = self._key(row[0], row[2])
            if row_key != key:
                continue
            sku = str(row[3] or "").strip()
            if not sku:
                continue
            found = LearningMapEntry(
                supplier_article=_normalize_text(row[0]),
                supplier_name=str(row[1] or "").strip(),
                canonical_supplier=_normalize_text(row[2]),
                sku=sku,
                master_name=str(row[4] or "").strip(),
                confidence=float(row[5] or 0.0),
                created_at=str(row[6] or ""),
                updated_at=str(row[7] or ""),
                confirmed_by=str(row[8] or "").strip(),
            )
            break

        wb.close()
        return found

    def upsert(
        self,
        *,
        supplier_article: object,
        supplier_name: object,
        canonical_supplier: object,
        sku: object,
        master_name: object,
        confidence: float,
        confirmed_by: object,
    ) -> LearningMapEntry:
        key = self._key(supplier_article, canonical_supplier)
        sku_value = str(sku or "").strip()
        if not key[0]:
            raise ValueError("SupplierArticle is required for learning map")
        if not key[1]:
            raise ValueError("CanonicalSupplier is required for learning map")
        if not sku_value:
            raise ValueError("SKU is required for learning map")

        now = self._now_iso()
        wb, ws = self._load_workbook()

        target_row: int | None = None
        created_at = now
        for row_idx in range(2, ws.max_row + 1):
            row_key = self._key(
                ws.cell(row=row_idx, column=1).value,
                ws.cell(row=row_idx, column=3).value,
            )
            if row_key == key:
                target_row = row_idx
                existing_created_at = str(ws.cell(row=row_idx, column=7).value or "").strip()
                if existing_created_at:
                    created_at = existing_created_at
                break

        if target_row is None:
            target_row = ws.max_row + 1

        ws.cell(row=target_row, column=1).value = key[0]
        ws.cell(row=target_row, column=2).value = str(supplier_name or "").strip()
        ws.cell(row=target_row, column=3).value = key[1]
        ws.cell(row=target_row, column=4).value = sku_value
        ws.cell(row=target_row, column=5).value = str(master_name or "").strip()
        ws.cell(row=target_row, column=6).value = float(confidence)
        ws.cell(row=target_row, column=7).value = created_at
        ws.cell(row=target_row, column=8).value = now
        ws.cell(row=target_row, column=9).value = str(confirmed_by or "").strip() or "MANUAL"

        wb.save(self.file_path)
        wb.close()

        return LearningMapEntry(
            supplier_article=key[0],
            supplier_name=str(supplier_name or "").strip(),
            canonical_supplier=key[1],
            sku=sku_value,
            master_name=str(master_name or "").strip(),
            confidence=float(confidence),
            created_at=created_at,
            updated_at=now,
            confirmed_by=str(confirmed_by or "").strip() or "MANUAL",
        )

    def confirm_review_item(
        self,
        item: object,
        *,
        sku: str,
        master_name: str,
        confidence: float,
        confirmed_by: str,
    ) -> LearningMapEntry:
        supplier_article = getattr(item, "supplier_article", None)
        supplier_name = getattr(item, "source_name", None)
        canonical_supplier = canonical_supplier_value(item)
        return self.upsert(
            supplier_article=supplier_article,
            supplier_name=supplier_name,
            canonical_supplier=canonical_supplier,
            sku=sku,
            master_name=master_name,
            confidence=confidence,
            confirmed_by=confirmed_by,
        )
