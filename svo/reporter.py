from pathlib import Path
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .models import ArrivalItem, MasterItem


class Reporter:
    """Writes matching results to RESULT.xlsx."""

    @staticmethod
    def _normalize_header(value: object) -> str:
        return str(value or "").strip().upper()

    @staticmethod
    def _set_cell_value(ws, row: int, column: int, value: object) -> bool:
        cell = ws.cell(row=row, column=column)
        if isinstance(cell, MergedCell):
            return False
        cell.value = value
        return True

    @classmethod
    def _find_sales_header_row(cls, ws) -> int | None:
        required_headers = {"ВОРОНЕЖ", "КРАСНОДАР 1", "КРАСНОДАР 2"}
        explicit_name_headers = {"НАИМЕНОВАНИЕ", "NAME", "PRODUCT", "PRODUCT NAME", "SOURCE_NAME"}

        for row_number in range(1, min(ws.max_row, 12) + 1):
            normalized = {
                cls._normalize_header(ws.cell(row=row_number, column=col).value)
                for col in range(1, ws.max_column + 1)
            }
            if required_headers.issubset(normalized):
                return row_number
            if required_headers.intersection(normalized) and explicit_name_headers.intersection(normalized):
                return row_number
        return None

    def _load_master_name_by_sku(self, master_file: str | Path | None) -> dict[str, object]:
        if master_file is None:
            return {}

        wb = load_workbook(filename=master_file, data_only=True)
        ws = wb.active

        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
        header_idx = {self._normalize_header(value): idx + 1 for idx, value in enumerate(headers) if value is not None}
        sku_col = header_idx.get("SKU")
        name_col = header_idx.get("MASTER_NAME")
        if sku_col is None or name_col is None:
            return {}

        names: dict[str, object] = {}
        for row in range(2, ws.max_row + 1):
            sku_value = ws.cell(row=row, column=sku_col).value
            if sku_value is None:
                continue
            sku = str(sku_value).strip()
            if not sku:
                continue
            names[sku] = ws.cell(row=row, column=name_col).value

        return names

    @staticmethod
    def _format_master_candidate(candidate: MasterItem | None) -> str:
        if candidate is None:
            return ""

        parts = [
            candidate.category,
            candidate.brand,
            candidate.variant,
            candidate.volume,
        ]
        text = " | ".join(str(part).strip() for part in parts if str(part or "").strip())
        if not text:
            return candidate.sku
        return f"{candidate.sku} | {text}"

    def write(
        self,
        items: list[ArrivalItem],
        output_file: str | Path,
        arrival_date: str | None = None,
    ) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "RESULT"

        ws.append(["REPORT", "SVO Match Engine"])
        ws.append(["ARRIVAL_DATE", arrival_date if arrival_date is not None else ""])
        ws.append(["METADATA", f"arrival_date={arrival_date}"])
        ws.append([])

        ws.append([
            "SOURCE_NAME",
            "CATEGORY",
            "BRAND",
            "VARIANT",
            "VOLUME",
            "SKU",
            "MASTER_NAME",
            "STATUS",
        ])

        for item in items:
            ws.append([
                item.source_name,
                item.category,
                item.brand,
                item.variant,
                item.volume,
                item.sku,
                item.master_name,
                item.status,
            ])

        Path(output_file).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_file)

        # SALES: preserve calculated shipped quantity in the final workbook.
        # SalesLoader calculates item.shipped_qty = D + G + H.
        if items:
            out_wb = load_workbook(output_file)
            out_ws = out_wb.active

            headers = [str(c.value or "").strip() for c in out_ws[1]]
            shipped_col = None
            for i, header in enumerate(headers, start=1):
                if header.casefold() in ("?????????", "shipped_qty", "shipped"):
                    shipped_col = i
                    break

            if shipped_col is None:
                shipped_col = out_ws.max_column + 1
                out_ws.cell(row=1, column=shipped_col, value="????????")

            shipped_by_row = {
                int(item.row_number): (getattr(item, "shipped_qty", 0) or 0)
                for item in items
                if getattr(item, "row_number", None) is not None
            }

            for excel_row, shipped_qty in shipped_by_row.items():
                if excel_row <= out_ws.max_row:
                    out_ws.cell(
                        row=excel_row,
                        column=shipped_col,
                        value=shipped_qty,
                    )

            out_wb.save(output_file)


    def write_matched_arrival(
        self,
        source_arrival_file: str | Path,
        items: list[ArrivalItem],
        output_file: str | Path,
    ) -> None:
        """Clones ARRIVAL workbook and appends match result columns."""
        self.write_matched_document(source_arrival_file, items, output_file)

    def write_matched_document(
        self,
        source_document_file: str | Path,
        items: list[ArrivalItem],
        output_file: str | Path,
    ) -> None:
        """Clones a source workbook and appends match result columns."""
        wb = load_workbook(filename=source_document_file, data_only=False)
        ws = wb.active

        header_row = self._find_sales_header_row(ws) or 1
        start_col = ws.max_column + 1
        headers = [
            "MATCH_STATUS",
            "MATCH_SKU",
            "MATCH_MASTER_NAME",
            "MATCH_CONFIDENCE",
            "MATCH_REASONS",
            "Отгружено",
        ]

        for offset, header in enumerate(headers):
            self._set_cell_value(ws, header_row, start_col + offset, header)

        items_by_row = {item.row_number: item for item in items}
        max_row = ws.max_row

        for row_number in range(header_row + 1, max_row + 1):
            item = items_by_row.get(row_number)
            if item is None:
                continue

            reasons = ",".join(item.review_reasons) if item.review_reasons else ""
            self._set_cell_value(ws, row_number, start_col + 0, item.status)
            self._set_cell_value(ws, row_number, start_col + 1, item.sku or "")
            self._set_cell_value(ws, row_number, start_col + 2, item.master_name or "")
            self._set_cell_value(ws, row_number, start_col + 3, item.confidence)
            self._set_cell_value(ws, row_number, start_col + 4, reasons)
            ws.cell(row=row_number, column=start_col + 5, value=getattr(item, "shipped_qty", 0) or 0)

        Path(output_file).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_file)

    def write_review_report(
        self,
        items: list[ArrivalItem],
        output_file: str | Path,
    ) -> int:
        wb = Workbook()
        ws = wb.active
        ws.title = "REVIEW_REPORT"
        ws.append([
            "ROW_NUMBER",
            "SOURCE_NAME",
            "BEST_MASTER_CANDIDATE",
            "SCORE",
            "REVIEW_REASONS",
        ])

        review_count = 0
        for item in items:
            if item.status != "REVIEW":
                continue

            best_candidate = item.candidates[0] if item.candidates else None
            ws.append([
                item.row_number,
                item.source_name,
                self._format_master_candidate(best_candidate),
                item.confidence,
                ",".join(item.review_reasons),
            ])
            review_count += 1

        Path(output_file).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_file)
        return review_count

    def finalize_result_by_master(
        self,
        result_file: str | Path,
        master_items: list[MasterItem],
        master_file: str | Path | None = None,
    ) -> None:
        wb = load_workbook(filename=result_file, data_only=False)
        ws = wb.active

        sales_header_row = self._find_sales_header_row(ws)
        header_row = sales_header_row or 1

        header_cells = [ws.cell(row=header_row, column=col).value for col in range(1, ws.max_column + 1)]
        header_to_col: dict[str, int] = {}
        for idx, value in enumerate(header_cells, start=1):
            normalized = self._normalize_header(value)
            if normalized:
                header_to_col[normalized] = idx

        match_status_col = header_to_col.get("MATCH_STATUS")
        match_sku_col = header_to_col.get("MATCH_SKU")
        match_master_name_col = header_to_col.get("MATCH_MASTER_NAME")
        match_reasons_col = header_to_col.get("MATCH_REASONS")
        match_confidence_col = header_to_col.get("MATCH_CONFIDENCE")

        if match_status_col is None or match_sku_col is None:
            wb.save(result_file)
            return

        master_by_sku = {item.sku: item for item in master_items}
        master_name_by_sku = self._load_master_name_by_sku(master_file) if master_file is not None else {}

        # Final RESULT <- MASTER mapping. Values are copied directly with no text assembly.
        master_value_by_result_header = {
            "SKU": lambda item, _name: item.sku,
            "ТИП ТОВАРА": lambda item, _name: item.category,
            "БРЕНД": lambda item, _name: item.brand,
            "VARIANT": lambda item, _name: item.variant,
            "ОБЪЕМ": lambda item, _name: item.volume,
            "НАИМЕНОВАНИЕ": lambda item, name: name if str(name or "").strip() else item.master_name,
        }

        # SALES source files can have one or more leading blank/title rows.
        # The final canonical result must expose the actual business header in row 1.
        sklad1_col = header_to_col.get("СКЛАД 1")
        is_sales_result = sales_header_row is not None
        if is_sales_result and header_row > 1:
            leading_rows = header_row - 1
            if all(
                all(ws.cell(row=row_number, column=col).value is None for col in range(1, ws.max_column + 1))
                for row_number in range(1, header_row)
            ):
                ws.delete_rows(1, leading_rows)
                header_row = 1

                # Rebuild all column locations after deleting leading rows.
                header_cells = [ws.cell(row=header_row, column=col).value for col in range(1, ws.max_column + 1)]
                header_to_col = {}
                for idx, value in enumerate(header_cells, start=1):
                    normalized = self._normalize_header(value)
                    if normalized:
                        header_to_col[normalized] = idx

                match_status_col = header_to_col.get("MATCH_STATUS")
                match_sku_col = header_to_col.get("MATCH_SKU")
                match_master_name_col = header_to_col.get("MATCH_MASTER_NAME")
                match_reasons_col = header_to_col.get("MATCH_REASONS")
                match_confidence_col = header_to_col.get("MATCH_CONFIDENCE")

        if is_sales_result:
            canonical_headers = ("№", "SKU", "НАИМЕНОВАНИЕ", "ТИП ТОВАРА", "БРЕНД", "VARIANT", "ОБЪЕМ")
            for col, title in enumerate(canonical_headers, start=1):
                self._set_cell_value(ws, header_row, col, title)

            header_cells = [ws.cell(row=header_row, column=col).value for col in range(1, ws.max_column + 1)]
            header_to_col = {}
            for idx, value in enumerate(header_cells, start=1):
                normalized = self._normalize_header(value)
                if normalized:
                    header_to_col[normalized] = idx

        actionable_rows: list[int] = []
        matched_rows_by_sku: dict[str, list[list[object]]] = {}
        matched_unknown_rows: list[list[object]] = []
        review_rows: list[list[object]] = []

        max_col = ws.max_column
        for row_number in range(header_row + 1, ws.max_row + 1):
            status_value = str(ws.cell(row=row_number, column=match_status_col).value or "").strip().upper()
            if status_value == "OUT_OF_SCOPE":
                continue
            if status_value not in {"MATCH", "REVIEW"}:
                fallback_match_sku = str(ws.cell(row=row_number, column=match_sku_col).value or "").strip()
                fallback_master_name = (
                    str(ws.cell(row=row_number, column=match_master_name_col).value or "").strip()
                    if match_master_name_col is not None
                    else ""
                )
                fallback_reasons = (
                    str(ws.cell(row=row_number, column=match_reasons_col).value or "").strip()
                    if match_reasons_col is not None
                    else ""
                )
                fallback_confidence = (
                    ws.cell(row=row_number, column=match_confidence_col).value
                    if match_confidence_col is not None
                    else None
                )
                if fallback_match_sku:
                    status_value = "MATCH"
                elif fallback_master_name or fallback_reasons or fallback_confidence is not None:
                    status_value = "REVIEW"

            if status_value not in {"MATCH", "REVIEW"}:
                continue

            actionable_rows.append(row_number)
            row_values = [ws.cell(row=row_number, column=col).value for col in range(1, max_col + 1)]
            row_values[match_status_col - 1] = status_value

            if status_value == "MATCH":
                sku = str(ws.cell(row=row_number, column=match_sku_col).value or "").strip()
                master = master_by_sku.get(sku)
                if master is not None:
                    master_name = master_name_by_sku.get(master.sku)

                    for result_header, getter in master_value_by_result_header.items():
                        target_col = header_to_col.get(result_header)
                        if target_col is None:
                            continue
                        row_values[target_col - 1] = getter(master, master_name)

                    row_values[match_sku_col - 1] = master.sku

                    matched_rows_by_sku.setdefault(master.sku, []).append(row_values)
                else:
                    matched_unknown_rows.append(row_values)
            else:
                review_rows.append(row_values)

        ordered_rows: list[list[object]] = []
        for master in master_items:
            ordered_rows.extend(matched_rows_by_sku.get(master.sku, []))
        ordered_rows.extend(matched_unknown_rows)
        ordered_rows.extend(review_rows)

        match_number = 1
        for row_values in ordered_rows:
            status_value = str(row_values[match_status_col - 1] or "").strip().upper()
            number_col = header_to_col.get("№")
            if number_col is not None:
                if status_value == "MATCH":
                    row_values[number_col - 1] = match_number
                    match_number += 1
                else:
                    row_values[number_col - 1] = None

            if status_value == "REVIEW":
                # Keep original source/business columns intact for REVIEW rows.
                # Only technical match fields must stay empty except factual reasons.
                row_values[match_sku_col - 1] = ""
                if match_master_name_col is not None:
                    row_values[match_master_name_col - 1] = ""
                if match_confidence_col is not None:
                    row_values[match_confidence_col - 1] = ""

        target_rows = list(actionable_rows)
        if len(target_rows) < len(ordered_rows):
            start_row = target_rows[-1] + 1 if target_rows else 2
            target_rows.extend(range(start_row, start_row + (len(ordered_rows) - len(target_rows))))
        elif len(target_rows) > len(ordered_rows):
            for stale_row in target_rows[len(ordered_rows):]:
                for col in range(1, max_col + 1):
                    self._set_cell_value(ws, stale_row, col, None)
            target_rows = target_rows[: len(ordered_rows)]

        for target_row, row_values in zip(target_rows, ordered_rows):
            for col, value in enumerate(row_values, start=1):
                self._set_cell_value(ws, target_row, col, value)

        wb.save(result_file)

    @staticmethod
    def _is_technical_header(header: object) -> bool:
        text = str(header or "").strip().upper()
        if not text:
            return False

        exact_names = {
            "MATCH_STATUS",
            "MATCH_SKU",
            "MATCH_MASTER_NAME",
            "MATCH_CONFIDENCE",
            "MATCH_REASONS",
            "MATCH_SCORE",
            "MATCH_TYPE",
            "MASTER_INDEX",
            "REVIEW_REASON",
            "NORMALIZED_NAME",
            "SOURCE_NAME",
            "INTERNAL_ID",
        }
        if text in exact_names:
            return True

        prefixes = (
            "DEBUG",
            "TMP",
            "TECH",
            "INTERNAL",
        )
        return any(text.startswith(prefix) for prefix in prefixes)

    def _clear_technical_columns(self, worksheet) -> None:
        for col in range(1, worksheet.max_column + 1):
            header = worksheet.cell(row=1, column=col).value
            if not self._is_technical_header(header):
                continue

            for row in range(1, worksheet.max_row + 1):
                worksheet.cell(row=row, column=col, value=None)
