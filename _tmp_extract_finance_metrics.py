from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

wb = load_workbook('output/FINANCE_RESULT.xlsx', data_only=True)
ws = wb['FINANCE_LINEAGE']
summary: dict[str, object] = {}
for r in range(1, ws.max_row + 1):
    key = ws.cell(r, 1).value
    value = ws.cell(r, 2).value
    if key in {'SKU with stock', 'Finance-complete SKU', 'Inventory alias rows applied'}:
        summary[str(key)] = value

lines = [
    f"SKU with stock={summary.get('SKU with stock', '')}",
    f"Finance-complete SKU={summary.get('Finance-complete SKU', '')}",
    f"Inventory alias rows applied={summary.get('Inventory alias rows applied', '')}",
]
Path('_tmp_finance_alias_metrics.txt').write_text('\n'.join(lines), encoding='utf-8')
print('\n'.join(lines))
