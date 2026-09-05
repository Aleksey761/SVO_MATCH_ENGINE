from pathlib import Path
from openpyxl import load_workbook

path = Path('output/EXACT_MATCH_FAILURES.xlsx')
out = Path('_tmp_exact_match_failures_summary.txt')
wb = load_workbook(path, data_only=True)
ws = wb['SUMMARY']
lines = []
for r in range(2, ws.max_row + 1):
    left = ws.cell(row=r, column=1).value
    right = ws.cell(row=r, column=2).value
    lines.append(f"{left}: {right}")
out.write_text('\n'.join(lines), encoding='utf-8')
