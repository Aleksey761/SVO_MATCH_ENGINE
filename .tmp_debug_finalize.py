from pathlib import Path
from openpyxl import load_workbook
from svo.loader import Loader
from svo.sales_loader import SalesLoader
from svo.normalizer import Normalizer
from svo.matcher import Matcher
from svo.reporter import Reporter

base = Path('.')
loader = Loader()
master_file, arrival_file, sales_file, arrival_date, sales_date = loader.discover_workbooks(base / 'data', require_sales=True)
master = loader.load_master(master_file)
items = SalesLoader().load(sales_file)
normalizer = Normalizer()
for item in items:
    normalizer.normalize(item)
Matcher(master).match_all(items)

pre = base / 'output' / '_tmp_pre_finalize_debug.xlsx'
post = base / 'output' / '_tmp_post_finalize_debug.xlsx'
Reporter().write_matched_document(sales_file, items, pre)
post.write_bytes(pre.read_bytes())
Reporter().finalize_result_by_master(post, master, master_file=master_file)

ws_pre = load_workbook(pre, data_only=True).active
pre_headers = [str(ws_pre.cell(1, c).value or '').strip().upper() for c in range(1, ws_pre.max_column + 1)]
status_col = pre_headers.index('MATCH_STATUS') + 1
sku_col = pre_headers.index('MATCH_SKU') + 1

print('PRE statuses for 201-204:')
for r in (201, 202, 203, 204):
    print(r, ws_pre.cell(r, status_col).value, ws_pre.cell(r, sku_col).value, ws_pre.cell(r, 2).value, ws_pre.cell(r, 3).value)

ws_post = load_workbook(post, data_only=True).active
print('POST rows 195-204:')
for r in range(195, 205):
    print(r, ws_post.cell(r, 1).value, ws_post.cell(r, 2).value, ws_post.cell(r, 3).value)

pre.unlink(missing_ok=True)
post.unlink(missing_ok=True)
