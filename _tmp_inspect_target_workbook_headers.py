from openpyxl import load_workbook


for path in ["data/MASTER.xlsx", "output/MASTER_DATASET.xlsx", "output/SALES_MATCH_FINAL.xlsx"]:
    wb = load_workbook(path, data_only=True)
    print("FILE|" + path)
    print("SHEETS|" + "|".join(wb.sheetnames))
    for ws in wb.worksheets:
        headers = ["" if ws.cell(1, c).value is None else str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
        print("HEADERS|" + ws.title + "|" + "|".join(headers))