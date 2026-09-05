from openpyxl import load_workbook
from pathlib import Path

path = Path(r"output\MASTER_DATASET.xlsx")

if not path.exists():
    raise FileNotFoundError(f"MASTER not found: {path}")

wb = load_workbook(path, data_only=True)
ws = wb.active

print("MASTER SHEET =", ws.title)
print("ROWS =", ws.max_row)
print()

count = 0

for row in ws.iter_rows(min_row=2, values_only=True):
    sku = row[0]
    category = row[1]
    brand = row[2]
    variant = row[3]
    volume = row[4]
    aroma = row[5]
    master_name = row[9]

    if str(brand or "").strip().upper() == "GILAR":
        count += 1
        print(
            f"{sku} | {category} | {brand} | {variant} | "
            f"{volume} | {aroma} | {master_name}"
        )

print()
print("GILAR ROWS =", count)
