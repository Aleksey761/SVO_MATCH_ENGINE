from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
CORRECTION_MAP_PATH = OUTPUT_DIR / "NAME_CORRECTION_MAP.xlsx"
NORMALIZATION_REPORT_PATH = OUTPUT_DIR / "NAME_NORMALIZATION_REPORT.xlsx"
MASTER_DATASET_PATH = OUTPUT_DIR / "MASTER_DATASET.xlsx"
AUDIT_OUTPUT_PATH = OUTPUT_DIR / "NAME_CORRECTION_AUDIT.xlsx"
SUMMARY_OUTPUT_PATH = OUTPUT_DIR / "NAME_CORRECTION_AUDIT_SUMMARY.txt"


@dataclass(frozen=True)
class MasterRecord:
    sku: str
    master_name: str


def _clean(value: object) -> str:
    return str(value or "").strip()


def _normalized_text(value: str) -> str:
    text = _clean(value).upper().replace("Ё", "Е")
    for token in [",", ".", ";", ":", "(", ")", "[", "]", "{", "}", "\"", "'", "/", "\\", "-", "_"]:
        text = text.replace(token, " ")
    return " ".join(text.split())


def _load_sheet_rows(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, object]]]:
    wb = load_workbook(path, data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    headers = [_clean(cell.value) for cell in ws[1]]
    rows: list[dict[str, object]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        payload = {headers[index]: values[index] for index in range(len(headers))}
        if any(_clean(value) for value in payload.values()):
            rows.append(payload)
    return headers, rows


def _load_master_records(path: Path) -> tuple[list[MasterRecord], dict[str, set[str]], dict[str, set[str]]]:
    _, rows = _load_sheet_rows(path)
    records: list[MasterRecord] = []
    sku_to_names: dict[str, set[str]] = defaultdict(set)
    name_to_skus: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        sku = _clean(row.get("SKU"))
        master_name = _clean(row.get("MASTER_NAME"))
        if not sku and not master_name:
            continue
        records.append(MasterRecord(sku=sku, master_name=master_name))
        if sku:
            sku_to_names[sku].add(master_name)
        if master_name:
            name_to_skus[master_name].add(sku)
    return records, sku_to_names, name_to_skus


def _score_candidate(source_name: str, master_name: str) -> float:
    left = _normalized_text(source_name)
    right = _normalized_text(master_name)
    if not left or not right:
        return 0.0
    sequence = SequenceMatcher(None, left, right).ratio()
    left_tokens = set(left.split())
    right_tokens = set(right.split())
    union = left_tokens | right_tokens
    token_overlap = len(left_tokens & right_tokens) / len(union) if union else 0.0
    return max(sequence, 0.65 * sequence + 0.35 * token_overlap)


def _classify_unresolved(source_name: str, ranked: list[tuple[float, MasterRecord]]) -> tuple[str, str, str]:
    if not ranked:
        return "", "", "new_product"

    top_score, top_record = ranked[0]
    second_score = ranked[1][0] if len(ranked) > 1 else 0.0
    if top_score >= 0.88:
        return top_record.master_name, top_record.sku, "typo_fix_needed"
    if top_score >= 0.70 and abs(top_score - second_score) <= 0.03:
        return top_record.master_name, top_record.sku, "ambiguous"
    if top_score >= 0.70:
        return top_record.master_name, top_record.sku, "manual_check"

    source_tokens = set(_normalized_text(source_name).split())
    if source_tokens and any(token in {"SVO", "GILAR", "GRASS", "SYNERGETIC"} for token in source_tokens):
        return "", "", "missing_master"
    return "", "", "new_product"


def build_audit() -> dict[str, int]:
    _, correction_rows = _load_sheet_rows(CORRECTION_MAP_PATH, "NAME_CORRECTION_MAP")
    _, corrected_rows = _load_sheet_rows(NORMALIZATION_REPORT_PATH, "CORRECTED_NAMES")
    _, unresolved_rows = _load_sheet_rows(NORMALIZATION_REPORT_PATH, "UNRESOLVED_NAMES")
    master_records, sku_to_names, name_to_skus = _load_master_records(MASTER_DATASET_PATH)

    wrong_name_to_pairs: dict[str, set[tuple[str, str]]] = defaultdict(set)
    wrong_name_to_rule_signatures: dict[str, set[tuple[str, str, str, str]]] = defaultdict(set)
    approved: list[list[object]] = []
    conflicts: list[list[object]] = []
    invalid: list[list[object]] = []

    for row in correction_rows:
        wrong_name = _clean(row.get("Wrong_Name"))
        correct_name = _clean(row.get("Correct_MASTER_NAME"))
        sku = _clean(row.get("SKU"))
        source = _clean(row.get("Source"))
        rule_type = _clean(row.get("Rule_Type"))
        evidence = _clean(row.get("Evidence"))
        status = _clean(row.get("Status"))

        wrong_name_to_pairs[wrong_name].add((correct_name, sku))
        wrong_name_to_rule_signatures[wrong_name].add((correct_name, sku, source, rule_type))

    conflict_wrong_names = {
        wrong_name
        for wrong_name, pairs in wrong_name_to_pairs.items()
        if len({pair[1] for pair in pairs if pair[1]}) > 1 or len(pairs) > 1
    }

    for wrong_name, signatures in wrong_name_to_rule_signatures.items():
        distinct_targets = {(correct_name, sku) for correct_name, sku, _source, _rule_type in signatures}
        if len(distinct_targets) > 1:
            conflict_wrong_names.add(wrong_name)

    corrected_index = {
        (_clean(row.get("BEFORE_NAME")), _clean(row.get("MASTER_NAME")), _clean(row.get("SKU")), _clean(row.get("Source"))): row
        for row in corrected_rows
    }

    for row in correction_rows:
        wrong_name = _clean(row.get("Wrong_Name"))
        correct_name = _clean(row.get("Correct_MASTER_NAME"))
        sku = _clean(row.get("SKU"))
        source = _clean(row.get("Source"))
        rule_type = _clean(row.get("Rule_Type"))
        evidence = _clean(row.get("Evidence"))
        status = _clean(row.get("Status"))

        notes: list[str] = []
        master_exists = correct_name in name_to_skus
        sku_exists = sku in sku_to_names
        sku_matches_name = master_exists and sku_exists and correct_name in sku_to_names.get(sku, set())
        corrected_seen = (wrong_name, correct_name, sku, source) in corrected_index
        if master_exists:
            notes.append("master_exists")
        else:
            notes.append("missing_master_name")
        if sku_exists:
            notes.append("sku_exists")
        else:
            notes.append("missing_sku")
        if sku_matches_name:
            notes.append("sku_master_match")
        else:
            notes.append("sku_master_mismatch")
        if corrected_seen:
            notes.append("seen_in_corrected_report")

        record = [wrong_name, correct_name, sku, source, rule_type, evidence, status, "; ".join(notes)]
        if wrong_name in conflict_wrong_names:
            conflicts.append(record)
        elif not (master_exists and sku_exists and sku_matches_name):
            invalid.append(record)
        else:
            approved.append(record)

    unresolved_review: list[list[object]] = []
    for row in unresolved_rows:
        source_name = _clean(row.get("BEFORE_NAME")) or _clean(row.get("AFTER_NAME"))
        source = _clean(row.get("Source"))
        reason = _clean(row.get("CHANGE_REASON"))
        ranked = sorted(
            ((_score_candidate(source_name, record.master_name), record) for record in master_records),
            key=lambda item: item[0],
            reverse=True,
        )[:3]
        possible_master_name, possible_sku, decision = _classify_unresolved(source_name, ranked)
        unresolved_review.append([
            source_name,
            source,
            possible_master_name,
            possible_sku,
            reason,
            decision,
        ])

    wb = Workbook()

    ws_approved = wb.active
    ws_approved.title = "APPROVED_CORRECTIONS"
    ws_approved.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status", "CheckNotes"])
    for row in approved:
        ws_approved.append(row)

    ws_conflicts = wb.create_sheet("CONFLICTS")
    ws_conflicts.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status", "CheckNotes"])
    for row in conflicts:
        ws_conflicts.append(row)

    ws_invalid = wb.create_sheet("INVALID")
    ws_invalid.append(["Wrong_Name", "Correct_MASTER_NAME", "SKU", "Source", "Rule_Type", "Evidence", "Status", "CheckNotes"])
    for row in invalid:
        ws_invalid.append(row)

    ws_summary = wb.create_sheet("SUMMARY")
    ws_summary.append(["Metric", "Value"])
    summary_rows = [
        ("Total corrections", len(correction_rows)),
        ("Valid", len(approved)),
        ("Conflicts", len(conflicts)),
        ("Invalid", len(invalid)),
        ("Unresolved review rows", len(unresolved_review)),
    ]
    for row in summary_rows:
        ws_summary.append(list(row))

    ws_unresolved = wb.create_sheet("UNRESOLVED_REVIEW")
    ws_unresolved.append(["SOURCE_NAME", "SOURCE", "Possible_MASTER_NAME", "Possible_SKU", "Reason", "Decision"])
    for row in unresolved_review:
        ws_unresolved.append(row)

    for sheet in wb.worksheets:
        for column_cells in sheet.columns:
            width = max(len(_clean(cell.value)) for cell in column_cells) if column_cells else 0
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 12), 60)

    AUDIT_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(AUDIT_OUTPUT_PATH)

    decision_counts = Counter(_clean(row[-1]) for row in unresolved_review)
    summary_lines = [
        f"AUDIT_FILE={AUDIT_OUTPUT_PATH}",
        f"TOTAL_CORRECTIONS={len(correction_rows)}",
        f"VALID={len(approved)}",
        f"CONFLICTS={len(conflicts)}",
        f"INVALID={len(invalid)}",
        f"UNRESOLVED_ROWS={len(unresolved_review)}",
    ]
    for key in sorted(decision_counts):
        summary_lines.append(f"UNRESOLVED_DECISION|{key}|{decision_counts[key]}")
    SUMMARY_OUTPUT_PATH.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    print(f"AUDIT_FILE={AUDIT_OUTPUT_PATH}")
    print(f"TOTAL_CORRECTIONS={len(correction_rows)}")
    print(f"VALID={len(approved)}")
    print(f"CONFLICTS={len(conflicts)}")
    print(f"INVALID={len(invalid)}")
    print(f"UNRESOLVED_ROWS={len(unresolved_review)}")
    for key in sorted(decision_counts):
        print(f"UNRESOLVED_DECISION|{key}|{decision_counts[key]}")

    return {
        "total": len(correction_rows),
        "valid": len(approved),
        "conflicts": len(conflicts),
        "invalid": len(invalid),
        "unresolved": len(unresolved_review),
    }


if __name__ == "__main__":
    build_audit()