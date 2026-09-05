from pathlib import Path

from openpyxl import load_workbook

from svo.loader import Loader
from svo.matcher import Matcher
from svo.normalizer import Normalizer
from svo.reporter import Reporter
from svo.sales_loader import SalesLoader

base = Path('.')
loader = Loader()
master_file, _arrival_file, sales_file, _arrival_date, _sales_date = loader.discover_workbooks(base / 'data', require_sales=True)
master = loader.load_master(master_file)
items = SalesLoader().load(sales_file)
normalizer = Normalizer()
for item in items:
    normalizer.normalize(item)
Matcher(master).match_all(items)

tmp = base / 'output' / '_tmp_pre.xlsx'
Reporter().write_matched_document(sales_file, items, tmp)
ws = load_workbook(tmp, data_only=True).active
headers = [str(ws.cell(1, c).value or '').strip().upper() for c in range(1, ws.max_column + 1)]
status_col = headers.index('MATCH_STATUS') + 1
sku_col = headers.index('MATCH_SKU') + 1

missing = []
for r in range(2, ws.max_row + 1):
    status = str(ws.cell(r, status_col).value or '').strip().upper()
    if status not in {'MATCH', 'REVIEW'}:
        row_values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(value is not None and str(value).strip() != '' for value in row_values):
            missing.append((r, ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, sku_col).value, status))

print('missing_count', len(missing))
print('missing_rows', missing)

tmp.unlink(missing_ok=True)
