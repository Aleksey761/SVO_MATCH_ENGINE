from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook


FEATURE_PATH = Path("output/FEATURE_EXTRACTION_V2.xlsx")
MASTER_PATH = Path("output/MASTER_DATASET.xlsx")
OUTPUT_PATH = Path("output/VARIANT_TRUTH_CHECK.xlsx")


FAILED_ROWS = [75, 102, 110, 111, 117, 118, 128, 176]


ROW_CLASSIFICATION = {
    75: {
        "Result": "CONFIRMED_ALIAS",
        "SupplierToken": "РОМАШКА Daisy",
        "MasterVariant": "Baby",
        "SKU": "SKU-291",
        "Evidence": "РОМАШКА Daisy -> Baby -> SKU-291",
        "ConfirmedRule": "РОМАШКА Daisy -> Baby",
    },
    102: {
        "Result": "CONFIRMED_UNIQUE",
        "SupplierToken": "для цв. и бел",
        "MasterVariant": "Color",
        "SKU": "SKU-212",
        "Evidence": "для цв. и бел -> Color -> SKU-212",
        "ConfirmedRule": "для цв. и бел -> Color",
    },
    110: {
        "Result": "NO_EVIDENCE",
        "SupplierToken": "для белых",
        "MasterVariant": "",
        "SKU": "",
        "Evidence": "для белых -> no MASTER variant in this ProductType+Brand+Volume family -> no SKU",
        "ConfirmedRule": "",
    },
    111: {
        "Result": "CONFIRMED_UNIQUE",
        "SupplierToken": "для цветных",
        "MasterVariant": "Color",
        "SKU": "SKU-223",
        "Evidence": "для цветных -> Color -> SKU-223",
        "ConfirmedRule": "для цветных -> Color",
    },
    117: {
        "Result": "CONFIRMED_ALIAS",
        "SupplierToken": "ПОДСНЕЖНИК",
        "MasterVariant": "Snowdrop",
        "SKU": "SKU-235",
        "Evidence": "ПОДСНЕЖНИК -> Snowdrop -> SKU-235",
        "ConfirmedRule": "ПОДСНЕЖНИК -> Snowdrop",
    },
    118: {
        "Result": "CONFIRMED_UNIQUE",
        "SupplierToken": "для цв. и бел",
        "MasterVariant": "Color",
        "SKU": "SKU-234",
        "Evidence": "для цв. и бел -> Color -> SKU-234",
        "ConfirmedRule": "для цв. и бел -> Color",
    },
    128: {
        "Result": "CONFIRMED_UNIQUE",
        "SupplierToken": "д\\цветного",
        "MasterVariant": "Color",
        "SKU": "SKU-245",
        "Evidence": "д\\цветного -> Color -> SKU-245",
        "ConfirmedRule": "д\\цветного -> Color",
    },
    176: {
        "Result": "CONFIRMED_ALIAS",
        "SupplierToken": "MENTOL",
        "MasterVariant": "Mint",
        "SKU": "SKU-018",
        "Evidence": "MENTOL -> Mint -> SKU-018",
        "ConfirmedRule": "MENTOL -> Mint",
    },
}


def norm(value: object) -> str:
    return str(value or "").strip()


def row_map(ws):
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rownum = int(row[0])
        rows[rownum] = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
    return rows


def same_pattern_elsewhere(master_ws, variant: str) -> str:
    if not variant:
        return "NO"
    target = variant.lower()
    count = 0
    for row in master_ws.iter_rows(min_row=2, values_only=True):
        if len(row) > 3 and norm(row[3]).lower() == target:
            count += 1
            if count > 1:
                return "YES"
    return "YES" if count == 1 else "NO"


def main() -> None:
    feature_wb = load_workbook(FEATURE_PATH, data_only=True, read_only=True)
    master_wb = load_workbook(MASTER_PATH, data_only=True, read_only=True)

    features = row_map(feature_wb["FEATURES_V2"])
    raw_rows = row_map(feature_wb["RAW"])
    master_comp = defaultdict(list)
    comp_ws = feature_wb["MASTER_COMPARISON"]
    for row in comp_ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        master_comp[int(row[0])].append(
            {
                "SKU": norm(row[1]),
                "MASTER_NAME": norm(row[2]),
                "MASTER_Variant": norm(row[6]),
            }
        )

    out_wb = Workbook()
    ws_cases = out_wb.active
    ws_cases.title = "CASES"
    ws_master = out_wb.create_sheet("MASTER_OPTIONS")
    ws_evidence = out_wb.create_sheet("EVIDENCE")
    ws_summary = out_wb.create_sheet("SUMMARY")

    ws_cases.append([
        "ROW",
        "SupplierName",
        "ProductType",
        "Brand",
        "Volume",
        "CurrentParsedVariant",
        "Result",
        "SupplierToken",
        "ExpectedVariantFromMASTER",
        "Evidence",
    ])

    ws_master.append([
        "ROW",
        "SKU",
        "MASTER_NAME",
        "MASTER Aroma/Variant",
        "SupplierName",
        "ProductType",
        "Brand",
        "Volume",
    ])

    ws_evidence.append([
        "ROW",
        "Supplier token",
        "MASTER Variant",
        "SKU",
        "Evidence",
    ])

    results = Counter()
    confirmed_rules = []
    variant_presence = []

    for rownum in FAILED_ROWS:
        feature = features[rownum]
        raw = raw_rows[rownum]
        classification = ROW_CLASSIFICATION[rownum]
        results[classification["Result"]] += 1
        if classification["ConfirmedRule"]:
            confirmed_rules.append(classification["ConfirmedRule"])

        variant_presence.append((rownum, same_pattern_elsewhere(master_wb.active, classification["MasterVariant"]), classification["MasterVariant"]))

        ws_cases.append([
            rownum,
            norm(raw.get("SupplierName")),
            norm(feature.get("ProductType")),
            norm(feature.get("Brand")),
            norm(feature.get("Volume")),
            norm(feature.get("Aroma/Variant")),
            classification["Result"],
            classification["SupplierToken"],
            classification["MasterVariant"],
            classification["Evidence"],
        ])

        for candidate in master_comp[rownum]:
            ws_master.append([
                rownum,
                candidate["SKU"],
                candidate["MASTER_NAME"],
                candidate["MASTER_Variant"],
                norm(raw.get("SupplierName")),
                norm(feature.get("ProductType")),
                norm(feature.get("Brand")),
                norm(feature.get("Volume")),
            ])

        ws_evidence.append([
            rownum,
            classification["SupplierToken"],
            classification["MasterVariant"] or "NO_MATCH",
            classification["SKU"] or "NO_SKU",
            classification["Evidence"],
        ])

    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["ROWS", len(FAILED_ROWS)])
    ws_summary.append(["CONFIRMED_UNIQUE", results.get("CONFIRMED_UNIQUE", 0)])
    ws_summary.append(["CONFIRMED_ALIAS", results.get("CONFIRMED_ALIAS", 0)])
    ws_summary.append(["AMBIGUOUS", results.get("AMBIGUOUS", 0)])
    ws_summary.append(["NO_EVIDENCE", results.get("NO_EVIDENCE", 0)])
    ws_summary.append([])
    ws_summary.append(["CONFIRMED_RULES"])
    for rule in list(dict.fromkeys(confirmed_rules)):
        ws_summary.append([rule])

    ws_summary.append([])
    ws_summary.append(["Variant", "Occurs Elsewhere In MASTER?", "Row"])
    for rownum, occurs, variant in variant_presence:
        ws_summary.append([variant or "", occurs, rownum])

    for ws in (ws_cases, ws_master, ws_evidence, ws_summary):
        for column_cells in ws.columns:
            values = [cell.value for cell in column_cells if cell.value is not None]
            if not values:
                continue
            width = min(max(len(str(value)) for value in values) + 2, 80)
            ws.column_dimensions[column_cells[0].column_letter].width = width

    out_wb.save(OUTPUT_PATH)

    print(f"ROWS = {len(FAILED_ROWS)}")
    print(f"CONFIRMED_UNIQUE = {results.get('CONFIRMED_UNIQUE', 0)}")
    print(f"CONFIRMED_ALIAS = {results.get('CONFIRMED_ALIAS', 0)}")
    print(f"AMBIGUOUS = {results.get('AMBIGUOUS', 0)}")
    print(f"NO_EVIDENCE = {results.get('NO_EVIDENCE', 0)}")
    print("CONFIRMED_RULES =")
    for rule in dict.fromkeys(confirmed_rules):
        print(rule)


if __name__ == "__main__":
    main()