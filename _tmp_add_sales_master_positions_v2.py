from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent

MASTER_SOURCE = ROOT / "data" / "MASTER.xlsx"
MASTER_V2 = ROOT / "output" / "MASTER_V2.xlsx"
DATASET_SOURCE = ROOT / "output" / "MASTER_DATASET.xlsx"
SALES_SOURCE = ROOT / "output" / "SALES_MATCH_FINAL.xlsx"
SALES_V2 = ROOT / "output" / "SALES_MATCH_FINAL_V2.xlsx"


@dataclass(frozen=True)
class Position:
    sku: str
    category: str
    brand: str
    variant: str
    volume: str
    master_name: str


POSITIONS = (
    Position("SKU-293", "ЖМС", "SVO", "BABY", "1 л", "ЖМС SVO Baby 1 л"),
    Position(
        "SKU-294",
        "Порошок",
        "BOSSFIX",
        "BABY",
        "400 г",
        "Порошок стиральный BOSSFIX Baby 400 г",
    ),
    Position(
        "SKU-296",
        "Порошок",
        "SVO",
        "BABY",
        "2,7 кг",
        "Порошок стиральный SVO Baby 2,7 кг",
    ),
)

SALES_UPDATES = {
    181: ("MATCH", "SKU-293", "ЖМС SVO Baby 1 л"),
    190: ("MATCH", "SKU-294", "Порошок стиральный BOSSFIX Baby 400 г"),
    219: ("MATCH", "SKU-296", "Порошок стиральный SVO Baby 2,7 кг"),
}


def normalize(value: object) -> str:
    return "" if value is None else str(value).strip()


def header_index_map(ws) -> dict[str, int]:
    return {normalize(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}


def require_headers(ws, names: list[str]) -> dict[str, int]:
    mapping = header_index_map(ws)
    missing = [name for name in names if name not in mapping]
    if missing:
        raise RuntimeError(f"Missing headers in {ws.title}: {', '.join(missing)}")
    return {name: mapping[name] for name in names}


def collect_column_values(ws, column: int) -> list[str]:
    return [normalize(ws.cell(row, column).value) for row in range(2, ws.max_row + 1)]


def first_empty_row(ws) -> int:
    row = ws.max_row + 1
    while row > 1:
        if any(ws.cell(row - 1, col).value is not None for col in range(1, ws.max_column + 1)):
            break
        row -= 1
    return row


def append_master_rows(ws, positions: tuple[Position, ...]) -> None:
    columns = require_headers(ws, ["SKU", "Категория", "Бренд", "Аромат", "Объем"])
    row = first_empty_row(ws)
    for pos in positions:
        ws.cell(row, columns["SKU"]).value = pos.sku
        ws.cell(row, columns["Категория"]).value = pos.category
        ws.cell(row, columns["Бренд"]).value = pos.brand
        ws.cell(row, columns["Аромат"]).value = pos.variant
        ws.cell(row, columns["Объем"]).value = pos.volume
        row += 1


def append_dataset_rows(ws, positions: tuple[Position, ...]) -> None:
    columns = require_headers(
        ws,
        ["SKU", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "AROMA", "C7", "C8", "C9", "MASTER_NAME"],
    )
    row = first_empty_row(ws)
    for pos in positions:
        ws.cell(row, columns["SKU"]).value = pos.sku
        ws.cell(row, columns["CATEGORY"]).value = pos.category
        ws.cell(row, columns["BRAND"]).value = pos.brand
        ws.cell(row, columns["VARIANT"]).value = pos.variant
        ws.cell(row, columns["VOLUME"]).value = pos.volume
        ws.cell(row, columns["AROMA"]).value = pos.variant
        ws.cell(row, columns["C7"]).value = ""
        ws.cell(row, columns["C8"]).value = ""
        ws.cell(row, columns["C9"]).value = ""
        ws.cell(row, columns["MASTER_NAME"]).value = pos.master_name
        row += 1


def update_sales_rows(ws) -> None:
    columns = require_headers(ws, ["MATCH_STATUS", "MATCH_SKU", "MATCH_MASTER_NAME"])
    for target_row, (match_status, match_sku, match_master_name) in SALES_UPDATES.items():
        ws.cell(target_row, columns["MATCH_STATUS"]).value = match_status
        ws.cell(target_row, columns["MATCH_SKU"]).value = match_sku
        ws.cell(target_row, columns["MATCH_MASTER_NAME"]).value = match_master_name


def count_duplicates(values: list[str]) -> dict[str, int]:
    counts = Counter(value for value in values if value)
    return {value: count for value, count in counts.items() if count > 1}


def run_prechecks() -> tuple[list[str], dict[str, str]]:
    results: list[str] = []
    details: dict[str, str] = {}

    master_wb = load_workbook(MASTER_SOURCE, data_only=True)
    master_ws = master_wb.active
    master_cols = require_headers(master_ws, ["SKU"])

    dataset_wb = load_workbook(DATASET_SOURCE, data_only=True)
    dataset_ws = dataset_wb.active
    dataset_cols = require_headers(dataset_ws, ["SKU", "MASTER_NAME"])

    sales_wb = load_workbook(SALES_SOURCE, data_only=True)
    sales_ws = sales_wb.active
    require_headers(sales_ws, ["MATCH_STATUS", "MATCH_SKU", "MATCH_MASTER_NAME"])

    master_skus = collect_column_values(master_ws, master_cols["SKU"])
    dataset_skus = collect_column_values(dataset_ws, dataset_cols["SKU"])
    dataset_names = collect_column_values(dataset_ws, dataset_cols["MASTER_NAME"])

    for sku in ("SKU-293", "SKU-294", "SKU-296"):
        free = sku not in master_skus and sku not in dataset_skus
        results.append(f"PRECHECK|{sku}_FREE|{'PASS' if free else 'FAIL'}")
        details[f"{sku}_FREE"] = "PASS" if free else "FAIL"

    sku_295_master = sum(1 for sku in master_skus if sku == "SKU-295")
    sku_295_dataset = sum(1 for sku in dataset_skus if sku == "SKU-295")
    sku_295_unchanged = sku_295_master == 1 and sku_295_dataset == 1
    results.append(f"PRECHECK|SKU-295_UNCHANGED|{'PASS' if sku_295_unchanged else 'FAIL'}|MASTER={sku_295_master}|DATASET={sku_295_dataset}")
    details["SKU-295_UNCHANGED"] = "PASS" if sku_295_unchanged else "FAIL"

    for row_number in SALES_UPDATES:
        ok = row_number <= sales_ws.max_row
        results.append(f"PRECHECK|SALES_ROW_{row_number}_EXISTS|{'PASS' if ok else 'FAIL'}")
        details[f"SALES_ROW_{row_number}_EXISTS"] = "PASS" if ok else "FAIL"

    master_post_sku_dupes = count_duplicates(master_skus + [pos.sku for pos in POSITIONS])
    dataset_post_sku_dupes = count_duplicates(dataset_skus + [pos.sku for pos in POSITIONS])
    sku_dupe_ok = not master_post_sku_dupes and not dataset_post_sku_dupes
    results.append(
        f"PRECHECK|NO_DUPLICATE_SKU_AFTER_INSERT|{'PASS' if sku_dupe_ok else 'FAIL'}|MASTER_COUNT={len(master_post_sku_dupes)}|DATASET_COUNT={len(dataset_post_sku_dupes)}"
    )
    details["NO_DUPLICATE_SKU_AFTER_INSERT"] = "PASS" if sku_dupe_ok else "FAIL"

    post_name_dupes = count_duplicates(dataset_names + [pos.master_name for pos in POSITIONS])
    name_dupe_ok = not post_name_dupes
    results.append(
        f"PRECHECK|NO_DUPLICATE_MASTER_NAME_AFTER_INSERT|{'PASS' if name_dupe_ok else 'FAIL'}|COUNT={len(post_name_dupes)}"
    )
    details["NO_DUPLICATE_MASTER_NAME_AFTER_INSERT"] = "PASS" if name_dupe_ok else "FAIL"

    return results, details


def compute_sales_metrics(ws) -> list[str]:
    columns = require_headers(ws, ["MATCH_STATUS", "MATCH_SKU", "MATCH_MASTER_NAME"])
    total_rows = max(ws.max_row - 1, 0)
    match_total = 0
    review_total = 0
    empty_sku_match = 0
    manual_mapping_mismatch = 0
    status_col = columns["MATCH_STATUS"]
    sku_col = columns["MATCH_SKU"]
    master_name_col = columns["MATCH_MASTER_NAME"]
    for row in range(2, ws.max_row + 1):
        status = normalize(ws.cell(row, status_col).value)
        sku = normalize(ws.cell(row, sku_col).value)
        master_name = normalize(ws.cell(row, master_name_col).value)
        if status == "MATCH":
            match_total += 1
            if not sku:
                empty_sku_match += 1
        if status == "REVIEW":
            review_total += 1
        if (sku and not master_name) or (master_name and not sku):
            manual_mapping_mismatch += 1

    metrics = [
        f"TOTAL_DATA_ROWS|{total_rows}",
        f"MATCH_TOTAL|{match_total}",
        f"REVIEW_TOTAL|{review_total}",
        f"EMPTY_SKU_MATCH|{empty_sku_match}",
        f"MANUAL_MAPPING_MISMATCH|{manual_mapping_mismatch}",
    ]

    for target_row in sorted(SALES_UPDATES):
        status = normalize(ws.cell(target_row, status_col).value)
        sku = normalize(ws.cell(target_row, sku_col).value)
        master_name = normalize(ws.cell(target_row, master_name_col).value)
        metrics.append(f"ROW_STATE|{target_row}|{status}|{sku}|{master_name}")

    return metrics


def main() -> int:
    precheck_lines, precheck_state = run_prechecks()
    for line in precheck_lines:
        print(line)

    if any(value != "PASS" for value in precheck_state.values()):
        print("WRITE_STATUS|SKIPPED")
        return 1

    master_wb = load_workbook(MASTER_SOURCE)
    master_ws = master_wb.active
    append_master_rows(master_ws, POSITIONS)
    master_wb.save(MASTER_V2)

    dataset_wb = load_workbook(DATASET_SOURCE)
    dataset_ws = dataset_wb.active
    append_dataset_rows(dataset_ws, POSITIONS)
    dataset_wb.save(DATASET_SOURCE)

    sales_wb = load_workbook(SALES_SOURCE)
    sales_ws = sales_wb.active
    update_sales_rows(sales_ws)
    sales_wb.save(SALES_V2)

    print("WRITE_STATUS|DONE")
    print(f"WRITTEN_PATH|{MASTER_V2}")
    print(f"WRITTEN_PATH|{DATASET_SOURCE}")
    print(f"WRITTEN_PATH|{SALES_V2}")

    for line in compute_sales_metrics(load_workbook(SALES_V2, data_only=True).active):
        print(line)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())