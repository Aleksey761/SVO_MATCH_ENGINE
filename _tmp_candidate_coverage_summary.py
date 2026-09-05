from pathlib import Path
from openpyxl import load_workbook

path = Path('output/CANDIDATE_COVERAGE_ANALYSIS.xlsx')
wb = load_workbook(path, data_only=True)
ws = wb['SUMMARY']
for r in range(2, ws.max_row + 1):
    print(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
