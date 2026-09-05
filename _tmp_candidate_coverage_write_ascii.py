from pathlib import Path
from openpyxl import load_workbook

path = Path('output/CANDIDATE_COVERAGE_ANALYSIS.xlsx')
out = Path('_tmp_candidate_coverage_summary_ascii.txt')
wb = load_workbook(path, data_only=True)
ws = wb['SUMMARY']
lines = []
for r in range(2, ws.max_row + 1):
    lines.append(f"{ws.cell(row=r, column=1).value}: {ws.cell(row=r, column=2).value}")
out.write_text('\n'.join(lines), encoding='utf-8')
