
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from tests.test_dataset_builder import _create_master_workbook, _create_match_workbook
from svo.dataset_builder import DatasetBuilder

with TemporaryDirectory() as tmp:
    root=Path(tmp)
    data_dir=root/"data"; output_dir=root/"output"
    data_dir.mkdir(); output_dir.mkdir()
    _create_master_workbook(data_dir/"MASTER.xlsx", [{"sku":"SKU-001","category":"C1","brand":"B1","variant":"V1","volume":"1L"}])
    match_path=_create_match_workbook(output_dir/"ARRIVAL_MATCH_2026-07-15.xlsx",
                                      [{"sku":"SKU-999","warehouse":[5,0,0,0,0,0,0,0,0,0]}])
    wb=load_workbook(match_path, data_only=True)
    ws=wb.active
    print("MASTER path:", data_dir/"MASTER.xlsx")
    print("MATCH path:", match_path)
    print("MATCH sheet:", ws.title)
    print("MATCH dimensions:", ws.max_row, ws.max_column)
    print("MATCH header:", [ws.cell(1,c).value for c in range(1, ws.max_column+1)])
    count=0
    samples=[]
    for r in range(2, ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1, min(ws.max_column, 16)+1)]
        if vals[11] is not None and str(vals[11]).strip().upper()=="MATCH":
            count+=1
            if len(samples)<5: samples.append(vals)
    print("MATCH rows with MATCH_STATUS:", count)
    print("Samples:", samples)
    b=DatasetBuilder()
    master=b._load_authoritative_master(data_dir/"MASTER.xlsx")
    print("Authoritative MASTER count:", len(master), [m.sku for m in master])
    print("Collected ARRIVAL files:", b._collect_match_files(None, "ARRIVAL_MATCH_*.xlsx", output_dir))
    result=b.build(input_dir=data_dir, output_file=output_dir/"MASTER_DATASET.xlsx", print_reports=False)
    print("BUILDER validation:", result["validation"])
