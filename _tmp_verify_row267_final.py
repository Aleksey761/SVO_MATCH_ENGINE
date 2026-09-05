from openpyxl import load_workbook

ws = load_workbook("output/SALES_MATCH_FINAL.xlsx", data_only=True).active
headers = ["" if ws.cell(1, c).value is None else str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
idx = {h: i + 1 for i, h in enumerate(headers)}
status_col = idx["MATCH_STATUS"]
sku_col = idx["MATCH_SKU"]
master_col = idx["MATCH_MASTER_NAME"]

r = 267
print("ROW267_VERIFY|" + str(r) + "|" + str(ws.cell(r, status_col).value) + "|" + str(ws.cell(r, sku_col).value) + "|" + str(ws.cell(r, master_col).value))
