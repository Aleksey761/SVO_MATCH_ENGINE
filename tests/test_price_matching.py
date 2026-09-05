from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from svo.engine import Engine
from svo.models import ArrivalItem


@pytest.fixture(autouse=True)
def _isolate_output_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.chdir(tmp_path)


def _write_master(path: Path, rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append([
        "SKU",
        "CATEGORY",
        "BRAND",
        "VARIANT",
        "VOLUME",
        "AROMA",
        "X1",
        "X2",
        "X3",
        "MASTER_NAME",
    ])
    for row in rows:
        ws.append(
            [
                row["sku"],
                row["category"],
                row["brand"],
                row["variant"],
                row["volume"],
                row.get("aroma", row["variant"]),
                "",
                "",
                "",
                row["master_name"],
            ]
        )
    wb.save(path)
    return path


def _write_price(path: Path, rows: list[list[object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE"
    ws.append(["ProductName", "UnitCost", "RetailPrice", "WholesalePrice", "MarketplacePrice"])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _single_master(tmp_path: Path) -> Path:
    return _write_master(
        tmp_path / "MASTER_TEST.xlsx",
        [
            {
                "sku": "SKU-APPLE",
                "category": "Шампунь",
                "brand": "SVO",
                "variant": "APPLE",
                "volume": "1 Л",
                "master_name": "Shampoo SVO APPLE 1 L",
            }
        ],
    )


def test_price_matching_exact_match(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo apple 1 l", 10, 20, None, None]],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    assert result["match"] == 1
    row = result["rows"][0]
    assert row["MatchStatus"] == "MATCH"
    assert row["SKU"] == "SKU-APPLE"
    assert row["MASTER_NAME"] == "Shampoo SVO APPLE 1 L"
    assert row["UnitCost"] == Decimal("10")
    assert row["RetailPrice"] == Decimal("20")


def test_price_matching_normalized_match(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["   svo   SHAMPOO   APPLE   1000 ml   ", 10, 20, None, None]],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    assert result["rows"][0]["MatchStatus"] == "MATCH"
    assert result["rows"][0]["SKU"] == "SKU-APPLE"


def test_price_matching_alias_match(tmp_path: Path):
    master_file = _write_master(
        tmp_path / "MASTER_TEST.xlsx",
        [
            {
                "sku": "SKU-WILD",
                "category": "Шампунь",
                "brand": "SVO",
                "variant": "WILD BERRIES",
                "volume": "1 Л",
                "master_name": "Shampoo SVO WILD BERRIES 1 L",
            }
        ],
    )
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo wildberries 1 l", 11, 21, None, None]],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    assert result["rows"][0]["MatchStatus"] == "MATCH"
    assert result["rows"][0]["SKU"] == "SKU-WILD"


def test_price_matching_marketing_word_removal(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo parfume plus apple 1 l", 10, 20, None, None]],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    assert result["rows"][0]["MatchStatus"] == "MATCH"
    assert result["rows"][0]["SKU"] == "SKU-APPLE"


def test_price_matching_unknown_product_goes_to_review(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["mystery foobar 999", 10, 20, None, None]],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    row = result["rows"][0]
    assert row["MatchStatus"] == "REVIEW"
    assert row["SKU"] == ""


def test_price_matching_ambiguous_product_goes_to_review(tmp_path: Path):
    master_file = _write_master(
        tmp_path / "MASTER_TEST.xlsx",
        [
            {
                "sku": "SKU-APPLE",
                "category": "Шампунь",
                "brand": "SVO",
                "variant": "APPLE",
                "volume": "1 Л",
                "master_name": "Shampoo SVO APPLE 1 L",
            },
            {
                "sku": "SKU-LIME",
                "category": "Шампунь",
                "brand": "SVO",
                "variant": "LIME",
                "volume": "1 Л",
                "master_name": "Shampoo SVO LIME 1 L",
            },
        ],
    )
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo 1 l", 10, 20, None, None]],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    row = result["rows"][0]
    assert row["MatchStatus"] == "REVIEW"
    assert "MULTIPLE_MATCH" in row["ReviewReason"]


def test_price_matching_allows_duplicate_names_with_same_sku_and_updates_history(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [
            ["SVO shampoo apple 1 l", 10, 20, None, None],
            ["SVO shampoo apple 1 l", 11, 21, None, None],
        ],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    assert result["price_rows"] == 2
    assert result["price_history_created"] == 2
    assert result["price_history_updated"] == 1
    assert result["price_history_skipped"] == 0


def test_price_matching_identical_duplicate_prices_are_skipped_by_history(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [
            ["SVO shampoo apple 1 l", 10, 20, None, None],
            ["SVO shampoo apple 1 l", 10, 20, None, None],
        ],
    )

    result = Engine().run_price_matching(master_file=master_file, price_file=price_file)

    assert result["price_rows"] == 2
    assert len(result["rows"]) == 2
    assert all(row["MatchStatus"] == "MATCH" for row in result["rows"])
    assert result["price_history_created"] == 1
    assert result["price_history_updated"] == 0
    assert result["price_history_skipped"] == 1


def test_price_matching_raises_only_for_different_sku_with_same_canonical_name():
    engine = Engine()
    first = ArrivalItem(row_number=10, source_name="SVO shampoo apple 1 l")
    first.sku = "SKU-001"
    first.unit_cost = Decimal("10")
    first.retail_price = Decimal("20")

    second = ArrivalItem(row_number=11, source_name="SVO shampoo apple 1 l")
    second.sku = "SKU-002"
    second.unit_cost = Decimal("11")
    second.retail_price = Decimal("21")

    with pytest.raises(ValueError, match="different SKU"):
        engine._validate_price_name_sku_conflicts([first, second])


def test_name_normalization_report_does_not_copy_previous_unresolved_rows(tmp_path: Path):
    output_dir = tmp_path / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "NAME_NORMALIZATION_REPORT.xlsx"

    wb = Workbook()
    ws_corrected = wb.active
    ws_corrected.title = "CORRECTED_NAMES"
    ws_corrected.append(["Source", "BEFORE_NAME", "AFTER_NAME", "MASTER_NAME", "SKU", "CHANGE_REASON"])
    ws_unresolved = wb.create_sheet("UNRESOLVED_NAMES")
    ws_unresolved.append(["Source", "BEFORE_NAME", "AFTER_NAME", "MASTER_NAME", "SKU", "CHANGE_REASON"])
    ws_unresolved.append(["PRICE", "mystery foobar 999", "MYSTERY FOOBAR 999", "", "", "stale"])
    ws_unresolved.append(["PRICE", "Unknown product without master", "UNKNOWN PRODUCT WITHOUT MASTER", "", "", "stale"])
    ws_summary = wb.create_sheet("SUMMARY")
    ws_summary.append(["Metric", "Value"])
    wb.save(report_path)

    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo apple 1 l", 10, 20, None, None]],
    )

    Engine().run_price_matching(master_file=master_file, price_file=price_file)

    report = load_workbook(report_path, data_only=True)
    unresolved_sheet = report["UNRESOLVED_NAMES"]
    unresolved_names = {
        str(unresolved_sheet.cell(row=row, column=2).value or "").strip()
        for row in range(2, unresolved_sheet.max_row + 1)
    }

    assert "mystery foobar 999" not in unresolved_names
    assert "Unknown product without master" not in unresolved_names


def test_price_matching_rejects_missing_price(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo apple 1 l", 10, None, None, None]],
    )

    with pytest.raises(ValueError, match="Missing required RetailPrice"):
        Engine().run_price_matching(master_file=master_file, price_file=price_file)


def test_price_matching_rejects_negative_price(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo apple 1 l", -10, 20, None, None]],
    )

    with pytest.raises(ValueError, match="Negative UnitCost"):
        Engine().run_price_matching(master_file=master_file, price_file=price_file)


def test_price_matching_rejects_invalid_decimal(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRICE.xlsx",
        [["SVO shampoo apple 1 l", "abc", 20, None, None]],
    )

    with pytest.raises(ValueError, match="Invalid Decimal"):
        Engine().run_price_matching(master_file=master_file, price_file=price_file)
