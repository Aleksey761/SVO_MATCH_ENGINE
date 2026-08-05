from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook

import runner


def _create_master(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-APPLE", "Шампунь", "SVO", "APPLE", "1 л", "Шампунь SVO APPLE 1 л"])
    ws.append(["SKU-LIME", "Шампунь", "SVO", "LIME", "1 л", "Шампунь SVO LIME 1 л"])
    wb.save(path)
    return path


def _create_price(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE"
    ws.append(["Supplier article", "Product name", "Free stock", "Distributor price"])
    ws.append(["A-100", "SVO shampoo apple 1 l", 42, "35,50"])
    ws.append(["A-200", "SVO shampoo lime 1 l", 12, "39,00"])
    wb.save(path)
    return path


def _create_sales(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append(["НАИМЕНОВАНИЕ", "SKU", "Тип товара", "Бренд", "Variant", "Объем"])
    ws.append(["SVO SHAMPUN APPLE 1 л", None, None, None, None, None])
    ws.append(["SVO SHAMPUN LIME 1 л", None, None, None, None, None])
    wb.save(path)
    return path


def _create_inventory(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "REVISION"
    ws.append([
        "Наименование",
        "Остаток на складе (шт) на 20.02.2024",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        "Остаток на складе (шт) на 10.07.24",
    ])
    ws.append(["SVO SHAMPUN APPLE 1 л", 20, 1, 0, 0, 1, 18])
    ws.append(["SVO SHAMPUN LIME 1 л", 15, 0, 1, 0, 1, 13])
    wb.save(path)
    return path


def _create_master_dataset(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER_DATASET"
    ws.append(["SKU", "MASTER_NAME", "CATEGORY", "BRAND", "ARRIVAL_QUANTITY", "SALES_QUANTITY", "NET_MOVEMENT", "ARRIVAL_COUNT"])
    ws.append(["SKU-APPLE", "Шампунь SVO APPLE 1 л", "Шампунь", "SVO", 1, 1, 0, 1])
    ws.append(["SKU-LIME", "Шампунь SVO LIME 1 л", "Шампунь", "SVO", 1, 1, 0, 1])
    wb.save(path)


def _sheet_as_dict(ws) -> dict[str, object]:
    values: dict[str, object] = {}
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row=row, column=1).value
        val = ws.cell(row=row, column=2).value
        if key is not None:
            values[str(key)] = val
    return values


def _normalize(value: object) -> str:
    return " ".join(str(value or "").strip().upper().split())


def _expected_summary(output_dir: Path) -> dict[str, object]:
    expected: dict[str, object] = {}

    wb_metrics = load_workbook(output_dir / "MASTER_DATASET.xlsx", data_only=True)
    ws_metrics = wb_metrics.active
    expected["MASTER rows"] = sum(1 for row in range(2, ws_metrics.max_row + 1) if str(ws_metrics.cell(row=row, column=1).value or "").strip())

    wb_match = load_workbook(output_dir / "price_match.xlsx", data_only=True)
    ws_match = wb_match.active
    headers = [_normalize(ws_match.cell(row=1, column=col).value) for col in range(1, ws_match.max_column + 1)]
    status_col = headers.index("STATUS") + 1
    original_col = headers.index("ORIGINALPRODUCTNAME") + 1

    price_rows = 0
    matched = 0
    for row in range(2, ws_match.max_row + 1):
        if not str(ws_match.cell(row=row, column=original_col).value or "").strip():
            continue
        price_rows += 1
        if _normalize(ws_match.cell(row=row, column=status_col).value) == "MATCH":
            matched += 1
    expected["PRICE rows"] = price_rows
    expected["Matched"] = matched

    wb_review = load_workbook(output_dir / "price_review.xlsx", data_only=True)
    ws_review = wb_review.active
    expected["Review"] = max(0, ws_review.max_row - 1)

    wb_change = load_workbook(output_dir / "price_change_report.xlsx", data_only=True)
    ws_change = wb_change.active
    change_headers = [_normalize(ws_change.cell(row=1, column=col).value) for col in range(1, ws_change.max_column + 1)]
    status_change_col = change_headers.index("STATUS") + 1
    updated = 0
    has_skipped = False
    skipped = 0
    for row in range(2, ws_change.max_row + 1):
        status = _normalize(ws_change.cell(row=row, column=status_change_col).value)
        if status == "UPDATED":
            updated += 1
        elif status == "SKIPPED":
            skipped += 1
            has_skipped = True
    expected["Updated prices"] = updated
    expected["Skipped"] = skipped if has_skipped else "N/A"

    analytics_lines = (output_dir / "MATCH_ANALYTICS.txt").read_text(encoding="utf-8").splitlines()
    in_new = False
    new_count = 0
    for line in analytics_lines:
        norm = _normalize(line)
        if norm == "NEW PRODUCTS FOR MASTER":
            in_new = True
            continue
        if in_new and norm.startswith("----------------------------------------"):
            in_new = False
            continue
        if in_new and line.strip().startswith("-"):
            if line.strip().lower() == "- none":
                continue
            new_count += 1
    expected["Coverage %"] = f"{matched / price_rows * 100.0:.2f} %" if price_rows else "N/A"
    expected["New products"] = new_count

    return expected


def _expected_review(output_dir: Path) -> dict[str, object]:
    wb = load_workbook(output_dir / "price_review.xlsx", data_only=True)
    ws = wb.active
    headers = [_normalize(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)]
    reason_col = headers.index("REASON") + 1
    recommendation_col = headers.index("RECOMMENDATION") + 1 if "RECOMMENDATION" in headers else None

    expected = {
        "NO_MATCH": 0,
        "MULTIPLE_CANDIDATES": 0,
        "PRODUCT_TYPE_MISMATCH": 0,
        "MANUAL_REVIEW": 0,
    }
    for row in range(2, ws.max_row + 1):
        reason = _normalize(ws.cell(row=row, column=reason_col).value)
        recommendation = _normalize(ws.cell(row=row, column=recommendation_col).value) if recommendation_col else ""
        if "PRODUCT_TYPE_MISMATCH" in reason or "PRODUCT TYPE" in reason or "PRODUCTTYPE" in reason:
            expected["PRODUCT_TYPE_MISMATCH"] += 1
        elif "AMBIGUOUS" in reason or "MULTIPLE" in reason:
            expected["MULTIPLE_CANDIDATES"] += 1
        elif "NO MASTER" in reason or "NO_MATCH" in reason:
            expected["NO_MATCH"] += 1
        elif "MANUAL" in reason or "MANUAL" in recommendation:
            expected["MANUAL_REVIEW"] += 1
        else:
            expected["MANUAL_REVIEW"] += 1
    return expected


def _expected_price(output_dir: Path) -> dict[str, object]:
    wb = load_workbook(output_dir / "price_change_report.xlsx", data_only=True)
    ws = wb.active
    headers = [_normalize(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)]
    status_col = headers.index("STATUS") + 1
    expected = {
        "New prices": 0,
        "Updated prices": 0,
        "Unchanged prices": 0,
    }
    has_skipped = False
    for row in range(2, ws.max_row + 1):
        status = _normalize(ws.cell(row=row, column=status_col).value)
        if status == "NEW":
            expected["New prices"] += 1
        elif status == "UPDATED":
            expected["Updated prices"] += 1
        elif status == "SKIPPED":
            expected["Unchanged prices"] += 1
            has_skipped = True
    if not has_skipped:
        expected["Unchanged prices"] = "N/A"
    return expected


def test_dashboard_matches_real_outputs_after_full_run(monkeypatch):
    workspace = Path("tests") / "integration_workspace" / "dashboard_full_pipeline"
    if workspace.exists():
        shutil.rmtree(workspace)
    workspace.mkdir(parents=True, exist_ok=True)
    workspace = workspace.resolve()

    master = _create_master((workspace / "MASTER.xlsx").resolve())
    price = _create_price((workspace / "PRICE.xlsx").resolve())
    sales = _create_sales((workspace / "SALES_2026-07-15.xlsx").resolve())
    inventory = _create_inventory((workspace / "REVISION_2026-07-15.xlsx").resolve())

    monkeypatch.chdir(workspace)

    code = runner.main(
        [
            "--mode",
            "full",
            "--master",
            str(master),
            "--price",
            str(price),
            "--sales",
            str(sales),
            "--inventory",
            str(inventory),
        ]
    )
    assert code == 0

    output_dir = Path("output").resolve()
    dashboard = load_workbook(output_dir / "DASHBOARD.xlsx", data_only=True)
    summary_ws = dashboard["Summary"]
    actual = _sheet_as_dict(summary_ws)
    expected = _expected_summary(output_dir)

    for key, value in expected.items():
        assert actual[key] == value, f"Summary mismatch for {key}: {actual[key]} != {value}"

    assert actual["Run date"] != "N/A"
    assert actual["Run time"] != "N/A"
    assert actual["Status 1"] == "✓ Dashboard created"
    assert actual["Status 2"] == "✓ Sources loaded"
    assert actual["Status 3"] == "✓ Metrics verified"

    review_actual = _sheet_as_dict(dashboard["Review"])
    review_expected = _expected_review(output_dir)
    for key, value in review_expected.items():
        assert review_actual[key] == value, f"Review mismatch for {key}: {review_actual[key]} != {value}"

    price_actual = _sheet_as_dict(dashboard["Price"])
    price_expected = _expected_price(output_dir)
    for key, value in price_expected.items():
        assert price_actual[key] == value, f"Price mismatch for {key}: {price_actual[key]} != {value}"

    run_ws = dashboard["Run"]
    run_values = _sheet_as_dict(run_ws)
    source_master = str(run_values["Source MASTER"])
    source_price = str(run_values["Source PRICE"])
    assert source_master.replace("\\", "/").endswith("/output/MASTER_DATASET.xlsx")
    assert source_price.replace("\\", "/").endswith("/output/price_match.xlsx")
    assert "pytest-" not in source_master.lower()
    assert "pytest-" not in source_price.lower()
    assert str(run_values["Duration"]).endswith(" sec")
