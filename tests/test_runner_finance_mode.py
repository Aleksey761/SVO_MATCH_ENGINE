from argparse import Namespace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from runner import _execute_finance_pipeline, _load_inventory_stock_by_sku_exact_alias


def _write_master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER_DATASET"
    ws.append(["SKU", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "AROMA", "C7", "C8", "C9", "MASTER_NAME"])
    ws.append(["SKU-1", "Cat", "Brand", "Var1", "1 Л", "Var1", "", "", "", "Master One"])
    ws.append(["SKU-2", "Cat", "Brand", "Var2", "1 Л", "Var2", "", "", "", "Master Two"])
    wb.save(path)


def _write_price_match(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matched"
    ws.append([
        "SKU",
        "MASTER_NAME",
        "SupplierArticle",
        "OriginalProductName",
        "CanonicalProductName",
        "UnitCost",
        "RetailPrice",
        "StockQty",
        "SourceUnit",
        "PackQty",
        "OriginalPrice",
        "ImportedAt",
        "SourceFile",
        "Status",
    ])
    ws.append(["SKU-1", "Master One", "A1", "Item 1", "ITEM 1", 10, 20, 5, None, None, 10, "", "", "MATCH"])
    ws.append(["SKU-2", "Master Two", "A2", "Item 2", "ITEM 2", 8, 12, 0, None, None, 8, "", "", "MATCH"])
    wb.save(path)


def _write_sales(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    ws.append(["№", "SKU", "Тип товара", "Бренд", "Variant", "Объем", "", "", "", "Склад 1", "", "Склад 2", "", "ВСЕГО", "Ср Годн", "", "СУММА"])
    ws.append(["", "", "", "", "", "", "", "", "", "КОР", "ШТ", "КОР", "ШТ", "", "", "ЦЕНА", "СУММА"])
    ws.append([1, "SKU-1", "", "", "", "", "", "", "", 0, 0, 0, 0, 3, "", "", 60])
    wb.save(path)


def _write_inventory(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "REV"
    ws.append(["", "", "", "", "", "", "", "", "", "", "Склад 1", "", "Склад 2", "", "ВСЕГО", "Ср Годн", "ЦЕНА", "СУММА"])
    ws.append(["", "", "", "", "", "", "", "", "", "", "КОР", "ШТ", "КОР", "ШТ", "", "", "", ""])
    ws.append([1, "", "Master One", "", "", "", "", "", "", "", 0, 3, 0, 2, 5, "2026-06-01", 88, 440])
    ws.append([2, "", "Master Two", "", "", "", "", "", "", "", 0, 9, 0, 0, 9, "2026-06-01", 88, 792])
    wb.save(path)


def _write_inventory_alias_map(
    path: Path,
    rows: list[tuple[str, str, str, str, str, str, str, str]],
    *,
    status: str = "APPROVED_READY",
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "APPROVED_CANDIDATES"
    ws.append([
        "InventoryName",
        "MASTER_NAME",
        "SKU",
        "ProductType",
        "Brand",
        "Volume",
        "Variant",
        "Evidence",
        "Status",
    ])
    for row in rows:
        ws.append([*row, status])
    wb.save(path)


def _write_inventory_with_alias_cases(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "REV"
    ws.append(["", "", "", "", "", "", "", "", "", "", "Склад 1", "", "Склад 2", "", "ВСЕГО", "Ср Годн", "ЦЕНА", "СУММА"])
    ws.append(["", "", "", "", "", "", "", "", "", "", "КОР", "ШТ", "КОР", "ШТ", "", "", "", ""])
    approved_names = [
        ("Alias One", 11),
        ("Alias Two", 12),
        ("Alias Three", 13),
        ("Alias Four", 14),
        ("Alias Five", 15),
        ("Alias Six", 16),
        ("Alias Seven", 17),
        ("Alias Eight", 18),
    ]
    pending_names = [
        ("Needs Review One", 21),
        ("Needs Review Two", 22),
        ("Needs Review Three", 23),
    ]
    unknown_names = [("Unknown Inventory Name", 99)]
    for idx, (name, total) in enumerate(approved_names + pending_names + unknown_names, start=1):
        ws.append([idx, "", name, "", "", "", "", "", "", "", 0, total, 0, 0, total, "2026-06-01", 1, total])
    wb.save(path)


def test_finance_mode_generates_result(tmp_path: Path):
    master_file = tmp_path / "MASTER_DATASET.xlsx"
    price_match_file = tmp_path / "price_match.xlsx"
    sales_file = tmp_path / "SALES_MATCH.xlsx"
    inventory_file = tmp_path / "REVISION 06.04.26.xlsx"
    output_file = tmp_path / "FINANCE_RESULT.xlsx"

    _write_master(master_file)
    _write_price_match(price_match_file)
    _write_sales(sales_file)
    _write_inventory(inventory_file)

    args = Namespace(
        master=str(master_file),
        price=str(price_match_file),
        sales=str(sales_file),
        inventory=str(inventory_file),
        output=str(output_file),
    )

    result = _execute_finance_pipeline(args)

    assert output_file.exists()
    assert result["total_sku"] == 2
    assert result["with_sales"] == 1
    assert result["with_cost"] == 2

    wb = load_workbook(output_file, data_only=True)
    ws = wb["FINANCE_RESULT"]
    rows = [
        [ws.cell(r, c).value for c in range(1, 13)]
        for r in range(2, ws.max_row + 1)
    ]
    row_by_sku = {row[0]: row for row in rows}

    sku1 = row_by_sku["SKU-1"]
    assert sku1[2] == 3  # SalesQty
    assert float(sku1[4]) == 60.0  # Revenue = 3 * 20
    assert float(sku1[6]) == 30.0  # COGS = 3 * 10
    assert float(sku1[7]) == 30.0  # GrossProfit
    assert float(sku1[8]) == 50.0  # MarginPct
    assert float(sku1[10]) == 50.0  # StockValue = 5 * 10
    assert sku1[11] == 0  # Отгружено is not sourced from this inventory structure.

    sku2 = row_by_sku["SKU-2"]
    assert sku2[2] == 0
    assert float(sku2[4]) == 0.0
    assert float(sku2[8]) == 0.0


def test_inventory_alias_stock_loader_applies_only_exact_approved_aliases(tmp_path: Path):
    inventory_file = tmp_path / "REVISION 06.04.26.xlsx"
    _write_inventory_with_alias_cases(inventory_file)

    approved_aliases = {
        "Alias One": "SKU-1",
        "Alias Two": "SKU-2",
        "Alias Three": "SKU-3",
        "Alias Four": "SKU-4",
        "Alias Five": "SKU-5",
        "Alias Six": "SKU-6",
        "Alias Seven": "SKU-7",
        "Alias Eight": "SKU-8",
    }

    stock_by_sku, applied_rows = _load_inventory_stock_by_sku_exact_alias(
        inventory_file,
        alias_by_inventory_name=approved_aliases,
    )

    assert applied_rows == 8
    assert stock_by_sku == {
        "SKU-1": 11,
        "SKU-2": 12,
        "SKU-3": 13,
        "SKU-4": 14,
        "SKU-5": 15,
        "SKU-6": 16,
        "SKU-7": 17,
        "SKU-8": 18,
    }
    assert "Needs Review One" not in stock_by_sku
    assert "Needs Review Two" not in stock_by_sku
    assert "Needs Review Three" not in stock_by_sku
    assert "Unknown Inventory Name" not in stock_by_sku


def test_finance_mode_uses_exact_inventory_alias_map_without_fuzzy_fallback(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.chdir(tmp_path)
    output_dir = tmp_path / "output"
    output_dir.mkdir()

    master_file = tmp_path / "MASTER_DATASET.xlsx"
    price_match_file = output_dir / "price_match.xlsx"
    sales_file = output_dir / "SALES_MATCH.xlsx"
    inventory_file = tmp_path / "REVISION 06.04.26.xlsx"
    output_file = output_dir / "FINANCE_RESULT.xlsx"
    alias_map_file = output_dir / "INVENTORY_ALIAS_MAP.xlsx"

    _write_master(master_file)
    _write_price_match(price_match_file)
    _write_sales(sales_file)

    wb = Workbook()
    ws = wb.active
    ws.title = "REV"
    ws.append(["", "", "", "", "", "", "", "", "", "", "Склад 1", "", "Склад 2", "", "ВСЕГО", "Ср Годн", "ЦЕНА", "СУММА"])
    ws.append(["", "", "", "", "", "", "", "", "", "", "КОР", "ШТ", "КОР", "ШТ", "", "", "", ""])
    ws.append([1, "", "Alias SKU Two", "", "", "", "", "", "", "", 0, 0, 0, 0, 9, "2026-06-01", 1, 9])
    ws.append([2, "", "Unapproved Similar Name", "", "", "", "", "", "", "", 0, 0, 0, 0, 25, "2026-06-01", 1, 25])
    wb.save(inventory_file)

    _write_inventory_alias_map(
        alias_map_file,
        [("Alias SKU Two", "Master Two", "SKU-2", "Cat", "Brand", "1 Л", "Var2", "approved exact alias")],
    )

    args = Namespace(
        master=str(master_file),
        price=str(price_match_file),
        sales=str(sales_file),
        inventory=str(inventory_file),
        output=str(output_file),
    )

    result = _execute_finance_pipeline(args)

    assert result["inventory_alias_rows_applied"] == 1
    assert result["with_stock"] == 2

    finance_wb = load_workbook(output_file, data_only=True)
    finance_ws = finance_wb["FINANCE_RESULT"]
    rows = {
        str(finance_ws.cell(r, 1).value): [finance_ws.cell(r, c).value for c in range(1, 13)]
        for r in range(2, finance_ws.max_row + 1)
    }
    assert rows["SKU-2"][9] == 9
    assert rows["SKU-1"][9] == 5
