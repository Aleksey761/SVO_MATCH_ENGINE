from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import runner
from svo.engine import Engine


@pytest.fixture(autouse=True)
def _isolate_output_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.chdir(tmp_path)


def _write_master(path: Path, rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append([
        "SKU",
        "CATEGORY",
        "BRAND",
        "VARIANT",
        "VOLUME",
        "AROMA",
        "X1",
        "X2",
        "X3",
        "MASTER_NAME",
    ])
    for row in rows:
        ws.append(
            [
                row["sku"],
                row["category"],
                row["brand"],
                row["variant"],
                row["volume"],
                row.get("aroma", row["variant"]),
                "",
                "",
                "",
                row["master_name"],
            ]
        )
    wb.save(path)
    return path


def _write_price(path: Path, rows: list[list[object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE"
    ws.append(["Supplier article", "ProductName", "StockQty", "UnitCost", "RetailPrice", "SourceUnit"])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _single_master(tmp_path: Path) -> Path:
    return _write_master(
        tmp_path / "MASTER_TEST.xlsx",
        [
            {
                "sku": "SKU-APPLE",
                "category": "Шампунь",
                "brand": "SVO",
                "variant": "APPLE",
                "volume": "1 Л",
                "master_name": "Shampoo SVO APPLE 1 L",
            }
        ],
    )


def test_price_match_output_workbook_created_with_required_columns(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRC.xlsx",
        [["A-100", "SVO shampoo apple 1 l", 42, 100, 150, "шт"]],
    )
    output_file = tmp_path / "output" / "price_match.xlsx"

    result = engine.run_price_matching(master_file=master_file, price_file=price_file, output_file=output_file)

    assert Path(result["price_match_output"]).exists()
    wb = load_workbook(result["price_match_output"], data_only=True)
    ws = wb["Matched"]
    headers = [ws.cell(row=1, column=col).value for col in range(1, 15)]
    assert headers == [
        "SKU",
        "MASTER_NAME",
        "SupplierArticle",
        "OriginalProductName",
        "CanonicalProductName",
        "UnitCost",
        "RetailPrice",
        "StockQty",
        "SourceUnit",
        "PackQty",
        "OriginalPrice",
        "ImportedAt",
        "SourceFile",
        "Status",
    ]


def test_price_change_report_output_has_production_columns(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    first = _write_price(tmp_path / "PRC_1.xlsx", [["A-100", "SVO shampoo apple 1 l", 42, 100, 150, "шт"]])
    second = _write_price(tmp_path / "PRC_2.xlsx", [["A-100", "SVO shampoo apple 1 l", 42, 120, 150, "шт"]])

    engine.run_price_matching(master_file=master_file, price_file=first, output_file=tmp_path / "output" / "price_match.xlsx")
    result = engine.run_price_matching(master_file=master_file, price_file=second, output_file=tmp_path / "output" / "price_match.xlsx")

    wb = load_workbook(result["price_change_report"], data_only=True)
    ws = wb["PRICE_CHANGES"]
    headers = [ws.cell(row=1, column=col).value for col in range(1, 14)]
    assert headers == [
        "SKU",
        "MASTER_NAME",
        "PreviousPrice",
        "NewPrice",
        "Difference",
        "PercentChange",
        "SourceUnit",
        "PackQty",
        "OriginalPrice",
        "ValidFrom",
        "ImportedAt",
        "SourceFile",
        "Status",
    ]


def test_price_review_output_generated(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    price_file = _write_price(
        tmp_path / "PRC_REVIEW.xlsx",
        [
            ["A-100", "SVO shampoo apple 1 l", 42, 100, 150, "шт"],
            ["A-404", "Unknown product 999", 5, 20, 30, "шт"],
        ],
    )

    result = engine.run_price_matching(master_file=master_file, price_file=price_file, output_file=tmp_path / "output" / "price_match.xlsx")

    report = Path(result["price_review_report"])
    assert report.exists()
    wb = load_workbook(report, data_only=True)
    ws = wb["PRICE_REVIEW"]
    headers = [ws.cell(row=1, column=col).value for col in range(1, 8)]
    assert headers == [
        "SupplierArticle",
        "OriginalProductName",
        "CanonicalProductName",
        "SKU",
        "MASTER_NAME",
        "Reason",
        "Recommendation",
    ]
    assert ws.max_row >= 2


def test_output_parameter_used_by_runner(tmp_path: Path):
    master_file = _single_master(tmp_path)
    price_file = _write_price(tmp_path / "PRC.xlsx", [["A-100", "SVO shampoo apple 1 l", 42, 100, 150, "шт"]])

    target_output = tmp_path / "output" / "price_match.xlsx"
    exit_code = runner.main(
        [
            "--mode",
            "price",
            "--master",
            str(master_file),
            "--price",
            str(price_file),
            "--output",
            str(target_output),
        ]
    )

    assert exit_code == 0
    assert target_output.exists()


def test_console_summary_printed_for_price_mode(tmp_path: Path, capsys):
    master_file = _single_master(tmp_path)
    price_file = _write_price(tmp_path / "PRC.xlsx", [["A-100", "SVO shampoo apple 1 l", 42, 100, 150, "шт"]])

    target_output = tmp_path / "output" / "price_match.xlsx"
    exit_code = runner.main(
        [
            "--mode",
            "price",
            "--master",
            str(master_file),
            "--price",
            str(price_file),
            "--output",
            str(target_output),
        ]
    )

    out = capsys.readouterr().out
    assert exit_code == 0
    assert "PRICE IMPORT SUMMARY" in out
    assert "MASTER rows" in out
    assert "PRICE rows" in out
    assert "New prices" in out
    assert "Updated prices" in out
    assert "Skipped" in out
    assert "price_match.xlsx" in out
    assert "price_change_report.xlsx" in out
    assert "price_review.xlsx" in out
    assert "Completed successfully" in out
