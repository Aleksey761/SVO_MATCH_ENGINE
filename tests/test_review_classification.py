from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook

import runner


ALLOWED_ROOT_CAUSES = {
    "MULTIPLE_CANDIDATES",
    "NO_MATCH",
    "PRODUCT_TYPE_MISMATCH",
    "BRAND_MISMATCH",
    "VOLUME_MISMATCH",
    "AROMA_MISMATCH",
    "MISSING_ALIAS",
    "OCR_OR_TYPOS",
    "NEW_PRODUCT",
    "OTHER",
}


def _create_master(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-APPLE", "Шампунь", "SVO", "APPLE", "1 л", "Шампунь SVO APPLE 1 л"])
    ws.append(["SKU-LIME", "Шампунь", "SVO", "LIME", "1 л", "Шампунь SVO LIME 1 л"])
    wb.save(path)
    return path


def _create_price(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE"
    ws.append(["Supplier article", "Product name", "Free stock", "Distributor price"])
    ws.append(["A-100", "SVO shampoo apple 1 l", 42, "35,50"])
    ws.append(["A-200", "SVO shampoo 1 l", 12, "39,00"])
    ws.append(["A-300", "Unknown product 999", 5, "19,00"])
    wb.save(path)
    return path


def _create_sales(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append(["НАИМЕНОВАНИЕ", "SKU", "Тип товара", "Бренд", "Variant", "Объем"])
    ws.append(["SVO SHAMPUN APPLE 1 л", None, None, None, None, None])
    wb.save(path)
    return path


def _create_inventory(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "REVISION"
    ws.append([
        "Наименование",
        "Остаток на складе (шт) на 20.02.2024",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        "Остаток на складе (шт) на 10.07.24",
    ])
    ws.append(["SVO SHAMPUN APPLE 1 л", 20, 1, 0, 0, 1, 18])
    wb.save(path)
    return path


def _sheet_rows(path: Path) -> list[dict[str, object]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(row=1, column=col).value or "").strip() for col in range(1, ws.max_column + 1)]
    rows: list[dict[str, object]] = []
    for row in range(2, ws.max_row + 1):
        payload: dict[str, object] = {}
        for index, header in enumerate(headers, start=1):
            payload[header] = ws.cell(row=row, column=index).value
        rows.append(payload)
    return rows


def test_review_classification_reports_after_full_run(monkeypatch, capsys):
    workspace = (Path("tests") / "integration_workspace" / "review_classification_full").resolve()
    if workspace.exists():
        shutil.rmtree(workspace)
    workspace.mkdir(parents=True, exist_ok=True)

    master = _create_master(workspace / "MASTER.xlsx")
    price = _create_price(workspace / "PRICE.xlsx")
    sales = _create_sales(workspace / "SALES_2026-07-15.xlsx")
    inventory = _create_inventory(workspace / "REVISION_2026-07-15.xlsx")

    monkeypatch.chdir(workspace)

    code = runner.main(
        [
            "--mode",
            "full",
            "--master",
            str(master),
            "--price",
            str(price),
            "--sales",
            str(sales),
            "--inventory",
            str(inventory),
        ]
    )
    assert code == 0

    stdout = capsys.readouterr().out
    assert "========== REVIEW ANALYSIS ==========" in stdout
    assert "Total REVIEW:" in stdout

    output_dir = workspace / "output"
    classification_rows = _sheet_rows(output_dir / "REVIEW_CLASSIFICATION.xlsx")
    statistics_rows = _sheet_rows(output_dir / "REVIEW_STATISTICS.xlsx")
    multiple_rows = _sheet_rows(output_dir / "MULTIPLE_CANDIDATES_ANALYSIS.xlsx")
    multiple_summary_rows = _sheet_rows(output_dir / "MULTIPLE_CANDIDATES_SUMMARY.xlsx")
    review_rows = _sheet_rows(output_dir / "price_review.xlsx")

    assert len(classification_rows) == len(review_rows)

    review_index = {
        (str(row["SupplierArticle"] or "").strip(), str(row["OriginalProductName"] or "").strip()): row
        for row in review_rows
    }

    for row in classification_rows:
        key = (str(row["SupplierArticle"] or "").strip(), str(row["OriginalName"] or "").strip())
        assert key in review_index
        assert row["ReviewReason"] == review_index[key]["Reason"]
        assert row["RootCause"] in ALLOWED_ROOT_CAUSES

    derived_counts: dict[str, int] = {}
    for row in classification_rows:
        root_cause = str(row["RootCause"] or "").strip()
        derived_counts[root_cause] = derived_counts.get(root_cause, 0) + 1

    statistics_counts = [(str(row["RootCause"] or "").strip(), int(row["Count"] or 0)) for row in statistics_rows]
    assert statistics_counts == sorted(statistics_counts, key=lambda item: (-item[1], item[0]))

    for root_cause, count in statistics_counts:
        assert derived_counts[root_cause] == count

    assert sum(count for _, count in statistics_counts) == len(review_rows)

    multiple_from_classification = [
        row
        for row in classification_rows
        if str(row["RootCause"] or "").strip() == "MULTIPLE_CANDIDATES"
    ]
    assert len(multiple_rows) == len(multiple_from_classification)

    for row in multiple_rows:
        c1 = row["Candidate1Score"]
        c2 = row["Candidate2Score"]
        diff = row["ScoreDifference12"]
        auto = str(row["AutoMatchPossible"] or "").strip().upper()
        recommendation = str(row["Recommendation"] or "").strip()
        assert auto in {"YES", "NO"}
        expected_recommendation = "Increase auto-match threshold candidate" if auto == "YES" else "Keep in REVIEW"
        assert recommendation == expected_recommendation
        if c1 is not None and c2 is not None:
            expected_diff = round(float(c1) - float(c2), 2)
            assert round(float(diff), 2) == expected_diff
            expected_auto = "YES" if expected_diff >= 20.0 else "NO"
            assert auto == expected_auto

    range_counts = {
        str(row["ScoreDifference range"] or "").strip(): int(row["Count"] or 0)
        for row in multiple_summary_rows
    }
    assert list(range_counts.keys()) == ["0-5", "5-10", "10-20", "20-30", "30+"]

    calculated = {"0-5": 0, "5-10": 0, "10-20": 0, "20-30": 0, "30+": 0}
    for row in multiple_rows:
        diff = row["ScoreDifference12"]
        if diff is None:
            continue
        value = float(diff)
        if 0.0 <= value < 5.0:
            calculated["0-5"] += 1
        elif 5.0 <= value < 10.0:
            calculated["5-10"] += 1
        elif 10.0 <= value < 20.0:
            calculated["10-20"] += 1
        elif 20.0 <= value < 30.0:
            calculated["20-30"] += 1
        elif value >= 30.0:
            calculated["30+"] += 1

    assert range_counts == calculated

    assert "Total MULTIPLE_CANDIDATES:" in stdout
    assert "Resolvable automatically:" in stdout
    assert "Still require REVIEW:" in stdout