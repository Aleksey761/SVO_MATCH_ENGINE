from pathlib import Path

from openpyxl import Workbook, load_workbook

from svo.inventory_input_reader import InventoryInputReader
from svo.inventory_reconciliation import InventoryReconciliation, build_inventory_reconciliation


def _create_master(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "SKU",
        "CATEGORY",
        "BRAND",
        "VARIANT",
        "VOLUME",
        "AROMA",
        "C7",
        "C8",
        "C9",
        "MASTER_NAME",
    ])
    ws.append(["SKU-1", "Cat", "Brand", "Var", "1 л", "Var", "", "", "", "Master One"])
    ws.append(["SKU-2", "Cat", "Brand", "Var", "1 л", "Var", "", "", "", "Master Two"])
    wb.save(path)
    return path


def _create_result(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESULT"
    ws.append(["REPORT", "SVO Match Engine"])
    ws.append(["ARRIVAL_DATE", "10.07.24"])
    ws.append(["METADATA", "arrival_date=10.07.24"])
    ws.append([])
    ws.append(["SOURCE_NAME", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "SKU", "MASTER_NAME", "STATUS"])
    ws.append(["Item A", "Cat", "Brand", "V1", "1 л", "SKU-1", "", "MATCH"])
    ws.append(["Item B", "Cat", "Brand", "V2", "1 л", "", "", "REVIEW"])
    ws.append(["Item C", "Cat", "Brand", "V3", "1 л", "SKU-1", "", "MATCH"])
    ws.append(["Item D", "Cat", "Brand", "V4", "1 л", "SKU-2", "", "MATCH"])
    wb.save(path)
    return path


def _create_inventory(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Наименование",
        "Остаток на складе (шт) на 20.02.2024",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        "Остаток на складе (шт) на 10.07.24",
    ])
    ws.append(["Item A", 100, 10, 5, 5, 2, 74])
    ws.append(["Item B", 70, 4, 3, 3, 1, 62])
    ws.append(["Item C", 50, 5, 4, 1, 3, 38])
    ws.append(["Item D", 200, 20, 10, 5, 4, 160])
    wb.save(path)
    return path


def _create_inventory_with_custom_opening_date(path: Path, opening_date: str, closing_date: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Наименование",
        f"Остаток на складе (шт) на {opening_date}",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        f"Остаток на складе (шт) на {closing_date}",
    ])
    ws.append(["Item X", 30, 3, 2, 1, 1, 25])
    wb.save(path)
    return path


def _create_inventory_without_dated_stock_header(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Наименование",
        "Остаток на складе (шт)",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
    ])
    ws.append(["Item Z", 10, 1, 1, 1, 1])
    wb.save(path)
    return path


def _create_inventory_with_sku(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "SKU",
        "Наименование",
        "Остаток на складе (шт) на 20.02.2024",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        "Остаток на складе (шт) на 10.07.24",
    ])
    ws.append(["SKU-2", "Item A alias", 100, 10, 5, 5, 2, 74])
    wb.save(path)
    return path


def _create_inventory_with_unknown_sku(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "SKU",
        "Наименование",
        "Остаток на складе (шт) на 20.02.2024",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        "Остаток на складе (шт) на 10.07.24",
    ])
    ws.append(["SKU-999", "Item A", 100, 10, 5, 5, 2, 74])
    wb.save(path)
    return path


def _create_result_ambiguous_name(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESULT"
    ws.append(["REPORT", "SVO Match Engine"])
    ws.append(["ARRIVAL_DATE", "10.07.24"])
    ws.append(["METADATA", "arrival_date=10.07.24"])
    ws.append([])
    ws.append(["SOURCE_NAME", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "SKU", "MASTER_NAME", "STATUS"])
    ws.append(["Item A", "Cat", "Brand", "V1", "1 л", "SKU-1", "Master One", "MATCH"])
    ws.append(["Item A", "Cat", "Brand", "V2", "1 л", "SKU-2", "Master Two", "MATCH"])
    wb.save(path)
    return path


def _create_inventory_shuffled(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Наименование",
        "Остаток на складе (шт) на 20.02.2024",
        "Воронеж",
        "Краснодар 1",
        "Краснодар 2",
        "розница",
        "Остаток на складе (шт) на 10.07.24",
    ])
    ws.append(["Item D", 200, 20, 10, 5, 4, 160])
    ws.append(["Item B", 70, 4, 3, 3, 1, 62])
    ws.append(["Item A", 100, 10, 5, 5, 2, 74])
    ws.append(["Item C", 50, 5, 4, 1, 3, 38])
    wb.save(path)
    return path


def _stock_rows(path: Path) -> list[list[object]]:
    wb = load_workbook(path, data_only=True)
    stock = wb["STOCK_RESULT"]
    return [
        [stock.cell(r, c).value for c in range(1, 10)]
        for r in range(2, stock.max_row + 1)
    ]


def test_inventory_reconciliation_business_formulas():
    calc = InventoryReconciliation._calculate_stock_qty(
        opening_stock_qty=100,
        receipt_qty=0,
        supplier_return_qty=0,
        sales_qty=22,
    )
    variance = InventoryReconciliation._calculate_variance_qty(
        actual_stock_qty=74,
        calculated_stock_qty=calc,
    )

    assert calc == 78
    assert variance == -4


def test_inventory_input_reader_maps_returns_and_retail_sales(tmp_path: Path):
    inventory_file = _create_inventory(tmp_path / "REVISION 10.07.24.xlsx")

    quantities = InventoryInputReader().read_inventory_quantities(inventory_file)

    assert quantities[0] == {
        "source_name": "Item A",
        "opening": 100,
        "receipt": 0,
        "supplier_return": 20,
        "sales": 2,
        "actual": 74,
    }
    assert quantities[2] == {
        "source_name": "Item C",
        "opening": 50,
        "receipt": 0,
        "supplier_return": 10,
        "sales": 3,
        "actual": 38,
    }


def test_inventory_input_reader_supports_dynamic_opening_date_header(tmp_path: Path):
    inventory_file = _create_inventory_with_custom_opening_date(
        tmp_path / "REVISION 06.04.26.xlsx",
        opening_date="15.03.2026",
        closing_date="06.04.26",
    )

    quantities = InventoryInputReader().read_inventory_quantities(inventory_file)

    assert len(quantities) == 1
    assert quantities[0]["opening"] == 30
    assert quantities[0]["actual"] == 25
    assert quantities[0]["supplier_return"] == 6
    assert quantities[0]["sales"] == 1


def test_inventory_input_reader_requires_dated_stock_header(tmp_path: Path):
    inventory_file = _create_inventory_without_dated_stock_header(tmp_path / "REVISION_BAD.xlsx")

    try:
        InventoryInputReader().read_inventory_quantities(inventory_file)
        assert False, "Expected ValueError for missing dated stock header"
    except ValueError as exc:
        assert "with date header" in str(exc)


def test_inventory_reconciliation_builds_stock_and_variance_sheets(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result(tmp_path / "RESULT 10.07.24.xlsx")
    inventory_file = _create_inventory(tmp_path / "REVISION 10.07.24.xlsx")

    output_file = tmp_path / "STOCK_RECONCILIATION_10.07.24.xlsx"
    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    assert output_file.exists()
    assert summary["matched_rows"] == 2
    assert summary["unmatched_rows"] == 1

    wb = load_workbook(output_file, data_only=True)
    assert wb.sheetnames == ["STOCK_RESULT", "VARIANCE_REPORT"]

    stock = wb["STOCK_RESULT"]
    variance = wb["VARIANCE_REPORT"]

    expected_headers = [
        "SKU",
        "Наименование",
        "Начальный остаток",
        "Приход",
        "Возврат поставщику",
        "Продажи",
        "Расчетный остаток",
        "Фактический остаток",
        "Расхождение",
    ]
    assert [stock.cell(1, c).value for c in range(1, 10)] == expected_headers
    assert [variance.cell(1, c).value for c in range(1, 10)] == expected_headers

    stock_rows = [
        [stock.cell(r, c).value for c in range(1, 10)]
        for r in range(2, stock.max_row + 1)
    ]
    # SKU-1 aggregated from Item A + Item C.
    assert stock_rows[0] == ["SKU-1", "Master One", 150, 0, 30, 5, 115, 112, -3]
    assert stock_rows[1] == ["SKU-2", "Master Two", 200, 0, 35, 4, 161, 160, -1]

    variance_rows = [
        [variance.cell(r, c).value for c in range(1, 10)]
        for r in range(2, variance.max_row + 1)
    ]
    assert len(variance_rows) == 2
    # Sorted by absolute variance descending.
    assert variance_rows[0][0] == "SKU-1"
    assert variance_rows[0][-1] == -3
    assert variance_rows[1][0] == "SKU-2"
    assert variance_rows[1][-1] == -1


def test_inventory_reconciliation_is_stable_when_inventory_row_order_changes(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result(tmp_path / "RESULT 10.07.24.xlsx")

    output_base = tmp_path / "STOCK_RECONCILIATION_BASE.xlsx"
    output_shuffled = tmp_path / "STOCK_RECONCILIATION_SHUFFLED.xlsx"

    build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=_create_inventory(tmp_path / "REVISION_BASE.xlsx"),
        output_file=output_base,
    )
    build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=_create_inventory_shuffled(tmp_path / "REVISION_SHUFFLED.xlsx"),
        output_file=output_shuffled,
    )

    assert _stock_rows(output_base) == _stock_rows(output_shuffled)


def test_inventory_reconciliation_handles_missing_inventory_sku(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "RESULT"
    ws.append(["REPORT", "SVO Match Engine"])
    ws.append(["ARRIVAL_DATE", "10.07.24"])
    ws.append(["METADATA", "arrival_date=10.07.24"])
    ws.append([])
    ws.append(["SOURCE_NAME", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "SKU", "MASTER_NAME", "STATUS"])
    ws.append(["Ghost Item", "Cat", "Brand", "V0", "1 л", "SKU-404", "", "MATCH"])
    result_file = tmp_path / "RESULT_MISSING.xlsx"
    wb.save(result_file)

    inventory_file = _create_inventory(tmp_path / "REVISION_PRESENT.xlsx")
    output_file = tmp_path / "STOCK_RECONCILIATION_MISSING.xlsx"

    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    assert summary["missing_inventory_sku"] == 1
    assert summary["matched_rows"] == 1

    rows = _stock_rows(output_file)
    assert rows[0][0] == "SKU-404"
    assert rows[0][2:] == [0, 0, 0, 0, 0, 0, 0]


def test_inventory_reconciliation_reports_duplicate_sku_usage(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result(tmp_path / "RESULT_DUPLICATE.xlsx")
    inventory_file = _create_inventory(tmp_path / "REVISION_DUPLICATE.xlsx")
    output_file = tmp_path / "STOCK_RECONCILIATION_DUPLICATE.xlsx"

    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    # SKU-1 is matched by two RESULT rows (Item A and Item C), which should be detected.
    assert summary["duplicate_result_sku"] >= 1
    # The same SKU receives quantities from multiple inventory rows and is aggregated.
    assert summary["duplicate_inventory_sku"] >= 1


def test_inventory_reconciliation_prefers_exact_inventory_sku(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result(tmp_path / "RESULT_SKU_FIRST.xlsx")
    inventory_file = _create_inventory_with_sku(tmp_path / "REVISION_SKU.xlsx")
    output_file = tmp_path / "STOCK_RECONCILIATION_SKU_FIRST.xlsx"

    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    assert summary["matched_rows"] == 2
    rows = _stock_rows(output_file)
    row_by_sku = {row[0]: row for row in rows}
    # Inventory row source name is alias, but explicit SKU maps deterministically to SKU-2.
    assert row_by_sku["SKU-2"][2:9] == [100, 0, 20, 2, 78, 74, -4]


def test_inventory_reconciliation_uses_unique_name_fallback(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result(tmp_path / "RESULT_NAME_FALLBACK.xlsx")
    inventory_file = _create_inventory(tmp_path / "REVISION_NAME_FALLBACK.xlsx")
    output_file = tmp_path / "STOCK_RECONCILIATION_NAME_FALLBACK.xlsx"

    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    assert summary["ambiguous_inventory_source"] == 0
    assert summary["matched_rows"] == 2


def test_inventory_reconciliation_reports_ambiguous_name_without_silent_sku_choice(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result_ambiguous_name(tmp_path / "RESULT_AMBIGUOUS.xlsx")
    inventory_file = _create_inventory(tmp_path / "REVISION_AMBIGUOUS.xlsx")
    output_file = tmp_path / "STOCK_RECONCILIATION_AMBIGUOUS.xlsx"

    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    assert summary["ambiguous_inventory_source"] >= 1
    # No SKU should be picked silently from ambiguous source-name mapping.
    rows = _stock_rows(output_file)
    for row in rows:
        assert row[2:9] == [0, 0, 0, 0, 0, 0, 0]


def test_inventory_reconciliation_reports_missing_inventory_sku_link(tmp_path: Path):
    master_file = _create_master(tmp_path / "MASTER_TEST.xlsx")
    result_file = _create_result(tmp_path / "RESULT_MISSING_LINK.xlsx")
    inventory_file = _create_inventory_with_unknown_sku(tmp_path / "REVISION_UNKNOWN_SKU.xlsx")
    output_file = tmp_path / "STOCK_RECONCILIATION_UNKNOWN_SKU.xlsx"

    summary = build_inventory_reconciliation(
        master_file=master_file,
        result_file=result_file,
        inventory_file=inventory_file,
        output_file=output_file,
    )

    assert summary["inventory_without_result_sku"] >= 1
