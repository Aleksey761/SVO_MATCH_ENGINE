from pathlib import Path

from openpyxl import Workbook, load_workbook

from svo.dataset_builder import DatasetBuilder, DatasetRecord


def _create_master_workbook(path: Path, rows: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume"])
    for row in rows:
        ws.append([
            row.get("sku", ""),
            row.get("category", ""),
            row.get("brand", ""),
            row.get("variant", ""),
            row.get("volume", ""),
        ])
    wb.save(path)
    return path


def _create_match_workbook(path: Path, rows: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MATCH"
    ws.append(
        [
            "NAME",
            "W1",
            "W2",
            "W3",
            "W4",
            "W5",
            "W6",
            "W7",
            "W8",
            "W9",
            "W10",
            "MATCH_STATUS",
            "MATCH_SKU",
            "MATCH_MASTER_NAME",
            "MATCH_CONFIDENCE",
            "MATCH_REASONS",
        ]
    )

    for row in rows:
        warehouse = row.get("warehouse", [0] * 10)
        ws.append(
            [row.get("name", "ITEM")]
            + warehouse
            + [
                row.get("status", "MATCH"),
                row.get("sku", ""),
                row.get("master_name", ""),
                row.get("confidence", 100),
                row.get("reasons", ""),
            ]
        )

    wb.save(path)
    return path


def test_dataset_record_total_movement():
    record = DatasetRecord(
        sku="SKU-1",
        master_name="M1",
        category="Cat",
        brand="Brand",
        aroma="Aroma",
        arrival_quantity=12,
        sales_quantity=8,
    )
    assert record.total_movement == 20


def test_dataset_builder_dataset_totals(tmp_path: Path):
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()

    _create_master_workbook(
        data_dir / "MASTER.xlsx",
        [
            {"sku": "SKU-001", "category": "C1", "brand": "B1", "variant": "V1", "volume": "1L"},
            {"sku": "SKU-002", "category": "C2", "brand": "B2", "variant": "V2", "volume": "2L"},
            {"sku": "SKU-003", "category": "C3", "brand": "B3", "variant": "V3", "volume": "3L"},
        ],
    )
    _create_match_workbook(
        output_dir / "ARRIVAL_MATCH_2026-07-15.xlsx",
        [
            {"sku": "SKU-001", "warehouse": [5, 5, 0, 0, 0, 0, 0, 0, 0, 0]},
            {"sku": "SKU-002", "warehouse": [3, 2, 0, 0, 0, 0, 0, 0, 0, 0]},
        ],
    )
    _create_match_workbook(
        output_dir / "SALES_MATCH_2026-07-16.xlsx",
        [
            {"sku": "SKU-002", "warehouse": [4, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
            {"sku": "SKU-003", "warehouse": [1, 1, 0, 0, 0, 0, 0, 0, 0, 0]},
        ],
    )

    result = DatasetBuilder().build(
        input_dir=data_dir,
        output_file=output_dir / "MASTER_DATASET.xlsx",
        print_reports=False,
    )

    validation = result["validation"]
    summary = result["summary"]

    assert validation["master_records"] == 3
    assert validation["dataset_records"] == 3
    assert validation["arrival_linked"] == 2
    assert validation["sales_linked"] == 2
    assert validation["missing_master_sku"] == 0
    assert validation["duplicate_sku"] == 0
    assert validation["integrity"] == "OK"

    assert summary["arrival_qty"] == 15
    assert summary["sales_qty"] == 6
    assert summary["master_sku"] == 3
    assert summary["arrival_linked"] == 2
    assert summary["sales_linked"] == 2
    assert round(summary["coverage_pct"], 2) == 100.00


def test_duplicate_sku_detection(tmp_path: Path):
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()

    _create_master_workbook(
        data_dir / "MASTER.xlsx",
        [
            {"sku": "SKU-001", "category": "C1", "brand": "B1", "variant": "V1", "volume": "1L"},
            {"sku": "SKU-001", "category": "C1", "brand": "B1", "variant": "V1", "volume": "1L"},
            {"sku": "SKU-002", "category": "C2", "brand": "B2", "variant": "V2", "volume": "2L"},
        ],
    )

    result = DatasetBuilder().build(
        input_dir=data_dir,
        output_file=output_dir / "MASTER_DATASET.xlsx",
        print_reports=False,
    )

    assert result["validation"]["duplicate_sku"] == 1
    assert result["validation"]["integrity"] == "FAILED"


def test_missing_master_sku_detection(tmp_path: Path):
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()

    _create_master_workbook(
        data_dir / "MASTER.xlsx",
        [{"sku": "SKU-001", "category": "C1", "brand": "B1", "variant": "V1", "volume": "1L"}],
    )

    _create_match_workbook(
        output_dir / "ARRIVAL_MATCH_2026-07-15.xlsx",
        [{"sku": "SKU-999", "warehouse": [5, 0, 0, 0, 0, 0, 0, 0, 0, 0]}],
    )

    result = DatasetBuilder().build(
        input_dir=data_dir,
        output_file=output_dir / "MASTER_DATASET.xlsx",
        print_reports=False,
    )

    assert result["validation"]["missing_master_sku"] == 1
    assert result["validation"]["integrity"] == "FAILED"


def test_coverage_calculation(tmp_path: Path):
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()

    _create_master_workbook(
        data_dir / "MASTER.xlsx",
        [
            {"sku": "SKU-001", "category": "C1", "brand": "B1", "variant": "V1", "volume": "1L"},
            {"sku": "SKU-002", "category": "C2", "brand": "B2", "variant": "V2", "volume": "2L"},
            {"sku": "SKU-003", "category": "C3", "brand": "B3", "variant": "V3", "volume": "3L"},
            {"sku": "SKU-004", "category": "C4", "brand": "B4", "variant": "V4", "volume": "4L"},
        ],
    )

    _create_match_workbook(
        output_dir / "ARRIVAL_MATCH_2026-07-15.xlsx",
        [
            {"sku": "SKU-001", "warehouse": [1, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
            {"sku": "SKU-004", "warehouse": [1, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
        ],
    )

    result = DatasetBuilder().build(
        input_dir=data_dir,
        output_file=output_dir / "MASTER_DATASET.xlsx",
        print_reports=False,
    )

    assert round(result["summary"]["coverage_pct"], 2) == 50.00


def test_load_dataset_from_workbook(tmp_path: Path):
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()

    _create_master_workbook(
        data_dir / "MASTER.xlsx",
        [{"sku": "SKU-001", "category": "Shampoo", "brand": "SVO", "variant": "AQUA", "volume": "1L"}],
    )
    _create_match_workbook(
        output_dir / "ARRIVAL_MATCH_2026-07-15.xlsx",
        [{"sku": "SKU-001", "warehouse": [2, 0, 0, 0, 0, 0, 0, 0, 0, 0]}],
    )

    output_file = output_dir / "MASTER_DATASET.xlsx"
    DatasetBuilder().build(input_dir=data_dir, output_file=output_file, print_reports=False)

    records = DatasetBuilder().load_dataset(output_file)
    assert len(records) == 1
    assert records[0].sku == "SKU-001"
    assert records[0].arrival_quantity == 2
    assert records[0].last_arrival_date == "15.07.2026"

    wb = load_workbook(output_file, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == DatasetBuilder.HEADERS
