from pathlib import Path

from openpyxl import Workbook, load_workbook

from runner import _write_master_dataset_from_master
from svo.loader import Loader


def _create_master_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume"])
    ws.append(["SKU-001", "Shampoo", "SVO", "AQUA", "1 L"])
    ws.append(["SKU-002", "Conditioner", "SVO", "ROSE", "2,7 L"])
    wb.save(path)
    return path


def test_write_master_dataset_is_loader_compatible(tmp_path: Path):
    master_file = _create_master_source(tmp_path / "MASTER.xlsx")
    dataset_file = tmp_path / "MASTER_DATASET.xlsx"

    written = _write_master_dataset_from_master(master_file, dataset_file)

    assert written == 2
    assert dataset_file.exists()

    wb = load_workbook(dataset_file, data_only=True)
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, 11)]
    assert headers == ["SKU", "CATEGORY", "BRAND", "VARIANT", "VOLUME", "AROMA", "C7", "C8", "C9", "MASTER_NAME"]

    records = Loader().load_master(dataset_file)
    assert len(records) == 2
    assert records[0].sku == "SKU-001"
    assert records[0].master_name == "Shampoo SVO AQUA 1 L"
    assert records[1].sku == "SKU-002"
    assert records[1].master_name == "Conditioner SVO ROSE 2,7 L"
