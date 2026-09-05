from pathlib import Path
import json

from openpyxl import Workbook, load_workbook

from svo.dashboard_builder import DashboardBuilder, RunContext


def _write_price_match(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matched"
    ws.append([
        "SKU",
        "MASTER_NAME",
        "SupplierArticle",
        "OriginalProductName",
        "CanonicalProductName",
        "UnitCost",
        "RetailPrice",
        "StockQty",
        "SourceUnit",
        "PackQty",
        "OriginalPrice",
        "ImportedAt",
        "SourceFile",
        "Status",
    ])
    ws.append(["SKU-1", "Master One", "A-1", "P1", "P1", 10, 20, 1, None, None, 10, "2026-08-05T10:11:12+00:00", "PRICE.xlsx", "MATCH"])
    ws.append(["SKU-2", "Master Two", "A-2", "P2", "P2", 10, 20, 1, None, None, 10, "2026-08-05T10:11:12+00:00", "PRICE.xlsx", "MATCH"])
    wb.save(path)


def _write_price_review(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE_REVIEW"
    ws.append(["SupplierArticle", "OriginalProductName", "CanonicalProductName", "SKU", "MASTER_NAME", "Reason", "Recommendation"])
    ws.append(["A-10", "N1", "N1", "", "", "No MASTER match", "Manual review required"])
    ws.append(["A-11", "N2", "N2", "", "", "Ambiguous match", "Manual review required"])
    ws.append(["A-12", "N3", "N3", "", "", "ProductType mismatch", "Manual review required"])
    ws.append(["A-13", "N4", "N4", "", "", "Manual review required", "Manual review required"])
    wb.save(path)


def _write_price_change(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE_CHANGES"
    ws.append(["SKU", "MASTER_NAME", "PreviousPrice", "NewPrice", "Difference", "PercentChange", "SourceUnit", "PackQty", "OriginalPrice", "ValidFrom", "ImportedAt", "SourceFile", "Status"])
    ws.append(["SKU-1", "Master One", None, 10, None, None, None, None, 10, "2026-08-05T10:11:12+00:00", "2026-08-05T10:11:12+00:00", "PRICE.xlsx", "NEW"])
    ws.append(["SKU-1", "Master One", 10, 11, 1, 10, None, None, 11, "2026-08-06T10:11:12+00:00", "2026-08-06T10:11:12+00:00", "PRICE.xlsx", "UPDATED"])
    ws.append(["SKU-1", "Master One", 11, 11, 0, 0, None, None, 11, "2026-08-07T10:11:12+00:00", "2026-08-07T10:11:12+00:00", "PRICE.xlsx", "SKIPPED"])
    wb.save(path)


def _write_master_dataset(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER_DATASET"
    ws.append(["SKU", "MASTER_NAME", "CATEGORY", "BRAND", "ARRIVAL_QUANTITY", "SALES_QUANTITY", "NET_MOVEMENT", "ARRIVAL_COUNT"])
    ws.append(["SKU-1", "Master One", "Cat", "Brand", 1, 2, -1, 1])
    ws.append(["SKU-2", "Master Two", "Cat", "Brand", 1, 2, -1, 1])
    ws.append(["SKU-3", "Master Three", "Cat", "Brand", 1, 2, -1, 1])
    wb.save(path)


def _write_master_audit(path: Path) -> None:
    path.write_text('{"gate":"master","status":"PASSED","errors":0}', encoding="utf-8")


def _write_match_analytics(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "========================================",
                "MATCH ANALYTICS",
                "----------------------------------------",
                "DATE              : 06.04.2026",
                "FILE              : PRICE.xlsx",
                "TOTAL ROWS        : 10",
                "MATCH             : 9",
                "REVIEW            : 1",
                "COVERAGE          : 90.00%",
                "----------------------------------------",
                "NEW PRODUCTS FOR MASTER",
                "- ROW 10: New product [NO_MATCH]",
                "----------------------------------------",
            ]
        ),
        encoding="utf-8",
    )


def _sheet_as_dict(ws) -> dict[str, object]:
    data: dict[str, object] = {}
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row=row, column=1).value
        value = ws.cell(row=row, column=2).value
        if key is not None:
            data[str(key)] = value
    return data


def _build_run_context(out: Path, *, status: str = "SUCCESS") -> RunContext:
    return RunContext(
        master_dataset_path=out / "MASTER_DATASET.xlsx",
        price_match_path=out / "price_match.xlsx",
        price_review_path=out / "price_review.xlsx",
        price_change_report_path=out / "price_change_report.xlsx",
        business_metrics_path=out / "BUSINESS_METRICS.xlsx",
        match_analytics_path=out / "MATCH_ANALYTICS.txt",
        master_audit_path=out / "MASTER_AUDIT.json",
        output_directory=out,
        run_date="2026-08-05",
        run_time="10:11:12",
        duration="5 sec",
        source_master="MASTER.xlsx",
        source_price="PRICE.xlsx",
        version="runner-full",
        status=status,
    )


def test_dashboard_builder_reads_existing_files_and_populates_sheets(tmp_path: Path):
    out = tmp_path / "output"
    out.mkdir(parents=True, exist_ok=True)

    _write_price_match(out / "price_match.xlsx")
    _write_price_review(out / "price_review.xlsx")
    _write_price_change(out / "price_change_report.xlsx")
    _write_master_dataset(out / "MASTER_DATASET.xlsx")
    _write_master_audit(out / "MASTER_AUDIT.json")
    _write_match_analytics(out / "MATCH_ANALYTICS.txt")

    dashboard_path = DashboardBuilder(_build_run_context(out)).build()

    assert dashboard_path.exists()
    wb = load_workbook(dashboard_path, data_only=True)
    assert wb.sheetnames == ["Summary", "Review", "Price", "Run"]

    summary = _sheet_as_dict(wb["Summary"])
    assert summary["MASTER rows"] == 3
    assert summary["PRICE rows"] == 2
    assert summary["Matched"] == 2
    assert summary["Review"] == 4
    assert summary["New products"] == 1
    assert summary["Updated prices"] == 1
    assert summary["Skipped"] == 1
    assert summary["Coverage %"] == "100.00 %"
    assert summary["Run date"] == "2026-08-05"
    assert summary["Run time"] == "10:11:12"
    assert summary["Validation"] is None
    assert summary["Status 1"] == "✓ Dashboard created"
    assert summary["Status 2"] == "✓ Sources loaded"
    assert summary["Status 3"] == "⚠ Metrics mismatch"

    review = _sheet_as_dict(wb["Review"])
    assert review["NO_MATCH"] == 1
    assert review["MULTIPLE_CANDIDATES"] == 1
    assert review["PRODUCT_TYPE_MISMATCH"] == 1
    assert review["MANUAL_REVIEW"] == 1

    price = _sheet_as_dict(wb["Price"])
    assert price["New prices"] == 1
    assert price["Updated prices"] == 1
    assert price["Unchanged prices"] == 1

    run = _sheet_as_dict(wb["Run"])
    assert run["Source MASTER"] == "MASTER.xlsx"
    assert run["Source PRICE"] == "PRICE.xlsx"
    assert run["Version"] == "runner-full"
    assert run["Duration"] == "5 sec"
    assert run["Status"] == "SUCCESS"

    audit_payload = json.loads((out / "MASTER_AUDIT.json").read_text(encoding="utf-8"))
    assert audit_payload["dashboard_warnings"] == [
        "Dashboard metrics mismatch: Matched(2) + Review(4) + Skipped(1) != PRICE rows(2)"
    ]


def test_dashboard_builder_handles_missing_files_with_na(tmp_path: Path, caplog):
    out = tmp_path / "output"
    out.mkdir(parents=True, exist_ok=True)

    caplog.set_level("DEBUG")
    missing_context = _build_run_context(out, status="N/A")
    dashboard_path = DashboardBuilder(missing_context).build()

    assert dashboard_path.exists()
    wb = load_workbook(dashboard_path, data_only=True)

    summary = _sheet_as_dict(wb["Summary"])
    assert summary["MASTER rows"] == "N/A"
    assert summary["PRICE rows"] == "N/A"
    assert summary["Coverage %"] == "N/A"
    assert summary["Status 2"] == "⚠ Sources missing"
    assert summary["Status 3"] == "✓ Metrics verified"

    review = _sheet_as_dict(wb["Review"])
    assert review["NO_MATCH"] == "N/A"

    price = _sheet_as_dict(wb["Price"])
    assert price["New prices"] == "N/A"

    run = _sheet_as_dict(wb["Run"])
    assert run["Status"] == "N/A"

    warnings = "\n".join(record.getMessage() for record in caplog.records)
    assert "Dashboard source file is missing" in warnings
    assert "Dashboard sources" in warnings
