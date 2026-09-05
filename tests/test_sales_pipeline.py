from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from svo.engine import Engine


@pytest.fixture(autouse=True)
def _isolate_output_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.chdir(tmp_path)


def _create_master_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-1", "Шампунь", "SVO", "AQUA", "1 л", "Шампунь SVO AQUA 1 л"])
    wb.save(path)
    return path


def _create_sales_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append([None, None, None, None, None, None, None, None])
    ws.append([None, "Наименование", None, "Воронеж", None, None, "Краснодар 1", "Краснодар 2"])
    ws.append([None, "SVO SHAMPUN AQUA 1 л", None, 5, None, None, 3, 2])
    wb.save(path)
    return path


def test_run_sales_produces_sales_match_output(tmp_path: Path):
    master_file = _create_master_workbook(tmp_path / "MASTER.xlsx")
    sales_file = _create_sales_workbook(tmp_path / "SALES_2026-07-15.xlsx")

    result = Engine().run_sales(
        master_file=master_file,
        sales_file=sales_file,
        output_file=tmp_path / "SALES_MATCH_15.07.2026.xlsx",
    )

    assert result["document"] == "SALES"
    assert result["rows"] == 1
    assert result["match"] == 1
    assert result["review"] == 0
    assert result["sales_date"] == "15.07.2026"
    assert result["output"] == str(tmp_path / "SALES_MATCH_15.07.2026.xlsx")
    assert Path(result["output"]).exists()

    wb = load_workbook(tmp_path / "SALES_MATCH_15.07.2026.xlsx")
    ws = wb.active
    assert ws.title == "SALES"
    headers = [str(value or "").strip().upper() for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    assert "НАИМЕНОВАНИЕ" in headers
    business_sku_col = headers.index("SKU") + 1
    status_col = headers.index("MATCH_STATUS") + 1
    match_sku_col = headers.index("MATCH_SKU") + 1
    assert ws.cell(row=2, column=status_col).value == "MATCH"
    assert ws.cell(row=2, column=business_sku_col).value == "SKU-1"
    assert ws.cell(row=2, column=match_sku_col).value == "SKU-1"

    review_report = load_workbook(tmp_path / "REVIEW_REPORT.xlsx")
    review_ws = review_report.active
    assert review_ws.title == "REVIEW_REPORT"
    assert review_ws.max_row == 1
    assert result["review_report"] == str(tmp_path / "REVIEW_REPORT.xlsx")
    assert result["review_report_rows"] == 0


def test_run_sales_writes_review_diagnostics_report(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-1", "Шампунь", "SVO", "AQUA", "1 л", "Шампунь SVO AQUA 1 л"])
    ws.append(["SKU-2", "Шампунь", "SVO", "LIME", "1 л", "Шампунь SVO LIME 1 л"])
    master_file = tmp_path / "MASTER.xlsx"
    wb.save(master_file)

    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append(["НАИМЕНОВАНИЕ", "SKU", "Тип товара", "Бренд", "Variant", "Объем"])
    ws.append(["SVO SHAMPUN 1 л", None, None, None, None, None])
    sales_file = tmp_path / "SALES_2026-07-15.xlsx"
    wb.save(sales_file)

    result = Engine().run_sales(
        master_file=master_file,
        sales_file=sales_file,
        output_file=tmp_path / "SALES_MATCH_15.07.2026.xlsx",
    )

    assert result["review"] == 1
    assert result["review_report"] == str(tmp_path / "REVIEW_REPORT.xlsx")
    assert result["review_report_rows"] == 1

    review_report = load_workbook(tmp_path / "REVIEW_REPORT.xlsx")
    review_ws = review_report.active
    assert review_ws.title == "REVIEW_REPORT"
    assert review_ws.max_row == 2
    assert review_ws.cell(row=1, column=1).value == "ROW_NUMBER"
    assert review_ws.cell(row=1, column=2).value == "SOURCE_NAME"
    assert review_ws.cell(row=1, column=3).value == "BEST_MASTER_CANDIDATE"
    assert review_ws.cell(row=1, column=4).value == "SCORE"
    assert review_ws.cell(row=1, column=5).value == "REVIEW_REASONS"
    assert review_ws.cell(row=2, column=2).value == "SVO SHAMPUN 1 л"
    assert review_ws.cell(row=2, column=3).value == "SKU-1 | Шампунь | SVO | AQUA | 1 л"
    assert float(review_ws.cell(row=2, column=4).value) > 100.0
    assert review_ws.cell(row=2, column=5).value == "MULTIPLE_MATCH"


def test_run_sales_finalizes_result_order_by_master_and_moves_review_last(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-2", "Шампунь", "SVO", "LIME", "1 л", "Шампунь SVO LIME 1 л"])
    ws.append(["SKU-1", "Шампунь", "SVO", "AQUA", "1 л", "Шампунь SVO AQUA 1 л"])
    master_file = tmp_path / "MASTER.xlsx"
    wb.save(master_file)

    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append(["НАИМЕНОВАНИЕ", "SKU", "Тип товара", "Бренд", "Variant", "Объем"])
    ws.append(["SVO SHAMPUN AQUA 1 л", None, None, None, None, None])
    ws.append(["SVO SHAMPUN 1 л", None, None, None, None, None])
    ws.append(["SVO SHAMPUN LIME 1 л", None, None, None, None, None])
    sales_file = tmp_path / "SALES_2026-07-15.xlsx"
    wb.save(sales_file)

    Engine().run_sales(
        master_file=master_file,
        sales_file=sales_file,
        output_file=tmp_path / "SALES_MATCH_15.07.2026.xlsx",
    )

    wb = load_workbook(tmp_path / "SALES_MATCH_15.07.2026.xlsx")
    ws = wb.active
    headers = [str(value or "").strip().upper() for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    assert "НАИМЕНОВАНИЕ" in headers
    status_col = headers.index("MATCH_STATUS") + 1
    business_sku_col = headers.index("SKU") + 1
    sku_col = headers.index("MATCH_SKU") + 1

    ordered_rows = []
    for row_number in range(2, ws.max_row + 1):
        status = str(ws.cell(row=row_number, column=status_col).value or "").strip().upper()
        if status in {"MATCH", "REVIEW"}:
            ordered_rows.append((
                status,
                str(ws.cell(row=row_number, column=sku_col).value or "").strip(),
                ws.cell(row=row_number, column=business_sku_col).value,
            ))

    assert ordered_rows[0][0] == "MATCH"
    assert ordered_rows[0][1] == "SKU-2"
    assert ordered_rows[0][2] == "SKU-2"

    assert ordered_rows[1][0] == "MATCH"
    assert ordered_rows[1][1] == "SKU-1"
    assert ordered_rows[1][2] == "SKU-1"

    assert ordered_rows[2][0] == "REVIEW"


def test_run_sales_preserves_source_name_for_review_rows(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-1", "Шампунь", "SVO", "AQUA", "1 л", "Шампунь SVO AQUA 1 л"])
    ws.append(["SKU-2", "Шампунь", "SVO", "LIME", "1 л", "Шампунь SVO LIME 1 л"])
    master_file = tmp_path / "MASTER.xlsx"
    wb.save(master_file)

    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append(["НАИМЕНОВАНИЕ", "SKU", "Тип товара", "Бренд", "Variant", "Объем"])
    source_name = "SVO SHAMPUN 1 л"
    ws.append([source_name, None, None, None, None, None])
    sales_file = tmp_path / "SALES_2026-07-15.xlsx"
    wb.save(sales_file)

    Engine().run_sales(
        master_file=master_file,
        sales_file=sales_file,
        output_file=tmp_path / "SALES_MATCH_15.07.2026.xlsx",
    )

    wb = load_workbook(tmp_path / "SALES_MATCH_15.07.2026.xlsx", data_only=True)
    ws = wb.active
    headers = [str(value or "").strip().upper() for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    name_col = headers.index("НАИМЕНОВАНИЕ") + 1
    status_col = headers.index("MATCH_STATUS") + 1
    match_sku_col = headers.index("MATCH_SKU") + 1
    match_master_col = headers.index("MATCH_MASTER_NAME") + 1
    match_conf_col = headers.index("MATCH_CONFIDENCE") + 1
    match_reasons_col = headers.index("MATCH_REASONS") + 1

    assert ws.cell(row=2, column=status_col).value == "REVIEW"
    assert ws.cell(row=2, column=name_col).value == source_name
    assert ws.cell(row=2, column=match_sku_col).value in (None, "")
    assert ws.cell(row=2, column=match_master_col).value in (None, "")
    assert ws.cell(row=2, column=match_conf_col).value in (None, "")
    assert ws.cell(row=2, column=match_reasons_col).value == "MULTIPLE_MATCH"
