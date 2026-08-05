from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from svo.engine import Engine


def _write_master(path: Path, rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append([
        "SKU",
        "CATEGORY",
        "BRAND",
        "VARIANT",
        "VOLUME",
        "AROMA",
        "X1",
        "X2",
        "X3",
        "MASTER_NAME",
    ])
    for row in rows:
        ws.append(
            [
                row["sku"],
                row["category"],
                row["brand"],
                row["variant"],
                row["volume"],
                row.get("aroma", row["variant"]),
                "",
                "",
                "",
                row["master_name"],
            ]
        )
    wb.save(path)
    return path


def _write_price(path: Path, rows: list[list[object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE"
    ws.append(["ProductName", "UnitCost", "RetailPrice", "SourceUnit"])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _single_master(tmp_path: Path) -> Path:
    return _write_master(
        tmp_path / "MASTER_TEST.xlsx",
        [
            {
                "sku": "SKU-APPLE",
                "category": "Шампунь",
                "brand": "SVO",
                "variant": "APPLE",
                "volume": "1 Л",
                "master_name": "Shampoo SVO APPLE 1 L",
            }
        ],
    )


def test_price_history_integration_first_import(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    price_file = _write_price(tmp_path / "PRICE_1.xlsx", [["SVO shampoo apple 1 l", 100, 150, "шт"]])

    result = engine.run_price_matching(master_file=master_file, price_file=price_file)

    assert result["price_history_created"] == 1
    assert result["price_history_updated"] == 0
    assert result["price_history_skipped"] == 0
    assert len(engine.price_history.records) == 1
    assert engine.price_history.records[0].Status == "ACTIVE"


def test_price_history_integration_second_identical_import(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    first_price = _write_price(tmp_path / "PRICE_1.xlsx", [["SVO shampoo apple 1 l", 100, 150, "шт"]])
    second_price = _write_price(tmp_path / "PRICE_2.xlsx", [["SVO shampoo apple 1 l", 100, 150, "шт"]])

    engine.run_price_matching(master_file=master_file, price_file=first_price)
    result = engine.run_price_matching(master_file=master_file, price_file=second_price)

    assert result["price_history_created"] == 0
    assert result["price_history_updated"] == 0
    assert result["price_history_skipped"] == 1
    assert len(engine.price_history.records) == 1
    assert engine.price_history.records[0].Status == "ACTIVE"


def test_price_history_integration_changed_price_closes_previous(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    first_price = _write_price(tmp_path / "PRICE_1.xlsx", [["SVO shampoo apple 1 l", 100, 150, "шт"]])
    second_price = _write_price(tmp_path / "PRICE_2.xlsx", [["SVO shampoo apple 1 l", 120, 150, "шт"]])

    engine.run_price_matching(master_file=master_file, price_file=first_price)
    result = engine.run_price_matching(master_file=master_file, price_file=second_price)

    assert result["price_history_created"] == 1
    assert result["price_history_updated"] == 1
    assert result["price_history_skipped"] == 0

    assert len(engine.price_history.records) == 2
    assert engine.price_history.records[0].Status == "CLOSED"
    assert engine.price_history.records[1].Status == "ACTIVE"


def test_price_history_integration_package_import(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    price_file = _write_price(tmp_path / "PRICE_PACK.xlsx", [["SVO shampoo apple 1 l", 1566, 1800, "упак (9 шт)"]])

    engine.run_price_matching(master_file=master_file, price_file=price_file)

    record = engine.price_history.records[0]
    assert record.UnitCost == Decimal("174")
    assert record.OriginalPrice == Decimal("1566")
    assert record.PackQty == 9


def test_price_change_report_generation(tmp_path: Path):
    engine = Engine()
    master_file = _single_master(tmp_path)
    first_price = _write_price(tmp_path / "PRICE_1.xlsx", [["SVO shampoo apple 1 l", 100, 150, "шт"]])
    second_price = _write_price(tmp_path / "PRICE_2.xlsx", [["SVO shampoo apple 1 l", 120, 150, "шт"]])

    engine.run_price_matching(master_file=master_file, price_file=first_price)
    result = engine.run_price_matching(master_file=master_file, price_file=second_price)

    report_path = Path(result["price_change_report"])
    assert report_path.exists()

    wb = load_workbook(report_path, data_only=True)
    ws = wb.active

    headers = [ws.cell(row=1, column=col).value for col in range(1, 14)]
    assert headers == [
        "SKU",
        "MASTER_NAME",
        "PreviousPrice",
        "NewPrice",
        "Difference",
        "PercentChange",
        "SourceUnit",
        "PackQty",
        "OriginalPrice",
        "ValidFrom",
        "ImportedAt",
        "SourceFile",
        "Status",
    ]

    assert ws.max_row == 2
    assert ws.cell(row=2, column=1).value == "SKU-APPLE"
    assert Decimal(str(ws.cell(row=2, column=3).value)) == Decimal("100")
    assert Decimal(str(ws.cell(row=2, column=4).value)) == Decimal("120")
    assert Decimal(str(ws.cell(row=2, column=5).value)) == Decimal("20")
    assert ws.cell(row=2, column=13).value == "UPDATED"
