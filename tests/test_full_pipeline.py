from pathlib import Path

from openpyxl import Workbook, load_workbook

import runner


def _create_master(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["SKU", "Category", "Brand", "Variant", "Volume", "MASTER_NAME"])
    ws.append(["SKU-APPLE", "Шампунь", "SVO", "APPLE", "1 л", "Шампунь SVO APPLE 1 л"])
    wb.save(path)
    return path


def _create_price(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRICE"
    ws.append(["Supplier article", "Product name", "Free stock", "Distributor price"])
    ws.append(["A-100", "SVO shampoo apple 1 l", 42, "35,50"])
    wb.save(path)
    return path


def _create_sales(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES"
    ws.append(["НАИМЕНОВАНИЕ", "SKU", "Тип товара", "Бренд", "Variant", "Объем"])
    ws.append(["SVO SHAMPUN APPLE 1 л", None, None, None, None, None])
    wb.save(path)
    return path


def _create_broken_inventory(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "REVISION"
    ws.append(["WRONG", "HEADERS", "ONLY"])
    ws.append(["x", "y", "z"])
    wb.save(path)
    return path


def test_full_pipeline_generates_outputs_summary_and_errors(tmp_path: Path, monkeypatch):
    master = _create_master(tmp_path / "MASTER.xlsx")
    price = _create_price(tmp_path / "PRICE.xlsx")
    sales = _create_sales(tmp_path / "SALES_2026-07-15.xlsx")
    inventory = _create_broken_inventory(tmp_path / "REVISION_2026-07-15.xlsx")

    monkeypatch.chdir(tmp_path)

    exit_code = runner.main(
        [
            "--mode",
            "full",
            "--master",
            str(master),
            "--price",
            str(price),
            "--sales",
            str(sales),
            "--inventory",
            str(inventory),
        ]
    )

    assert exit_code == 0

    output_dir = tmp_path / "output"
    expected_files = [
        output_dir / "MASTER_MATCH.xlsx",
        output_dir / "PRICE_MATCH.xlsx",
        output_dir / "PRICE_HISTORY.xlsx",
        output_dir / "PRICE_CHANGE_REPORT.xlsx",
        output_dir / "PRICE_REVIEW.xlsx",
        output_dir / "SALES_MATCH.xlsx",
        output_dir / "STOCK_MATCH.xlsx",
        output_dir / "FINAL_REPORT.xlsx",
        output_dir / "ERRORS.xlsx",
        output_dir / "RUN_SUMMARY.txt",
    ]
    for file_path in expected_files:
        assert file_path.exists(), f"Missing file: {file_path}"

    summary_text = (output_dir / "RUN_SUMMARY.txt").read_text(encoding="utf-8")
    assert "MASTER" in summary_text
    assert "PRICE" in summary_text
    assert "SALES" in summary_text
    assert "STOCK" in summary_text
    assert "Execution time" in summary_text

    errors_wb = load_workbook(output_dir / "ERRORS.xlsx", data_only=True)
    errors_ws = errors_wb.active
    headers = [errors_ws.cell(row=1, column=idx).value for idx in range(1, 6)]
    assert headers == ["Module", "Row", "Object", "Reason", "Details"]
    assert errors_ws.max_row >= 2

    modules = [
        str(errors_ws.cell(row=row, column=1).value or "").strip().upper()
        for row in range(2, errors_ws.max_row + 1)
    ]
    assert "STOCK" in modules

    final_wb = load_workbook(output_dir / "FINAL_REPORT.xlsx", data_only=True)
    assert final_wb.sheetnames == [
        "Dashboard",
        "Price Changes",
        "Review",
        "New Products",
        "Sales",
        "Stock",
        "Analytics",
        "Audit",
    ]

    dashboard = final_wb["Dashboard"]
    assert dashboard.cell(row=1, column=1).value == "KPI"
    assert dashboard.cell(row=1, column=2).value == "Value"
