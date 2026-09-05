from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook


INPUT_MATCH = Path("output/SALES_MATCH.xlsx")
INPUT_MANUAL = Path("output/SALES_MATCH_MANUAL.xlsx")
INPUT_FINAL = Path("output/SALES_MATCH_FINAL.xlsx")
INPUT_MAP_REVIEWED = Path("output/SALES_MANUAL_MAPPING_REVIEWED.xlsx")
INPUT_MAP_REVIEW12 = Path("output/SALES_REVIEW_12_MANUAL_MAPPING.xlsx")
OUTPUT_PATH = Path("output/SALES_FINAL_CONTROL_CHECK.xlsx")


def norm_header(value: object) -> str:
    text = "" if value is None else str(value)
    cleaned = "".join(ch for ch in text.upper().strip() if ch.isalnum())
    return cleaned


def find_column(headers: Sequence[object], aliases: Iterable[str]) -> Optional[int]:
    alias_norm = {norm_header(a) for a in aliases}
    for idx, value in enumerate(headers, start=1):
        if norm_header(value) in alias_norm:
            return idx
    return None


def read_table(path: Path) -> Dict[str, object]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    headers = [ws.cell(1, c).value for c in range(1, max_col + 1)]
    source_row_col = find_column(headers, ["SOURCE_ROW", "SOURCE ROW", "SOURCEROW"])
    source_name_col = find_column(headers, ["SOURCE_NAME", "SOURCE NAME", "NAME", "НАИМЕНОВАНИЕ"])
    match_status_col = find_column(headers, ["MATCH_STATUS", "STATUS"])
    match_sku_col = find_column(headers, ["MATCH_SKU", "SKU", "MATCHED_SKU"])
    match_master_name_col = find_column(headers, ["MATCH_MASTER_NAME", "MASTER_NAME", "MATCH_NAME"])
    match_reasons_col = find_column(headers, ["MATCH_REASONS", "MATCH_REASON", "REASONS"])

    rows: List[Dict[str, object]] = []
    source_rows_resolved: List[int] = []

    for r in range(2, max_row + 1):
        row_values = [ws.cell(r, c).value for c in range(1, max_col + 1)]
        if all(v is None or str(v).strip() == "" for v in row_values):
            continue

        if source_row_col is not None:
            raw_source_row = ws.cell(r, source_row_col).value
            try:
                source_row = int(raw_source_row) if raw_source_row is not None and str(raw_source_row).strip() != "" else r
            except Exception:
                source_row = r
        else:
            source_row = r

        source_rows_resolved.append(source_row)

        rows.append(
            {
                "sheet_row": r,
                "source_row": source_row,
                "source_name": ws.cell(r, source_name_col).value if source_name_col else None,
                "match_status": ws.cell(r, match_status_col).value if match_status_col else None,
                "match_sku": ws.cell(r, match_sku_col).value if match_sku_col else None,
                "match_master_name": ws.cell(r, match_master_name_col).value if match_master_name_col else None,
                "match_reasons": ws.cell(r, match_reasons_col).value if match_reasons_col else None,
                "row_values": row_values,
            }
        )

    source_counter = Counter(source_rows_resolved)
    duplicates = sorted((k, v) for k, v in source_counter.items() if v > 1)

    return {
        "path": str(path),
        "headers": headers,
        "rows": rows,
        "data_row_count": len(rows),
        "source_row_min": min(source_rows_resolved) if source_rows_resolved else None,
        "source_row_max": max(source_rows_resolved) if source_rows_resolved else None,
        "source_row_unique_count": len(source_counter),
        "source_row_duplicates": duplicates,
        "has_explicit_source_row": source_row_col is not None,
    }


def get_manual_map_source_rows(path: Path) -> set[int]:
    wb = load_workbook(path, data_only=True)
    approved: set[int] = set()
    for ws in wb.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        if max_row < 2 or max_col == 0:
            continue
        headers = [ws.cell(1, c).value for c in range(1, max_col + 1)]
        source_row_col = find_column(headers, ["SOURCE_ROW", "SOURCE ROW", "SOURCEROW"])
        decision_col = find_column(headers, ["DECISION", "STATUS", "ACTION"])
        if source_row_col is None or decision_col is None:
            continue
        for r in range(2, max_row + 1):
            decision = ws.cell(r, decision_col).value
            if decision is None:
                continue
            decision_norm = str(decision).strip().upper()
            if decision_norm not in {"APPROVE", "NO_MASTER"}:
                continue
            source_raw = ws.cell(r, source_row_col).value
            try:
                source_row = int(source_raw)
            except Exception:
                continue
            approved.add(source_row)
    return approved


def row_signature(row: Dict[str, object], compare_len: int) -> Tuple[object, ...]:
    values = row["row_values"]
    sig = tuple(values[i] if i < len(values) else None for i in range(compare_len))
    return sig


def main() -> None:
    match = read_table(INPUT_MATCH)
    manual = read_table(INPUT_MANUAL)
    final = read_table(INPUT_FINAL)

    map_rows = get_manual_map_source_rows(INPUT_MAP_REVIEWED) | get_manual_map_source_rows(INPUT_MAP_REVIEW12)

    match_rows = match["rows"]
    manual_rows = manual["rows"]
    final_rows = final["rows"]

    match_presence = {int(row["source_row"]) for row in match_rows}
    manual_presence = {int(row["source_row"]) for row in manual_rows}

    new_rows = [row for row in final_rows if int(row["source_row"]) not in match_presence]

    expected_extra = max(int(final["data_row_count"]) - int(match["data_row_count"]), 0)
    if len(new_rows) < expected_extra:
        taken_sheet_rows = {int(r["sheet_row"]) for r in new_rows}
        appended = [row for row in final_rows if int(row["sheet_row"]) > int(match["data_row_count"]) + 1]
        for row in appended:
            if int(row["sheet_row"]) in taken_sheet_rows:
                continue
            new_rows.append(row)
            taken_sheet_rows.add(int(row["sheet_row"]))
            if len(new_rows) >= expected_extra:
                break

    new_rows = sorted(new_rows, key=lambda x: int(x["sheet_row"]))

    new_rows_out: List[Dict[str, object]] = []
    for row in new_rows:
        appeared_in = "SALES_MATCH_MANUAL" if int(row["source_row"]) in manual_presence else "SALES_MATCH_FINAL"
        new_rows_out.append(
            {
                "source_row": row["source_row"],
                "source_name": row["source_name"],
                "match_status": row["match_status"],
                "match_sku": row["match_sku"],
                "match_master_name": row["match_master_name"],
                "appeared_in_file": appeared_in,
                "sheet_row": row["sheet_row"],
            }
        )

    review_rows = [
        row
        for row in final_rows
        if row.get("match_status") is not None and str(row.get("match_status")).strip().upper() == "REVIEW"
    ]

    match_by_source: Dict[int, Dict[str, object]] = {}
    for row in match_rows:
        src = int(row["source_row"])
        match_by_source.setdefault(src, row)

    final_by_source: Dict[int, Dict[str, object]] = {}
    for row in final_rows:
        src = int(row["source_row"])
        final_by_source.setdefault(src, row)

    compare_len = min(len(match["headers"]), len(final["headers"]))
    changed_outside_map: List[Dict[str, object]] = []
    for src in sorted(match_presence & {int(r["source_row"]) for r in final_rows}):
        if src in map_rows:
            continue
        m_row = match_by_source.get(src)
        f_row = final_by_source.get(src)
        if m_row is None or f_row is None:
            continue
        if row_signature(m_row, compare_len) != row_signature(f_row, compare_len):
            changed_outside_map.append(
                {
                    "source_row": src,
                    "match_status_before": m_row.get("match_status"),
                    "match_status_after": f_row.get("match_status"),
                    "match_sku_before": m_row.get("match_sku"),
                    "match_sku_after": f_row.get("match_sku"),
                }
            )

    wb_out = Workbook()
    ws_summary = wb_out.active
    ws_summary.title = "SUMMARY"
    ws_new = wb_out.create_sheet("NEW_ROWS")
    ws_review = wb_out.create_sheet("REVIEW_3")

    ws_summary.append(["KEY", "VALUE"])

    def add_file_metrics(prefix: str, data: Dict[str, object]) -> None:
        ws_summary.append([f"{prefix}_DATA_ROW_COUNT", data["data_row_count"]])
        ws_summary.append([f"{prefix}_SOURCE_ROW_MIN", data["source_row_min"]])
        ws_summary.append([f"{prefix}_SOURCE_ROW_MAX", data["source_row_max"]])
        ws_summary.append([f"{prefix}_SOURCE_ROW_UNIQUE_COUNT", data["source_row_unique_count"]])
        dups = data["source_row_duplicates"]
        dup_txt = "; ".join(f"{k}:{v}" for k, v in dups) if dups else "NONE"
        ws_summary.append([f"{prefix}_SOURCE_ROW_DUPLICATES", dup_txt])

    add_file_metrics("SALES_MATCH", match)
    add_file_metrics("SALES_MATCH_MANUAL", manual)
    add_file_metrics("SALES_MATCH_FINAL", final)

    ws_summary.append(["ROW_COUNT_DIFF_FINAL_MINUS_MATCH", int(final["data_row_count"]) - int(match["data_row_count"])])
    ws_summary.append(["FINAL_ROWS_NOT_IN_MATCH_COUNT", len(new_rows_out)])
    ws_summary.append(["FINAL_ROWS_NOT_IN_MATCH_SOURCE_ROWS", ", ".join(str(r["source_row"]) for r in new_rows_out) if new_rows_out else "NONE"])
    ws_summary.append(["FINAL_REVIEW_ROW_COUNT", len(review_rows)])
    ws_summary.append(["MANUAL_MAP_APPROVE_OR_NO_MASTER_SOURCE_ROWS", len(map_rows)])
    ws_summary.append(["CHANGED_EXISTING_ROWS_OUTSIDE_MANUAL_MAP_COUNT", len(changed_outside_map)])
    changed_src_txt = ", ".join(str(r["source_row"]) for r in changed_outside_map) if changed_outside_map else "NONE"
    ws_summary.append(["CHANGED_EXISTING_ROWS_OUTSIDE_MANUAL_MAP_SOURCE_ROWS", changed_src_txt])

    ws_summary.append([])
    ws_summary.append(["CHANGED_OUTSIDE_MAP_DETAIL", ""])
    ws_summary.append(["SOURCE_ROW", "STATUS_BEFORE -> STATUS_AFTER | SKU_BEFORE -> SKU_AFTER"])
    if changed_outside_map:
        for item in changed_outside_map:
            ws_summary.append(
                [
                    item["source_row"],
                    f"{item['match_status_before']} -> {item['match_status_after']} | {item['match_sku_before']} -> {item['match_sku_after']}",
                ]
            )
    else:
        ws_summary.append(["NONE", "NONE"])

    ws_new.append(
        [
            "SOURCE_ROW",
            "SOURCE_NAME",
            "MATCH_STATUS",
            "MATCH_SKU",
            "MATCH_MASTER_NAME",
            "APPEARED_IN_FILE",
            "SHEET_ROW_IN_FINAL",
        ]
    )
    for row in new_rows_out:
        ws_new.append(
            [
                row["source_row"],
                row["source_name"],
                row["match_status"],
                row["match_sku"],
                row["match_master_name"],
                row["appeared_in_file"],
                row["sheet_row"],
            ]
        )

    ws_review.append(
        [
            "SOURCE_ROW",
            "SOURCE_NAME",
            "MATCH_STATUS",
            "MATCH_SKU",
            "MATCH_MASTER_NAME",
            "MATCH_REASONS",
            "SHEET_ROW_IN_FINAL",
        ]
    )
    for row in review_rows:
        ws_review.append(
            [
                row["source_row"],
                row["source_name"],
                row["match_status"],
                row["match_sku"],
                row["match_master_name"],
                row["match_reasons"],
                row["sheet_row"],
            ]
        )

    wb_out.save(OUTPUT_PATH)

    print(f"OUTPUT|{OUTPUT_PATH}")
    print(f"DATA_ROWS|MATCH|{match['data_row_count']}")
    print(f"DATA_ROWS|MANUAL|{manual['data_row_count']}")
    print(f"DATA_ROWS|FINAL|{final['data_row_count']}")
    print(f"NEW_ROWS_COUNT|{len(new_rows_out)}")
    print(f"REVIEW_ROWS_COUNT|{len(review_rows)}")
    print(f"CHANGED_OUTSIDE_MAP_COUNT|{len(changed_outside_map)}")


if __name__ == "__main__":
    main()
