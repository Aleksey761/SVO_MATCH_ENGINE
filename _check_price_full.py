from openpyxl import load_workbook

def norm(x):
    return str(x or "").strip()

p = load_workbook("output/price_match.xlsx", data_only=True).active
f = load_workbook("output/_FULL_RESULT_FOR_STOCK.xlsx", data_only=True).active

price = {
    norm(p.cell(r, 1).value)
    for r in range(2, p.max_row + 1)
    if norm(p.cell(r, 1).value)
}

full = {
    norm(f.cell(r, 6).value)
    for r in range(2, f.max_row + 1)
    if norm(f.cell(r, 6).value)
}

print("PRICE_UNIQUE =", len(price))
print("FULL_RESULT_SKU_UNIQUE =", len(full))
print("PRICE_TO_FULL =", len(price & full))
print("PRICE_NOT_IN_FULL =", len(price - full))
print("FULL_NOT_IN_PRICE =", len(full - price))

print("MISSING_PRICE:")
for sku in sorted(price - full):
    print(sku)

print("FULL_NOT_IN_PRICE_LIST:")
for sku in sorted(full - price):
    print(sku)
