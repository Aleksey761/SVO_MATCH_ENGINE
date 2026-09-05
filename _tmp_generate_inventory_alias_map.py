from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook


def main() -> None:
    validation_file = Path("output") / "INVENTORY_ALIAS_VALIDATION.xlsx"
    draft_file = Path("output") / "INVENTORY_ALIAS_MAP_DRAFT.xlsx"
    output_file = Path("output") / "INVENTORY_ALIAS_MAP.xlsx"

    validation_wb = load_workbook(validation_file, data_only=True)
    validation_ws = validation_wb["VALIDATION"]
    validation_idx = {str(validation_ws.cell(1, c).value or "").strip(): c for c in range(1, validation_ws.max_column + 1)}

    approved_names: set[str] = set()
    evidence_by_name: dict[str, str] = {}
    for row in range(2, validation_ws.max_row + 1):
        inventory_name = str(validation_ws.cell(row, validation_idx["InventoryName"]).value or "").strip()
        status = str(validation_ws.cell(row, validation_idx["Status"]).value or "").strip().upper()
        evidence = str(validation_ws.cell(row, validation_idx["Evidence"]).value or "").strip()
        if inventory_name and status == "APPROVED_READY":
            approved_names.add(inventory_name)
            evidence_by_name[inventory_name] = evidence

    draft_wb = load_workbook(draft_file, data_only=True)
    draft_ws = draft_wb["APPROVED_CANDIDATES"]
    draft_idx = {str(draft_ws.cell(1, c).value or "").strip(): c for c in range(1, draft_ws.max_column + 1)}

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "APPROVED_CANDIDATES"
    ws_out.append([
        "InventoryName",
        "MASTER_NAME",
        "SKU",
        "ProductType",
        "Brand",
        "Volume",
        "Variant",
        "Evidence",
        "Status",
    ])

    written = 0
    for row in range(2, draft_ws.max_row + 1):
        inventory_name = str(draft_ws.cell(row, draft_idx["InventoryName"]).value or "").strip()
        if inventory_name not in approved_names:
            continue
        ws_out.append([
            inventory_name,
            str(draft_ws.cell(row, draft_idx["MASTER_NAME"]).value or "").strip(),
            str(draft_ws.cell(row, draft_idx["SKU"]).value or "").strip(),
            str(draft_ws.cell(row, draft_idx["ProductType"]).value or "").strip(),
            str(draft_ws.cell(row, draft_idx["Brand"]).value or "").strip(),
            str(draft_ws.cell(row, draft_idx["Volume"]).value or "").strip(),
            str(draft_ws.cell(row, draft_idx["Variant"]).value or "").strip(),
            evidence_by_name.get(inventory_name, str(draft_ws.cell(row, draft_idx["Evidence"]).value or "").strip()),
            "APPROVED_READY",
        ])
        written += 1

    ws_summary = wb_out.create_sheet("SUMMARY")
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Alias rows", written])

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_file)
    print(f"Generated: {output_file}")
    print(f"Alias rows: {written}")


if __name__ == "__main__":
    main()
