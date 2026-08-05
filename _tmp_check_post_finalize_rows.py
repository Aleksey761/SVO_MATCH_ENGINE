from pathlib import Path

from openpyxl import load_workbook

from svo.loader import Loader
from svo.matcher import Matcher
from svo.normalizer import Normalizer
from svo.reporter import Reporter
from svo.sales_loader import SalesLoader

base = Path('.')
loader = Loader()
master_file, _arrival_file, sales_file, _arrival_date, _sales_date = loader.discover_workbooks(base / 'data', require_sales=True)
master = loader.load_master(master_file)
items = SalesLoader().load(sales_file)
normalizer = Normalizer()
for item in items:
    normalizer.normalize(item)
Matcher(master).match_all(items)

pre = base / 'output' / '_tmp_pre.xlsx'
post = base / 'output' / '_tmp_post.xlsx'
Reporter().write_matched_document(sales_file, items, pre)
post.write_bytes(pre.read_bytes())
Reporter().finalize_result_by_master(post, master, master_file=master_file)

ws = load_workbook(post, data_only=True).active
for r in range(195, 205):
    print(r, ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value)

pre.unlink(missing_ok=True)
post.unlink(missing_ok=True)
