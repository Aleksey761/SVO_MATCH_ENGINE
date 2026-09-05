from openpyxl import Workbook, load_workbook
from pathlib import Path
import json

base = Path(".")
review_path = base / "output" / "SALES_PRODUCT_TYPE_REVIEW.xlsx"
master_path = base / "output" / "MASTER_DATASET.xlsx"
out_path = base / "output" / "SALES_MANUAL_MAPPING.xlsx"
report_path = base / "_tmp_sales_manual_mapping_report.json"

review_wb = load_workbook(review_path, data_only=True)
review_ws = review_wb["REVIEW_79"]

master_wb = load_workbook(master_path, data_only=True)
master_ws = master_wb.active

review_headers = [
    "" if review_ws.cell(1, c).value is None else str(review_ws.cell(1, c).value).strip()
    for c in range(1, review_ws.max_column + 1)
]
master_headers = [
    "" if master_ws.cell(1, c).value is None else str(master_ws.cell(1, c).value).strip()
    for c in range(1, master_ws.max_column + 1)
]

review_idx = {h: i + 1 for i, h in enumerate(review_headers) if h}
master_idx = {h: i + 1 for i, h in enumerate(master_headers) if h}

# Required review fields from task phrasing.
required_review = ["НАИМЕНОВАНИЕ", "POSSIBLE_MASTER_NAME", "POSSIBLE_SKU"]
missing_review = [h for h in required_review if h not in review_idx]
if missing_review:
    raise RuntimeError(f"Missing required columns in REVIEW_79: {missing_review}. Available: {review_headers}")

# Master sheet field detection with exact known aliases only.
master_name_col = None
for cand in ["MASTER_NAME", "НАИМЕНОВАНИЕ", "NAME"]:
    if cand in master_idx:
        master_name_col = cand
        break
if master_name_col is None:
    raise RuntimeError(f"Could not locate master name column. Available: {master_headers}")

master_sku_col = None
for cand in ["SKU", "MASTER_SKU", "КОД", "АРТИКУЛ"]:
    if cand in master_idx:
        master_sku_col = cand
        break
if master_sku_col is None:
    raise RuntimeError(f"Could not locate master SKU column. Available: {master_headers}")

# Build exact pair lookup from master data.
master_pairs = {}
for r in range(2, master_ws.max_row + 1):
    name_val = master_ws.cell(r, master_idx[master_name_col]).value
    sku_val = master_ws.cell(r, master_idx[master_sku_col]).value
    if name_val is None or sku_val is None:
        continue
    name_s = str(name_val).strip()
    sku_s = str(sku_val).strip()
    if not name_s or not sku_s:
        continue
    key = (name_s, sku_s)
    if key not in master_pairs:
        # Keep exact text from master dataset.
        master_pairs[key] = (name_val, sku_val)

out_wb = Workbook()
out_ws = out_wb.active
out_ws.title = "MANUAL_MAPPING"
out_headers = ["SOURCE_ROW", "SOURCE_NAME", "MASTER_NAME", "SKU", "DECISION", "COMMENT"]
out_ws.append(out_headers)

approve_rows = []
counts = {"TOTAL": 0, "APPROVE": 0, "REVIEW": 0, "NO_MASTER": 0}

source_row_col = review_idx.get("SOURCE_ROW")
name_col = review_idx["НАИМЕНОВАНИЕ"]
pos_name_col = review_idx["POSSIBLE_MASTER_NAME"]
pos_sku_col = review_idx["POSSIBLE_SKU"]

for r in range(2, review_ws.max_row + 1):
    counts["TOTAL"] += 1

    source_row_val = review_ws.cell(r, source_row_col).value if source_row_col else (r - 1)
    source_name_val = review_ws.cell(r, name_col).value
    possible_name_val = review_ws.cell(r, pos_name_col).value
    possible_sku_val = review_ws.cell(r, pos_sku_col).value

    source_name_s = "" if source_name_val is None else str(source_name_val)
    pos_name_s = "" if possible_name_val is None else str(possible_name_val).strip()
    pos_sku_s = "" if possible_sku_val is None else str(possible_sku_val).strip()

    master_name_out = ""
    sku_out = ""
    comment = ""

    if pos_name_s and pos_sku_s and (pos_name_s, pos_sku_s) in master_pairs:
        decision = "APPROVE"
        master_name_exact, sku_exact = master_pairs[(pos_name_s, pos_sku_s)]
        master_name_out = master_name_exact
        sku_out = sku_exact
        counts["APPROVE"] += 1
        approve_rows.append([
            source_row_val,
            source_name_s,
            "" if master_name_exact is None else str(master_name_exact),
            "" if sku_exact is None else str(sku_exact),
        ])
    elif (not pos_name_s) and (not pos_sku_s):
        decision = "NO_MASTER"
        counts["NO_MASTER"] += 1
    else:
        decision = "REVIEW"
        counts["REVIEW"] += 1

    out_ws.append([
        source_row_val,
        source_name_s,
        master_name_out,
        sku_out,
        decision,
        comment,
    ])

if counts["TOTAL"] != 79:
    raise RuntimeError(f"Expected 79 rows in REVIEW_79, got {counts['TOTAL']}")

summary_ws = out_wb.create_sheet("SUMMARY")
summary_ws.append(["TOTAL", counts["TOTAL"]])
summary_ws.append(["APPROVE", counts["APPROVE"]])
summary_ws.append(["REVIEW", counts["REVIEW"]])
summary_ws.append(["NO_MASTER", counts["NO_MASTER"]])

out_wb.save(out_path)

report = {
    "output_path": out_path.as_posix(),
    "counts": counts,
    "approve_rows": approve_rows,
    "review_headers": review_headers,
    "master_headers": master_headers,
    "master_name_col": master_name_col,
    "master_sku_col": master_sku_col,
}
report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

print("OUTPUT_PATH|" + out_path.as_posix())
print("TOTAL|" + str(counts["TOTAL"]))
print("APPROVE|" + str(counts["APPROVE"]))
print("REVIEW|" + str(counts["REVIEW"]))
print("NO_MASTER|" + str(counts["NO_MASTER"]))
print("REPORT_PATH|" + report_path.as_posix())
