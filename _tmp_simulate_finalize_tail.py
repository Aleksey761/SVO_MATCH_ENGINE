from pathlib import Path

from openpyxl import load_workbook

from svo.loader import Loader
from svo.matcher import Matcher
from svo.normalizer import Normalizer
from svo.reporter import Reporter
from svo.sales_loader import SalesLoader


def norm(value: object) -> str:
    return str(value or "").strip().upper()


base = Path('.')
loader = Loader()
master_file, _arrival_file, sales_file, _arrival_date, _sales_date = loader.discover_workbooks(base / 'data', require_sales=True)
master_items = loader.load_master(master_file)
items = SalesLoader().load(sales_file)
normalizer = Normalizer()
for item in items:
    normalizer.normalize(item)
Matcher(master_items).match_all(items)

tmp = base / 'output' / '_tmp_pre.xlsx'
Reporter().write_matched_document(sales_file, items, tmp)
ws = load_workbook(tmp, data_only=False).active

header_cells = [cell.value for cell in ws[1]]
header_to_col = {}
for idx, value in enumerate(header_cells, start=1):
    normalized = norm(value)
    if normalized:
        header_to_col[normalized] = idx

match_status_col = header_to_col['MATCH_STATUS']
match_sku_col = header_to_col['MATCH_SKU']
master_by_sku = {item.sku: item for item in master_items}
master_name_by_sku = Reporter()._load_master_name_by_sku(master_file)
master_value_by_result_header = {
    'SKU': lambda item, _name: item.sku,
    'ТИП ТОВАРА': lambda item, _name: item.category,
    'БРЕНД': lambda item, _name: item.brand,
    'VARIANT': lambda item, _name: item.variant,
    'ОБЪЕМ': lambda item, _name: item.volume,
    'НАИМЕНОВАНИЕ': lambda item, name: name,
}

sklad1_col = header_to_col.get('СКЛАД 1')
if sklad1_col is not None and sklad1_col > 7:
    for col, title in enumerate(('№', 'SKU', 'Наименование', 'Тип товара', 'Бренд', 'Variant', 'Объем'), start=1):
        ws.cell(row=1, column=col, value=title)

    header_cells = [cell.value for cell in ws[1]]
    header_to_col = {}
    for idx, value in enumerate(header_cells, start=1):
        normalized = norm(value)
        if normalized:
            header_to_col[normalized] = idx

actionable_rows = []
matched_rows_by_sku = {}
review_rows = []
max_col = ws.max_column
for row_number in range(2, ws.max_row + 1):
    status_value = norm(ws.cell(row=row_number, column=match_status_col).value)
    if status_value not in {'MATCH', 'REVIEW'}:
        continue
    actionable_rows.append(row_number)
    row_values = [ws.cell(row=row_number, column=col).value for col in range(1, max_col + 1)]
    row_values[match_status_col - 1] = status_value
    if status_value == 'MATCH':
        sku = str(ws.cell(row=row_number, column=match_sku_col).value or '').strip()
        master = master_by_sku.get(sku)
        if master is not None:
            master_name = master_name_by_sku.get(master.sku)
            for result_header, getter in master_value_by_result_header.items():
                target_col = header_to_col.get(result_header)
                if target_col is not None:
                    row_values[target_col - 1] = getter(master, master_name)
            row_values[match_sku_col - 1] = master.sku
            matched_rows_by_sku.setdefault(master.sku, []).append(row_values)
    else:
        review_rows.append(row_values)

ordered_rows = []
for master in master_items:
    ordered_rows.extend(matched_rows_by_sku.get(master.sku, []))
ordered_rows.extend(review_rows)

match_number = 1
number_col = header_to_col.get('№')
for row_values in ordered_rows:
    status_value = norm(row_values[match_status_col - 1])
    if number_col is not None:
        if status_value == 'MATCH':
            row_values[number_col - 1] = match_number
            match_number += 1
        else:
            row_values[number_col - 1] = None
    if status_value != 'MATCH':
        for header_name in ('SKU', 'ТИП ТОВАРА', 'БРЕНД', 'VARIANT', 'ОБЪЕМ', 'НАИМЕНОВАНИЕ'):
            target_col = header_to_col.get(header_name)
            if target_col is not None:
                row_values[target_col - 1] = None

for row_no, row_values in zip(actionable_rows[-4:], ordered_rows[-4:]):
    print('target', row_no, 'status', norm(row_values[match_status_col - 1]), 'col1', row_values[0], 'col2', row_values[1], 'col3', row_values[2], 'tech_sku', row_values[match_sku_col - 1])

tmp.unlink(missing_ok=True)
