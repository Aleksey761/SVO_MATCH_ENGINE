from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


WORKSPACE = Path("f:/SVO/AI/SVO_MATCH_ENGINE")
SOURCE_FILE = WORKSPACE / "output" / "FINANCE_RESULT.xlsx"
PRICE_MATCH_FILE = WORKSPACE / "output" / "price_match.xlsx"
OUTPUT_FILE = WORKSPACE / "output" / "INVENTORY_ALIAS_FINANCE_BEFORE_AFTER.xlsx"
REPORT_FILE = WORKSPACE / "output" / "INVENTORY_ALIAS_FINANCE_BEFORE_AFTER.txt"
BASELINE_WITH_STOCK = 79
BASELINE_FINANCE_COMPLETE = 63


def _normalize_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    return int(str(value).strip())


def _load_summary_metrics(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, data_only=True)
    if "FINANCE_LINEAGE" not in workbook.sheetnames:
        raise KeyError("Sheet FINANCE_LINEAGE not found in output/FINANCE_RESULT.xlsx")

    sheet = workbook["FINANCE_LINEAGE"]
    metrics: dict[str, int] = {}
    wanted = {
        "SKU with stock",
        "Finance-complete SKU",
        "Inventory alias rows applied",
        "Total SKU",
        "Missing MASTER for PRICE SKU",
    }
    for row_idx in range(1, sheet.max_row + 1):
        key = sheet.cell(row=row_idx, column=1).value
        if key in wanted:
            metrics[str(key)] = _normalize_int(sheet.cell(row=row_idx, column=2).value)

    missing = [name for name in wanted if name not in metrics]
    if missing:
        raise KeyError(f"Missing FINANCE_LINEAGE summary metrics: {', '.join(missing)}")
    return metrics


def _load_price_match_count(path: Path) -> int:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [str(sheet.cell(row=1, column=col).value or "").strip() for col in range(1, sheet.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
    status_col = idx.get("Status")
    if status_col is None:
        raise KeyError("Status column not found in output/price_match.xlsx")

    return sum(
        1
        for row_idx in range(2, sheet.max_row + 1)
        if str(sheet.cell(row=row_idx, column=status_col).value or "").strip().upper() == "MATCH"
    )


def _write_summary_workbook(path: Path, metrics: dict[str, int], *, price_match_count: int) -> None:
    after_with_stock = metrics["SKU with stock"]
    after_finance_complete = metrics["Finance-complete SKU"]
    alias_rows = metrics["Inventory alias rows applied"]
    regressions = after_with_stock < BASELINE_WITH_STOCK or after_finance_complete < BASELINE_FINANCE_COMPLETE

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SUMMARY"
    sheet.append(["Metric", "Value"])
    sheet.append(["Alias applied", "YES" if alias_rows > 0 else "NO"])
    sheet.append(["Stock improved", "YES" if after_with_stock > BASELINE_WITH_STOCK else "NO"])
    sheet.append(["Finance complete improved", "YES" if after_finance_complete > BASELINE_FINANCE_COMPLETE else "NO"])
    sheet.append(["Regressions", "YES" if regressions else "NO"])
    sheet.append(["WITH_STOCK BEFORE", BASELINE_WITH_STOCK])
    sheet.append(["WITH_STOCK AFTER", after_with_stock])
    sheet.append(["FINANCE_COMPLETE BEFORE", BASELINE_FINANCE_COMPLETE])
    sheet.append(["FINANCE_COMPLETE AFTER", after_finance_complete])
    sheet.append(["Inventory alias rows applied", alias_rows])
    sheet.append(["PRICE_MATCH MATCH count", price_match_count])
    sheet.append(["PRICE_MATCH scope warning", "CURRENT_SCOPE_IS_NOT_168_BASELINE" if price_match_count != 168 else "NO"])
    sheet.append(["Total SKU", metrics["Total SKU"]])
    sheet.append(["Missing MASTER for PRICE SKU", metrics["Missing MASTER for PRICE SKU"]])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    if not SOURCE_FILE.exists():
        raise FileNotFoundError("output/FINANCE_RESULT.xlsx was not created")

    metrics = _load_summary_metrics(SOURCE_FILE)
    price_match_count = _load_price_match_count(PRICE_MATCH_FILE)
    _write_summary_workbook(OUTPUT_FILE, metrics, price_match_count=price_match_count)

    regressions = (
        metrics["SKU with stock"] < BASELINE_WITH_STOCK
        or metrics["Finance-complete SKU"] < BASELINE_FINANCE_COMPLETE
    )
    lines = [
        f"AFTER_WITH_STOCK={metrics['SKU with stock']}",
        f"AFTER_FINANCE_COMPLETE={metrics['Finance-complete SKU']}",
        f"AFTER_ALIAS_ROWS={metrics['Inventory alias rows applied']}",
        f"AFTER_TOTAL_SKU={metrics['Total SKU']}",
        f"AFTER_MISSING_MASTER={metrics['Missing MASTER for PRICE SKU']}",
        f"PRICE_MATCH_MATCH_COUNT={price_match_count}",
        f"REGRESSIONS={'YES' if regressions else 'NO'}",
        f"WORKBOOK_CREATED={'YES' if OUTPUT_FILE.exists() else 'NO'}",
    ]
    REPORT_FILE.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()