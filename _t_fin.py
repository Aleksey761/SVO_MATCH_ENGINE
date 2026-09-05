from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook


root = Path("f:/SVO/AI/SVO_MATCH_ENGINE")
source = root / "output" / "FINANCE_RESULT.xlsx"
target = root / "output" / "INVENTORY_ALIAS_FINANCE_BEFORE_AFTER.xlsx"
report = root / "output" / "INVENTORY_ALIAS_FINANCE_BEFORE_AFTER.txt"

before_with_stock = 79
before_finance_complete = 63

wb = load_workbook(source, data_only=True)
ws = wb["FINANCE_LINEAGE"]

needed = [
    "SKU with stock",
    "Finance-complete SKU",
    "Inventory alias rows applied",
    "Total SKU",
    "Missing MASTER for PRICE SKU",
]
metrics: dict[str, int] = {}
for row in ws.iter_rows(values_only=True):
    if row and row[0] in needed:
        metrics[str(row[0])] = int(row[1] or 0)

after_with_stock = metrics["SKU with stock"]
after_finance_complete = metrics["Finance-complete SKU"]
alias_rows = metrics["Inventory alias rows applied"]
regressions = after_with_stock < before_with_stock or after_finance_complete < before_finance_complete

out_wb = Workbook()
out_ws = out_wb.active
out_ws.title = "SUMMARY"
out_ws.append(["Metric", "Value"])
out_ws.append(["Alias applied", "YES" if alias_rows > 0 else "NO"])
out_ws.append(["Stock improved", "YES" if after_with_stock > before_with_stock else "NO"])
out_ws.append(["Finance complete improved", "YES" if after_finance_complete > before_finance_complete else "NO"])
out_ws.append(["Regressions", "YES" if regressions else "NO"])
out_ws.append(["WITH_STOCK BEFORE", before_with_stock])
out_ws.append(["WITH_STOCK AFTER", after_with_stock])
out_ws.append(["FINANCE_COMPLETE BEFORE", before_finance_complete])
out_ws.append(["FINANCE_COMPLETE AFTER", after_finance_complete])
out_ws.append(["Inventory alias rows applied", alias_rows])
out_ws.append(["Total SKU", metrics["Total SKU"]])
out_ws.append(["Missing MASTER for PRICE SKU", metrics["Missing MASTER for PRICE SKU"]])
out_wb.save(target)

lines = [
    f"AFTER_WITH_STOCK={after_with_stock}",
    f"AFTER_FINANCE_COMPLETE={after_finance_complete}",
    f"AFTER_ALIAS_ROWS={alias_rows}",
    f"AFTER_TOTAL_SKU={metrics['Total SKU']}",
    f"AFTER_MISSING_MASTER={metrics['Missing MASTER for PRICE SKU']}",
    f"REGRESSIONS={'YES' if regressions else 'NO'}",
    f"WORKBOOK_CREATED={'YES' if target.exists() else 'NO'}",
]
report.write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines))