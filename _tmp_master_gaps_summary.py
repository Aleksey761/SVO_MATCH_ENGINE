from pathlib import Path
from openpyxl import load_workbook

path = Path('output/MASTER_GAPS.xlsx')
wb = load_workbook(path, data_only=True)
ws = wb['SUMMARY']
for r in range(1, ws.max_row + 1):
    left = ws.cell(row=r, column=1).value
    right = ws.cell(row=r, column=2).value
    if left is None and right is None:
        continue
    print(f"{left}: {right}")
