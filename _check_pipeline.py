from openpyxl import load_workbook

def n(x):
    return str(x or "").strip()

p = load_workbook("output/price_match.xlsx", data_only=True).active
f = load_workbook("output/_FULL_RESULT_FOR_STOCK.xlsx", data_only=True).active
s = load_workbook("output/STOCK_MATCH.xlsx", data_only=True)["STOCK_RESULT"]

price = {n(p.cell(r,1).value) for r in range(2,p.max_row+1) if n(p.cell(r,1).value)}
full = {n(f.cell(r,6).value) for r in range(2,f.max_row+1) if n(f.cell(r,6).value)}
stock = {n(s.cell(r,1).value) for r in range(2,s.max_row+1) if n(s.cell(r,1).value)}

print("PRICE =", len(price))
print("FULL_RESULT =", len(full))
print("STOCK =", len(stock))
print("PRICE_TO_FULL =", len(price & full))
print("FULL_TO_STOCK =", len(full & stock))
print("PRICE_TO_STOCK =", len(price & stock))
print("PRICE_MISSING_FULL =", len(price - full))
print("FULL_MISSING_STOCK =", len(full - stock))
print("PRICE_MISSING_STOCK =", len(price - stock))

print("PRICE_MISSING_FULL_LIST:")
print(",".join(sorted(price-full)))

print("FULL_MISSING_STOCK_LIST:")
print(",".join(sorted(full-stock)))

